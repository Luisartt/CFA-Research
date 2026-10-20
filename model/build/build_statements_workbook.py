"""Build model/BIMBOA_statements_2021-2030_v1.xlsx.

Full as-reported consolidated statements of Grupo Bimbo (MXN millions):
income statement, OCI, balance sheet, cash flow and equity roll-forward,
FY2021A-FY2025A typed from the audited filings, plus empty FY2026E-FY2030E
columns for the team to project and balance. Subtotals are live formulas in
every column; the Checks sheet tests that the statements tie to each other.

Sources (PDF page numbers):
  IA25 = filings/annual/IA25_GB_EEFF_V10.pdf  (FY2023-FY2025; BS p.12, IS p.13,
         OCI p.14, equity p.15, CF p.16)
  RA23 = filings/bmv/Reporte Anual 2023_XBRL_VF.PDF  (FY2021-FY2022; BS p.256,
         IS p.257, OCI p.258, equity p.259, CF p.260)
Rule (data/adjustments.md): the latest report that presents a year wins, so
FY2023 comes from IA25 and FY2021-FY2022 from RA23 (restated for Ricolino).
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "model" / "BIMBOA_statements_2021-2030_v1.xlsx"

HIST = ["FY2021A", "FY2022A", "FY2023A", "FY2024A", "FY2025A"]
FCST = ["FY2026E", "FY2027E", "FY2028E", "FY2029E", "FY2030E"]
YEARS = HIST + FCST
COLS = [get_column_letter(i) for i in range(2, 12)]  # B..K
HC, FC = COLS[:5], COLS[5:]
TAG_COL, SRC_COL = "L", "M"

SRC = {
    "IS": "RA23 p.257 (FY21-22); IA25 p.13 (FY23-25)",
    "OCI": "RA23 p.258 (FY21-22); IA25 p.14 (FY23-25)",
    "BS": "RA23 p.256 (FY21-22); IA25 p.12 (FY23-25)",
    "CF": "RA23 p.260 (FY21-22); IA25 p.16 (FY23-25)",
    "EQ": "RA23 p.259 (FY21-22); IA25 p.15 (FY23-25)",
}

# Row spec: (key, label, kind, payload, note)
#   kind "h"   section header, "i" typed input (payload = 5 historical values),
#   "f"   formula in every year column (payload = template with {c} = column,
#         {p} = previous column, and {key} = row number of another line),
#   "fh"  formula in history only, forecast left as input (payload = template),
#   "b"   blank spacer.
IS_ROWS = [
    ("h1", "Continuing operations", "h", None, ""),
    ("sales", "Net sales", "i", [338792, 398706, 399879, 408335, 426952], ""),
    ("cogs", "Cost of sales", "i", [157703, 193325, 194369, 193407, 203465], "Positive = expense"),
    ("gp", "Gross profit", "f", "={c}{sales}-{c}{cogs}", ""),
    ("dist", "Distribution and selling expenses", "i", [124422, 140455, 139477, 147628, 154866], ""),
    ("admin", "Administrative expenses", "i", [23147, 26172, 25860, 27536, 29023], ""),
    ("integ", "Integration expenses", "i", [724, 367, 288, 317, 541], ""),
    ("oth", "Other (income) expenses, net", "i", [216, -15309, 4430, 6298, 4911],
     "FY2022 includes the non-cash gain on multi-employer pension plans (MEPP)"),
    ("opex", "Total general expenses", "f", "=SUM({c}{dist}:{c}{oth})", ""),
    ("ebit", "Operating income (EBIT)", "f", "={c}{gp}-{c}{opex}", ""),
    ("b1", "", "b", None, ""),
    ("h2", "Comprehensive financing result", "h", None, ""),
    ("intexp", "Interest expense", "i", [7823, 8049, 10006, 13100, 14464], "Includes IFRS 16 lease interest"),
    ("intinc", "Interest income", "i", [-363, -740, -809, -924, -897], "Negative = income (as in the filing)"),
    ("fx", "Foreign exchange loss, net", "i", [538, 995, 1323, 506, 384], ""),
    ("mon", "Monetary position loss (gain)", "i", [-25, -34, -73, -103, 50], ""),
    ("cif", "Total comprehensive financing result", "f", "=SUM({c}{intexp}:{c}{mon})", ""),
    ("assoc", "Share of results of associates", "i", [247, 452, 316, 464, -2], ""),
    ("pbt", "Income before taxes, continuing operations", "f", "={c}{ebit}-{c}{cif}+{c}{assoc}", ""),
    ("tax", "Income taxes", "i", [8726, 14381, 8386, 6797, 7197], ""),
    ("nicont", "Net income from continuing operations", "f", "={c}{pbt}-{c}{tax}", ""),
    ("disc", "Net income from discontinued operations, after tax", "i", [1254, 16988, -16, -194, 0],
     "Ricolino (sold 1-Nov-2022)"),
    ("ni", "Consolidated net income", "f", "={c}{nicont}+{c}{disc}", ""),
    ("b2", "", "b", None, ""),
    ("h3", "Attribution", "h", None, ""),
    ("nci", "Non-controlling interest", "i", [1466, 1575, 1445, 1499, 1813],
     "FY21-22 include discontinued NCI (30; 402)"),
    ("nictrl", "Controlling interest", "f", "={c}{ni}-{c}{nci}", ""),
    ("shares", "Weighted-average shares (millions)", "i", [4487.268, 4448.693, 4418.137, 4357.250, 4310.531],
     "Filing reports thousands; basic = diluted"),
    ("eps", "Basic EPS (MXN)", "f", "=IFERROR({c}{nictrl}/{c}{shares},0)", ""),
    ("b3", "", "b", None, ""),
    ("h4", "Memo", "h", None, ""),
    ("da", "D&A, continuing (segment note)", "i", [16126, 18038, 18929, 23051, 24838],
     "RA23 p.358-359; IA25 p.85-86 (data/financials.csv)"),
    ("ebitda", "EBITDA (EBIT + D&A)", "f", "={c}{ebit}+{c}{da}", ""),
]

OCI_ROWS = [
    ("ni", "Consolidated net income", "i", [17382, 48485, 16922, 14043, 12946], "Must equal IS consolidated net income"),
    ("b1", "", "b", None, ""),
    ("h1", "Items that will not be reclassified to profit or loss", "h", None, ""),
    ("eqv", "Valuation of equity financial instruments", "i", [-112, -32, -236, -43, -118], ""),
    ("act", "Actuarial gains (losses) on defined-benefit plans", "i", [1543, 800, -349, 802, -433], ""),
    ("tax1", "Income tax", "i", [-432, -207, 164, -540, 369], ""),
    ("sub1", "Subtotal", "f", "=SUM({c}{eqv}:{c}{tax1})", ""),
    ("h2", "Items that may be reclassified to profit or loss", "h", None, ""),
    ("nih", "Net investment hedge", "i", [-723, 667, 1100, -3407, 2096], ""),
    ("cta", "Translation of foreign operations", "i", [2114, -9862, -15635, 15551, -10811], ""),
    ("cfh", "Cash-flow hedges, net change", "i", [2988, -2213, 938, 287, -1800], ""),
    ("tax2", "Income tax", "i", [-937, 142, -772, 1164, -348], ""),
    ("sub2", "Subtotal", "f", "=SUM({c}{nih}:{c}{tax2})", ""),
    ("oci", "Total other comprehensive income", "f", "={c}{sub1}+{c}{sub2}", ""),
    ("ci", "Consolidated comprehensive income", "f", "={c}{ni}+{c}{oci}", ""),
    ("cinci", "  Attributable to non-controlling interest", "i", [1470, 1155, 1540, 1677, 2012], ""),
    ("cictrl", "  Attributable to controlling interest", "f", "={c}{ci}-{c}{cinci}", ""),
]

BS_ROWS = [
    ("h1", "Current assets", "h", None, ""),
    ("cash", "Cash and cash equivalents", "i", [8748, 12313, 6353, 8057, 8535], "Must equal CF ending cash"),
    ("ar", "Trade and other receivables, net", "i", [27170, 32185, 33269, 36659, 35782], ""),
    ("inv", "Inventories", "i", [13710, 16989, 16120, 18832, 17151], ""),
    ("prep", "Prepaid expenses", "i", [2296, 928, 1378, 1371, 1798], ""),
    ("derca", "Derivative financial instruments", "i", [1293, 38, 119, 1417, 60], ""),
    ("coll", "Collateral deposits on derivatives", "i", [0, 879, 1519, 577, 459], ""),
    ("hfs", "Assets held for sale", "i", [194, 138, 156, 267, 886], ""),
    ("tca", "Total current assets", "f", "=SUM({c}{cash}:{c}{hfs})", ""),
    ("ppe", "Property, plant and equipment, net", "i", [103891, 113505, 129156, 155376, 157423], ""),
    ("rou", "Right-of-use assets, net", "i", [30754, 28477, 25848, 34220, 32757], ""),
    ("assoc", "Investment in associates", "i", [3891, 4398, 5031, 7531, 7018], ""),
    ("dernca", "Derivative financial instruments", "i", [1962, 2005, 468, 784, 657], ""),
    ("dta", "Deferred income taxes", "i", [7861, 3962, 7313, 6884, 6908], ""),
    ("intang", "Intangible assets, net", "i", [56965, 59062, 55101, 65808, 62947], ""),
    ("gw", "Goodwill", "i", [74565, 68085, 62000, 75100, 72568], ""),
    ("otha", "Other assets, net", "i", [4340, 4800, 4271, 3921, 4856], ""),
    ("ta", "TOTAL ASSETS", "f", "={c}{tca}+SUM({c}{ppe}:{c}{otha})", ""),
    ("b1", "", "b", None, ""),
    ("h2", "Current liabilities", "h", None, ""),
    ("std", "Current portion of long-term debt", "i", [10625, 6435, 12932, 4862, 12635], ""),
    ("ap", "Trade payables", "i", [35752, 44058, 40409, 40815, 39074], ""),
    ("othp", "Other payables and accrued liabilities", "i", [24102, 24198, 22072, 27280, 26395], ""),
    ("leasec", "Lease liabilities", "i", [5793, 5930, 5751, 7140, 7076], ""),
    ("relp", "Payables to related parties", "i", [1527, 1240, 1412, 1260, 1154], ""),
    ("taxp", "Income taxes payable", "i", [708, 4604, 665, 1535, 1059], ""),
    ("ptu", "Employee profit sharing (PTU)", "i", [1692, 1811, 1787, 1752, 1592], ""),
    ("dercl", "Derivative financial instruments", "i", [169, 1458, 2164, 511, 1588], ""),
    ("othst", "Other short-term liabilities", "i", [392, 0, 0, 0, 0], ""),
    ("tcl", "Total current liabilities", "f", "=SUM({c}{std}:{c}{othst})", ""),
    ("ltd", "Long-term debt", "i", [82230, 77619, 97003, 146043, 141028], ""),
    ("leasenc", "Lease liabilities", "i", [25356, 23255, 21064, 28661, 27714], ""),
    ("derncl", "Derivative financial instruments", "i", [67, 590, 1888, 689, 2314], ""),
    ("empb", "Employee benefits", "i", [30712, 9382, 7217, 6635, 6507], "FY2022 drop = MEPP settlement"),
    ("dtl", "Deferred income taxes", "i", [7087, 7575, 9224, 8891, 8415], ""),
    ("othlt", "Other long-term liabilities", "i", [9822, 12007, 12886, 13032, 12999], ""),
    ("tl", "TOTAL LIABILITIES", "f", "={c}{tcl}+SUM({c}{ltd}:{c}{othlt})", ""),
    ("b2", "", "b", None, ""),
    ("h3", "Stockholders' equity", "h", None, ""),
    ("cap", "Capital stock", "i", [4021, 3984, 3946, 3887, 3866], ""),
    ("re", "Retained earnings", "i", [73384, 111486, 118718, 122763, 127596], ""),
    ("perp", "Other equity financial instrument (perpetual notes)", "i", [8867, 8098, 0, 0, 0],
     "Reclassified to debt in 2023"),
    ("ctab", "Cumulative translation adjustment", "i", [10297, 1132, -13832, -868, -10377], ""),
    ("actb", "Actuarial gains on defined-benefit plans", "i", [696, 1152, 734, 1315, 1008], ""),
    ("eqvb", "Valuation of equity financial instruments", "i", [-742, -690, -857, -1011, -1132], ""),
    ("cfhb", "Unrealized (losses) gains on cash-flow hedges", "i", [577, -1051, -387, -142, -1449], ""),
    ("ctrl", "Controlling interest", "f", "=SUM({c}{cap}:{c}{cfhb})", ""),
    ("ncib", "Non-controlling interest", "i", [4506, 3491, 3306, 1754, 743], ""),
    ("te", "TOTAL EQUITY", "f", "={c}{ctrl}+{c}{ncib}", ""),
    ("tle", "TOTAL LIABILITIES AND EQUITY", "f", "={c}{tl}+{c}{te}", ""),
    ("b3", "", "b", None, ""),
    ("chk", "Check: total assets - (liabilities + equity)", "f", "={c}{ta}-{c}{tle}", "Must be 0"),
]

CF_ROWS = [
    ("h1", "Operating activities", "h", None, ""),
    ("pbtc", "Income before taxes, continuing operations", "i", [24854, 45878, 25324, 21034, 20143], "Must equal IS"),
    ("pbtd", "Income (loss) before taxes, discontinued operations", "i", [1499, 23516, -23, -277, 0], ""),
    ("pbt", "Income before taxes", "f", "={c}{pbtc}+{c}{pbtd}", ""),
    ("h1a", "Adjustments for:", "h", None, ""),
    ("da", "Depreciation and amortization", "i", [16375, 18282, 18929, 23051, 24838],
     "FY21-22 include discontinued D&A (249; 244)"),
    ("ppel", "Loss (gain) on sale of property, plant and equipment", "i", [-157, 650, 751, -192, 68], ""),
    ("cas", "Write-off of PP&E due to casualty", "i", [379, 0, 0, 0, 0], ""),
    ("assoc", "Share of results of associates", "i", [-247, -452, -316, -464, 2], ""),
    ("imp", "Impairment of long-lived assets", "i", [694, 1046, 383, 249, 445], ""),
    ("dgain", "Gain on disposal of discontinued operation", "i", [0, -22038, 0, 277, 0], ""),
    ("mepp", "Multi-employer pension plans and other long-term liabilities", "i", [-2247, -19010, 0, -982, 27], ""),
    ("svc", "Current service cost", "i", [1128, 1013, 837, 942, 909], ""),
    ("intexp", "Interest expense", "i", [7884, 8049, 10006, 13100, 14464], ""),
    ("intinc", "Interest income", "i", [-373, -740, -809, -924, -897], ""),
    ("stl", "Short-term and low-value lease expenses", "i", [2183, 2519, 3968, 4526, 4964], ""),
    ("h1b", "Changes in assets and liabilities:", "h", None, ""),
    ("dar", "Trade and other receivables", "i", [666, -6647, -4206, -319, 601], ""),
    ("dinv", "Inventories", "i", [-2320, -4163, -1078, -1104, 901], ""),
    ("dprep", "Prepaid expenses", "i", [-365, 1261, -625, 189, -594], ""),
    ("dap", "Trade payables", "i", [8286, 9920, -851, -2825, -1108], ""),
    ("doth", "Other payables and accrued liabilities", "i", [-1025, -3395, 390, -3827, -560], ""),
    ("drel", "Payables to related parties", "i", [209, 287, -172, 153, 105], ""),
    ("taxpd", "Income taxes paid", "i", [-7578, -11824, -13831, -6472, -8955], ""),
    ("dcoll", "Collateral deposits on derivatives", "i", [6, -1271, -640, 942, 118], ""),
    ("dptu", "Employee profit sharing (PTU)", "i", [675, 122, -21, -36, -159], ""),
    ("demp", "Employee benefits and social security", "i", [-2567, -754, -2637, -2608, -2595], ""),
    ("ddisc", "Dividends receivable from discontinued operation", "i", [0, -879, 0, 0, 0], ""),
    ("stlp", "Short-term and low-value lease payments", "i", [-2183, -2519, -3968, -4526, -4964], ""),
    ("cfo", "Net cash from operating activities", "f", "=SUM({c}{pbt}:{c}{stlp})", ""),
    ("b1", "", "b", None, ""),
    ("h2", "Investing activities", "h", None, ""),
    ("capex", "Purchases of property, plant and equipment", "i", [-20671, -28669, -34754, -29402, -22530], ""),
    ("acq", "Business acquisitions and NCI, net of cash acquired", "i", [-10637, -6520, -6548, -5988, -11294], ""),
    ("divst", "Proceeds from sale of discontinued operation", "i", [0, 25797, 0, 0, 0], "Ricolino"),
    ("ppes", "Proceeds from sale of PP&E", "i", [882, 20, 152, 984, 824], ""),
    ("ins", "Insurance recoveries", "i", [201, 272, 0, 0, 0], ""),
    ("intg", "Purchases of intangible assets, net of disposals", "i", [-622, -798, -918, -772, -311], ""),
    ("drights", "Increase in distribution rights in structured entities", "i", [-77, 3, -129, -705, -52], ""),
    ("otha", "Other assets", "i", [-1433, -794, -577, 455, -860], ""),
    ("divr", "Dividends received from associates and others", "i", [59, 954, 86, 207, 158], ""),
    ("intr", "Interest received", "i", [373, 740, 809, 924, 897], ""),
    ("ncip", "Purchase of non-controlling interest", "i", [0, 0, 0, 0, -1468], "Bimbo Colombia 40%"),
    ("assc", "Capital contributions to associates, net", "i", [-534, -127, -561, -1841, -16], ""),
    ("cfi", "Net cash used in investing activities", "f", "=SUM({c}{capex}:{c}{assc})", ""),
    ("b2", "", "b", None, ""),
    ("h3", "Financing activities", "h", None, ""),
    ("borr", "Proceeds from borrowings, net of issuance costs", "i", [38924, 51670, 136638, 79111, 69267], ""),
    ("repay", "Repayment of borrowings", "i", [-33535, -55542, -109847, -56495, -55477], ""),
    ("intp", "Interest paid", "i", [-6781, -6407, -7436, -8376, -11533], "Debt interest; lease interest is inside lease payments"),
    ("perpy", "Yield on perpetual notes (equity instrument)", "i", [-621, -649, -395, 0, 0], ""),
    ("perpr", "Repurchase of perpetual notes", "i", [0, -344, -78, 0, 0], ""),
    ("divp", "Dividends paid", "i", [-4636, -5885, -3549, -4234, -4451], "Parent + NCI"),
    ("leasep", "Lease payments", "i", [-5372, -6385, -6278, -7072, -8060], "Principal + interest"),
    ("derp", "Payments on derivative financial instruments", "i", [-1690, 0, -1655, -1889, -1034], ""),
    ("derc", "Collections on derivative financial instruments", "i", [1496, 418, 2090, 692, 633], ""),
    ("buyb", "Share repurchases", "i", [-1901, -2568, -3586, -4433, -1253], ""),
    ("cff", "Net cash from (used in) financing activities", "f", "=SUM({c}{borr}:{c}{buyb})", ""),
    ("b3", "", "b", None, ""),
    ("fxadj", "Effect of exchange rates and inflation on cash", "i", [279, -472, -835, 631, -715], ""),
    ("net", "Net increase (decrease) in cash", "f", "={c}{cfo}+{c}{cfi}+{c}{cff}+{c}{fxadj}", ""),
    ("beg", "Cash at beginning of year", "beg", [9268, 8748, 12313, 6353, 8057],
     "FY2021 typed (FY2020 close); later years = prior-year ending cash"),
    ("end", "Cash at end of year", "f", "={c}{beg}+{c}{net}", "Must equal BS cash"),
]

EQ_ROWS = [
    ("h1", "Controlling interest", "h", None, ""),
    ("o1", "Opening balance", "open", [83713, 97100, 124111, 108322, 125944],
     "FY2021 typed (FY2020 close); later years = prior-year closing"),
    ("py", "Yield on perpetual notes", "i", [-621, -649, -395, 0, 0], ""),
    ("ptx", "Tax effect on perpetual notes", "i", [58, -229, -20, 0, 0], ""),
    ("prep", "Repurchase of perpetual notes", "i", [0, -377, -78, 0, 0], ""),
    ("prcl", "Reclassification of perpetual notes to debt", "i", [0, 0, -8844, 0, 0], ""),
    ("div1", "Dividends declared", "i", [-4502, -5791, -3458, -4125, -4316], ""),
    ("bb", "Share repurchases", "i", [-1901, -2568, -3586, -4433, -1253], ""),
    ("own", "Changes in ownership interest", "i", [0, 0, 0, 0, -752], "Colombia NCI purchase"),
    ("ni1", "Net income (controlling)", "i", [15916, 46910, 15477, 12544, 11133], "Must equal IS controlling NI"),
    ("oci1", "Other comprehensive income (controlling)", "i", [4437, -10285, -14885, 13636, -11244], ""),
    ("c1", "Closing balance, controlling", "f", "=SUM({c}{o1}:{c}{oci1})", "Must equal BS controlling interest"),
    ("b1", "", "b", None, ""),
    ("h2", "Non-controlling interest", "h", None, ""),
    ("o2", "Opening balance", "open", [4298, 4506, 3491, 3306, 1754],
     "FY2021 typed (FY2020 close); later years = prior-year closing"),
    ("se", "Consolidation of structured entities", "i", [-1099, -2011, -1851, -3120, -2240], ""),
    ("div2", "Dividends declared", "i", [-163, -94, -91, -109, -135], ""),
    ("own2", "Acquisitions, disposals and ownership changes", "i", [0, -65, 217, 0, -648],
     "FY22 discontinued (-65); FY23 acquisitions (+217); FY25 acquisitions 68 less decrease 716"),
    ("ni2", "Net income (NCI)", "i", [1466, 1575, 1445, 1499, 1813], "Must equal IS NCI"),
    ("oci2", "Other comprehensive income (NCI)", "i", [4, -420, 95, 178, 199], ""),
    ("c2", "Closing balance, NCI", "f", "=SUM({c}{o2}:{c}{oci2})", "Must equal BS NCI"),
    ("b2", "", "b", None, ""),
    ("tot", "Total equity, closing", "f", "={c}{c1}+{c}{c2}", "Must equal BS total equity"),
]

BLUE = Font(name="Arial", size=9, color="0000FF")
BLACK = Font(name="Arial", size=9, color="000000")
BOLD = Font(name="Arial", size=9, bold=True)
HDR = Font(name="Arial", size=9, bold=True, color="FFFFFF")
TITLE = Font(name="Arial", size=12, bold=True)
NOTE = Font(name="Arial", size=8, italic=True, color="595959")
FILL_HDR = PatternFill("solid", fgColor="1F3864")
FILL_HIST = PatternFill("solid", fgColor="D9E1F2")
FILL_FCST = PatternFill("solid", fgColor="FCE4D6")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_SEC = PatternFill("solid", fgColor="F2F2F2")
TOP = Border(top=Side(style="thin"))
NUM = '#,##0;(#,##0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
TOTAL_KEYS = {"gp", "ebit", "pbt", "ni", "nicont", "ta", "tl", "te", "tle", "tca", "tcl", "ctrl",
              "cfo", "cfi", "cff", "end", "oci", "ci", "c1", "c2", "tot", "opex", "cif", "sub1", "sub2", "net"}


def build_sheet(wb, name, title, subtitle, rows, src):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = subtitle
    ws["A2"].font = NOTE
    hdr = 4
    ws.cell(hdr, 1, "MXN millions").font = HDR
    for j, y in enumerate(YEARS):
        c = ws.cell(hdr, 2 + j, y)
        c.font = HDR
        c.alignment = Alignment(horizontal="center")
    ws[f"{TAG_COL}{hdr}"] = "Tag"
    ws[f"{SRC_COL}{hdr}"] = "Source / note"
    for col in ["A"] + COLS + [TAG_COL, SRC_COL]:
        ws[f"{col}{hdr}"].fill = FILL_HDR
        ws[f"{col}{hdr}"].font = HDR
    ws.cell(hdr - 1, 2, "Actual (audited)").fill = FILL_HIST
    ws.merge_cells(start_row=hdr - 1, start_column=2, end_row=hdr - 1, end_column=6)
    ws.cell(hdr - 1, 7, "Forecast - team projects and balances").fill = FILL_FCST
    ws.merge_cells(start_row=hdr - 1, start_column=7, end_row=hdr - 1, end_column=11)
    for c in (ws.cell(hdr - 1, 2), ws.cell(hdr - 1, 7)):
        c.font = BOLD
        c.alignment = Alignment(horizontal="center")

    rownum = {}
    r = hdr + 1
    for key, *_ in rows:
        rownum[key] = r
        r += 1

    for key, label, kind, payload, note in rows:
        r = rownum[key]
        a = ws.cell(r, 1, label)
        if kind == "h":
            a.font = BOLD
            for col in ["A"] + COLS:
                ws[f"{col}{r}"].fill = FILL_SEC
            continue
        if kind == "b":
            continue
        a.font = BOLD if key in TOTAL_KEYS else BLACK
        fmt = NUM2 if key == "eps" else ('#,##0.000' if key == "shares" else NUM)
        for j, col in enumerate(COLS):
            cell = ws[f"{col}{r}"]
            prev = COLS[j - 1] if j else None
            if kind == "i":
                if j < 5:
                    cell.value = payload[j]
                    cell.font = BLUE
                else:
                    cell.fill = FILL_IN
            elif kind in ("beg", "open"):
                if j == 0:
                    cell.value = payload[0]
                    cell.font = BLUE
                else:
                    close_key = "end" if kind == "beg" else ("c1" if key == "o1" else "c2")
                    cell.value = f"={prev}{rownum[close_key]}"
                    cell.font = BLACK
            elif kind == "f":
                cell.value = payload.format(c=col, p=prev, **rownum)
                cell.font = BOLD if key in TOTAL_KEYS else BLACK
            cell.number_format = fmt
            if key in TOTAL_KEYS:
                cell.border = TOP
        if kind == "i":
            tag = "[sourced]"
        elif kind in ("beg", "open"):
            tag = "[sourced] FY21; [calc] after"
        else:
            tag = "[calc]"
        if key in ("da",) and name == "IS":
            srcnote = note
        else:
            srcnote = (src + ("; " + note if note else "")) if kind != "f" else note
        ws[f"{TAG_COL}{r}"] = tag
        ws[f"{SRC_COL}{r}"] = srcnote
        ws[f"{TAG_COL}{r}"].font = NOTE
        ws[f"{SRC_COL}{r}"].font = NOTE

    ws.column_dimensions["A"].width = 52
    for col in COLS:
        ws.column_dimensions[col].width = 11
    ws.column_dimensions[TAG_COL].width = 12
    ws.column_dimensions[SRC_COL].width = 70
    ws.freeze_panes = ws.cell(hdr + 1, 2)
    return ws, rownum


def main():
    wb = Workbook()
    readme = wb.active
    readme.title = "README"

    _, is_r = build_sheet(wb, "IS", "Grupo Bimbo - Consolidated income statement",
                          "Continuing operations; FY2021 restated for Ricolino (discontinued from 2022). Costs by function, as reported.",
                          IS_ROWS, SRC["IS"])
    _, oci_r = build_sheet(wb, "OCI", "Grupo Bimbo - Consolidated statement of other comprehensive income",
                           "As reported. Positive = gain.", OCI_ROWS, SRC["OCI"])
    _, bs_r = build_sheet(wb, "BS", "Grupo Bimbo - Consolidated balance sheet (31 December)",
                          "As reported, every line of the face of the statement.", BS_ROWS, SRC["BS"])
    _, cf_r = build_sheet(wb, "CF", "Grupo Bimbo - Consolidated cash flow statement",
                          "Indirect method, as reported. Sign as in the filing: negative = cash outflow. FY2021-22 include Ricolino until its sale.",
                          CF_ROWS, SRC["CF"])
    _, eq_r = build_sheet(wb, "Equity", "Grupo Bimbo - Equity roll-forward (statement of changes in equity)",
                          "Columns of the statement of changes in equity condensed to controlling and non-controlling totals.",
                          EQ_ROWS, SRC["EQ"])

    # Checks sheet
    ck = wb.create_sheet("Checks")
    ck["A1"] = "Integrity checks - every line must be 0 in every year"
    ck["A1"].font = TITLE
    ck["A2"] = "Historical columns are 0 by construction. In the forecast, a non-zero value tells you which link is still missing."
    ck["A2"].font = NOTE
    ck.cell(4, 1, "Check (difference, MXN m)").font = HDR
    ck.cell(4, 1).fill = FILL_HDR
    for j, y in enumerate(YEARS):
        c = ck.cell(4, 2 + j, y)
        c.font = HDR
        c.fill = FILL_HDR
        c.alignment = Alignment(horizontal="center")
    checks = [
        ("BS: total assets - (liabilities + equity)", "=BS!{c}%d-BS!{c}%d" % (bs_r["ta"], bs_r["tle"])),
        ("CF ending cash - BS cash", "=CF!{c}%d-BS!{c}%d" % (cf_r["end"], bs_r["cash"])),
        ("CF income before taxes (continuing) - IS", "=CF!{c}%d-IS!{c}%d" % (cf_r["pbtc"], is_r["pbt"])),
        ("OCI net income - IS consolidated net income", "=OCI!{c}%d-IS!{c}%d" % (oci_r["ni"], is_r["ni"])),
        ("Equity: controlling net income - IS", "=Equity!{c}%d-IS!{c}%d" % (eq_r["ni1"], is_r["nictrl"])),
        ("Equity: NCI net income - IS", "=Equity!{c}%d-IS!{c}%d" % (eq_r["ni2"], is_r["nci"])),
        ("Equity: OCI (controlling + NCI) - OCI statement", "=Equity!{c}%d+Equity!{c}%d-OCI!{c}%d" % (eq_r["oci1"], eq_r["oci2"], oci_r["oci"])),
        ("Equity: OCI to NCI - OCI statement", "=Equity!{c}%d-(OCI!{c}%d-IS!{c}%d)" % (eq_r["oci2"], oci_r["cinci"], is_r["nci"])),
        ("Equity: closing controlling - BS", "=Equity!{c}%d-BS!{c}%d" % (eq_r["c1"], bs_r["ctrl"])),
        ("Equity: closing NCI - BS", "=Equity!{c}%d-BS!{c}%d" % (eq_r["c2"], bs_r["ncib"])),
        ("Equity: share repurchases - CF", "=Equity!{c}%d-CF!{c}%d" % (eq_r["bb"], cf_r["buyb"])),
    ]
    tie_outs = [
        ("IS net sales vs filing", is_r["sales"], [338792, 398706, 399879, 408335, 426952]),
        ("IS gross profit vs filing", is_r["gp"], [181089, 205381, 205510, 214928, 223487]),
        ("IS operating income vs filing", is_r["ebit"], [32580, 53696, 35455, 33149, 34146]),
        ("IS income before taxes vs filing", is_r["pbt"], [24854, 45878, 25324, 21034, 20143]),
        ("IS consolidated net income vs filing", is_r["ni"], [17382, 48485, 16922, 14043, 12946]),
        ("IS controlling net income vs filing", is_r["nictrl"], [15916, 46910, 15477, 12544, 11133]),
    ]
    tie_bs = [
        ("BS total current assets vs filing", "tca", [53411, 63470, 58914, 67180, 64671]),
        ("BS total assets vs filing", "ta", [337640, 347764, 348102, 416804, 409805]),
        ("BS total current liabilities vs filing", "tcl", [80760, 89734, 87192, 85155, 90573]),
        ("BS total liabilities vs filing", "tl", [236034, 220162, 236474, 289106, 289550]),
        ("BS total equity vs filing", "te", [101606, 127602, 111628, 127698, 120255]),
    ]
    tie_cf = [
        ("CF operating vs filing", "cfo", [45776, 38851, 31411, 39907, 47753]),
        ("CF investing vs filing", "cfi", [-32459, -9122, -42440, -36138, -34652]),
        ("CF financing vs filing", "cff", [-14116, -25692, 5904, -2696, -11908]),
        ("OCI total vs filing", "oci", [4441, -10705, -14790, 13814, -11045]),
    ]
    r = 5
    ck.cell(r, 1, "Links between statements (all years)").font = BOLD
    r += 1
    first_check = r
    for label, f in checks:
        ck.cell(r, 1, label).font = BLACK
        for j, col in enumerate(COLS):
            c = ck[f"{col}{r}"]
            c.value = f.format(c=col)
            c.number_format = NUM
        r += 1
    r += 1
    ck.cell(r, 1, "Subtotals vs the audited filing (history only)").font = BOLD
    r += 1
    rows_hist = [(l, "IS", row, v) for l, row, v in tie_outs]
    rows_hist += [(l, "BS", bs_r[k], v) for l, k, v in tie_bs]
    rows_hist += [(l, "CF", cf_r[k], v) for l, k, v in tie_cf[:3]]
    rows_hist += [(tie_cf[3][0], "OCI", oci_r["oci"], tie_cf[3][2])]
    for label, sh, row, vals in rows_hist:
        ck.cell(r, 1, label).font = BLACK
        for j, col in enumerate(HC):
            c = ck[f"{col}{r}"]
            c.value = f"={sh}!{col}{row}-{vals[j]}"
            c.number_format = NUM
        r += 1
    last_check = r - 1
    r += 1
    ck.cell(r, 1, "Sum of absolute differences").font = BOLD
    for col in COLS:
        c = ck[f"{col}{r}"]
        c.value = f"=SUMPRODUCT(ABS({col}{first_check}:{col}{last_check}))"
        c.number_format = NUM
        c.font = BOLD
    r += 1
    ck.cell(r, 1, "Status").font = BOLD
    for col in COLS:
        c = ck[f"{col}{r}"]
        c.value = f'=IF(ROUND({col}{r-1},0)=0,"OK","CHECK")'
        c.font = BOLD
        c.alignment = Alignment(horizontal="center")
    ck.column_dimensions["A"].width = 52
    for col in COLS:
        ck.column_dimensions[col].width = 11
    ck.freeze_panes = "B5"

    # README
    lines = [
        ("Grupo Bimbo (BIMBOA) - Consolidated financial statements FY2021A-FY2030E", TITLE),
        ("MXN millions. Built 2026-10-19 by model/build/build_statements_workbook.py.", NOTE),
        ("", None),
        ("What is here", BOLD),
        ("IS, OCI, BS, CF and Equity: every line of the face of the audited statements, FY2021A-FY2025A.", None),
        ("Checks: links between the statements and tie-outs of every subtotal to the filing. History: all 0.", None),
        ("", None),
        ("Sources (PDF page numbers)", BOLD),
        ("FY2023-FY2025: filings/annual/IA25_GB_EEFF_V10.pdf - BS p.12, IS p.13, OCI p.14, equity p.15, CF p.16.", None),
        ("FY2021-FY2022: filings/bmv/Reporte Anual 2023_XBRL_VF.PDF - BS p.256, IS p.257, OCI p.258, equity p.259, CF p.260.", None),
        ("Rule: the latest report that presents a year wins (data/adjustments.md). FY2021 is restated for the Ricolino disposal.", None),
        ("", None),
        ("Colour code", BOLD),
        ("Blue number = typed from the filing [sourced]. Black = formula [calc].", None),
        ("Yellow cell (FY2026E-FY2030E) = empty input for the team to project.", None),
        ("", None),
        ("What is already wired in the forecast columns", BOLD),
        ("Every subtotal and total; CF beginning cash = prior-year ending cash; equity opening = prior-year closing.", None),
        ("", None),
        ("What the team must link to balance FY2026E-FY2030E", BOLD),
        ("1. IS net income (controlling and NCI) -> Equity roll-forward; IS income before taxes -> CF first line; IS NI -> OCI.", None),
        ("2. Changes in working-capital lines of the BS -> CF 'changes in assets and liabilities' (sign: asset up = cash out).", None),
        ("3. Capex and D&A -> PP&E and intangibles roll-forward; debt issued/repaid -> debt balances; lease payments -> lease liabilities.", None),
        ("4. Dividends and buybacks -> both Equity and CF financing.", None),
        ("5. BS cash = CF ending cash. When all links are in, Checks!B-K shows OK in every year.", None),
        ("", None),
        ("Notes", BOLD),
        ("Signs follow each filing: IS expenses positive; interest income negative inside the financing result; CF outflows negative.", None),
        ("Lease interest is paid inside the CF line 'Lease payments' (data/adjustments.md, adjustment 6).", None),
        ("The condensed model inputs (data/financials.csv) group these lines; totals agree with this workbook.", None),
    ]
    for i, (t, f) in enumerate(lines, start=1):
        c = readme.cell(i, 1, t)
        if f:
            c.font = f
    readme.column_dimensions["A"].width = 130

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
