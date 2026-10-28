"""Build data/macro/nominal_gdp_growth_2016-2030.xlsx.

Nominal GDP growth in local currency for the countries in Grupo Bimbo's
regions, aggregated per region the way Bimbo reports ex-FX growth: each
country's local-currency growth, weighted at the prior year's exchange rate
(constant currency). Levels, CPI and FX are typed as stored [sourced]; every
growth, weight and regional figure is a live Excel formula [calc].

Sources:
  data/capitaliq/capiq-sales-drivers-{M1..M5,X1,E1..E5}_2026-10-24.md
  BIMBOA/research/Informacion_BIMBO_macro.txt (ESP/PRT/FRA/ITA levels;
      FX average per USD, history lines 1024-1052, forecast after line 1126)

Run from the repo root:
  python model/build/build_macro_gdp.py data/macro/nominal_gdp_growth_2016-2030.xlsx
"""
import re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import CellIsRule

REG = {'MEX': 'Mexico', 'USA': 'North America', 'CAN': 'North America'}
for c in 'ARG BRA CHL COL CRI ECU SLV GTM HND NIC PAN PRY PER URY VEN'.split():
    REG[c] = 'Latin America'
for c in 'GBR CHE ROU SRB SVN HRV MNE MKD RUS UKR TUR KAZ CHN IND KOR ZAF MAR TUN ESP PRT FRA ITA'.split():
    REG[c] = 'EAA'
num = re.compile(r'^-?[\d,]+(\.\d+)?$')
data = {'nominalGDPLCU': {}, 'consumerPriceInflation': {}}


def parse(path, only=None):
    for ln, line in enumerate(open(path, encoding='utf-8'), 1):
        if 'Derived' in line:
            continue
        cells = [c.strip() for c in re.split(r'\t|\|', line) if c.strip()]
        if not cells or cells[0] not in REG or (only and cells[0] not in only):
            continue
        for fld in data:
            fi = next((i for i, c in enumerate(cells) if c.startswith(fld)), None)
            if fi is None:
                continue
            while fi + 1 < len(cells) and cells[fi + 1] == fld:  # M5 repeats the field name
                fi += 1
            yi = next((i for i, c in enumerate(cells) if i > fi and re.fullmatch(r'20\d\d', c)), None)
            if yi is None:
                continue
            vi = next((i for i in range(yi + 1, len(cells)) if num.match(cells[i])), None)
            if vi is None:
                continue
            key = (cells[0], int(cells[yi]))
            if key in data[fld]:
                continue  # first block wins; duplicates checked identical
            asof = next((c for c in reversed(cells) if re.search(r'\d\d/\d\d/20\d\d', c)), '')
            data[fld][key] = dict(v=float(cells[vi].replace(',', '')), unit=cells[fi + 1],
                                  flag=cells[vi + 1] if vi + 1 < len(cells) else '', asof=asof,
                                  src=path.split('/')[-1], line=ln)


for b in 'M1 M2 M3 M4 M5 X1 E1 E2 E3 E4 E5'.split():
    parse(f'data/capitaliq/capiq-sales-drivers-{b}_2026-10-24.md')
parse('BIMBOA/research/Informacion_BIMBO_macro.txt', only={'ESP', 'PRT', 'FRA', 'ITA'})
cpi, gdp = data['consumerPriceInflation'], data['nominalGDPLCU']
missing = [(i, y) for i in REG for y in range(2015, 2031) if (i, y) not in cpi]
print('CPI missing', len(missing), missing[:30])

UNV = {('KAZ', 2026), ('KOR', 2026), ('MAR', 2026), ('UKR', 2026), ('UKR', 2027), ('SVN', 2025)}
B = Font(bold=True)
H = PatternFill('solid', fgColor='DDEBF7')
U = PatternFill('solid', fgColor='FCE4D6')
wb = Workbook()

n = wb.active
n.title = 'Notes'
for row in [
    ['Nominal GDP growth, local currency, 2016-2030 (Grupo Bimbo regions)'], [],
    ['Built', '2026-10-24 by the AI for the team (logged in docs/context/ai-use-log.md)'],
    ['Levels', 'S&P Capital IQ nominalGDPLCU, entered as values exactly as stored [sourced]; file and line in Data columns H-I'],
    ['Sources', 'data/capitaliq/capiq-sales-drivers-{M1..M5,X1,E1..E5}_2026-10-24.md; ESP/PRT/FRA/ITA from BIMBOA/research/Informacion_BIMBO_macro.txt (its "Derived" growth rows, computed by ChatIQ, are not used)'],
    ['Growth', 'Live formula in Data column J: level_t / level_t-1 - 1 [calc]. Growth_view links to Data; nothing is pasted as a value'],
    ['Duplicates', 'PER, URY, VEN 2015-2025 appear in M5 and X1 with identical values; the M5 row is cited'],
    ['[unverified]', 'KAZ/KOR/MAR 2026 (likely series break; 2027-2030 consistent); UKR 2026-2027 (gap vs real GDP + CPI); VEN all years (unusable); SVN 2025 (flag conflict). Shaded orange'],
    ['Unit notes', 'HRV labelled knB but euro-scale, one unit throughout so growth is unaffected; PRY 2026 flagged ACTUAL (label wrong, value plausible); India on calendar years'],
    ['FX', 'Local currency per USD, annual average, as stored in BIMBOA/research/Informacion_BIMBO_macro.txt [sourced]; USD = 1; VES and MKD not in source'],
    ['Regional_agg', 'Same logic as Bimbo ex-FX growth: each country grows in local currency; countries are added at the prior-year exchange rate, so weight_t = GDP USD t-1 / region total t-1 and region growth_t = sum(weight x growth) [calc]. All consolidated countries included (team decision 2026-10-25). VEN excluded because Bimbo has not consolidated it since 1-Jun-2017; MKD excluded for lack of FX (toggle column C). Weights are GDP, not Bimbo sales (Bimbo discloses sales by region only)'],
    ['IAS29_test', 'IAS 29 para 3(e): one indicator of hyperinflation is cumulative inflation over three years approaching or exceeding 100%. Computed with live formulas from the Capital IQ CPI annual averages in the same blocks [calc]. Annual-average CPI is a proxy; IAS 29 practice uses index levels'],
]:
    n.append(row)
n['A1'].font = Font(bold=True, size=13)
n.column_dimensions['A'].width = 14
n.column_dimensions['B'].width = 140

d = wb.create_sheet('Data')
d.append(['iso3', 'region', 'year', 'nominal_gdp_lcu', 'unit', 'flag', 'as_of', 'source_file', 'source_line', 'growth', 'tag'])
for c in d[1]:
    c.font, c.fill = B, H
pos, r = {}, 2
for iso in REG:
    for y in range(2015, 2031):
        g = gdp[(iso, y)]
        d.append([iso, REG[iso], y, g['v'], g['unit'], g['flag'], g['asof'], g['src'], g['line']])
        pos[(iso, y)] = r
        d.cell(r, 4).number_format = '#,##0.00'
        if y == 2015:
            d.cell(r, 11, '[sourced] level only')
        else:
            d.cell(r, 10, f'=IF(D{r-1}=0,"",D{r}/D{r-1}-1)').number_format = '0.00%'
            unv = iso == 'VEN' or (iso, y) in UNV
            d.cell(r, 11, '[unverified]' if unv else '[calc]')
            if unv:
                for c in range(1, 12):
                    d.cell(r, c).fill = U
        r += 1
d.freeze_panes = 'A2'
d.auto_filter.ref = f'A1:K{r-1}'
for i, w in enumerate([6, 14, 6, 14, 8, 22, 12, 38, 8, 9, 20], 1):
    d.column_dimensions[L(i)].width = w

v = wb.create_sheet('Growth_view')
v.append(['iso3', 'region'] + list(range(2016, 2031)))
for c in v[1]:
    c.font, c.fill = B, H
for k, iso in enumerate(REG, 2):
    v.cell(k, 1, iso)
    v.cell(k, 2, REG[iso])
    for j, y in enumerate(range(2016, 2031), 3):
        c = v.cell(k, j, f'=Data!J{pos[(iso, y)]}')
        c.number_format = '0.0%'
        if iso == 'VEN' or (iso, y) in UNV:
            c.fill = U
v.freeze_panes = 'C2'
v.column_dimensions['B'].width = 14

t = wb.create_sheet('IAS29_test')
t.append(['CPI annual average [sourced]: Capital IQ consumerPriceInflation, % stored / 100. Source lines in column S'])
t.append(['iso3', 'region'] + list(range(2015, 2031)) + ['source'])
for c in t[2]:
    c.font, c.fill = B, H
for k, iso in enumerate(REG, 3):
    t.cell(k, 1, iso)
    t.cell(k, 2, REG[iso])
    for j, y in enumerate(range(2015, 2031), 3):
        if (iso, y) in cpi:
            t.cell(k, j, cpi[(iso, y)]['v'] / 100).number_format = '0.0%'
    lines = sorted({(cpi[(iso, y)]['src'], cpi[(iso, y)]['line']) for y in range(2015, 2031) if (iso, y) in cpi}, key=lambda s: s[1])
    t.cell(k, 19, '; '.join(sorted({s for s, _ in lines})) + (f' L{lines[0][1]}-L{lines[-1][1]}' if lines else ''))
top = len(REG) + 5
t.cell(top - 1, 1, 'Cumulative 3-year inflation ending in each year [calc] = (1+CPI t-2)(1+CPI t-1)(1+CPI t) - 1; orange = 100% or more (IAS 29 para 3(e))').font = B
t.cell(top, 1, 'iso3').font = B
t.cell(top, 2, 'region').font = B
for j, y in enumerate(range(2017, 2031), 5):
    t.cell(top, j, y).font = B
for k, iso in enumerate(REG):
    rr, src = top + 1 + k, 3 + k
    t.cell(rr, 1, iso)
    t.cell(rr, 2, REG[iso])
    for j in range(5, 19):
        a, b, c = L(j - 2), L(j - 1), L(j)
        t.cell(rr, j, f'=IF(COUNT({a}{src}:{c}{src})<3,"",(1+{a}{src})*(1+{b}{src})*(1+{c}{src})-1)').number_format = '0%'
t.conditional_formatting.add(f'E{top+1}:R{top+len(REG)}', CellIsRule(operator='greaterThanOrEqual', formula=['1'], fill=U))
t.column_dimensions['B'].width = 14

# FX: local currency per USD, annual average. History 2015-2025 and forecast 2026-2030.
FXSRC = 'BIMBOA/research/Informacion_BIMBO_macro.txt'
fx, fxline = {}, {}
lines = open(FXSRC, encoding='utf-8').read().split('\n')


def fx_block(start, years):
    for ln in range(start, len(lines)):
        cells = lines[ln].split('\t')
        if cells[0].startswith('Cierre'):
            return
        if re.fullmatch(r'[A-Z]{3}', cells[0]) and len(cells) > len(years):
            for y, raw in zip(years, cells[1:]):
                fx[(cells[0], y)] = float(raw.replace(',', '').rstrip('~c'))
                fxline[(cells[0], y)] = ln + 1


fx_block(lines.index('3A - Historia 2015-2025') + 2, list(range(2015, 2026)))
fx_block(lines.index('3C - Pronósticos 2026-2030 (Capital IQ)') + 2, list(range(2026, 2031)))
CUR = {'MEX': 'MXN', 'USA': 'USD', 'CAN': 'CAD', 'ARG': 'ARS', 'BRA': 'BRL', 'CHL': 'CLP', 'COL': 'COP',
       'CRI': 'CRC', 'ECU': 'USD', 'SLV': 'USD', 'GTM': 'GTQ', 'HND': 'HNL', 'NIC': 'NIO', 'PAN': 'USD',
       'PRY': 'PYG', 'PER': 'PEN', 'URY': 'UYU', 'VEN': 'VES', 'GBR': 'GBP', 'CHE': 'CHF', 'ROU': 'RON',
       'SRB': 'RSD', 'SVN': 'EUR', 'HRV': 'EUR', 'MNE': 'EUR', 'MKD': 'MKD', 'RUS': 'RUB', 'UKR': 'UAH',
       'TUR': 'TRY', 'KAZ': 'KZT', 'CHN': 'CNY', 'IND': 'INR', 'KOR': 'KRW', 'ZAF': 'ZAR', 'MAR': 'MAD',
       'TUN': 'TND', 'ESP': 'EUR', 'PRT': 'EUR', 'FRA': 'EUR', 'ITA': 'EUR'}
f = wb.create_sheet('FX')
f.append(['Local currency per USD, annual average [sourced]: ' + FXSRC + ' (2015-2025 Banxico FIX for MXN, and CAD/GBP/EUR from 2019, else Capital IQ; 2026-2030 Capital IQ forecast). USD = 1 by definition'])
f.append(['currency'] + list(range(2015, 2031)) + ['source lines'])
for c in f[2]:
    c.font, c.fill = B, H
fxrow = {}
for k, cur in enumerate(sorted(set(CUR.values())), 3):
    fxrow[cur] = k
    f.cell(k, 1, cur)
    for j, y in enumerate(range(2015, 2031), 2):
        if cur == 'USD':
            f.cell(k, j, 1)
        elif (cur, y) in fx:
            f.cell(k, j, fx[(cur, y)])
    if cur not in ('USD',) and (cur, 2015) in fx:
        f.cell(k, 18, f'L{fxline[(cur, 2015)]} (2015-2025), L{fxline[(cur, 2026)]} (2026-2030)')
    elif cur != 'USD':
        f.cell(k, 18, 'not in source: no FX')
f.column_dimensions['R'].width = 30

# Regional aggregate, constant currency at prior-year FX:
# region growth_t = sum(GDP_i,t-1 in USD x growth_i,t) / sum(GDP_i,t-1 in USD), included countries only.
EXCL = {'VEN': 'excluded: not consolidated by Bimbo since 1-Jun-2017 (fair-value instrument) [sourced] wiki bimbo-history-2014-2020; GDP also unusable',
        'MKD': 'excluded: no MKD FX in source, cannot weight [unverified]'}
g = wb.create_sheet('Regional_agg')
g.append(['Regional nominal GDP growth, local currency, constant prior-year FX (same logic as Bimbo ex-FX growth) [calc]. include = 1/0 (team can toggle). GDP USD bn = level (x1000 when stored in trillions: CHL, COL, PRY, KOR) / FX'])
YRS = list(range(2015, 2031))
g.append(['iso3', 'region', 'include', 'note'] + [f'GDP USD bn {y}' for y in YRS])
for c in g[2]:
    c.font, c.fill = B, H
order = [i for reg in ('Mexico', 'North America', 'Latin America', 'EAA') for i in REG if REG[i] == reg]
grow = {}
for k, iso in enumerate(order, 3):
    grow[iso] = k
    g.cell(k, 1, iso)
    g.cell(k, 2, REG[iso])
    g.cell(k, 3, 0 if iso in EXCL else 1)
    g.cell(k, 4, EXCL.get(iso, ''))
    fr = fxrow[CUR[iso]]
    for j, y in enumerate(YRS, 5):
        fxc = f"FX!{L(2 + y - 2015)}{fr}"
        scale = '*1000' if gdp[(iso, y)]['unit'].endswith('T') else ''  # trillions -> billions
        g.cell(k, j, f'=IF(N({fxc})=0,"",Data!D{pos[(iso, y)]}{scale}/{fxc})').number_format = '#,##0.0'
last = 2 + len(order)
# weights block (prior-year weights used for each growth year)
wtop = last + 3
g.cell(wtop - 1, 1, 'Weight inside its region used for year t = included GDP USD t-1 / region total t-1 [calc]').font = B
g.cell(wtop, 1, 'iso3').font = B
for j, y in enumerate(YRS[1:], 5):
    g.cell(wtop, j, y).font = B
span = {}
for reg in ('Mexico', 'North America', 'Latin America', 'EAA'):
    rows_ = [grow[i] for i in order if REG[i] == reg]
    span[reg] = (min(rows_), max(rows_))
for iso in order:
    rr = wtop + grow[iso] - 2
    a, b = span[REG[iso]]
    g.cell(rr, 1, iso)
    g.cell(rr, 2, REG[iso])
    for j, y in enumerate(YRS[1:], 5):
        pc = L(5 + (y - 1 - 2015))  # GDP USD column of year t-1
        g.cell(rr, j, f'=IF($C{grow[iso]}=0,0,N({pc}{grow[iso]})/SUMPRODUCT($C${a}:$C${b},{pc}${a}:{pc}${b}))').number_format = '0.0%'
wlast = wtop + len(order)
rtop = wlast + 3
g.cell(rtop - 1, 1, 'Regional nominal GDP growth, local currency, constant FX [calc] = sum(weight x Data growth)').font = B
g.cell(rtop, 1, 'region').font = B
for j, y in enumerate(YRS[1:], 5):
    g.cell(rtop, j, y).font = B
for n_, reg in enumerate(('Mexico', 'North America', 'Latin America', 'EAA'), 1):
    rr = rtop + n_
    g.cell(rr, 1, reg).font = B
    a, b = span[reg]
    for j, y in enumerate(YRS[1:], 5):
        terms = '+'.join(f'{L(j)}{wtop + grow[i] - 2}*N(Data!J{pos[(i, y)]})' for i in order if REG[i] == reg)
        g.cell(rr, j, f'={terms}').number_format = '0.0%'
g.freeze_panes = 'E3'
g.column_dimensions['D'].width = 40

# Historical ratio: Bimbo ex-FX net sales growth by region / regional nominal GDP growth (local currency).
# Bimbo gives ex-FX growth by region in a table only from FY2023 (wiki bimbo-income-statement-fy21-fy25 §3b).
BIMBO = {  # [sourced] 4Q23 release p.2; 4Q24 release p.2; RA25 p.132
    'Mexico': (0.115, 0.040, 0.024), 'North America': (0.067, -0.046, -0.040),
    'EAA': (0.235, 0.079, 0.117), 'Latin America': (0.169, 0.100, 0.101)}
HY = (2023, 2024, 2025)
q = wb.create_sheet('Ratio_history')
q.append(['Historical ratio = Bimbo ex-FX net sales growth / regional nominal GDP growth in local currency (Regional_agg) [calc]. Team picks the forward ratio; this is the evidence'])
q.append(['region'] + [f'Bimbo ex-FX {y} [sourced]' for y in HY] + [f'GDP {y} [calc]' for y in HY]
         + [f'ratio {y} [calc]' for y in HY] + ['Bimbo CAGR 23-25', 'GDP CAGR 23-25', 'ratio of CAGRs [calc]'])
for c in q[2]:
    c.font, c.fill = B, H
for n_, reg in enumerate(('Mexico', 'North America', 'Latin America', 'EAA'), 1):
    rr = 2 + n_
    q.cell(rr, 1, reg).font = B
    for k, y in enumerate(HY):
        q.cell(rr, 2 + k, BIMBO[reg][k]).number_format = '0.0%'
        q.cell(rr, 5 + k, f'=Regional_agg!{L(5 + y - 2016)}{rtop + n_}').number_format = '0.0%'
        q.cell(rr, 8 + k, f'=IF({L(5 + k)}{rr}=0,"",{L(2 + k)}{rr}/{L(5 + k)}{rr})').number_format = '0.00x'
    q.cell(rr, 11, f'=((1+B{rr})*(1+C{rr})*(1+D{rr}))^(1/3)-1').number_format = '0.0%'
    q.cell(rr, 12, f'=((1+E{rr})*(1+F{rr})*(1+G{rr}))^(1/3)-1').number_format = '0.0%'
    q.cell(rr, 13, f'=IF(L{rr}=0,"",K{rr}/L{rr})').number_format = '0.00x'
for k, txt in enumerate([
        'Sources Bimbo: 4Q23 release p.2 (FY2023), 4Q24 release p.2 (FY2024), RA25 p.132 (FY2025), via wiki bimbo-income-statement-fy21-fy25 §3b.',
        'FY2021-FY2022: Bimbo gives no ex-FX growth by region in a table (4Q22 release p.4); FY2021 also restated for Ricolino. Not used.',
        'Ex-FX growth is NOT organic: it includes acquisitions (EAA and Latin America mainly; US sourdough in North America 2025). Bimbo does not quantify the split (S2).',
        'FY2023 North America and group growth is held down by the 53rd week in FY2022 (group 10.7% reported vs 11.4% without it, 4Q23 release p.1).',
        'Latin America and EAA GDP include ARG and TUR at local-currency (hyperinflation) growth, weighted by prior-year GDP in USD.'], 8):
    q.cell(k, 1, txt)
q.column_dimensions['A'].width = 16

# North America ratio (team decision 2026-10-25): base = 1Q26-3Q26 ex-FX net sales growth vs US nominal GDP growth 2026.
NAQ = [  # quarter, net sales MXN m, prior-year quarter MXN m, ex-FX growth, source [sourced]
    ('1Q26', 40533, 46580, 0.007, '1T26 release p.3'),
    ('2Q26', 44368, 49124, 0.013, '2T26 release p.3'),
    ('3Q26', 44241, 47470, 0.010, '3T26 release p.3')]
na = wb.create_sheet('Ratio_NA')
na.append(['North America ratio: 1Q26-3Q26 ex-FX net sales growth / US nominal GDP growth 2026 (team decision 2026-10-25)'])
na.append(['quarter', 'net sales MXN m [sourced]', 'prior-year quarter MXN m [sourced]', 'growth in MXN [calc]', 'ex-FX growth [sourced]', 'source'])
for c in na[2]:
    c.font, c.fill = B, H
for k, (qq, s, p, gx, src) in enumerate(NAQ, 3):
    na.append([qq, s, p, f'=B{k}/C{k}-1', gx, src])
    na.cell(k, 4).number_format = na.cell(k, 5).number_format = '0.0%'
    na.cell(k, 2).number_format = na.cell(k, 3).number_format = '#,##0'
usa26 = pos[('USA', 2026)]
rows_na = [
    ('9M26 ex-FX growth [calc] (quarters weighted by prior-year quarter sales, the constant-currency basis)', '=SUMPRODUCT(C3:C5,E3:E5)/SUM(C3:C5)', '0.00%'),
    ('US nominal GDP growth 2026, local currency [calc] (Data, USA 2026)', f'=Data!J{usa26}', '0.00%'),
    ('North America ratio [calc]; adopted by the team as the FY2027E-FY2030E ratio [assumption]', '=B7/B8', '0.00x'),
]
for k, (lab, fml, fmt) in enumerate(rows_na, 7):
    na.cell(k, 1, lab).font = B
    na.cell(k, 2, fml).number_format = fmt
na.cell(11, 1, 'Implied North America ex-FX growth = ratio x US nominal GDP growth [calc]').font = B
for j, y in enumerate(range(2027, 2031), 2):
    na.cell(12, j, y).font = B
    na.cell(13, j, f'=Data!J{pos[("USA", y)]}').number_format = '0.0%'
    na.cell(14, j, f'=$B$9*{L(j)}13').number_format = '0.0%'
na.cell(13, 1, 'US nominal GDP growth [calc]')
na.cell(14, 1, 'North America ex-FX growth [calc]')
for k, txt in enumerate([
        'Why this base: North America ex-FX sales fell nine quarters in a row (4Q23-4Q25) and turned positive in 1Q26-3Q26; Bimbo cites US share gains and better price mix (1T26 p.3; 2T26 p.5). FY2023-FY2025 history (ratio -0.13) reflects the decline phase.',
        'Limits: three quarters only; 2Q26 helped by the football World Cup (2T26 p.5); 2026 laps a weak 2025 base; quarterly year-on-year growth is compared with a full-year 2026 GDP forecast (quarterly US GDP not in the data); Canada (about 7% of the region GDP weight) left out by choice.'], 16):
    na.cell(k, 1, txt)
na.column_dimensions['A'].width = 60
na.column_dimensions['F'].width = 18

# Mexico evidence for the team's ratio decision. Mexico is a peso business: growth in MXN = ex-FX growth.
MXQ = [  # quarter, net sales MXN m, prior-year quarter, growth as published, release page [sourced]
    ('3Q23', 36375, 33562, 0.084, '3T23 p.3'), ('4Q23', 36985, 34818, 0.062, '4T23 p.3'),
    ('1Q24', 37386, 36244, 0.032, '1T24 p.3'), ('2Q24', 37369, 35783, 0.044, '2T24 p.3'),
    ('3Q24', 38793, 36375, 0.066, '3T24 p.3'), ('4Q24', 37616, 36985, 0.017, '4T24 p.3'),
    ('1Q25', 38018, 37386, 0.017, '1T25 p.3'), ('2Q25', 38475, 37369, 0.030, '2T25 p.3'),
    ('3Q25', 38897, 38793, 0.003, '3T25 p.3'), ('4Q25', 39419, 37616, 0.048, '4T25 p.4'),
    ('1Q26', 39726, 38018, 0.045, '1T26 p.3'), ('2Q26', 40279, 38475, 0.047, '2T26 p.3'),
    ('3Q26', 40905, 38897, 0.052, '3T26 p.3')]
mx = wb.create_sheet('Evidence_MEX')
mx.append(['Mexico: evidence for the ratio (team decides). Peso business, so growth in MXN = ex-FX growth'])
mx.append(['A. Annual', 'Bimbo Mexico growth', 'Mexico nominal GDP growth [calc]', 'ratio [calc]', 'source'])
for c in mx[2]:
    c.font, c.fill = B, H
ANN = [(2022, '=130401/109089-1', '[calc] segment sales FY22 130,401 (RA23 p.358) / FY21 restated 109,089 (RA22 p.313)'),
       (2023, 0.115, '[sourced] 4Q23 release p.2'), (2024, 0.040, '[sourced] 4Q24 release p.2'),
       (2025, 0.024, '[sourced] RA25 p.132')]
for k, (y, val, src) in enumerate(ANN, 3):
    mx.cell(k, 1, y)
    mx.cell(k, 2, val).number_format = '0.0%'
    mx.cell(k, 3, f'=Data!J{pos[("MEX", y)]}').number_format = '0.0%'
    mx.cell(k, 4, f'=B{k}/C{k}').number_format = '0.00x'
    mx.cell(k, 5, src)
qt = 9
mx.cell(qt - 1, 1, 'B. Quarterly, year on year').font = B
for j, h in enumerate(['quarter', 'net sales MXN m [sourced]', 'prior-year quarter [sourced]', 'growth [calc]',
                       'growth as published [sourced]', 'source', 'trailing 4Q growth [calc]'], 1):
    c = mx.cell(qt, j, h)
    c.font, c.fill = B, H
for k, (qq, s, p, gp, src) in enumerate(MXQ, qt + 1):
    mx.cell(k, 1, qq)
    mx.cell(k, 2, s).number_format = '#,##0'
    mx.cell(k, 3, p).number_format = '#,##0'
    mx.cell(k, 4, f'=B{k}/C{k}-1').number_format = '0.0%'
    mx.cell(k, 5, gp).number_format = '0.0%'
    mx.cell(k, 6, src)
    if k >= qt + 4:
        mx.cell(k, 7, f'=SUM(B{k-3}:B{k})/SUM(C{k-3}:C{k})-1').number_format = '0.0%'
ql = qt + len(MXQ)
ct = ql + 2
mx.cell(ct, 1, 'C. Candidate ratios (evidence only; team picks one or another with a written reason)').font = B
cands = [('9M26 (1Q26-3Q26 growth / Mexico GDP 2026)', f'=(SUM(B{ql-2}:B{ql})/SUM(C{ql-2}:C{ql})-1)/Data!J{pos[("MEX", 2026)]}'),
         ('Trailing 4Q to 3Q26 / Mexico GDP 2026', f'=G{ql}/Data!J{pos[("MEX", 2026)]}'),
         ('Two-year 9M26 vs 9M24, annualised / GDP 2025-2026 annualised (strips the weak-2025 base)',
          f'=((SUM(B{ql-2}:B{ql})/SUM(B{qt+3}:B{qt+5}))^(1/2)-1)/(((1+Data!J{pos[("MEX", 2025)]})*(1+Data!J{pos[("MEX", 2026)]}))^(1/2)-1)'),
         ('FY2025 / GDP 2025', '=D6'),
         ('FY2023-FY2025 compound (backward-looking)', '=(((1+B4)*(1+B5)*(1+B6))^(1/3)-1)/(((1+C4)*(1+C5)*(1+C6))^(1/3)-1)')]
for j, y in enumerate(range(2027, 2031), 3):
    mx.cell(ct + 1, j, y).font = B
    mx.cell(ct + 2, j, f'=Data!J{pos[("MEX", y)]}').number_format = '0.0%'
mx.cell(ct + 1, 1, 'candidate').font = B
mx.cell(ct + 1, 2, 'ratio').font = B
mx.cell(ct + 2, 1, 'Mexico nominal GDP growth [calc]')
for k, (lab, fml) in enumerate(cands, ct + 3):
    mx.cell(k, 1, lab)
    mx.cell(k, 2, fml).number_format = '0.00x'
    for j in range(3, 7):
        mx.cell(k, j, f'=$B{k}*{L(j)}${ct + 2}').number_format = '0.0%'
nt = ct + 3 + len(cands) + 1
for k, txt in enumerate([
        'Bimbo on Mexico 2026: volume up and favourable mix, growth across categories and channels (1T26 p.4); mix and volume, partly helped by World Cup consumption (2T26 p.3); categories, channels, mix, positive volumes, execution, service and innovation (3T26 p.3) [sourced]. Volumes are not quantified.',
        'No acquisitions named in Mexico 2024-2026 (handoff, S2 block), so Mexico growth is close to organic [guidance, unverified].',
        'FY2022 is a high-inflation year and its base is restated for the Ricolino sale; shown for context.',
        'TEAM DECISION 2026-10-25: Mexico ratio = 9M26 candidate (same base as North America, for consistency) [assumption]. Compound and two-year candidates rejected by the team.'], nt):
    mx.cell(k, 1, txt)
mx.column_dimensions['A'].width = 44
mx.column_dimensions['F'].width = 12

# Consistency pass (team decision 2026-10-25): every region on the same 9M26 base.
# Latin America and EAA: 1Q26-3Q26 ex-FX growth vs regional GDP 2026 (Regional_agg). These include acquisitions.
RQ = {  # quarter: (net sales MXN m, prior-year quarter as shown in the 2026 release, ex-FX growth) [sourced] release p.3
    'Latin America': [('1Q26', 11545, 10902, 0.151, '1T26 p.3'), ('2Q26', 11620, 10735, 0.141, '2T26 p.3'),
                      ('3Q26', 12047, 10714, 0.152, '3T26 p.3')],
    'EAA': [('1Q26', 12631, 12042, 0.126, '1T26 p.3'), ('2Q26', 13188, 13436, 0.084, '2T26 p.3'),
            ('3Q26', 13371, 14370, 0.021, '3T26 p.3')]}
sm = wb.create_sheet('Ratios_9M26')
sm.append(['Ratios on one base for all regions: 1Q26-3Q26 ex-FX net sales growth / nominal GDP growth 2026 (team decision 2026-10-25, consistency pass)'])
sm.append(['region', 'quarter', 'net sales MXN m [sourced]', 'prior-year quarter [sourced]', 'ex-FX growth [sourced]', 'source'])
for c in sm[2]:
    c.font, c.fill = B, H
qrow = {}
r = 3
for reg, qs in RQ.items():
    qrow[reg] = r
    for qq, s, p, gx, src in qs:
        sm.append([reg, qq, s, p, gx, src])
        sm.cell(r, 5).number_format = '0.0%'
        r += 1
st = r + 1
sm.cell(st, 1, 'Summary').font = B
hdr = ['region', '9M26 ex-FX growth [calc]', 'GDP growth 2026 [calc]', 'GDP used', 'ratio [calc]'] + [f'implied {y} [calc]' for y in range(2027, 2031)]
for j, h in enumerate(hdr, 1):
    c = sm.cell(st + 1, j, h)
    c.font, c.fill = B, H
ragc = {y: L(5 + y - 2016) for y in range(2016, 2031)}  # Regional_agg column per year
regrow = {reg: rtop + n_ for n_, reg in enumerate(('Mexico', 'North America', 'Latin America', 'EAA'), 1)}
summary = [
    ('North America', '=Ratio_NA!B7', '=Ratio_NA!B8', 'US only (team choice)', {y: f'=Data!J{pos[("USA", y)]}' for y in range(2027, 2031)}),
    ('Mexico', f'=SUM(Evidence_MEX!B{ql-2}:B{ql})/SUM(Evidence_MEX!C{ql-2}:C{ql})-1', f'=Data!J{pos[("MEX", 2026)]}', 'Mexico', {y: f'=Data!J{pos[("MEX", y)]}' for y in range(2027, 2031)}),
]
for reg in ('Latin America', 'EAA'):
    a = qrow[reg]
    summary.append((reg, f'=SUMPRODUCT(D{a}:D{a+2},E{a}:E{a+2})/SUM(D{a}:D{a+2})',
                    f'=Regional_agg!{ragc[2026]}{regrow[reg]}', 'regional aggregate, prior-year FX',
                    {y: f'=Regional_agg!{ragc[y]}{regrow[reg]}' for y in range(2027, 2031)}))
for k, (reg, g9, gdp26, used, fwd) in enumerate(summary, st + 2):
    sm.cell(k, 1, reg).font = B
    sm.cell(k, 2, g9).number_format = '0.0%'
    sm.cell(k, 3, gdp26).number_format = '0.0%'
    sm.cell(k, 4, used)
    sm.cell(k, 5, f'=B{k}/C{k}').number_format = '0.00x'
    for j, y in enumerate(range(2027, 2031), 6):
        sm.cell(k, j, f'=$E{k}*({fwd[y][1:]})').number_format = '0.0%'
nt2 = st + 2 + len(summary) + 1
for k, txt in enumerate([
        'Not organic: 9M26 includes acquisitions whose sales Bimbo does not disclose: Stone House (North America, closed 2-Jun-2025, inorganic in 1Q26 and Apr-May 2026); Wickbold (Latin America, closed 6-Oct-2025, inorganic in all of 1Q26-3Q26, comparable from 4Q26); Don Don and Karamolegos (EAA, about Jun-2025, inorganic in 1Q26); Bonel (EAA, 7-Apr-2026, inorganic in 2Q26-3Q26) [sourced] wiki bimbo-acquisitions; releases 1T26 p.4, 2T26 p.4, 3T26 p.4.',
        'Prior-year quarters are the comparatives printed in the 2026 releases; some differ from the figures first reported in 2025 (e.g. Latin America 1Q25: 10,902 in 1T26 vs 11,176 in 1T25), so the 2026 releases are used for consistency with the published growth.',
        'EAA 2026 GDP includes [unverified] 2026 steps for KAZ, KOR, MAR (likely series breaks) and UKR; local-currency GDP includes hyperinflation for ARG (Latin America) and TUR (EAA) (see IAS29_test).',
        'Next (team plan): redo all four regions on Alan\'s organic-only rule and compare with this table.'], nt2):
    sm.cell(k, 1, txt)
sm.column_dimensions['A'].width = 16
sm.column_dimensions['D'].width = 30

# Organic pass (team decision 2026-10-25): acquisition points = 2026 quarter ex-FX growth minus the same
# quarter of 2025 (before the deal), quarter by quarter [assumption]. Organic = reported - adjust flag x points.
Q25 = {  # same quarter 2025 ex-FX growth [sourced] 1T25 p.3, 2T25 p.3, 3T25 p.3
    'North America': (-0.049, -0.046, -0.035), 'Mexico': (0.017, 0.030, 0.003),
    'Latin America': (0.052, 0.082, 0.091), 'EAA': (0.045, 0.068, 0.173)}
FLAG = {  # 1 = deduct the before/after points in that quarter (team can toggle)
    'North America': (0, 0, 0), 'Mexico': (0, 0, 0), 'Latin America': (1, 1, 1), 'EAA': (1, 1, 0)}
WHY = {
    'North America': 'Stone House (2-Jun-2025) inorganic in 1Q26 and Apr-May 2026, size not disclosed. Not deducted: the before/after points (+5.6, +5.9) are the turnaround Bimbo attributes to share gains and price mix (1T26 p.3, 2T26 p.5), not the deal.',
    'Mexico': 'No acquisitions named; growth is organic.',
    'Latin America': 'Wickbold (6-Oct-2025) inorganic in all of 1Q26-3Q26; 1Q25-3Q25 have no Wickbold. Deducted every quarter.',
    'EAA': 'Don Don, Karamolegos (~Jun-2025) inorganic in 1Q26 and Apr-May 2026; Bonel (7-Apr-2026) in 2Q26-3Q26. 3Q26 not deducted: its 2025 base (17.3%) already contains Don Don and Karamolegos, so the difference is negative and measures nothing.'}
src26 = {'North America': ('Ratio_NA', 3, 'C', 'E'), 'Mexico': ('Evidence_MEX', ql - 2, 'C', None),
         'Latin America': ('Ratios_9M26', qrow['Latin America'], 'D', 'E'), 'EAA': ('Ratios_9M26', qrow['EAA'], 'D', 'E')}
og = wb.create_sheet('Organic_pass')
og.append(['Organic pass: acquisition points per quarter = 2026 ex-FX growth - same quarter 2025 [assumption: organic growth in the deal quarters = the 2025 pace]; flag 1 = deduct'])
og.append(['region', 'quarter', 'prior-year quarter sales', 'reported ex-FX 2026', 'same quarter 2025 [sourced]',
           'before/after points [calc]', 'deduct? (1/0)', 'organic growth [calc]'])
for c in og[2]:
    c.font, c.fill = B, H
orow, r = {}, 3
for reg in ('North America', 'Mexico', 'Latin America', 'EAA'):
    orow[reg] = r
    sh, r0, pc, gc = src26[reg]
    for k, qq in enumerate(('1Q26', '2Q26', '3Q26')):
        og.cell(r, 1, reg)
        og.cell(r, 2, qq)
        og.cell(r, 3, f'={sh}!{pc}{r0 + k}').number_format = '#,##0'
        og.cell(r, 4, f'={sh}!{gc}{r0 + k}' if gc else f'={sh}!B{r0 + k}/{sh}!C{r0 + k}-1').number_format = '0.0%'
        og.cell(r, 5, Q25[reg][k]).number_format = '0.0%'
        og.cell(r, 6, f'=D{r}-E{r}').number_format = '0.0%'
        og.cell(r, 7, FLAG[reg][k])
        og.cell(r, 8, f'=D{r}-G{r}*F{r}').number_format = '0.0%'
        r += 1
ot = r + 1
for j, h in enumerate(['region', 'organic 9M26 [calc]', 'GDP 2026 [calc]', 'organic ratio [calc]', 'consistency ratio (Ratios_9M26)']
                      + [f'implied {y} [calc]' for y in range(2027, 2031)] + ['why'], 1):
    c = og.cell(ot, j, h)
    c.font, c.fill = B, H
for n_, (reg, g9, gdp26, used, fwd) in enumerate(summary):
    k = ot + 1 + n_
    a = orow[reg]
    srow = st + 2 + n_
    og.cell(k, 1, reg).font = B
    og.cell(k, 2, f'=SUMPRODUCT(C{a}:C{a+2},H{a}:H{a+2})/SUM(C{a}:C{a+2})').number_format = '0.0%'
    og.cell(k, 3, f'=Ratios_9M26!C{srow}').number_format = '0.0%'
    og.cell(k, 4, f'=B{k}/C{k}').number_format = '0.00x'
    og.cell(k, 5, f'=Ratios_9M26!E{srow}').number_format = '0.00x'
    for j, y in enumerate(range(2027, 2031), 6):
        og.cell(k, j, f'=$D{k}*({fwd[y][1:]})').number_format = '0.0%'
    og.cell(k, 10, WHY[reg])
og.cell(ot + 6, 1, 'Limits: the method assumes nothing else changed between the two quarters; any organic acceleration (price, volume, Argentina) is counted as acquisition. 2025 quarters also carry older deals (Pagnifique Sep-2024 and La Zarcerena 1Q25 in Latin America), so the 2025 pace is not fully organic either.')
og.column_dimensions['A'].width = 16
og.column_dimensions['J'].width = 90

# Long-history cross-check (team decision 2026-10-25): yearly ratios Bimbo growth / nominal GDP growth, as far back as
# each region's data allow. Reported growth, NOT organic (no history of acquired sales). Rule fixed in advance:
# a year's ratio is not computed when GDP growth is zero or negative (the ratio flips sign and means nothing).
xc = wb.create_sheet('Cross_check_history')
xc.append(['Long-history cross-check: Bimbo growth / nominal GDP growth per year. Reported growth incl. acquisitions; summary = median of valid years (not a compound rate)'])
# A. FY2022 by region from quarters, weighted by the prior-year quarter sales printed in the 2022 releases
Q22 = {  # region: [(prior-year quarter MXN m [sourced] 1T22-4T22 release, growth ex-FX/USD [sourced] wiki bimbo-quarterly-trend-4q21-2q26)]
    'North America': [(38963, 0.158), (42538, 0.168), (44028, 0.190), (50050, 0.219)],
    'Mexico': [(27661, 0.195), (26119, 0.216), (27867, 0.204), (29653, 0.174)],
    'Latin America': [(7374, 0.251), (7369, 0.377), (7830, 0.392), (8603, 0.334)],
    'EAA': [(7843, 0.211), (8289, 0.174), (8680, 0.233), (9384, 0.242)]}
xc.append(['A. FY2022 from quarters [calc]', 'prior 1Q21', 'g 1Q22', 'prior 2Q21', 'g 2Q22', 'prior 3Q21', 'g 3Q22', 'prior 4Q21', 'g 4Q22', 'FY2022 [calc]'])
for c in xc[2]:
    c.font, c.fill = B, H
fy22 = {}
for k, (reg, qs) in enumerate(Q22.items(), 3):
    xc.cell(k, 1, reg)
    for n_, (p, gq) in enumerate(qs):
        xc.cell(k, 2 + 2 * n_, p).number_format = '#,##0'
        xc.cell(k, 3 + 2 * n_, gq).number_format = '0.0%'
    xc.cell(k, 10, f'=(B{k}*C{k}+D{k}*E{k}+F{k}*G{k}+H{k}*I{k})/(B{k}+D{k}+F{k}+H{k})').number_format = '0.0%'
    fy22[reg] = f'J{k}'
xc.cell(7, 1, 'Sources: prior-year quarters 1T22 p.3, 2T22 p.3, 3T22 p.3, 4T22 p.5; growth: North America in US dollars through 2Q23, others ex-FX (narrative). 4Q22 North America includes a 53rd week.')
# B. North America in US dollars before FY2022 [calc] = segment sales MXN / average MXN per USD (FX sheet)
NAS = {2015: 116399, 2016: 135219, 2017: 137662, 2018: 143968, 2019: 144005, 2020: 176395}
xc.cell(9, 1, 'B. North America sales in USD m [calc] = segment sales MXN m [sourced] / MXN per USD average (FX)').font = B
xc.cell(10, 1, 'year').font = B
xc.cell(11, 1, 'segment sales MXN m')
xc.cell(12, 1, 'MXN per USD')
xc.cell(13, 1, 'sales USD m')
xc.cell(14, 1, 'growth in USD')
yrsB = list(range(2015, 2022))
for j, y in enumerate(yrsB, 2):
    xc.cell(10, j, y).font = B
    xc.cell(11, j, NAS.get(y, 176275)).number_format = '#,##0'
    xc.cell(12, j, f'=FX!{L(2 + y - 2015)}{fxrow["MXN"]}')
    xc.cell(13, j, f'={L(j)}11/{L(j)}12').number_format = '#,##0'
    if j > 2:
        xc.cell(14, j, f'={L(j)}13/{L(j-1)}13-1').number_format = '0.0%'
xc.cell(15, 1, 'Segment sales: RA18 p.373 (FY15-FY17, recast), RA20 p.286-287 (FY18-FY20; FY18 reclassified, overstates FY18 growth ~0.5 pt), RA21 p.304 (FY21 as first reported, same pre-Ricolino basis as FY20), via wiki bimbo-income-statement-fy14-fy20 and -fy21-fy25. Canada is translated through USD too.')
# C. Yearly growth and ratios
YC = list(range(2016, 2026))
MXG = {2016: '=81455/75597-1', 2017: '=90367/81455-1', 2018: 0.105, 2019: '=102688/100327-1', 2020: '=104593/102688-1',
       2021: '=118661/104593-1', 2022: f'={fy22["Mexico"]}', 2023: 0.115, 2024: 0.040, 2025: 0.024}
NAG = {y: f'={L(2 + y - 2015)}14' for y in range(2016, 2022)}
NAG.update({2022: f'={fy22["North America"]}', 2023: 0.067, 2024: -0.046, 2025: -0.040})
LAG = {2022: f'={fy22["Latin America"]}', 2023: 0.169, 2024: 0.100, 2025: 0.101}
EAG = {2022: f'={fy22["EAA"]}', 2023: 0.235, 2024: 0.079, 2025: 0.117}
GDPREF = {'Mexico': lambda y: f'=Data!J{pos[("MEX", y)]}', 'North America': lambda y: f'=Data!J{pos[("USA", y)]}',
          'Latin America': lambda y: f'=Regional_agg!{ragc[y]}{regrow["Latin America"]}', 'EAA': lambda y: f'=Regional_agg!{ragc[y]}{regrow["EAA"]}'}
c0 = 18
xc.cell(c0 - 1, 1, 'C. Yearly ratios. Bimbo growth: Mexico in pesos (= ex-FX), North America in USD, Latin America and EAA ex-FX').font = B
xc.cell(c0, 1, 'row').font = B
for j, y in enumerate(YC, 2):
    xc.cell(c0, j, y).font = B
xc.cell(c0, 13, 'median of valid years [calc]').font = B
xc.cell(c0, 14, 'valid years').font = B
xc.cell(c0, 15, '9M26 organic ratio (Organic_pass)').font = B
r = c0 + 1
for n_, (reg, gd) in enumerate((('North America', NAG), ('Mexico', MXG), ('Latin America', LAG), ('EAA', EAG))):
    xc.cell(r, 1, f'{reg}: Bimbo growth')
    xc.cell(r + 1, 1, f'{reg}: GDP growth')
    xc.cell(r + 2, 1, f'{reg}: ratio').font = B
    for j, y in enumerate(YC, 2):
        if y in gd:
            xc.cell(r, j, gd[y]).number_format = '0.0%'
            xc.cell(r + 1, j, GDPREF[reg](y)).number_format = '0.0%'
            xc.cell(r + 2, j, f'=IF({L(j)}{r+1}<=0,"n/a",{L(j)}{r}/{L(j)}{r+1})').number_format = '0.00x'
    xc.cell(r + 2, 13, f'=MEDIAN(B{r+2}:K{r+2})').number_format = '0.00x'
    xc.cell(r + 2, 14, f'=COUNT(B{r+2}:K{r+2})')
    xc.cell(r + 2, 15, f'=Organic_pass!D{ot + 1 + n_}').number_format = '0.00x'
    r += 4
for k, txt in enumerate([
        'Mexico growth sources: FY16-FY17 segment sales (RA18 p.373); FY18 company-reported +10.5% (4T18 p.3); FY19-FY20 RA20 p.286-287; FY21 RA21 p.304 vs FY20 (both include Ricolino); FY22 from quarters (Ricolino excluded); FY23-FY25 4Q23 p.2, 4Q24 p.2, RA25 p.132.',
        'Breaks to keep in mind: Ricolino (confectionery) inside Mexico until FY21, outside from FY22; China moved from Mexico to EAA in 2017 (FY15-FY17 recast); 53rd week in North America in FY15 and FY22; COVID 2020-2021; acquisitions throughout (e.g. Canada Bread 2014, Saputo Bakery 2015, East Balt/Bimbo QSR Oct-2017 in North America).',
        'Latin America and EAA before FY2022: no ex-FX growth by region in the 4Q14-4Q21 releases or the FY2018/FY2020 BMV annual reports (searched 2026-10-25); only scattered quarter remarks (e.g. EAA 4Q19 local currency +2%, 4T19 p.4). Not usable as a series.'], r + 1):
    xc.cell(k, 1, txt)
xc.column_dimensions['A'].width = 30

# Regional FX (team decision 2026-10-25, advisor: grow in local currency, consolidate in MXN, split business vs FX).
# Leg 1, local currency -> USD, per region: translation factor of the regional GDP basket used for the growth,
# factor_t = sum(w x (1 + g) x FX t-1 / FX t) / sum(w x (1 + g)), same prior-year weights as Regional_agg [calc].
# Leg 2, USD -> MXN, is applied in the statements workbook with the model's USD/MXN path.
FXY = list(range(2027, 2031))
rx = wb.create_sheet('Regional_FX')
rx.append(['Regional FX, local currency to USD [calc]: FX ratio = local currency per USD t-1 / t (above 1 = the currency gains vs USD); '
           'region factor = sum(weight x (1 + GDP growth) x FX ratio) / sum(weight x (1 + GDP growth)), weights and growth as in Regional_agg'])
rx.append(['iso3', 'region', 'currency'] + [f'FX ratio {y}' for y in FXY])
for c in rx[2]:
    c.font, c.fill = B, H
xrow = {}
for k, iso in enumerate(order, 3):
    xrow[iso] = k
    rx.cell(k, 1, iso)
    rx.cell(k, 2, REG[iso])
    rx.cell(k, 3, CUR[iso])
    fr = fxrow[CUR[iso]]
    for j, y in enumerate(FXY, 4):
        a, b = f"FX!{L(2 + y - 1 - 2015)}{fr}", f"FX!{L(2 + y - 2015)}{fr}"
        rx.cell(k, j, f'=IFERROR({a}/{b},0)').number_format = '0.0000'
ftop = 3 + len(order) + 2
rx.cell(ftop - 1, 1, 'Region factor, local currency to USD [calc] (Mexico = 1: peso business; North America = USA only, the team choice for GDP, so 1)').font = B
rx.cell(ftop, 1, 'region').font = B
for j, y in enumerate(FXY, 4):
    rx.cell(ftop, j, y).font = B
fxfac = {}
for n_, reg in enumerate(('Mexico', 'North America', 'Latin America', 'EAA'), 1):
    rr = ftop + n_
    fxfac[reg] = rr
    rx.cell(rr, 1, reg).font = B
    for j, y in enumerate(FXY, 4):
        if reg in ('Mexico', 'North America'):
            rx.cell(rr, j, 1).number_format = '0.0000'
            continue
        wc = L(5 + y - 2016)
        isos = [i for i in order if REG[i] == reg]
        num_ = '+'.join(f'Regional_agg!{wc}{wtop + grow[i] - 2}*(1+N(Data!J{pos[(i, y)]}))*{L(j)}{xrow[i]}' for i in isos)
        den_ = '+'.join(f'Regional_agg!{wc}{wtop + grow[i] - 2}*(1+N(Data!J{pos[(i, y)]}))' for i in isos)
        rx.cell(rr, j, f'=({num_})/({den_})').number_format = '0.0000'
rx.cell(ftop + 6, 1, 'Limits: weights are GDP, not Bimbo sales by country (not disclosed): China is the largest EAA weight and Brazil the largest Latin America weight. '
        'Annual-average FX; Bimbo translates hyperinflationary Argentina at the closing rate (IAS 29, EEFF25 p.23, p.26). Capital IQ FX forecasts [sourced] via sheet FX.')
rx.column_dimensions['A'].width = 16

# Terminal FX by purchasing power parity (team decision 2026-10-25): in perpetuity each currency moves against the USD
# by its inflation gap, PPP ratio = (1 + US CPI 2030) / (1 + currency CPI 2030). Dollarized countries = 1; the euro uses
# the GDP-weighted CPI of the euro countries in Bimbo's list (proxy for euro-area inflation). Region factor aggregated
# exactly as in Regional_FX for 2030 (same weights and GDP growth) [calc].
pp = wb.create_sheet('PPP_terminal')
pp.append(['Terminal FX, local currency to USD, by purchasing power parity 2030 [calc]: PPP ratio = (1 + US CPI) / (1 + currency CPI); '
           'below 1 = the currency loses vs USD. CPI from IAS29_test (Capital IQ) [sourced]. USD countries = 1; EUR = GDP-weighted CPI of the euro countries'])
pp.append(['iso3', 'region', 'currency', 'CPI 2030', 'currency CPI 2030', 'PPP ratio 2030'])
for c in pp[2]:
    c.font, c.fill = B, H
cpirow = {iso: k for k, iso in enumerate(REG, 3)}
cpic = L(3 + 2030 - 2015)            # IAS29_test column of 2030
gcol = L(5 + 2029 - 2015)            # Regional_agg GDP USD column of 2029 (weights for 2030)
prow = {}
for k, iso in enumerate(order, 3):
    prow[iso] = k
eur = [i for i in order if CUR[i] == 'EUR']
eur_cpi = '(' + '+'.join(f'Regional_agg!{gcol}{grow[i]}*IAS29_test!{cpic}{cpirow[i]}' for i in eur) + ')/(' + \
          '+'.join(f'Regional_agg!{gcol}{grow[i]}' for i in eur) + ')'
us = f'IAS29_test!{cpic}{cpirow["USA"]}'
for iso in order:
    k = prow[iso]
    pp.cell(k, 1, iso)
    pp.cell(k, 2, REG[iso])
    pp.cell(k, 3, CUR[iso])
    pp.cell(k, 4, f'=IAS29_test!{cpic}{cpirow[iso]}').number_format = '0.00%'
    if CUR[iso] == 'USD':
        pp.cell(k, 5, f'={us}')
    elif CUR[iso] == 'EUR':
        pp.cell(k, 5, '=' + eur_cpi)
    else:
        pp.cell(k, 5, f'=D{k}')
    pp.cell(k, 5).number_format = '0.00%'
    pp.cell(k, 6, f'=IFERROR((1+{us})/(1+E{k}),0)').number_format = '0.0000'
ptop = 3 + len(order) + 2
pp.cell(ptop - 1, 1, 'Region factor 2030, local currency to USD, PPP [calc] = sum(weight x (1 + GDP growth) x PPP ratio) / sum(weight x (1 + GDP growth)), as Regional_FX').font = B
pp.cell(ptop, 1, 'region').font = B
pp.cell(ptop, 4, 2030).font = B
for n_, reg in enumerate(('Mexico', 'North America', 'Latin America', 'EAA'), 1):
    rr = ptop + n_
    pp.cell(rr, 1, reg).font = B
    if reg in ('Mexico', 'North America'):
        pp.cell(rr, 4, 1).number_format = '0.0000'
        continue
    wc = L(5 + 2030 - 2016)
    isos = [i for i in order if REG[i] == reg]
    num_ = '+'.join(f'Regional_agg!{wc}{wtop + grow[i] - 2}*(1+N(Data!J{pos[(i, 2030)]}))*F{prow[i]}' for i in isos)
    den_ = '+'.join(f'Regional_agg!{wc}{wtop + grow[i] - 2}*(1+N(Data!J{pos[(i, 2030)]}))' for i in isos)
    pp.cell(rr, 4, f'=({num_})/({den_})').number_format = '0.0000'
pp.cell(ptop + 6, 1, 'Mexico and North America = 1 here: the USD/MXN PPP step is applied in the statements workbook '
        '((1 + Mexico CPI) / (1 + US CPI)). Limits: GDP weights, not Bimbo sales by country; annual-average CPI; one year (2030) as the long-run proxy.')
pp.column_dimensions['A'].width = 16
wb.save(sys.argv[1])

# Python-side replica for checking the formulas
usd = {(i, y): float(gdp[(i, y)]['v']) * (1000 if gdp[(i, y)]['unit'].endswith('T') else 1) / (1 if CUR[i] == 'USD' else fx[(CUR[i], y)])
       for i in REG for y in YRS if CUR[i] == 'USD' or (CUR[i], y) in fx}
print('Regional growth [calc] 2023..2030 and top weights 2027:')
for reg in ('Mexico', 'North America', 'Latin America', 'EAA'):
    inc = [i for i in order if REG[i] == reg and i not in EXCL]
    out = []
    for y in range(2023, 2031):
        tot = sum(usd[(i, y - 1)] for i in inc)
        out.append(sum(usd[(i, y - 1)] / tot * (float(gdp[(i, y)]['v']) / float(gdp[(i, y - 1)]['v']) - 1) for i in inc))
    tot = sum(usd[(i, 2026)] for i in inc)
    top = sorted(((usd[(i, 2026)] / tot, i) for i in inc), reverse=True)[:4]
    print(reg, ' '.join(f'{x*100:.1f}' for x in out), '|', ', '.join(f'{i} {w*100:.0f}%' for w, i in top))


def cum(iso, y):
    try:
        return (1 + cpi[(iso, y - 2)]['v'] / 100) * (1 + cpi[(iso, y - 1)]['v'] / 100) * (1 + cpi[(iso, y)]['v'] / 100) - 1
    except KeyError:
        return None


print('3-yr cumulative CPI 2023..2030 (countries reaching 60%+):')
for iso in REG:
    vals = [cum(iso, y) for y in range(2023, 2031)]
    if any(x is not None and x >= 0.6 for x in vals):
        print(iso, ' '.join('NA' if x is None else f'{x*100:.0f}' for x in vals))


print('Regional FX factor, local currency to USD [calc] 2027..2030:')
for reg in ('Latin America', 'EAA'):
    inc = [i for i in order if REG[i] == reg and i not in EXCL]
    out = []
    for y in FXY:
        tot = sum(usd[(i, y - 1)] for i in inc)
        w = {i: usd[(i, y - 1)] / tot for i in inc}
        g_ = {i: float(gdp[(i, y)]['v']) / float(gdp[(i, y - 1)]['v']) - 1 for i in inc}
        r_ = {i: 1 if CUR[i] == 'USD' else fx[(CUR[i], y - 1)] / fx[(CUR[i], y)] for i in inc}
        out.append(sum(w[i] * (1 + g_[i]) * r_[i] for i in inc) / sum(w[i] * (1 + g_[i]) for i in inc))
    print(reg, ' '.join(f'{x:.4f}' for x in out))



print('PPP region factor 2030, local currency to USD [calc]:')
_c = {(i, 2030): cpi[(i, 2030)]['v'] / 100 for i in REG if (i, 2030) in cpi}
_eur = [i for i in order if CUR[i] == 'EUR' and i not in EXCL]
_eurc = sum(usd[(i, 2029)] * _c[(i, 2030)] for i in _eur) / sum(usd[(i, 2029)] for i in _eur)
for reg in ('Latin America', 'EAA'):
    inc = [i for i in order if REG[i] == reg and i not in EXCL]
    tot = sum(usd[(i, 2029)] for i in inc)
    w = {i: usd[(i, 2029)] / tot for i in inc}
    g_ = {i: float(gdp[(i, 2030)]['v']) / float(gdp[(i, 2029)]['v']) - 1 for i in inc}
    cc = {i: _c[('USA', 2030)] if CUR[i] == 'USD' else (_eurc if CUR[i] == 'EUR' else _c[(i, 2030)]) for i in inc}
    r_ = {i: (1 + _c[('USA', 2030)]) / (1 + cc[i]) for i in inc}
    print(reg, f"{sum(w[i] * (1 + g_[i]) * r_[i] for i in inc) / sum(w[i] * (1 + g_[i]) for i in inc):.4f}")
