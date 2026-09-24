"""Build the pitch deck (.pptx) from pitch/outline.yaml on the team's own template.

Template: --template <file>, else pitch/template.pptx, else python-pptx's plain default.
Example slides already in the template are removed; its layouts, theme, fonts and colors
are kept. Every slide except the title gets a "Sources" footer (Challenge rule: slides
show their sources); speaker notes go in the notes field. Output (never overwritten):
pitch/<TICKER>_deck_v<N>.pptx.

Usage: python build_pptx.py --project <team folder> [--template <file>] [--list-layouts]
Console output is ASCII only.
"""

from __future__ import annotations

import argparse
import re
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import yaml
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.exc import PackageNotFoundError
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

EXIT_OK = 0
EXIT_INPUT_ERROR = 2
EXIT_IO_ERROR = 4
KINDS = ("title", "section", "content", "chart")
DEFAULT_MINUTES = 10.0
MAX_BULLETS = 6
MAX_BULLET_WORDS = 20
MARGIN = Inches(0.5)
FOOTER_HEIGHT = Inches(0.4)
GAP = Inches(0.15)
SOURCES_NAME = "Sources"
SOURCE_GREY = RGBColor(0x59, 0x59, 0x59)  # type: ignore[no-untyped-call]  # RGBColor.__new__ has no annotations
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
BODY_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT}


class DeckInputError(Exception):
    def __init__(self, messages: list[str]) -> None:
        super().__init__("; ".join(messages))
        self.messages = messages


@dataclass(frozen=True)
class SlideSpec:
    kind: str
    title: str
    subtitle: str = ""
    bullets: tuple[str, ...] = ()
    chart: str = ""
    notes: str = ""
    sources: str = ""
    speaker: str = ""
    minutes: float = 0.0


def ascii_safe(text: str) -> str:
    return text.encode("ascii", "replace").decode("ascii")


def _read_yaml(path: Path) -> Any:
    try:
        return yaml.safe_load(path.read_text(encoding="utf-8-sig"))
    except UnicodeDecodeError:
        raise DeckInputError([f"{path.name} is not saved as UTF-8; re-save it as UTF-8"]) from None
    except yaml.YAMLError as exc:
        raise DeckInputError([f"{path.name} is not valid YAML: {exc}"]) from None


def load_outline(path: Path, charts_dir: Path) -> tuple[dict[str, Any], list[SlideSpec], list[str]]:
    if not path.is_file():
        raise DeckInputError(["pitch/outline.yaml not found: run the pitch skill first"])
    doc = _read_yaml(path)
    if not isinstance(doc, dict) or not isinstance(doc.get("slides"), list) or not doc["slides"]:
        raise DeckInputError(["pitch/outline.yaml needs a non-empty 'slides:' list"])
    deck_raw = doc.get("deck")
    deck: dict[str, Any] = deck_raw if isinstance(deck_raw, dict) else {}
    problems: list[str] = []
    warnings: list[str] = []
    slides: list[SlideSpec] = []
    for n, raw in enumerate(doc["slides"], start=1):
        if not isinstance(raw, dict):
            problems.append(f"slide {n}: must be a mapping")
            continue
        kind, title = str(raw.get("kind", "content")), str(raw.get("title", "")).strip()
        where = f"slide {n} ({title or 'untitled'})"
        if kind not in KINDS:
            problems.append(f"{where}: kind must be one of {', '.join(KINDS)}")
        if not title:
            problems.append(f"slide {n}: title is missing")
        bullets = raw.get("bullets") or []
        if not isinstance(bullets, list) or not all(isinstance(b, str) for b in bullets):
            problems.append(f"{where}: bullets must be a list of text lines")
            bullets = []
        if len(bullets) > MAX_BULLETS:
            warnings.append(f"{where}: {len(bullets)} bullets; judges read slides fast - keep to {MAX_BULLETS} or fewer")
        if any(len(b.split()) > MAX_BULLET_WORDS for b in bullets):
            warnings.append(f"{where}: a bullet has more than {MAX_BULLET_WORDS} words; shorten it")
        chart = str(raw.get("chart") or "").strip()
        if kind == "chart" and not chart:
            problems.append(f"{where}: a chart slide needs 'chart:' (a PNG name in report/charts)")
        if chart and not (charts_dir / f"{chart}.png").is_file():
            problems.append(f"{where}: chart '{chart}' not found in report/charts (run the report skill's charts.py)")
        sources = str(raw.get("sources") or "").strip()
        if kind != "title" and not sources:
            problems.append(f"{where}: every slide except the title needs sources (Challenge rule)")
        minutes = raw.get("minutes", 0)
        if isinstance(minutes, bool) or not isinstance(minutes, (int, float)) or minutes < 0:
            problems.append(f"{where}: minutes must be a number >= 0")
            minutes = 0
        slides.append(SlideSpec(kind, title, str(raw.get("subtitle") or ""), tuple(bullets), chart,
                                str(raw.get("notes") or ""), sources, str(raw.get("speaker") or ""), float(minutes)))
    if problems:
        raise DeckInputError(problems)
    limit = deck.get("minutes", DEFAULT_MINUTES)
    total = sum(s.minutes for s in slides)
    if isinstance(limit, (int, float)) and total > limit:
        warnings.append(f"planned time is {total:g} minutes, over the {limit:g}-minute limit")
    return deck, slides, warnings


def _types(layout: Any) -> set[Any]:
    return {ph.placeholder_format.type for ph in layout.placeholders}


def _bodies(layout: Any) -> int:
    return sum(1 for ph in layout.placeholders if ph.placeholder_format.type in BODY_TYPES)


def pick_layout(prs: Any, kind: str, with_chart: bool) -> Any:
    """Choose a template layout by its placeholders (works with any team template)."""
    layouts = list(prs.slide_layouts)

    def has_title(layout: Any) -> bool:
        return bool(_types(layout) & TITLE_TYPES)

    def title_only(layout: Any) -> bool:
        return has_title(layout) and _bodies(layout) == 0 and PP_PLACEHOLDER.SUBTITLE not in _types(layout)

    preferences: list[Callable[[Any], bool]]
    if kind == "title":
        preferences = [lambda l: has_title(l) and PP_PLACEHOLDER.SUBTITLE in _types(l)]
    elif kind == "section":
        preferences = [lambda l: "section" in l.name.lower(), title_only]
    elif kind == "chart":
        preferences = [title_only]
    elif with_chart:
        preferences = [lambda l: has_title(l) and _bodies(l) == 2 and "comparison" not in l.name.lower(), title_only]
    else:
        preferences = [lambda l: has_title(l) and _bodies(l) == 1]
    for predicate in [*preferences, has_title]:
        for layout in layouts:
            if predicate(layout):
                return layout
    return layouts[0]


def _clear_slides(prs: Any) -> None:
    ids = prs.slides._sldIdLst
    for sld_id in list(ids):
        prs.part.drop_rel(sld_id.rId)
        ids.remove(sld_id)


def _remove(shape: Any) -> None:
    element = shape._element
    element.getparent().remove(element)


def _fill(text_frame: Any, lines: Sequence[str]) -> None:
    text_frame.paragraphs[0].text = lines[0] if lines else ""
    for line in lines[1:]:
        text_frame.add_paragraph().text = line


def _content_box(prs: Any, slide: Any) -> tuple[int, int, int, int]:
    title = slide.shapes.title
    top = int(title.top + title.height + GAP) if title is not None and title.top is not None else int(Inches(1.4))
    height = int(prs.slide_height - top - FOOTER_HEIGHT - GAP)
    return int(MARGIN), top, int(prs.slide_width - 2 * MARGIN), height


def _picture(slide: Any, path: Path, box: tuple[int, int, int, int]) -> None:
    left, top, width, height = box
    with Image.open(path) as image:
        w_px, h_px = image.size
    scale = min(width / w_px, height / h_px)
    w, h = int(w_px * scale), int(h_px * scale)
    slide.shapes.add_picture(str(path), Emu(left + (width - w) // 2), Emu(top + (height - h) // 2), Emu(w), Emu(h))


def _textbox_bullets(slide: Any, bullets: Sequence[str], box: tuple[int, int, int, int]) -> None:
    left, top, width, height = box
    frame = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height)).text_frame
    frame.word_wrap = True
    _fill(frame, bullets)
    for paragraph in frame.paragraphs:
        paragraph.font.size = Pt(16)
        props = paragraph._p.get_or_add_pPr()
        props.set("marL", str(int(Inches(0.3))))
        props.set("indent", str(-int(Inches(0.3))))
        props.append(props.makeelement(qn("a:buChar"), {"char": "•"}))  # bullet character (U+2022)


def _sources(prs: Any, slide: Any, text: str) -> None:
    box = slide.shapes.add_textbox(MARGIN, Emu(prs.slide_height - FOOTER_HEIGHT),
                                   Emu(prs.slide_width - 2 * MARGIN), Emu(FOOTER_HEIGHT - Inches(0.05)))
    box.name = SOURCES_NAME
    box.text_frame.word_wrap = True
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = f"Source: {text}"
    paragraph.font.size = Pt(9)
    paragraph.font.color.rgb = SOURCE_GREY


def _add_slide(prs: Any, spec: SlideSpec, charts_dir: Path) -> None:
    layout = pick_layout(prs, spec.kind, bool(spec.chart))
    slide = prs.slides.add_slide(layout)
    used: set[int] = set()
    if slide.shapes.title is not None:
        slide.shapes.title.text = spec.title
        used.add(slide.shapes.title.placeholder_format.idx)
    placeholders = sorted((ph for ph in slide.placeholders if ph.placeholder_format.idx not in used),
                          key=lambda ph: (ph.left or 0))
    subtitles = [ph for ph in placeholders if ph.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE]
    bodies = [ph for ph in placeholders if ph.placeholder_format.type in BODY_TYPES]
    if spec.kind == "title" and subtitles and spec.subtitle:
        subtitles[0].text = spec.subtitle
        used.add(subtitles[0].placeholder_format.idx)
    chart_path = charts_dir / f"{spec.chart}.png" if spec.chart else None
    if spec.kind in ("content", "chart"):
        if spec.bullets and bodies:
            _fill(bodies[0].text_frame, spec.bullets)
            used.add(bodies[0].placeholder_format.idx)
        if chart_path is not None:
            if spec.bullets and len(bodies) >= 2:
                right = bodies[1]
                box = (int(right.left), int(right.top), int(right.width), int(right.height))
                _picture(slide, chart_path, box)
            elif spec.bullets and not bodies:
                left, top, width, height = _content_box(prs, slide)
                half = width // 2
                _textbox_bullets(slide, spec.bullets, (left, top, half - int(GAP), height))
                _picture(slide, chart_path, (left + half, top, half, height))
            else:
                _picture(slide, chart_path, _content_box(prs, slide))
        elif spec.bullets and not bodies:
            _textbox_bullets(slide, spec.bullets, _content_box(prs, slide))
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx not in used:
            _remove(ph)
    if spec.sources:
        _sources(prs, slide, spec.sources)
    header = f"[{spec.speaker} - {spec.minutes:g} min] " if spec.speaker else ""
    slide.notes_slide.notes_text_frame.text = (header + spec.notes).strip()


def _ticker(project: Path) -> str:
    profile = project / "company-profile.yaml"
    try:
        doc = _read_yaml(profile) if profile.is_file() else None
    except DeckInputError:
        doc = None
    ticker = doc.get("company", {}).get("ticker") if isinstance(doc, dict) else None
    return re.sub(r"[^A-Za-z0-9]+", "-", str(ticker or "DECK")).strip("-") or "DECK"


def build(project: Path, template: Path | None) -> tuple[Path, list[str]]:
    _, slides, warnings = load_outline(project / "pitch" / "outline.yaml", project / "report" / "charts")
    try:
        prs = Presentation(str(template)) if template is not None else Presentation()
    except PackageNotFoundError:
        raise DeckInputError([f"{template} is not a valid .pptx file"]) from None
    _clear_slides(prs)
    for spec in slides:
        _add_slide(prs, spec, project / "report" / "charts")
    out_dir = project / "pitch"
    ticker = _ticker(project)
    pattern = re.compile(rf"^{re.escape(ticker)}_deck_v(\d+)\.pptx$")
    versions = [int(m.group(1)) for p in out_dir.glob("*.pptx") if (m := pattern.match(p.name))]
    target = out_dir / f"{ticker}_deck_v{max(versions, default=0) + 1}.pptx"
    prs.save(str(target))
    return target, warnings


def list_layouts(template: Path | None) -> None:
    prs = Presentation(str(template)) if template is not None else Presentation()
    for i, layout in enumerate(prs.slide_layouts):
        kinds = ", ".join(sorted(str(ph.placeholder_format.type).split(" ")[0]
                                 for ph in layout.placeholders))  # type: ignore[misc]  # pptx __iter__ typing quirk
        print(f"{i}: {ascii_safe(layout.name)} | placeholders: {ascii_safe(kinds)}")


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build the pitch deck from pitch/outline.yaml.")
    parser.add_argument("--project", default=".", help="Team project folder")
    parser.add_argument("--template", help="Team .pptx template (default: pitch/template.pptx)")
    parser.add_argument("--list-layouts", action="store_true", help="List the template's layouts and exit")
    args = parser.parse_args(argv)
    project = Path(args.project).resolve()
    template: Path | None = Path(args.template).resolve() if args.template else project / "pitch" / "template.pptx"
    if template is not None and not template.is_file():
        if args.template:
            print(f"[x] Template not found: {ascii_safe(args.template)}")
            return EXIT_INPUT_ERROR
        print("[warn] no template found at pitch/template.pptx: using a plain default (put your team template there)")
        template = None
    try:
        if args.list_layouts:
            list_layouts(template)
            return EXIT_OK
        target, warnings = build(project, template)
    except DeckInputError as exc:
        print("[x] Cannot build the deck. Fix these first:")
        for message in exc.messages:
            print("  - " + ascii_safe(message))
        return EXIT_INPUT_ERROR
    except OSError as exc:
        print(f"[x] Could not write the deck: {ascii_safe(exc.strerror or str(exc))}. Close it in PowerPoint and run again.")
        return EXIT_IO_ERROR
    for warning in warnings:
        print("[warn] " + ascii_safe(warning))
    print(f"[ok] deck -> pitch/{ascii_safe(target.name)}")
    print("Next: run audit_pptx.py, then open the deck in PowerPoint and check every slide.")
    return EXIT_OK


if __name__ == "__main__":
    raise SystemExit(main())
