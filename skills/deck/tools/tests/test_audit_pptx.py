from __future__ import annotations

from collections.abc import Callable
from pathlib import Path
from typing import Any

import pytest
from deck_fixtures import deck_project, write_template  # noqa: F401  (import first: sets sys.path)
from pptx import Presentation
from pptx.util import Inches, Pt

from audit_pptx import Finding, audit, main
from build_pptx import main as build


def _built(project: Path) -> Path:
    assert build(["--project", str(project)]) == 0
    return project / "pitch" / "ACME_deck_v1.pptx"


def _codes(findings: list[Finding]) -> set[str]:
    return {f.code for f in findings}


def _slides_with(findings: list[Finding], code: str) -> set[int]:
    return {f.slide for f in findings if f.code == code}


def _deck(project: Path, fill: Callable[[Any], None]) -> Path:
    """A deck on the plain default template whose slides are added by `fill`."""
    prs = Presentation()
    fill(prs)
    path = project / "pitch" / "ACME_deck_v1.pptx"
    prs.save(str(path))
    return path


def _slide(prs: Any, title: str, layout: int = 5) -> Any:
    slide = prs.slides.add_slide(prs.slide_layouts[layout])
    slide.shapes.title.text = title
    slide.notes_slide.notes_text_frame.text = "Real notes."
    return slide


def _box(slide: Any, text: str, top: float = 6.9, left: float = 0.5, width: float = 9.0, height: float = 0.4) -> Any:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.text_frame.text = text
    return box


def test_built_deck_has_no_errors(deck_project: Path) -> None:
    findings = audit(_built(deck_project), deck_project)
    assert not [f for f in findings if f.severity == "ERROR"], findings
    assert not _codes(findings) & {"overflow", "overlap", "leftover", "stale-target", "stale-wacc"}, findings


def test_planted_problems_are_found(deck_project: Path) -> None:
    path = _built(deck_project)
    prs = Presentation(str(path))
    extra = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only, no sources, no notes
    extra.shapes.title.text = "Our target price is MXN 40.0 with a WACC of 9.0%"
    box = extra.shapes.add_textbox(Inches(9), Inches(1), Inches(3), Inches(1))  # runs off a 10in slide
    box.text_frame.text = "Click to add text"
    box.text_frame.paragraphs[0].font.size = Pt(7)
    box.text_frame.paragraphs[0].font.name = "Comic Sans MS"
    prs.save(str(path))
    codes = _codes(audit(path, deck_project))
    for code in ("no-sources", "no-notes", "leftover", "off-slide", "small-font", "stale-target", "stale-wacc"):
        assert code in codes, code


def test_too_many_slides(deck_project: Path) -> None:
    path = _built(deck_project)
    prs = Presentation(str(path))
    for _ in range(12):
        prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(path))
    assert "slide-count" in _codes(audit(path, deck_project))


def test_report_file_and_ascii_output(deck_project: Path, capsys: pytest.CaptureFixture[str]) -> None:
    _built(deck_project)
    assert main(["--project", str(deck_project)]) == 0
    assert (deck_project / "pitch" / "deck-audit.md").is_file()
    out = capsys.readouterr().out
    assert out.isascii() and "ACME_deck_v1.pptx" in out


def test_fix_removes_empty_placeholders_and_off_palette_fonts(deck_project: Path) -> None:
    path = _built(deck_project)
    prs = Presentation(str(path))
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + empty body placeholder
    slide.shapes.title.text = "Extra"
    note = _box(slide, "A stray pasted note", top=2)
    note.text_frame.paragraphs[0].runs[0].font.name = "Comic Sans MS"
    prs.save(str(path))
    assert main(["--project", str(deck_project), "--fix"]) == 0
    fixed = Presentation(str(deck_project / "pitch" / "ACME_deck_v2.pptx"))
    last = fixed.slides[len(fixed.slides) - 1]
    assert all(ph.text_frame.text.strip() for ph in last.placeholders if ph.has_text_frame)
    fonts = {r.font.name for sh in last.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs}
    assert "Comic Sans MS" not in fonts
    assert (deck_project / "pitch" / "ACME_deck_v1.pptx").is_file()  # original kept


def test_fix_never_touches_theme_or_title_fonts(deck_project: Path) -> None:
    def fill(prs: Any) -> None:
        for n in range(8):
            slide = _slide(prs, f"Title {n}", layout=1)
            slide.shapes.title.text_frame.paragraphs[0].runs[0].font.name = "Montserrat"  # brand font, titles only
            body = slide.placeholders[1].text_frame
            body.text = "A long body text written in the theme font that dominates the character count " * 3
            theme = body.add_paragraph()
            theme.text = "Short theme-font line"
            theme.runs[0].font.name = "Calibri"  # the default theme's font, set explicitly
            stray = body.add_paragraph()
            stray.text = "Pasted line"
            stray.runs[0].font.name = "Comic Sans MS"
    _deck(deck_project, fill)
    assert main(["--project", str(deck_project), "--fix"]) == 0
    fixed = Presentation(str(deck_project / "pitch" / "ACME_deck_v2.pptx"))
    for slide in fixed.slides:
        assert slide.shapes.title.text_frame.paragraphs[0].runs[0].font.name == "Montserrat"
        paragraphs = slide.placeholders[1].text_frame.paragraphs
        assert paragraphs[1].runs[0].font.name == "Calibri"
        assert paragraphs[2].runs[0].font.name is None


def test_fix_without_changes_saves_no_new_version(deck_project: Path, capsys: pytest.CaptureFixture[str]) -> None:
    _built(deck_project)
    assert main(["--project", str(deck_project), "--fix"]) == 0
    assert "no mechanical fixes needed" in capsys.readouterr().out
    assert not (deck_project / "pitch" / "ACME_deck_v2.pptx").exists()


def test_no_deck_exits_2(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    (tmp_path / "pitch").mkdir()
    assert main(["--project", str(tmp_path)]) == 2
    assert "no deck" in capsys.readouterr().out


# --- sources, leftovers, notes -----------------------------------------------------------------

def test_sources_in_groups_tables_last_paragraphs_and_other_languages(deck_project: Path) -> None:
    def fill(prs: Any) -> None:
        _slide(prs, "Cover", layout=0)
        group = _slide(prs, "Grouped footer").shapes.add_group_shape()
        group.shapes.add_textbox(Inches(0.5), Inches(6.9), Inches(8), Inches(0.4)).text_frame.text = "Source: Bloomberg"
        table = _slide(prs, "Table source").shapes.add_table(2, 1, Inches(0.5), Inches(2), Inches(9), Inches(1)).table
        table.cell(0, 0).text = "Revenue 2025"
        table.cell(1, 0).text = "Source: team model"
        body = _box(_slide(prs, "Body"), "Revenue grows 6%", top=2, height=3)
        body.text_frame.add_paragraph().text = "Source: company filings"
        _box(_slide(prs, "Tesis"), "Fuente: Bloomberg")
        _box(_slide(prs, "Mercado"), "FUENTES: INEGI")
        _box(_slide(prs, "Tese"), "Fonte: B3")
        _box(_slide(prs, "Sourcing strategy"), "Sourcing costs fell 5% in 2025", top=2)
    findings = audit(_deck(deck_project, fill), deck_project)
    assert _slides_with(findings, "no-sources") == {8}  # only "Sourcing strategy"


def test_leftovers_are_case_aware_and_multilingual(deck_project: Path) -> None:
    texts = ["Crecemos en todos los segmentos", "TODO: add comps", "Haga clic para agregar texto",
             "Clique para adicionar texto", "Margins TBD", "Lorem ipsum dolor", "[Insert chart]",
             "Click to add notes", "Methodology to do list"]

    def fill(prs: Any) -> None:
        for n, text in enumerate(texts):
            slide = _slide(prs, f"Slide {n}")
            _box(slide, text, top=2)
            _box(slide, "Source: team")
    findings = audit(_deck(deck_project, fill), deck_project)
    assert _slides_with(findings, "leftover") == {2, 3, 4, 5, 6, 7, 8}


def test_notes_with_only_the_speaker_header_count_as_missing(deck_project: Path) -> None:
    findings = audit(_built(deck_project), deck_project)
    assert _slides_with(findings, "no-notes") == {5}  # the fixture's section slide: "[Eva - 0.2 min]" only


# --- overflow and overlap --------------------------------------------------------------------

def test_overflow_uses_the_inherited_placeholder_size(deck_project: Path) -> None:
    def fill(prs: Any) -> None:
        dense = _slide(prs, "Dense at the inherited 32 pt", layout=1)
        frame = dense.placeholders[1].text_frame
        frame.text = "Mexican packaged food demand keeps growing as urban households trade up to branded products"
        for _ in range(6):
            frame.add_paragraph().text = "Margins expand on lower input costs, better mix and plant utilisation"
        light = _slide(prs, "Two short bullets", layout=1)
        light.placeholders[1].text_frame.text = "Volume recovery"
    findings = audit(_deck(deck_project, fill), deck_project)
    assert _slides_with(findings, "overflow") == {1}


def test_overflow_of_vertical_text_swaps_width_and_height(deck_project: Path) -> None:
    """Vertical text (inherited from the layout) runs down the box: paragraphs are columns across its width."""
    def fill(prs: Any) -> None:
        vertical = next(item for item in prs.slide_layouts if item.name == "Vertical Title and Text")
        for count in (5, 3):  # a 2-inch-wide vertical body at 32 pt holds about 4 columns
            slide = prs.slides.add_slide(vertical)
            slide.shapes.title.text = "Thesis"
            body = slide.placeholders[1]
            body.left, body.top, body.width, body.height = Inches(0.5), Inches(0.3), Inches(2), Inches(6.4)
            body.text_frame.text = "Volume up"
            for _ in range(count - 1):
                body.text_frame.add_paragraph().text = "Volume up"
            slide.notes_slide.notes_text_frame.text = "Real notes."
    findings = audit(_deck(deck_project, fill), deck_project)
    assert _slides_with(findings, "overflow") == {1}


def test_fix_on_a_hand_made_deck_saves_numbered_copies(deck_project: Path, tmp_path_factory: pytest.TempPathFactory,
                                                      capsys: pytest.CaptureFixture[str]) -> None:
    def fill(prs: Any) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + empty body placeholder
        slide.shapes.title.text = "Hand made"
    inside = _deck(deck_project, fill).rename(deck_project / "pitch" / "Our deck.pptx")
    outside_dir = tmp_path_factory.mktemp("elsewhere")  # outside the project folder
    outside = outside_dir / "Our deck.pptx"
    outside.write_bytes(inside.read_bytes())

    assert main(["--project", str(deck_project), "--deck", str(inside), "--fix"]) == 0
    assert "-> pitch/Our deck_fixed1.pptx" in capsys.readouterr().out
    assert main(["--project", str(deck_project), "--deck", str(inside), "--fix"]) == 0
    assert "-> pitch/Our deck_fixed2.pptx" in capsys.readouterr().out
    assert inside.is_file() and (deck_project / "pitch" / "Our deck_fixed2.pptx").is_file()

    assert main(["--project", str(deck_project), "--deck", str(outside), "--fix"]) == 0
    out = capsys.readouterr().out
    assert out.isascii() and f"-> {outside_dir.resolve() / 'Our deck_fixed1.pptx'}" in out


def test_overlap_of_pictures_on_titles_and_of_the_sources_line(deck_project: Path) -> None:
    chart = deck_project / "report" / "charts" / "football-field.png"

    def fill(prs: Any) -> None:
        on_title = _slide(prs, "Picture over the title")
        on_title.shapes.add_picture(str(chart), Inches(1), Inches(0.5), Inches(4), Inches(2))
        _box(on_title, "Source: team")
        crowded = _slide(prs, "Footer over a shape")
        _box(crowded, "Key message box", top=6.5, height=0.8)
        _box(crowded, "Source: team")
        clean = _slide(prs, "Clean")
        clean.shapes.add_picture(str(chart), Inches(1), Inches(2), Inches(6), Inches(3))
        _box(clean, "Source: team")
    findings = audit(_deck(deck_project, fill), deck_project)
    assert _slides_with(findings, "overlap") == {1, 2}


# --- stale numbers ---------------------------------------------------------------------------

@pytest.mark.parametrize(("text", "code"), [
    ("Target price MXN 40.0", "stale-target"),
    ("precio objetivo MXN 40.0", "stale-target"),
    ("Precio objetivo: MXN 31,0", None),
    ("preço-alvo R$ 40,00", "stale-target"),
    ("preco-alvo R$ 31,00", None),
    ("TP: MXN 31.0", None),
    ("TP MXN 40", "stale-target"),
    ("TP (12m): MXN 31.0", None),
    ("Price target for 2027: MXN 31.0", None),
    ("Management targets 2027 revenue of MXN 5bn; we target a 20% margin", None),
    ("Target price implies 24% upside", None),
    ("WACC of 9,0%", "stale-wacc"),
    ("WACC 12.45%", None),
    ("WACC 12,4%", None),
])
def test_stale_numbers(deck_project: Path, text: str, code: str | None) -> None:
    def fill(prs: Any) -> None:
        slide = _slide(prs, "Valuation")
        _box(slide, text, top=2)
        _box(slide, "Source: team")
    codes = _codes(audit(_deck(deck_project, fill), deck_project)) & {"stale-target", "stale-wacc"}
    assert codes == ({code} if code else set()), text


def test_stale_target_in_a_table_cell(deck_project: Path) -> None:
    def fill(prs: Any) -> None:
        slide = _slide(prs, "Summary")
        table = slide.shapes.add_table(2, 2, Inches(0.5), Inches(2), Inches(9), Inches(1)).table
        table.cell(0, 0).text = "Target price"
        table.cell(1, 0).text = "Target price MXN 40.0"
        _box(slide, "Source: team")
    assert "stale-target" in _codes(audit(_deck(deck_project, fill), deck_project))


def test_unreadable_reference_files_are_reported_without_crashing(deck_project: Path,
                                                                  capsys: pytest.CaptureFixture[str]) -> None:
    _built(deck_project)
    (deck_project / "report" / "header.yaml").write_text("target_price: [1,\n", encoding="utf-8")
    (deck_project / "model" / "model-summary.json").write_text("{bad", encoding="utf-8")
    assert main(["--project", str(deck_project)]) == 0
    out = capsys.readouterr().out
    assert out.isascii() and "header.yaml" in out and "model-summary.json" in out
    (deck_project / "model" / "model-summary.json").write_text('{"valuation": null}', encoding="utf-8")
    assert main(["--project", str(deck_project)]) == 0
