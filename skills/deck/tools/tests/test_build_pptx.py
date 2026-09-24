from __future__ import annotations

import copy
from pathlib import Path
from typing import Any

import pytest
from deck_fixtures import OUTLINE, deck_project, write_deck_project  # noqa: F401  (import first: sets sys.path)
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from build_pptx import SOURCES_NAME, main, pick_layout


def _build(project: Path) -> Any:
    assert main(["--project", str(project)]) == 0
    return Presentation(str(project / "pitch" / "ACME_deck_v1.pptx"))


def test_one_slide_per_outline_entry_and_template_examples_removed(deck_project: Path) -> None:
    prs = _build(deck_project)
    assert len(prs.slides) == len(OUTLINE["slides"])
    titles = [s.shapes.title.text if s.shapes.title is not None else "" for s in prs.slides]
    assert "Example slide from the template" not in titles
    assert titles[1] == "Investment thesis"


def test_sources_footer_on_every_slide_but_the_title(deck_project: Path) -> None:
    prs = _build(deck_project)
    for i, slide in enumerate(prs.slides):
        if i == 0:
            continue
        footers = [sh for sh in slide.shapes if sh.name == SOURCES_NAME]
        assert len(footers) == 1, i
        assert footers[0].text_frame.text.startswith("Source: ")


def test_speaker_notes_in_the_notes_field(deck_project: Path) -> None:
    prs = _build(deck_project)
    notes = prs.slides[1].notes_slide.notes_text_frame.text
    assert "Ana" in notes and "Three pillars." in notes


def test_charts_placed_inside_the_slide(deck_project: Path) -> None:
    prs = _build(deck_project)
    pictures = [sh for s in prs.slides for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 2
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.left is None or sh.width is None:
                continue
            assert sh.left >= 0 and sh.top >= 0
            assert sh.left + sh.width <= prs.slide_width and sh.top + sh.height <= prs.slide_height


def test_bullets_land_on_the_slide(deck_project: Path) -> None:
    prs = _build(deck_project)
    text = " ".join(sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame)
    assert "Volume recovery" in text and "Margin expansion" in text


def test_no_empty_placeholders_left(deck_project: Path) -> None:
    prs = _build(deck_project)
    for slide in prs.slides:
        for ph in slide.placeholders:
            assert not ph.has_text_frame or ph.text_frame.text.strip(), ph.name


def test_versions_are_never_overwritten(deck_project: Path) -> None:
    _build(deck_project)
    assert main(["--project", str(deck_project)]) == 0
    assert (deck_project / "pitch" / "ACME_deck_v2.pptx").is_file()


def test_plain_default_when_no_template(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    project = write_deck_project(tmp_path, template=False)
    assert main(["--project", str(project)]) == 0
    out = capsys.readouterr().out
    assert out.isascii() and "no template" in out


def test_missing_sources_is_an_error(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    outline = copy.deepcopy(OUTLINE)
    outline["slides"][1]["sources"] = ""
    project = write_deck_project(tmp_path, outline=outline)
    assert main(["--project", str(project)]) == 2
    assert "sources" in capsys.readouterr().out


def test_missing_chart_is_an_error(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    project = write_deck_project(tmp_path, charts=False)
    assert main(["--project", str(project)]) == 2
    assert "revenue-margin" in capsys.readouterr().out


def test_long_outline_warns_about_time(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    outline = copy.deepcopy(OUTLINE)
    outline["slides"][1]["minutes"] = 9
    project = write_deck_project(tmp_path, outline=outline)
    assert main(["--project", str(project)]) == 0
    assert "minutes" in capsys.readouterr().out


def test_list_layouts(deck_project: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["--project", str(deck_project), "--list-layouts"]) == 0
    out = capsys.readouterr().out
    assert out.isascii() and "Title Slide" in out and "Two Content" in out


def test_layout_choice_on_the_default_template() -> None:
    prs = Presentation()
    assert pick_layout(prs, "title", False).name == "Title Slide"
    assert pick_layout(prs, "section", False).name == "Section Header"
    assert pick_layout(prs, "content", False).name == "Title and Content"
    assert pick_layout(prs, "content", True).name == "Two Content"
    assert pick_layout(prs, "chart", True).name == "Title Only"
