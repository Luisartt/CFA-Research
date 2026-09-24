"""Audit a pitch deck against the Challenge rules and basic presentation hygiene.

Structural checks only (no rendering): slide count for a 10-minute talk, a source on
every slide but the title ("Source:", "Fuente:" or "Fonte:" anywhere, groups and tables
included), speaker notes, leftover placeholder text, shapes off the slide, small fonts,
too many font families, likely text overflow (at the size the text really inherits),
pictures over titles or a sources line over other shapes, dense slides, and numbers that
disagree with model/model-summary.json and report/header.yaml.
--fix saves a NEW version with mechanical fixes only (empty placeholders removed, rare
pasted fonts in body text reset to the theme font; theme and title fonts are never
touched); nothing is saved when there is nothing to fix. Everything else is reported.
Writes pitch/deck-audit.md. Console output is ASCII only.

Usage: python audit_pptx.py --project <team folder> [--deck <file>] [--fix]
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import yaml
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.exc import PackageNotFoundError
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches

from build_pptx import Box, ascii_safe, decor_boxes, shape_box, vertical_text

EXIT_OK = 0
EXIT_NO_DECK = 2
EXIT_IO_ERROR = 4
MIN_SLIDES, MAX_SLIDES = 6, 14
MIN_FONT_PT, MIN_FOOTER_PT = 10, 8
MAX_WORDS = 60
MAX_FONT_FAMILIES = 2
DEFAULT_FONT_PT = 18
EMU_PER_PT = 12700
OVERLAP_TOLERANCE = Inches(0.05)
SOURCE_RE = re.compile(r"^\s*(sources?|fuentes?|fontes?)\s*:", re.IGNORECASE)
LEFTOVERS = (re.compile(r"\bTODO\b"), re.compile(r"\bTBD\b"), re.compile(r"click to add", re.IGNORECASE),
             re.compile(r"haga clic para agregar", re.IGNORECASE), re.compile(r"clique para adicionar", re.IGNORECASE),
             re.compile(r"lorem ipsum", re.IGNORECASE), re.compile(r"\[insert", re.IGNORECASE),
             re.compile(r"\bxx\.x\b", re.IGNORECASE))
NOTES_HEADER_RE = re.compile(r"^\s*\[[^\]]*\bmin\]")
TARGET_RE = re.compile(r"target price|price target|\bTP\b|precio objetivo|pre[cç]o-alvo", re.IGNORECASE)
TARGET_WINDOW = 20  # characters between the words and the number
NUMBER_RE = re.compile(r"\d+(?:[.,]\d+)*")
YEAR_RE = re.compile(r"(?:19|20)\d{2}")
WACC_RE = re.compile(r"WACC\D{0,15}?(\d+(?:[.,]\d+)?)\s*%", re.IGNORECASE)
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}


@dataclass(frozen=True)
class Finding:
    slide: int  # 0 = whole deck
    severity: str  # ERROR or WARN
    code: str
    message: str


def _texts(slide: Any) -> Iterator[tuple[Any, str]]:
    for shape in slide.shapes:
        if shape.has_text_frame:
            yield shape, shape.text_frame.text


def _frames(shapes: Any) -> Iterator[Any]:
    """Every text frame on a slide: shapes, shapes inside groups, and table cells."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _frames(shape.shapes)
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame
        elif shape.has_text_frame:
            yield shape.text_frame


def _load_reference(project: Path) -> tuple[dict[str, float], list[Finding]]:
    ref: dict[str, float] = {}
    problems: list[Finding] = []
    header = project / "report" / "header.yaml"
    if header.is_file():
        try:
            doc = yaml.safe_load(header.read_text(encoding="utf-8-sig"))
        except (yaml.YAMLError, UnicodeDecodeError):
            doc = None
            problems.append(Finding(0, "WARN", "reference", "report/header.yaml is unreadable: target check skipped"))
        if isinstance(doc, dict) and isinstance(doc.get("target_price"), (int, float)):
            ref["target"] = float(doc["target_price"])
    summary = project / "model" / "model-summary.json"
    if summary.is_file():
        try:
            data = json.loads(summary.read_text(encoding="utf-8-sig"))
        except (json.JSONDecodeError, UnicodeDecodeError):
            data = None
            problems.append(Finding(0, "WARN", "reference",
                                    "model/model-summary.json is unreadable: WACC check skipped"))
        valuation = data.get("valuation") if isinstance(data, dict) else None
        wacc = valuation.get("wacc") if isinstance(valuation, dict) else None
        if isinstance(wacc, (int, float)) and not isinstance(wacc, bool):
            ref["wacc_pct"] = float(wacc) * 100
    return ref, problems


# --- effective font size ---------------------------------------------------------------------

def _lvl_size(container: Any, level: int) -> float | None:
    """sz of a:lvlNpPr/a:defRPr in a list style (a:lstStyle, p:titleStyle, p:bodyStyle...)."""
    if container is None:
        return None
    props = container.find(qn(f"a:lvl{level}pPr"))
    run = props.find(qn("a:defRPr")) if props is not None else None
    size = run.get("sz") if run is not None else None
    return int(size) / 100 if size else None


def _inherited_size(shape: Any, level: int) -> float:
    """The size text inherits: the shape's list style, then its layout and master placeholders, then the
    master text styles; 18 pt when nothing says (plain text boxes)."""
    chain: list[Any] = [shape]
    if shape.is_placeholder:
        base = getattr(shape, "_base_placeholder", None)
        while base is not None and len(chain) < 4:
            chain.append(base)
            base = getattr(base, "_base_placeholder", None)
    for item in chain:
        size = _lvl_size(item._element.find(f"{qn('p:txBody')}/{qn('a:lstStyle')}"), level)
        if size:
            return size
    if shape.is_placeholder:
        styles = shape.part.slide_layout.slide_master._element.find(qn("p:txStyles"))
        kind = shape.placeholder_format.type
        name = "p:titleStyle" if kind in TITLE_TYPES else \
            "p:otherStyle" if kind in (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER) \
            else "p:bodyStyle"
        size = _lvl_size(styles.find(qn(name)) if styles is not None else None, level)
        if size:
            return size
    return DEFAULT_FONT_PT


def _font_scale(shape: Any) -> float:
    autofit = shape._element.find(f"{qn('p:txBody')}/{qn('a:bodyPr')}/{qn('a:normAutofit')}")
    scale = autofit.get("fontScale") if autofit is not None else None
    return int(scale) / 100000 if scale else 1.0


def _size(shape: Any, paragraph: Any, run: Any | None) -> float:
    if run is not None and run.font.size is not None:
        return float(run.font.size.pt)
    if paragraph.font.size is not None:
        return float(paragraph.font.size.pt)
    return _inherited_size(shape, min(paragraph.level, 8) + 1) * _font_scale(shape)


def _overflow(shape: Any, text: str) -> bool:
    if shape.width is None or shape.height is None or not text.strip():
        return False
    sizes = [_size(shape, p, r) for p in shape.text_frame.paragraphs for r in (p.runs or [None])]
    size = min(sizes) if sizes else DEFAULT_FONT_PT
    along, across = int(shape.width), int(shape.height)
    if vertical_text(shape):  # lines run down the box and stack across its width
        along, across = across, along
    chars_per_line = max(1.0, (along / EMU_PER_PT) / (size * 0.5))
    lines = sum(max(1, -(-len(p.text) // int(chars_per_line))) for p in shape.text_frame.paragraphs)
    capacity = (across / EMU_PER_PT) / (size * 1.2)
    return bool(lines > capacity + 0.5)


# --- numbers -----------------------------------------------------------------------------------

def _number(raw: str) -> float | None:
    """'31.0', '31,0', '1,234.5', '1.234,5', '3,100' -> float."""
    if "," in raw and "." in raw:
        decimal = "," if raw.rfind(",") > raw.rfind(".") else "."
        raw = raw.replace("." if decimal == "," else ",", "").replace(decimal, ".")
    elif "," in raw or "." in raw:
        sep = "," if "," in raw else "."
        parts = raw.split(sep)
        thousands = len(parts) > 2 or len(parts[-1]) == 3
        raw = raw.replace(sep, "") if thousands else raw.replace(sep, ".")
    try:
        return float(raw)
    except ValueError:
        return None


def _targets(text: str) -> Iterator[float]:
    """Numbers right after 'target price' words: not percentages, years or numbers glued to units (12m, 5bn)."""
    for key in TARGET_RE.finditer(text):
        for match in NUMBER_RE.finditer(text, key.end()):
            if match.start() - key.end() > TARGET_WINDOW:
                break
            after = text[match.end():]
            if after.lstrip().startswith("%") or YEAR_RE.fullmatch(match.group()) or after[:1].isalpha():
                continue
            if (value := _number(match.group())) is not None:
                yield value
            break


# --- overlap -----------------------------------------------------------------------------------

def _overlaps(a: Box, b: Box) -> bool:
    width, height = a.overlap(b)
    return width > OVERLAP_TOLERANCE and height > OVERLAP_TOLERANCE


def _overlap_findings(n: int, slide: Any, width: int, height: int) -> list[Finding]:
    findings: list[Finding] = []
    title = slide.shapes.title
    title_box = shape_box(title) if title is not None else None
    boxes = [(shape, box) for shape in slide.shapes if (box := shape_box(shape)) is not None]
    for shape, box in boxes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and title_box is not None and _overlaps(box, title_box):
            findings.append(Finding(n, "WARN", "overlap", f"picture '{shape.name}' overlaps the title"))
    for footer, footer_box in boxes:
        if not (footer.has_text_frame and SOURCE_RE.match(footer.text_frame.text)):
            continue
        others = [(s.name, b) for s, b in boxes if s is not footer]
        others += [("the template's logo/decoration", b) for b in decor_boxes(slide.slide_layout, width, height)]
        for name, box in others:
            if _overlaps(footer_box, box):
                findings.append(Finding(n, "WARN", "overlap", f"the sources line overlaps '{name}'"))
                break
    return findings


def audit(path: Path, project: Path) -> list[Finding]:
    prs = Presentation(str(path))
    reference, findings = _load_reference(project)
    count = len(prs.slides)
    if not MIN_SLIDES <= count <= MAX_SLIDES:
        findings.append(Finding(0, "WARN", "slide-count",
                                f"{count} slides; a 10-minute talk usually needs {MIN_SLIDES}-{MAX_SLIDES}"))
    families: Counter[str] = Counter()
    for n, slide in enumerate(prs.slides, start=1):
        frames = list(_frames(slide.shapes))
        all_text = " ".join(frame.text for frame in frames)
        if n > 1 and not any(SOURCE_RE.match(p.text) for frame in frames for p in frame.paragraphs):
            findings.append(Finding(n, "ERROR", "no-sources", "no 'Source:' line (Challenge rule: slides show sources)"))
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        if not NOTES_HEADER_RE.sub("", notes, count=1).strip():
            findings.append(Finding(n, "WARN", "no-notes", "no speaker notes"))
        for marker in LEFTOVERS:
            if found := marker.search(all_text):
                findings.append(Finding(n, "ERROR", "leftover", f"placeholder text left: '{found.group()}'"))
        if len(all_text.split()) > MAX_WORDS:
            findings.append(Finding(n, "WARN", "dense", f"{len(all_text.split())} words; aim for under {MAX_WORDS}"))
        for shape in slide.shapes:
            if None not in (shape.left, shape.top, shape.width, shape.height) and (
                    shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width
                    or shape.top + shape.height > prs.slide_height):
                findings.append(Finding(n, "ERROR", "off-slide", f"'{shape.name}' extends beyond the slide"))
        for shape, text in _texts(slide):
            is_footer = bool(SOURCE_RE.match(text))
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.name:
                        families[run.font.name] += len(run.text)
                    size = run.font.size or paragraph.font.size
                    if size is not None and size.pt < (MIN_FOOTER_PT if is_footer else MIN_FONT_PT):
                        findings.append(Finding(n, "WARN", "small-font", f"{size.pt:g} pt text in '{shape.name}'"))
            if _overflow(shape, text):
                findings.append(Finding(n, "WARN", "overflow", f"text may not fit in '{shape.name}'"))
        findings += _overlap_findings(n, slide, int(prs.slide_width or 0), int(prs.slide_height or 0))
        if "target" in reference:
            for value in _targets(all_text):
                if abs(value - reference["target"]) > 0.01 * reference["target"]:
                    findings.append(Finding(n, "WARN", "stale-target",
                                            f"target {value:g} differs from report/header.yaml ({reference['target']:g})"))
        if "wacc_pct" in reference:
            for match in WACC_RE.finditer(all_text):
                wacc = float(match.group(1).replace(",", "."))
                if abs(wacc - reference["wacc_pct"]) > 0.1:
                    findings.append(Finding(n, "WARN", "stale-wacc",
                                            f"WACC {match.group(1)}% differs from the model ({reference['wacc_pct']:.1f}%)"))
    if len(families) > MAX_FONT_FAMILIES:
        findings.append(Finding(0, "WARN", "fonts", "fonts used: " + ", ".join(sorted(families))))
    return findings


RARE_FONT_SHARE = 0.2


def _theme_fonts(prs: Any) -> set[str]:
    fonts: set[str] = set()
    for master in prs.slide_masters:
        theme = parse_xml(master.part.part_related_by(RT.THEME).blob)
        for tag in ("a:majorFont", "a:minorFont"):
            latin = theme.find(f".//{qn(tag)}/{qn('a:latin')}")
            if latin is not None and latin.get("typeface"):
                fonts.add(latin.get("typeface"))
    return fonts


def _is_title(shape: Any) -> bool:
    return bool(shape.is_placeholder and shape.placeholder_format.type in TITLE_TYPES)


def fix(path: Path, target: Path) -> int:
    """Mechanical fixes only; returns the number of changes. Saves to `target` only when there are changes.

    - removes empty placeholders;
    - a font used for less than 20% of the deck's text in body text is reset to the template's
      theme font (explicit font removed), so stray pasted fonts disappear. Theme fonts and any
      font used on titles (a brand font) are never touched.
    """
    prs = Presentation(str(path))
    families: Counter[str] = Counter()
    protected = _theme_fonts(prs)
    total = 0
    for slide in prs.slides:
        for shape, _ in _texts(slide):
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    total += len(run.text)
                    if run.font.name:
                        families[run.font.name] += len(run.text)
                        if _is_title(shape):
                            protected.add(run.font.name)
    rare = {name for name, chars in families.items()
            if chars < RARE_FONT_SHARE * total and name not in protected and not name.startswith("+")}
    changes = 0
    for slide in prs.slides:
        for ph in list(slide.placeholders):
            if ph.has_text_frame and not ph.text_frame.text.strip():
                ph._element.getparent().remove(ph._element)
                changes += 1
        for shape, _ in _texts(slide):
            if _is_title(shape):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.name in rare:
                        run.font.name = None
                        changes += 1
    if changes:
        prs.save(str(target))
    return changes


def _latest_deck(pitch: Path) -> Path | None:
    decks = sorted(pitch.glob("*_deck_v*.pptx"),
                   key=lambda p: int(m.group(1)) if (m := re.search(r"_v(\d+)\.pptx$", p.name)) else 0)
    return decks[-1] if decks else None


def _next_version(path: Path) -> Path:
    """<name>_deck_v<N+1>.pptx for built decks; otherwise the next free <stem>_fixed<N>.pptx (never overwrites)."""
    match = re.match(r"^(.*_deck_v)(\d+)\.pptx$", path.name)
    if not match:
        base = re.sub(r"_fixed\d+$", "", path.stem)
        pattern = re.compile(rf"^{re.escape(base)}_fixed(\d+)\.pptx$", re.IGNORECASE)
        used = [int(m.group(1)) for p in path.parent.glob("*.pptx") if (m := pattern.match(p.name))]
        return path.with_name(f"{base}_fixed{max(used, default=0) + 1}.pptx")
    stem = match.group(1)
    versions = [int(m.group(1)) for p in path.parent.glob(f"{stem}*.pptx")
                if (m := re.search(r"_v(\d+)\.pptx$", p.name))]
    return path.with_name(f"{stem}{max(versions) + 1}.pptx")


def _shown(path: Path, project: Path) -> str:
    """The path relative to the project when it is inside it, else absolute (ASCII-safe)."""
    try:
        return ascii_safe(path.relative_to(project).as_posix())
    except ValueError:
        return ascii_safe(str(path))


def write_report(project: Path, deck: Path, findings: Sequence[Finding]) -> Path:
    lines = [f"# Deck audit - {deck.name}", "",
             f"{sum(f.severity == 'ERROR' for f in findings)} errors, {sum(f.severity == 'WARN' for f in findings)} warnings.",
             "", "| Slide | Severity | Check | Finding |", "|---|---|---|---|"]
    lines += [f"| {f.slide or 'deck'} | {f.severity} | {f.code} | {f.message} |" for f in findings]
    lines += ["", "Structural checks only: open the deck in PowerPoint and look at every slide before presenting."]
    report = project / "pitch" / "deck-audit.md"
    report.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return report


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Audit the pitch deck (and optionally apply mechanical fixes).")
    parser.add_argument("--project", default=".", help="Team project folder")
    parser.add_argument("--deck", help="Deck to audit (default: newest pitch/<TICKER>_deck_v<N>.pptx)")
    parser.add_argument("--fix", action="store_true", help="Save a new version with mechanical fixes")
    args = parser.parse_args(argv)
    project = Path(args.project).resolve()
    deck = Path(args.deck).resolve() if args.deck else _latest_deck(project / "pitch")
    if deck is None or not deck.is_file():
        print("[x] no deck found: build it with build_pptx.py or pass --deck <file>")
        return EXIT_NO_DECK
    try:
        if args.fix:
            fixed = _next_version(deck)
            changes = fix(deck, fixed)
            if changes:
                print(f"[ok] {changes} mechanical fixes -> {_shown(fixed, project)}")
                deck = fixed
            else:
                print("[ok] no mechanical fixes needed; no new version saved")
        findings = audit(deck, project)
        report = write_report(project, deck, findings)
    except (PackageNotFoundError, ValueError, KeyError):
        print(f"[x] {ascii_safe(deck.name)} is not a valid .pptx file")
        return EXIT_NO_DECK
    except OSError as exc:
        print(f"[x] Could not write: {ascii_safe(exc.strerror or str(exc))}. Close the file and run again.")
        return EXIT_IO_ERROR
    errors = sum(f.severity == "ERROR" for f in findings)
    print(f"audited {ascii_safe(deck.name)}: {errors} errors, {len(findings) - errors} warnings")
    for f in findings:
        where = "deck" if f.slide == 0 else f"slide {f.slide}"
        print(f"[{'x' if f.severity == 'ERROR' else 'warn'}] {where}: {ascii_safe(f.message)}")
    print(f"[ok] report -> pitch/{report.name}")
    return EXIT_OK


if __name__ == "__main__":
    raise SystemExit(main())
