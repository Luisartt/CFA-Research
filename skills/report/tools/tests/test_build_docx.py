from __future__ import annotations

import copy
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Mm

from report_fixtures import HEADER, SECTIONS, report_project, write_report_project  # noqa: F401

from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from build_docx import (HEADER_PAGES, BUDGET_PAGES, ReportInputError, budget_rows, load_sections, looks_numeric,
                        main, parse_markdown, uncited)
from charts import main as draw_charts


def _build(project: Path) -> Path:
    draw_charts(["--project", str(project)])
    assert main(["--project", str(project)]) == 0
    return project / "report" / "ACME_report_v1.docx"


def _project_with(tmp_path: Path, **overrides: str) -> Path:
    sections = dict(SECTIONS)
    sections.update(overrides)
    return write_report_project(tmp_path, sections=sections)


def _run_main(project: Path, capsys: pytest.CaptureFixture[str], *extra: str) -> tuple[int, str]:
    code = main(["--project", str(project), *extra])
    out = capsys.readouterr().out
    assert out.isascii()
    return code, out


def test_builds_a4_body_with_header_sections_and_figures(report_project: Path) -> None:
    doc = Document(str(_build(report_project)))
    section = doc.sections[0]
    assert abs(section.page_width - Mm(210)) < Mm(1) and abs(section.page_height - Mm(297)) < Mm(1)  # type: ignore[operator]
    text = "\n".join(p.text for p in doc.paragraphs)
    for heading in ("Investment summary", "Business description", "Valuation", "Appendix"):
        assert heading in text
    assert "Figure 1: Revenue and EBITDA margin" in text
    assert len(doc.inline_shapes) == 1
    usable = section.page_width - section.left_margin - section.right_margin  # type: ignore[operator]
    assert all(shape.width <= usable for shape in doc.inline_shapes)


def test_header_block_has_the_required_fields(report_project: Path) -> None:
    doc = Document(str(_build(report_project)))
    header_text = " ".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
    for expected in ("Acme Alimentos", "BMV: ACME", "Consumer Staples / Packaged Foods", "BUY",
                     "MXN 25.00 (2026-10-14)", "MXN 31.00 (+24.0%)"):
        assert expected in header_text, expected


def test_markdown_tables_and_inline_formatting(report_project: Path) -> None:
    doc = Document(str(_build(report_project)))
    assert len(doc.tables) == 2  # header block + the valuation table
    assert doc.tables[1].rows[0].cells[0].text == "Method"
    bold_runs = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "BUY" in bold_runs


def test_appendix_starts_on_a_new_page(report_project: Path) -> None:
    doc = Document(str(_build(report_project)))
    appendix = next(p for p in doc.paragraphs if p.text == "Appendix")
    assert appendix.paragraph_format.page_break_before is True
    assert 'w:type="page"' not in doc.element.body.xml  # no empty page-break paragraph


def test_appendix_starting_with_a_table_still_breaks_the_page(tmp_path: Path) -> None:
    project = _project_with(tmp_path, **{"99-appendix.md": "| A | B |\n|---|---|\n| x | y |\n"})
    doc = Document(str(_build(project)))
    previous = doc.tables[-1]._tbl.getprevious()
    assert previous.tag == qn("w:p")
    assert previous.find(qn("w:pPr")).find(qn("w:pageBreakBefore")) is not None


def test_versions_are_never_overwritten(report_project: Path) -> None:
    _build(report_project)
    assert main(["--project", str(report_project)]) == 0
    assert (report_project / "report" / "ACME_report_v1.docx").is_file()
    assert (report_project / "report" / "ACME_report_v2.docx").is_file()


def test_check_mode_reports_budget_and_uncited(report_project: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["--project", str(report_project), "--check"]) == 0
    out = capsys.readouterr().out
    assert out.isascii()
    assert "investment-summary" in out and "pages" in out
    assert "uncited" in out and "Our DCF gives 31 per share" in out
    assert not list((report_project / "report").glob("*.docx"))


def test_check_mode_flags_over_budget(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    sections = copy.deepcopy(SECTIONS)
    sections["02-business-description.md"] = "# Business description\n\n" + ("word " * 600) + "\n"
    project = write_report_project(tmp_path, sections=sections)
    main(["--project", str(project), "--check"])
    assert "[over] business-description" in capsys.readouterr().out


def test_missing_header_exits_2(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    project = write_report_project(tmp_path, header=None)
    assert main(["--project", str(project)]) == 2
    assert "header.yaml" in capsys.readouterr().out


def test_bad_recommendation_exits_2(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    header = dict(HEADER)
    header["recommendation"] = "STRONG BUY"
    project = write_report_project(tmp_path, header=header)
    assert main(["--project", str(project)]) == 2
    assert "recommendation" in capsys.readouterr().out


def test_parser_blocks() -> None:
    blocks = parse_markdown("# Title\n\nPara one\nstill one.\n\n- a\n- b\n\n1. x\n\n![Cap](c.png)\n\n| A | B |\n|---|---|\n| 1 | 2 |\n")
    assert [b.kind for b in blocks] == ["heading", "paragraph", "bullet", "bullet", "number", "image", "table"]
    assert blocks[1].text == "Para one still one."
    assert blocks[6].rows == [["A", "B"], ["1", "2"]]


def test_budget_covers_the_seven_graded_sections() -> None:
    assert set(BUDGET_PAGES) == {"investment-summary", "business-description", "industry", "financial-analysis",
                                 "valuation", "risks", "esg"}
    assert sum(BUDGET_PAGES.values()) == 10


# --- headings and inline formatting ------------------------------------------------------------------

def test_headings_inherit_bold_and_use_report_sizes(tmp_path: Path) -> None:
    project = _project_with(tmp_path, **{"03-industry.md": "# Industry\n\n## Market\n\n### Share\n\nText.\n"})
    doc = Document(str(_build(project)))
    sizes = {"Industry": Pt(11), "Market": Pt(10.5), "Share": Pt(10)}
    seen = set()
    for paragraph in doc.paragraphs:
        if paragraph.text in sizes:
            seen.add(paragraph.text)
            assert paragraph.style is not None and paragraph.style.name.startswith("Heading")
            assert paragraph.style.font.bold is True
            for run in paragraph.runs:
                assert run.bold is None and run.italic is None  # inherit the bold heading style
                assert run.font.size == sizes[paragraph.text]
    assert seen == set(sizes)


def test_plain_runs_do_not_switch_bold_off(report_project: Path) -> None:
    doc = Document(str(_build(report_project)))
    assert all(run.bold in (None, True) and run.italic in (None, True) for p in doc.paragraphs for run in p.runs)


def test_bold_markers_in_a_table_header_are_not_doubled(tmp_path: Path) -> None:
    project = _project_with(tmp_path, **{"05-valuation.md": "| **Method** | Value |\n|---|---|\n| DCF | 31 |\n"})
    doc = Document(str(_build(project)))
    cell = doc.tables[1].rows[0].cells[0]
    assert cell.text == "Method" and all(run.bold for run in cell.paragraphs[0].runs)


# --- tables ------------------------------------------------------------------------------------------

def test_table_cells_have_tight_spacing(report_project: Path) -> None:
    doc = Document(str(_build(report_project)))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    assert paragraph.paragraph_format.space_after == Pt(0)
                    assert paragraph.paragraph_format.line_spacing == 1.0


def test_header_block_column_widths(report_project: Path) -> None:
    doc = Document(str(_build(report_project)))
    for row in doc.tables[0].rows:
        assert abs(row.cells[0].width - Mm(45)) < Mm(0.5)
        assert abs(row.cells[1].width - Mm(125)) < Mm(0.5)


def test_data_table_right_aligns_numbers_and_uses_horizontal_rules(report_project: Path) -> None:
    doc = Document(str(_build(report_project)))
    table = doc.tables[1]
    assert table.style is None or table.style.name != "Table Grid"

    def align(i: int, j: int) -> Any:
        return table.rows[i].cells[j].paragraphs[0].alignment

    assert align(1, 1) == WD_ALIGN_PARAGRAPH.RIGHT  # 31.0
    assert align(0, 1) == WD_ALIGN_PARAGRAPH.RIGHT  # header of a numeric column
    assert align(1, 0) != WD_ALIGN_PARAGRAPH.RIGHT  # DCF

    def edges(i: int) -> set[str]:
        tc_pr = table.rows[i].cells[0]._tc.tcPr
        borders = tc_pr.find(qn("w:tcBorders")) if tc_pr is not None else None
        return set() if borders is None else {child.tag.split("}")[1] for child in borders}

    assert edges(0) == {"top", "bottom"}
    assert edges(1) == set()
    assert edges(len(table.rows) - 1) == {"bottom"}


@pytest.mark.parametrize("text", ["31.0", "(1,234.5)", "+24.0%", "-3.2", "8.5x", "150bp", "$12", "1,200", "n.m.",
                                  "-", "**12.0**"])
def test_numeric_cells(text: str) -> None:
    assert looks_numeric(text)


@pytest.mark.parametrize("text", ["DCF", "Box", "Comps P/E", "2025 plan", "", "MXN bn"])
def test_text_cells(text: str) -> None:
    assert not looks_numeric(text)


# --- numbered lists ----------------------------------------------------------------------------------

def test_separate_numbered_lists_each_start_at_one(tmp_path: Path) -> None:
    project = _project_with(tmp_path, **{"03-industry.md": "# Industry\n\n1. a\n2. b\n\nBreak.\n\n1. c\n2. d\n"})
    doc = Document(str(_build(project)))
    items = {p.text: p for p in doc.paragraphs if p.text in {"a", "b", "c", "d"}}

    def num_id(text: str) -> int:
        paragraph: Any = items[text]._p
        return int(paragraph.pPr.numPr.numId.val)

    assert num_id("a") == num_id("b") and num_id("c") == num_id("d") and num_id("a") != num_id("c")
    numbering = doc.part.numbering_part.element
    style_num = doc.styles["List Number"].element.pPr.numPr.numId.val
    abstract = numbering.num_having_numId(style_num).abstractNumId.val
    for text in ("a", "c"):
        num = numbering.num_having_numId(num_id(text))
        assert num.abstractNumId.val == abstract
        overrides = num.lvlOverride_lst
        assert len(overrides) == 1 and overrides[0].ilvl == 0 and overrides[0].startOverride.val == 1


# --- Markdown robustness -----------------------------------------------------------------------------

@pytest.mark.parametrize("source", [
    "| A | B |\n|---|---|\n| 1 | 2 |",
    "A | B\n--- | ---\n1 | 2",
    "| A | B\n|---|---\n| 1 | 2",
    "| A | B |\n|--|-|\n| 1 | 2 |",
    "| A | B |\n|:--|--:|\n| 1 | 2 |",
])
def test_table_variants(source: str) -> None:
    assert [(b.kind, b.rows) for b in parse_markdown(source)] == [("table", [["A", "B"], ["1", "2"]])]


def test_table_escaped_pipe_is_a_literal_pipe() -> None:
    blocks = parse_markdown("| A | B |\n|---|---|\n| x \\| y | 2 |")
    assert blocks[0].rows == [["A", "B"], ["x | y", "2"]]


def test_pipe_in_a_paragraph_is_not_a_table() -> None:
    blocks = parse_markdown("Revenue | cost is a ratio.\nSecond line.")
    assert [b.kind for b in blocks] == ["paragraph"]


def test_paragraph_directly_before_and_after_a_table_is_kept() -> None:
    blocks = parse_markdown("Intro line.\n| A | B |\n|---|---|\n| 1 | 2 |\nAfter.")
    assert [b.kind for b in blocks] == ["paragraph", "table", "paragraph"]
    assert blocks[0].text == "Intro line." and blocks[2].text == "After."


def test_list_continuation_lines_join_the_item() -> None:
    blocks = parse_markdown("- item that wraps\n  onto the next line\n- second\n\n1. one\n   more\n")
    assert [(b.kind, b.text) for b in blocks] == [
        ("bullet", "item that wraps onto the next line"), ("bullet", "second"), ("number", "one more")]


def test_html_comments_are_dropped() -> None:
    source = "<!-- TODO: add data -->\nKept <!-- inline --> text.\n\n<!--\nmulti\nline\n-->\n# Title\n"
    blocks = parse_markdown(source)
    assert [(b.kind, b.text) for b in blocks] == [("paragraph", "Kept text."), ("heading", "Title")]


def test_comments_keep_line_numbers_for_errors() -> None:
    with pytest.raises(ReportInputError, match="line 5"):
        parse_markdown("<!--\na\nb\n-->\n#### Deep\n", name="x.md")


def test_unclosed_comment_is_an_error() -> None:
    with pytest.raises(ReportInputError, match="x.md: line 2: .*-->"):
        parse_markdown("Text.\n<!-- never closed\nmore\n", name="x.md")


def test_links_render_as_their_text() -> None:
    blocks = parse_markdown("See [INEGI](https://inegi.org.mx) data.\n\n- [a](b) item\n")
    assert [b.text for b in blocks] == ["See INEGI data.", "a item"]


@pytest.mark.parametrize("line, what", [
    ("#### Deep heading", "#### heading"),
    ("---", "horizontal rule"),
    ("***", "horizontal rule"),
    ("> quoted text", "blockquote"),
])
def test_unsupported_markdown_names_file_and_line(line: str, what: str) -> None:
    with pytest.raises(ReportInputError) as info:
        parse_markdown(f"# Ok\n\n{line}\n", name="03-industry.md")
    [message] = info.value.messages
    assert message.startswith("03-industry.md: line 3: ") and what in message
    assert "is not supported; use" in message


def test_unsupported_markdown_errors_are_collected(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    project = _project_with(tmp_path, **{"03-industry.md": "# Industry\n\n---\n\n> quote\n"})
    code, out = _run_main(project, capsys)
    assert code == 2
    assert "03-industry.md: line 3" in out and "03-industry.md: line 5" in out


# --- input errors never crash ------------------------------------------------------------------------

def test_invalid_yaml_header_exits_2(report_project: Path, capsys: pytest.CaptureFixture[str]) -> None:
    (report_project / "report" / "header.yaml").write_text("company: Acme\nticker: [ACME\n", encoding="utf-8")
    code, out = _run_main(report_project, capsys)
    assert code == 2 and "header.yaml" in out and "YAML" in out


def test_header_not_utf8_exits_2(report_project: Path, capsys: pytest.CaptureFixture[str]) -> None:
    path = report_project / "report" / "header.yaml"
    path.write_bytes(path.read_text(encoding="utf-8").replace("Acme Alimentos", "Acmé").encode("cp1252"))
    code, out = _run_main(report_project, capsys)
    assert code == 2 and "header.yaml" in out and "save it as UTF-8" in out


def test_section_not_utf8_exits_2(report_project: Path, capsys: pytest.CaptureFixture[str]) -> None:
    path = report_project / "report" / "sections" / "02-business-description.md"
    path.write_bytes("# Descripción\n\nMéxico".encode("cp1252"))
    code, out = _run_main(report_project, capsys, "--check")
    assert code == 2 and "02-business-description.md" in out and "save it as UTF-8" in out


@pytest.mark.parametrize("content", [b"<svg xmlns='http://www.w3.org/2000/svg'/>", b"not an image at all"])
def test_unreadable_image_exits_2_naming_the_file(report_project: Path, capsys: pytest.CaptureFixture[str],
                                                  content: bytes) -> None:
    draw_charts(["--project", str(report_project)])
    capsys.readouterr()
    (report_project / "report" / "charts" / "revenue-margin.png").write_bytes(content)
    code, out = _run_main(report_project, capsys)
    assert code == 2 and "01-investment-summary.md" in out and "revenue-margin.png" in out


def test_numeric_ticker_must_be_quoted(report_project: Path, capsys: pytest.CaptureFixture[str]) -> None:
    path = report_project / "report" / "header.yaml"
    path.write_text(path.read_text(encoding="utf-8").replace("ticker: ACME", "ticker: 0700"), encoding="utf-8")
    code, out = _run_main(report_project, capsys)
    assert code == 2 and "ticker" in out and "put the ticker in quotes" in out


# --- uncited check -----------------------------------------------------------------------------------

def _uncited(text: str) -> list[tuple[str, str]]:
    return uncited([(Path("03-industry.md"), parse_markdown(text))])


@pytest.mark.parametrize("text", [
    "Figure 1 shows the trend.",
    "We see 3Q25 as the trough and 1H26 as the recovery.",
    "Pillar 2 is margins; Table 3 has the detail.",
    "FY2025 closed well and 2026E looks better.",
    "Sales rose 5% (Bimbo, 2025a).",
    "Growth of 4% (INEGI, n.d.).",
    "Revenue was MXN 400bn (Grupo Bimbo 2025, p. 4).",
    "Margins reach 21% (team estimate).",
    "WACC is 12.5% (team model).",
    "Source: INEGI household survey 2024.",
])
def test_uncited_ignores_cited_or_label_numbers(text: str) -> None:
    assert _uncited(text) == []


@pytest.mark.parametrize("text", [
    "Margin rose to 19.8% (vs. 19.0% in 2021).",
    "Revenue grew 7% in 2025.",
    "Peers trade at 8x [1].",
])
def test_uncited_flags_unsourced_numbers(text: str) -> None:
    assert len(_uncited(text)) == 1


def test_uncited_flags_a_table_without_a_source_line() -> None:
    table = "| Method | Value |\n|---|---|\n| DCF | 31.0 |\n"
    flagged = _uncited(table + "\nOur view (team model).\n")
    assert len(flagged) == 1 and "table" in flagged[0][1].lower() and "Method" in flagged[0][1]
    assert _uncited(table + "\nSource: team model.\n") == []
    assert len(_uncited(table)) == 1  # table at the end of the section


# --- figures -----------------------------------------------------------------------------------------

def test_two_consecutive_images_sit_side_by_side(tmp_path: Path) -> None:
    project = _project_with(tmp_path, **{"03-industry.md": (
        "# Industry\n\n![Free cash flow](../charts/free-cash-flow.png)\n![Risk matrix](../charts/risk-matrix.png)\n")})
    doc = Document(str(_build(project)))
    section = doc.sections[0]
    usable = section.page_width - section.left_margin - section.right_margin  # type: ignore[operator]
    pair = next(t for t in doc.tables if "Figure" in t.rows[0].cells[0].text)
    assert len(pair.columns) == 2
    assert "w:tcBorders" not in pair._tbl.xml and "w:tblBorders" not in pair._tbl.xml
    assert pair.style is None or pair.style.name != "Table Grid"
    texts = [cell.text for cell in pair.rows[0].cells]
    assert "Figure 2: Free cash flow" in texts[0] and "Figure 3: Risk matrix" in texts[1]
    for cell in pair.rows[0].cells:
        assert len(cell._tc.xpath(".//pic:pic")) == 1
        assert cell.paragraphs[0].paragraph_format.keep_with_next is True
    shapes = list(doc.inline_shapes)
    assert len(shapes) == 3
    assert all(shape.width <= usable / 2 for shape in shapes[1:])


def test_a_full_width_image_is_never_squeezed_into_a_pair(tmp_path: Path) -> None:
    project = _project_with(tmp_path, **{"03-industry.md": (
        "# Industry\n\n![Football field](../charts/football-field.png)\n![Sensitivity](../charts/sensitivity.png)\n")})
    doc = Document(str(_build(project)))
    assert not any("Figure" in t.rows[0].cells[0].text for t in doc.tables)
    pictures = [p for p in doc.paragraphs if p._p.xpath(".//pic:pic")]
    assert len(pictures) == 3  # revenue chart in the fixture + the two above, each on its own


def test_single_image_is_centered_and_kept_with_its_caption(report_project: Path) -> None:
    doc = Document(str(_build(report_project)))
    picture = next(p for p in doc.paragraphs if p._p.xpath(".//pic:pic"))
    assert picture.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert picture.paragraph_format.keep_with_next is True


# --- --check output ----------------------------------------------------------------------------------

def test_check_columns_fit_the_longest_section_name(report_project: Path, capsys: pytest.CaptureFixture[str]) -> None:
    code, out = _run_main(report_project, capsys, "--check")
    assert code == 0
    lines = out.splitlines()
    header = next(line for line in lines if "words" in line)
    words_end = header.index("words") + len("words")
    names = ("investment-summary", "business-description", "valuation", "header block")
    rows = [line for line in lines if line[7:].startswith(names)]
    assert len(rows) == 4
    for row in rows:
        assert row[words_end] == " " and row[words_end - 1] in "0123456789-", row


def test_check_counts_the_header_block_and_says_conservative(report_project: Path,
                                                            capsys: pytest.CaptureFixture[str]) -> None:
    _, out = _run_main(report_project, capsys, "--check")
    assert HEADER_PAGES == 0.25 and "header block" in out
    total = sum(r[3] for r in budget_rows(load_sections(report_project / "report" / "sections"))) + HEADER_PAGES
    total_line = next(line for line in out.splitlines() if line.startswith("total estimated pages"))
    assert f"{total:.1f} of 10" in total_line and "(conservative estimate" in total_line
