"""Evaluate the model in Python and resolve cell addresses for Excel formulas."""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from typing import TypeAlias, Union

from openpyxl.utils import get_column_letter

from .expr import Expr
from .inputs import ModelInputs
from .placement import PERIOD_COL0, Placement, place
from .spec import Inputs, LayoutItem, Line, Observed


class ModelError(Exception):
    """A formula references something that does not exist or loops back on itself."""


@dataclass(frozen=True)
class ObservedCell:
    value: float
    comment: str


@dataclass(frozen=True)
class InputCell:
    value: float


@dataclass(frozen=True)
class FormulaCell:
    expr: Expr


@dataclass(frozen=True)
class BlankCell:
    pass


BLANK = BlankCell()
Cell: TypeAlias = Union[ObservedCell, InputCell, FormulaCell, BlankCell]


class Model:
    """The evaluated model; implements the expr.Context protocol."""

    def __init__(self, layout: Sequence[LayoutItem], inputs: ModelInputs) -> None:
        self.inputs = inputs
        self.placed = place(layout)
        self.lines: dict[str, Line] = {line.key: line for line, _ in self.placed.lines}
        self._cache: dict[tuple[str, int | None], float] = {}
        self._visiting: set[tuple[str, int | None]] = set()

    @property
    def h(self) -> int:
        return self.inputs.h

    @property
    def n(self) -> int:
        return self.inputs.n

    def has(self, key: str) -> bool:
        return key in self.lines

    def line(self, key: str) -> Line:
        try:
            return self.lines[key]
        except KeyError:
            raise ModelError(f"unknown line '{key}'") from None

    def sheet_of(self, key: str) -> str:
        return self._placement(key).sheet

    def row_of(self, key: str) -> int:
        return self._placement(key).row

    def cell(self, key: str, period: int | None) -> Cell:
        line = self.line(key)
        if line.scalar is not None:
            scalar = line.scalar
            return FormulaCell(scalar) if isinstance(scalar, Expr) else InputCell(float(scalar))
        if period is None:
            raise ModelError(f"'{key}' is a per-year line; a single-value formula must use At()")
        spec = line.hist if period < self.h else line.fcst
        if spec is None:
            return BLANK
        if isinstance(spec, Observed):
            return self._observed(key, period)
        if isinstance(spec, Inputs):
            return InputCell(spec.values[period - self.h])
        if period - spec.max_lag() < 0:
            return BLANK
        return FormulaCell(spec)

    def value(self, key: str, period: int | None) -> float:
        line = self.line(key)
        slot: int | None = None
        if line.scalar is None:
            if period is None:
                raise ModelError(f"'{key}' is a per-year line; a single-value formula must use At()")
            if not 0 <= period < self.n:
                raise ModelError(f"'{key}' referenced outside the model years (period {period})")
            slot = period
        cache_key = (key, slot)
        if cache_key in self._cache:
            return self._cache[cache_key]
        if cache_key in self._visiting:
            raise ModelError(f"circular reference through '{key}'")
        self._visiting.add(cache_key)
        try:
            result = self._compute(key, slot)
        finally:
            self._visiting.discard(cache_key)
        self._cache[cache_key] = result
        return result

    def address(self, key: str, period: int | None, from_sheet: str) -> str:
        placement = self._placement(key)
        if placement.col is not None:
            col = placement.col
        elif period is None:
            raise ModelError(f"'{key}' is a per-year line; a single-value formula must use At()")
        else:
            col = PERIOD_COL0 + period
        ref = f"{get_column_letter(col)}{placement.row}"
        return ref if placement.sheet == from_sheet else f"{placement.sheet}!{ref}"

    def range_address(self, key: str, first: int, last: int, from_sheet: str) -> str:
        placement = self._placement(key)
        if placement.col is not None:
            raise ModelError(f"'{key}' is a single value, not a range")
        start = f"{get_column_letter(PERIOD_COL0 + first)}{placement.row}"
        end = f"{get_column_letter(PERIOD_COL0 + last)}{placement.row}"
        span = f"{start}:{end}"
        return span if placement.sheet == from_sheet else f"{placement.sheet}!{span}"

    def evaluate_all(self) -> None:
        """Evaluate every cell once, so broken references and cycles fail before writing."""
        for key, line in self.lines.items():
            periods: Sequence[int | None] = [None] if line.is_scalar else range(self.n)
            for period in periods:
                self.value(key, period)

    def _placement(self, key: str) -> Placement:
        try:
            return self.placed.placements[key]
        except KeyError:
            raise ModelError(f"unknown line '{key}'") from None

    def _observed(self, key: str, period: int) -> ObservedCell:
        year = self.inputs.hist_years[period]
        point = self.inputs.history.get(key, {}).get(year)
        if point is None:
            return ObservedCell(0.0, f"{key} {year}: not reported, set to 0")
        page = f" p.{point.page}" if point.page else ""
        return ObservedCell(point.value, f"[{point.tag}] {point.source_doc}{page}")

    def _compute(self, key: str, slot: int | None) -> float:
        cell = self.cell(key, slot)
        if isinstance(cell, (ObservedCell, InputCell)):
            return cell.value
        if isinstance(cell, FormulaCell):
            return cell.expr.evaluate(self, slot)
        return 0.0  # blank cells read as zero, exactly as in Excel
