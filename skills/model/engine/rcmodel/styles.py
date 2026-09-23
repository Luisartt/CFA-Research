"""Cell styles for the model workbook.

Palette and number-format conventions adapted from research_analyst
tools/xlsx_builder.py (MIT License, Copyright (c) 2026 CFA Society Mexico -
AI for Finance): blue inputs, black formulas, green links.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT_NAME = "Aptos Narrow"
BLUE = "FF0000FF"
BLACK = "FF000000"
GREEN = "FF008000"
WHITE = "FFFFFFFF"
GREY = "FF7F7F7F"


def font(color: str, bold: bool = False) -> Font:
    return Font(name=FONT_NAME, color=color, bold=bold)


BAND_FILL = PatternFill("solid", fgColor="FF1F4E79")
SECTION_FILL = PatternFill("solid", fgColor="FFBDD7EE")
INPUT_FILL = PatternFill("solid", fgColor="FFFFF2CC")
ERROR_FILL = PatternFill("solid", fgColor="FFC00000")
WARN_FILL = PatternFill("solid", fgColor="FFFFC000")
BAND_FONT = Font(name=FONT_NAME, color=WHITE, bold=True, size=12)
SECTION_FONT = Font(name=FONT_NAME, color=BLACK, bold=True)
BOLD_FONT = Font(name=FONT_NAME, bold=True)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
NOTE_FONT = Font(name=FONT_NAME, color=GREY, italic=True)
WHITE_BOLD = Font(name=FONT_NAME, color=WHITE, bold=True)
TOP_BORDER = Border(top=Side(style="thin"))
RIGHT = Alignment(horizontal="right")
