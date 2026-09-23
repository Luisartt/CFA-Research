"""Write the evaluated model to an .xlsx with live formulas (openpyxl)."""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

from . import styles
from .checks import Check, CheckResult, Severity, overall_status
from .engine import BlankCell, Cell, InputCell, Model, ObservedCell
from .expr import At, Expr, Ref
from .placement import FIRST_ROW, HEADER_ROW, PERIOD_COL0
from .spec import Header, Line, Style
from .valuation import FootballSpec

SHEET_TITLES: dict[str, str] = {
    "Cover": "Model cover",
    "Drivers": "Drivers and assumptions",
    "IS": "Income statement",
    "BS": "Balance sheet",
    "CF": "Cash flow statement",
    "Schedules": "Supporting schedules",
    "Ratios": "Ratios",
    "WACC": "Cost of capital",
    "DCF": "Discounted cash flow",
    "Comps": "Comparable companies",
    "Sensitivity": "Sensitivity (value per share)",
    "Football": "Football field",
    "Checks": "Integrity checks",
}
SHEET_ORDER: tuple[str, ...] = tuple(SHEET_TITLES)
PERIOD_SHEETS = frozenset({"Drivers", "IS", "BS", "CF", "Schedules", "Ratios", "DCF", "Checks"})
AUTHOR = "research-challenge"
COVER_NOTES = (
    "Blue = input or reported figure (hover a reported figure for its source); black = formula; "
    "green = link to another sheet.",
    "Interest is charged on opening balances, so the model has no circular references.",
    "A revolving credit line keeps cash at or above the minimum cash balance.",
    "Valuation is as of the last fiscal year-end (no stub-period adjustment).",
)


@dataclass(frozen=True)
class BuildInfo:
    version: int
    filename: str
    built_on: date


def write_workbook(path: Path, model: Model, checks: Sequence[Check], results: Sequence[CheckResult],
                   football: FootballSpec | None, info: BuildInfo) -> None:
    workbook = Workbook()
    cover = workbook.active
    cover.title = "Cover"
    used = {placement.sheet for placement in model.placed.placements.values()} | {"Checks"}
    sheets: dict[str, Any] = {}
    for name in SHEET_ORDER[1:]:
        if name in used:
            sheets[name] = workbook.create_sheet(name)
            _frame(sheets[name], model, name)
    for sheet, row, header in model.placed.headers:
        _header(sheets[sheet], row, header, model)
    for line, row in model.placed.lines:
        _write_line(sheets[line.sheet], line, row, model)
    last_check_row = _write_checks(sheets["Checks"], model, checks)
    if football is not None:
        _football_chart(sheets["Football"], model, football)
    _cover(cover, model, info, results, last_check_row)
    workbook.calculation.fullCalcOnLoad = True
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _last_col(model: Model) -> int:
    return PERIOD_COL0 + model.n + 2


def _frame(ws: Any, model: Model, name: str) -> None:
    profile = model.inputs.profile
    last_col = _last_col(model)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 52
    for col in range(PERIOD_COL0, last_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 13
    if name == "Drivers":
        ws.column_dimensions[get_column_letter(PERIOD_COL0 + model.n + 1)].width = 48
    for col in range(2, last_col + 1):
        ws.cell(1, col).fill = styles.BAND_FILL
    ws.cell(1, 2, f"{profile.name} ({profile.ticker}) - {SHEET_TITLES[name]}").font = styles.BAND_FONT
    ws.cell(2, 2, f"{profile.currency} {profile.units} | A = actual, E = estimate | "
                  "blue = input, black = formula, green = link to another sheet").font = styles.NOTE_FONT
    if name in PERIOD_SHEETS:
        for period, year in enumerate(model.inputs.years):
            cell = ws.cell(HEADER_ROW, PERIOD_COL0 + period, year)
            cell.number_format = '0"A"' if period < model.h else '0"E"'
            cell.font = styles.BOLD_FONT
            cell.alignment = styles.RIGHT
        ws.freeze_panes = ws.cell(HEADER_ROW + 1, PERIOD_COL0)


def _header(ws: Any, row: int, header: Header, model: Model) -> None:
    for col in range(2, _last_col(model) + 1):
        ws.cell(row, col).fill = styles.SECTION_FILL
    ws.cell(row, 2, header.title).font = styles.SECTION_FONT
    for col, text in header.columns:
        cell = ws.cell(row, col, text)
        cell.font = styles.SECTION_FONT
        cell.alignment = styles.RIGHT


def _write_line(ws: Any, line: Line, row: int, model: Model) -> None:
    if line.new_row and line.label:
        ws.cell(row, 2, line.label).font = styles.font(styles.BLACK, bold=line.style is not Style.NORMAL)
    periods: Sequence[int | None] = [None] if line.is_scalar else range(model.n)
    for period in periods:
        col = line.col if period is None else PERIOD_COL0 + period
        _write_cell(ws, row, col, line, model.cell(line.key, period), model, period)
    if line.notes:
        start = line.col + 1 if line.is_scalar else PERIOD_COL0 + model.n
        for offset, text in enumerate(line.notes):
            if text:
                ws.cell(row, start + offset, text).font = styles.NOTE_FONT


def _write_cell(ws: Any, row: int, col: int, line: Line, cell: Cell, model: Model, period: int | None) -> None:
    if isinstance(cell, BlankCell):
        return
    target = ws.cell(row, col)
    color = styles.BLACK
    if isinstance(cell, ObservedCell):
        target.value = cell.value
        target.comment = Comment(cell.comment, AUTHOR)
        color = styles.BLUE
    elif isinstance(cell, InputCell):
        target.value = cell.value
        target.fill = styles.INPUT_FILL
        color = styles.BLUE
    else:
        target.value = "=" + cell.expr.render(model, period, ws.title)
        if _is_link(cell.expr, ws.title, model):
            color = styles.GREEN
    target.font = styles.font(color, bold=line.style is not Style.NORMAL)
    target.number_format = line.fmt.value
    if line.style is Style.TOTAL:
        target.border = styles.TOP_BORDER


def _is_link(expr: Expr, sheet: str, model: Model) -> bool:
    return isinstance(expr, (Ref, At)) and model.sheet_of(expr.key) != sheet


def _write_checks(ws: Any, model: Model, checks: Sequence[Check]) -> int:
    row = FIRST_ROW
    groups = (
        ("Checks by year", [c for c in checks if c.periods is not None]),
        ("Single-value checks (result in column C)", [c for c in checks if c.periods is None]),
    )
    for title, group in groups:
        if not group:
            continue
        if row > FIRST_ROW:
            row += 1
        _header(ws, row, Header("Checks", title), model)
        row += 1
        for check in group:
            suffix = " (warning only)" if check.severity is Severity.WARN else ""
            ws.cell(row, 2, check.label + suffix)
            for period in check.periods if check.periods is not None else (None,):
                col = PERIOD_COL0 if period is None else PERIOD_COL0 + period
                condition = check.cond.render(model, period, ws.title)
                ws.cell(row, col, f'=IF({condition},"OK","{check.severity.value}")').alignment = styles.RIGHT
            row += 1
    last = max(row - 1, FIRST_ROW)
    span = f"C{FIRST_ROW}:{get_column_letter(PERIOD_COL0 + model.n - 1)}{last}"
    ws.conditional_formatting.add(span, CellIsRule(operator="equal", formula=['"ERROR"'],
                                                   fill=styles.ERROR_FILL, font=styles.WHITE_BOLD))
    ws.conditional_formatting.add(span, CellIsRule(operator="equal", formula=['"WARN"'], fill=styles.WARN_FILL))
    return last


def _football_chart(ws: Any, model: Model, spec: FootballSpec) -> None:
    first, last = model.row_of(spec.first_row_key), model.row_of(spec.last_row_key)
    chart = BarChart()
    chart.type = "bar"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "Value per share by method"
    chart.add_data(Reference(ws, min_col=3, min_row=first, max_row=last), titles_from_data=False)
    chart.add_data(Reference(ws, min_col=5, min_row=first, max_row=last), titles_from_data=False)
    chart.set_categories(Reference(ws, min_col=2, min_row=first, max_row=last))
    chart.series[0].graphicalProperties.noFill = True
    chart.series[0].graphicalProperties.line.noFill = True
    chart.legend = None
    chart.height = 9
    chart.width = 18
    ws.add_chart(chart, f"G{FIRST_ROW}")


def _cover(ws: Any, model: Model, info: BuildInfo, results: Sequence[CheckResult], last_check_row: int) -> None:
    profile, years = model.inputs.profile, model.inputs.years
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 90
    ws.cell(2, 2, f"{profile.name} ({profile.ticker})").font = styles.TITLE_FONT
    facts = (
        ("Accounting framework", profile.framework),
        ("Currency and units", f"{profile.currency} {profile.units}"),
        ("Historical years", f"{years[0]}-{years[model.h - 1]}"),
        ("Forecast years", f"{years[model.h]}-{years[-1]}"),
        ("Model file", info.filename),
        ("Version", str(info.version)),
        ("Built on", info.built_on.isoformat()),
    )
    row = 4
    for label, value in facts:
        ws.cell(row, 2, label).font = styles.BOLD_FONT
        ws.cell(row, 3, value)
        row += 1
    row += 1
    span = f"Checks!$C${FIRST_ROW}:${get_column_letter(PERIOD_COL0 + model.n - 1)}${last_check_row}"
    ws.cell(row, 2, "Model status (live)").font = styles.BOLD_FONT
    status = ws.cell(row, 3, f'=IF(COUNTIF({span},"ERROR")>0,"CHECKS FAILING",'
                             f'IF(COUNTIF({span},"WARN")>0,"OK WITH WARNINGS","ALL CHECKS OK"))')
    status.font = styles.TITLE_FONT
    ws.conditional_formatting.add(f"C{row}", FormulaRule(formula=[f'$C${row}="CHECKS FAILING"'],
                                                         fill=styles.ERROR_FILL, font=styles.WHITE_BOLD))
    row += 1
    ws.cell(row, 2, "Status computed by Python at build").font = styles.BOLD_FONT
    ws.cell(row, 3, overall_status(results))
    row += 2
    ws.cell(row, 2, "How to read this model").font = styles.BOLD_FONT
    for note in COVER_NOTES:
        ws.cell(row, 3, note)
        row += 1
    if model.inputs.warnings:
        row += 1
        ws.cell(row, 2, "Engine warnings").font = styles.BOLD_FONT
        for warning in model.inputs.warnings:
            ws.cell(row, 3, warning)
            row += 1
