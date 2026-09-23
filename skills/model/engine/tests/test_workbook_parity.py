"""Recalculate the written workbook with an Excel engine and compare every formula with Python."""

from __future__ import annotations

import math
import numbers
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from rcmodel.checks import evaluate_checks, overall_status
from rcmodel.engine import FormulaCell
from rcmodel.placement import PERIOD_COL0
from rcmodel.writer import check_layout
from test_writer import build_all

formulas = pytest.importorskip("formulas")

KEY = re.compile(r"^'\[(?P<book>[^\]]+)\](?P<sheet>[^']+)'!(?P<cell>[A-Z]+[0-9]+)$")


def _scalar(result: object) -> object:
    """`formulas` wraps every cell result in a (1, 1) array; unwrap it."""
    import numpy as np

    return np.asarray(getattr(result, "value", result), dtype=object).ravel()[0]


def _to_float(value: object) -> float:
    """A number as float; text, booleans and Excel errors (#DIV/0!, #VALUE!, ...) as NaN."""
    if isinstance(value, numbers.Real) and not isinstance(value, bool):
        return float(value)
    return math.nan


def test_excel_values_match_python(project: Path) -> None:
    path, model, checks = build_all(project)
    solution = formulas.ExcelModel().loads(str(path)).finish().calculate()
    excel: dict[tuple[str, str], object] = {}
    for name, ranges in solution.items():
        match = KEY.match(str(name).upper())
        if match:
            excel[(match["sheet"], match["cell"])] = _scalar(ranges)

    formula_cells = finite = compared = 0
    for line, row in model.placed.lines:
        periods = [None] if line.is_scalar else list(range(model.n))
        for period in periods:
            if not isinstance(model.cell(line.key, period), FormulaCell):
                continue
            formula_cells += 1
            col = line.col if period is None else PERIOD_COL0 + period
            key = (line.sheet.upper(), f"{get_column_letter(col)}{row}")
            assert key in excel, f"{line.key} {period}: {key} missing from the recalculated workbook"
            python_value = model.value(line.key, period)
            excel_value = _to_float(excel[key])
            if not math.isfinite(python_value):
                assert not math.isfinite(excel_value), f"{line.key} {period}: Python {python_value}, Excel {excel[key]}"
                continue
            finite += 1
            assert excel_value == pytest.approx(python_value, rel=1e-9, abs=1e-6), f"{line.key} {period}"
            compared += 1
    assert formula_cells > 0
    assert compared == finite

    evaluated = evaluate_checks(model, checks)
    results = {(r.key, r.period): r.status for r in evaluated}
    seen = 0
    for row, check in check_layout(checks)[1]:
        for period in check.periods if check.periods is not None else (None,):
            col = PERIOD_COL0 if period is None else PERIOD_COL0 + period
            key = ("CHECKS", f"{get_column_letter(col)}{row}")
            assert key in excel, f"check {check.key} {period}: {key} missing"
            assert excel[key] == results[(check.key, period)], f"check {check.key} {period}"
            seen += 1
    assert seen == len(results)

    cover = load_workbook(path)["Cover"]
    live = next(c for r in cover.iter_rows() for c in r
                if isinstance(c.value, str) and c.value.startswith("=IF(COUNTIF("))
    assert excel[("COVER", live.coordinate)] == overall_status(evaluated)
