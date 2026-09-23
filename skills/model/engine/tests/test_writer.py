from __future__ import annotations

import zipfile
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from conftest import write_project
from rcmodel.assemble import assemble
from rcmodel.checks import build_checks, evaluate_checks, parity_checks
from rcmodel.engine import Model
from rcmodel.inputs import load_inputs
from rcmodel.writer import BuildInfo, write_workbook


def build_file(project: Path) -> tuple[Path, Model]:
    model, football = assemble(load_inputs(project))
    checks = [*build_checks(model.inputs), *parity_checks(model)]
    path = project / "model" / "TEST.xlsx"
    write_workbook(path, model, checks, evaluate_checks(model, checks), football,
                   BuildInfo(1, path.name, date(2026, 9, 23)))
    return path, model


def test_sheets_in_order(project: Path) -> None:
    path, _ = build_file(project)
    assert load_workbook(path).sheetnames == [
        "Cover", "Drivers", "IS", "BS", "CF", "Schedules", "Ratios",
        "WACC", "DCF", "Comps", "Sensitivity", "Football", "Checks",
    ]


def test_history_is_values_forecast_is_formulas(project: Path) -> None:
    path, model = build_file(project)
    ws = load_workbook(path)["IS"]
    row = model.row_of("revenue")
    assert ws.cell(row, 3 + model.h - 1).value == 13400
    forecast = ws.cell(row, 3 + model.h).value
    assert isinstance(forecast, str) and forecast.startswith("=") and "Drivers!" in forecast
    assert "sourced" in ws.cell(row, 3).comment.text


def test_inputs_are_blue_on_yellow(project: Path) -> None:
    path, model = build_file(project)
    cell = load_workbook(path)["Drivers"].cell(model.row_of("gross_margin"), 3 + model.h)
    assert cell.value == 0.41
    assert cell.font.color.rgb == "FF0000FF"
    assert cell.fill.fgColor.rgb == "FFFFF2CC"


def test_links_to_other_sheets_are_green(project: Path) -> None:
    path, model = build_file(project)
    cell = load_workbook(path)["BS"].cell(model.row_of("cash"), 3 + model.h)
    assert str(cell.value).startswith("=CF!")
    assert cell.font.color.rgb == "FF008000"


def test_checks_and_cover_formulas(project: Path) -> None:
    path, _ = build_file(project)
    workbook = load_workbook(path)
    check_formulas = [c.value for row in workbook["Checks"].iter_rows() for c in row
                      if isinstance(c.value, str) and c.value.startswith("=IF(")]
    assert check_formulas
    cover_text = [c.value for row in workbook["Cover"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any("COUNTIF(Checks!" in text for text in cover_text)
    assert "ALL CHECKS OK" in cover_text or "OK WITH WARNINGS" in cover_text


def test_workbook_recalculates_on_open(project: Path) -> None:
    path, _ = build_file(project)
    assert load_workbook(path).calculation.fullCalcOnLoad is True


def test_football_chart_is_embedded(project: Path) -> None:
    path, _ = build_file(project)
    with zipfile.ZipFile(path) as archive:
        assert any(name.startswith("xl/charts/chart") for name in archive.namelist())


def test_no_valuation_sheets_without_valuation(tmp_path: Path) -> None:
    path, _ = build_file(write_project(tmp_path, valuation=None))
    assert load_workbook(path).sheetnames == ["Cover", "Drivers", "IS", "BS", "CF", "Schedules", "Ratios", "Checks"]
