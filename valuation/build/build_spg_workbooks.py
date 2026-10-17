"""Generate the S&P Capital IQ Pro-layout WACC and DCF workbooks for Grupo Bimbo.

One source of truth: the canonical inputs of the project's model engine
    data/financials.csv, model/drivers.yaml, valuation/valuation.yaml
plus presentation extras in valuation/inputs/spg_inputs.yaml, and the engine's own
results in model/model-summary.json (used only to reconcile).

The layouts follow valuation/SPG_DCF_Analysis_v1.xlsm and valuation/SPG_WACCAnalysis_v1.xlsm
(those templates are never opened for writing). Every calculation is a live Excel formula;
only inputs are constants (blue font). Run from the project root:

    python -I valuation/build/build_spg_workbooks.py --project .

then recalculate with valuation/build/recalc_excel.ps1 and scan with scan_errors.py.
"""

from __future__ import annotations

import argparse
import datetime as dt
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import dcf_core as core  # noqa: E402

# ----------------------------------------------------------------------------- styles

NAVY = "1F3864"
F_TITLE = Font(name="Calibri", size=14, bold=True, color=NAVY)
F_SUB = Font(name="Calibri", size=10, italic=True, color="595959")
F_HDR = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_BOLD = Font(name="Calibri", size=10, bold=True)
F_NORM = Font(name="Calibri", size=10)
F_INPUT = Font(name="Calibri", size=10, color="0000FF")
F_LINK = Font(name="Calibri", size=10, color="00803C")
F_NOTE = Font(name="Calibri", size=9, italic=True, color="595959")
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_SUB = PatternFill("solid", fgColor="D9E1F2")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_KEY = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
TOP = Border(top=Side(style="thin", color="000000"))

FMT = {
    "money": '#,##0;(#,##0);"-"',
    "money1": '#,##0.0;(#,##0.0);"-"',
    "pct": '0.0%;(0.0%);"-"',
    "pct2": '0.00%;(0.00%);"-"',
    "mult": '0.00"x"',
    "price": '#,##0.00;(#,##0.00)',
    "factor": "0.0000",
    "num": "0.00",
    "num3": "0.000",
    "days": "0.0",
    "shares": "#,##0.0",
    "year": "0",
    "text": "@",
}

TAG_HELP = (
    "[sourced] from a filing or public source; [guidance] management statement; [assumption] team or "
    "assistant judgment (proposal until the team confirms); [calc] computed; [unverified] not yet tied to a source."
)


def safe_text(value: str) -> str:
    """Imported text is written as text (formula-injection guard)."""
    return "'" + value if value and value[0] in "=+-@" else value


class Sheet:
    """Small helper: writes cells, remembers named cells and builds absolute references."""

    def __init__(self, ws, label_col: int = 2):
        self.ws = ws
        self.label_col = label_col
        self.names: dict[str, str] = {}
        self.rows: dict[str, int] = {}

    def a(self, key: str) -> str:
        return self.names[key]

    def x(self, key: str, sheet: str | None = None) -> str:
        ref = self.names[key]
        return f"'{self.ws.title}'!{ref}" if sheet else ref

    def text(self, row: int, col: int, value: str, font=F_NORM, fill=None, wrap=False, fmt=None):
        c = self.ws.cell(row=row, column=col)
        c.value = safe_text(str(value))
        c.data_type = "s"
        c.font = font
        if fill:
            c.fill = fill
        if wrap:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        return c

    def put(self, row: int, col: int, value, fmt: str = "money", kind: str = "calc", key: str | None = None,
            bold: bool = False):
        c = self.ws.cell(row=row, column=col)
        if isinstance(value, str) and not value.startswith("="):
            c.value = safe_text(value)
            c.data_type = "s"
        else:
            c.value = value
        c.number_format = FMT.get(fmt, fmt)
        if kind == "input":
            c.font = Font(name="Calibri", size=10, color="0000FF", bold=bold)
            c.fill = FILL_IN
        elif kind == "link":
            c.font = Font(name="Calibri", size=10, color="00803C", bold=bold)
        else:
            c.font = Font(name="Calibri", size=10, bold=bold)
        if key:
            self.names[key] = f"${L(col)}${row}"
        return c

    def header(self, row: int, text: str, last_col: int):
        for col in range(self.label_col, last_col + 1):
            self.ws.cell(row=row, column=col).fill = FILL_HDR
        self.text(row, self.label_col, text, font=F_HDR, fill=FILL_HDR)

    def subheader(self, row: int, text: str, last_col: int):
        for col in range(self.label_col, last_col + 1):
            self.ws.cell(row=row, column=col).fill = FILL_SUB
        self.text(row, self.label_col, text, font=F_BOLD, fill=FILL_SUB)

    def note(self, row: int, col: int, text: str):
        return self.text(row, col, text, font=F_NOTE)


# ----------------------------------------------------------------------------- inputs

class Ctx:
    def __init__(self, project: Path):
        self.project = project
        self.case = core.build_case(project)
        self.extras = self.case.extra or {}
        self.summary = core.load_summary(project / "model" / "model-summary.json")
        self.val = core.load_yaml(project / "valuation" / "valuation.yaml")
        self.drv_doc = core.load_yaml(project / "model" / "drivers.yaml")
        self.py = core.value(self.case)                          # engine basis (reconciliation)
        self.py_pub = core.value(self.case, basis="published")    # published basis (headline)
        meta = self.extras.get("meta") or {}
        self.company = meta.get("company", "Grupo Bimbo, S.A.B. de C.V.")
        self.ticker = meta.get("ticker", "BMV:BIMBO A")
        self.valuation_date = str(meta.get("valuation_date", "2026-10-16"))
        self.currency = meta.get("currency", "MXN")
        self.units = meta.get("units", "millions")

    def drv(self, key: str) -> dict:
        return self.drv_doc["drivers"][key]

    def src(self, key: str, default: str = "") -> str:
        s = (self.extras.get("sources_short") or {}).get(key)
        return s or default


# ----------------------------------------------------------------------------- DCF workbook

DRIVER_ROWS = [
    ("gross_margin", "Gross margin", "pct"),
    ("opex_pct_revenue", "Operating expenses / revenue", "pct"),
    ("da_pct_revenue", "D&A / revenue", "pct"),
    ("capex_pct_revenue", "Capex / revenue", "pct"),
    ("lease_principal_pct_revenue", "Lease principal repaid / revenue (engine: = new leases)", "pct"),
    ("dso", "Receivable days (DSO, on revenue)", "days"),
    ("dio", "Inventory days (DIO, on cost of sales)", "days"),
    ("dpo", "Payable days (DPO, on cost of sales)", "days"),
    ("other_ca_pct_revenue", "Other current assets / revenue", "pct"),
    ("other_cl_pct_revenue", "Other current liabilities / revenue", "pct"),
    ("tax_rate", "Effective tax rate (on EBIT in the DCF)", "pct"),
]


def build_dcf(ctx: Ctx, out: Path) -> None:
    c = ctx.case
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    ws = wb.create_sheet("DCF")
    s = Sheet(ws)
    hist, fcst = c.hist_years, c.fcst_years
    h, n = len(hist), len(fcst)
    Y0 = 4                                   # first year column (D)
    col = {y: Y0 + i for i, y in enumerate(hist + fcst)}
    first_f, last_f = col[fcst[0]], col[fcst[-1]]
    last_h = col[hist[-1]]
    NOTE = last_f + 1
    LASTC = NOTE + 1
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["C"].width = 9
    for yy in hist + fcst:
        ws.column_dimensions[L(col[yy])].width = 11
    ws.column_dimensions[L(NOTE)].width = 70
    ws.sheet_view.showGridLines = False

    def yref(row_key: str, year: int) -> str:
        return f"{L(col[year])}{s.rows[row_key]}"

    r = 2
    s.text(r, 2, "Discounted Cash Flow Analysis", font=F_TITLE)
    r += 1
    s.text(r, 2, f"{ctx.company} ({ctx.ticker}) | valuation date {ctx.valuation_date} | "
                 f"{ctx.currency} {ctx.units}, IFRS | layout after S&P Capital IQ Pro 'DCF Analysis'", font=F_SUB)
    r += 1
    s.text(r, 2, "Blue on yellow = input (change here); black = formula; green = link to another sheet. "
                 "Every input carries a tag and source in 'Inputs_Sources'. " + TAG_HELP, font=F_NOTE)
    r += 2

    # --- key outputs (filled with references after the cells exist)
    s.header(r, "KEY OUTPUTS - PUBLISHED BASIS (leases grow with the business; exit on peer-EV claims). "
                "Engine basis kept for reconciliation ('Engine_Recon')", LASTC)
    r += 1
    key_out_row = r
    r += 14

    # --- cost of capital
    s.header(r, "ASSUMPTIONS - COST OF CAPITAL (from valuation/valuation.yaml; full build in "
                f"BIMBOA_WACC_v{getattr(ctx, 'version', '2')}.xlsx)", LASTC)
    r += 1
    w = c.wacc_in
    wacc_rows = [
        ("rf", "Risk-free rate", w["risk_free"], "pct2", "input", "risk_free"),
        ("erp", "Equity risk premium (mature market)", w["equity_risk_premium"], "pct2", "input", "equity_risk_premium"),
        ("crp", "Country risk premium", w["country_risk_premium"], "pct2", "input", "country_risk_premium"),
        ("beta_u", "Unlevered beta", w["beta_unlevered"], "num3", "input", "beta_unlevered"),
        ("de", "Target debt / equity (market values, debt incl. leases)", w["target_debt_to_equity"], "num3", "input",
         "target_debt_to_equity"),
        ("kd", "Pre-tax cost of debt (MXN)", w["pre_tax_cost_of_debt"], "pct2", "input", "pre_tax_cost_of_debt"),
        ("tax_m", "Marginal tax rate (WACC)", w["tax_rate"], "pct", "input", "tax_rate"),
    ]
    for key, label, v, fmt, kind, src_key in wacc_rows:
        s.text(r, 2, label)
        s.put(r, 4, v, fmt, kind, key=key)
        s.note(r, 5, ctx.src(src_key, "valuation/valuation.yaml"))
        r += 1
    calc_rows = [
        ("beta_l", "Relevered beta = unlevered x (1 + (1 - tax) x D/E)", f"={s.a('beta_u')}*(1+(1-{s.a('tax_m')})*{s.a('de')})", "num3"),
        ("ke", "Cost of equity = rf + beta x ERP + CRP", f"={s.a('rf')}+{{bl}}*{s.a('erp')}+{s.a('crp')}", "pct2"),
        ("kd_after", "After-tax cost of debt", f"={s.a('kd')}*(1-{s.a('tax_m')})", "pct2"),
        ("we", "Equity weight = 1 / (1 + D/E)", f"=1/(1+{s.a('de')})", "pct"),
        ("wd", "Debt weight = D/E / (1 + D/E)", f"={s.a('de')}/(1+{s.a('de')})", "pct"),
    ]
    for key, label, formula, fmt in calc_rows:
        if "{bl}" in formula:
            formula = formula.replace("{bl}", s.a("beta_l"))
        s.text(r, 2, label)
        s.put(r, 4, formula, fmt, key=key)
        r += 1
    s.text(r, 2, "WACC", font=F_BOLD)
    s.put(r, 4, f"={s.a('we')}*{s.a('ke')}+{s.a('wd')}*{s.a('kd_after')}", "pct2", key="wacc", bold=True)
    s.note(r, 5, "Same formula as the engine's WACC sheet; the reconciliation is in 'Engine_Recon'.")
    r += 2

    # --- discounting and terminal inputs
    s.header(r, "ASSUMPTIONS - DISCOUNTING, TERMINAL VALUE AND SCENARIO SHIFTS", LASTC)
    r += 1
    disc_rows = [
        ("offset", "Mid-year convention offset (0.5 = mid-year, 0 = end of year)", 0.5 if c.mid_year else 0.0, "num", "input",
         "Engine convention: cash flow t is discounted t - offset years from 31-Dec-2025."),
        ("stub", "Years since last fiscal year-end (31-Dec-2025 to the valuation date)", c.stub, "num3", "input",
         "Roll-forward of equity value to today at the cost of equity (engine approximation)."),
        ("g", "Terminal growth (g), nominal MXN", c.g, "pct2", "input", ctx.src("growth", "valuation/valuation.yaml")),
        ("exit_mult", "Exit EV / EBITDA multiple (trailing, applied to final-year EBITDA)", c.exit_multiple, "mult", "input",
         ctx.src("exit_ev_ebitda", "valuation/valuation.yaml")),
        ("lt_gdp", "Long-term nominal GDP growth, MXN (cap for g)", float(ctx.val["terminal"]["lt_nominal_gdp_growth"]),
         "pct2", "input", ctx.src("lt_nominal_gdp_growth", "valuation/valuation.yaml")),
        ("g_shift", "Scenario: revenue growth shift added to every forecast year (0 = base case)", 0.0, "pct2", "input",
         "Used by the reverse DCF (Goal Seek this cell). Keep 0 for the base case."),
        ("gm_shift", "Scenario: gross margin shift, all forecast years (e.g. -0.01 = -100 bp wheat shock)", 0.0, "pct2",
         "input", "Climate / input-cost scenario lever. Keep 0 for the base case."),
        ("capex_shift", "Scenario: capex / revenue shift, all forecast years (e.g. +0.005 = +50 bp)", 0.0, "pct2",
         "input", "Reinvestment lever (debt / FCF pillar). Keep 0 for the base case."),
    ]
    for key, label, v, fmt, kind, note in disc_rows:
        s.text(r, 2, label)
        s.put(r, 4, v, fmt, kind, key=key)
        s.note(r, 5, note)
        r += 1
    r += 1

    # --- year header helper
    def year_header(row: int, title: str):
        s.header(row, title, LASTC)
        for yy in hist + fcst:
            cc = ws.cell(row=row, column=col[yy])
            cc.value = f"FY{yy}{'A' if yy in hist else 'E'}"
            cc.font = F_HDR
            cc.alignment = Alignment(horizontal="right")
        ws.cell(row=row, column=NOTE).value = "Tag / source / note"
        ws.cell(row=row, column=NOTE).font = F_HDR

    # --- drivers
    year_header(r, "FORECAST DRIVERS (model/drivers.yaml; history = implied values)")
    r += 1
    seg_keys = c.segments
    seg_labels = {sg["key"]: sg.get("label", sg["key"]) for sg in (ctx.drv_doc.get("revenue_segments") or [])}
    if seg_keys:
        for sg in seg_keys:
            key = f"drv_seg_{sg}"
            s.rows[key] = r
            s.text(r, 2, f"{seg_labels.get(sg, sg)} revenue growth")
            for i, yy in enumerate(fcst):
                s.put(r, col[yy], ctx.drv(f"seg_{sg}_growth")["values"][i], "pct", "input")
            d = ctx.drv(f"seg_{sg}_growth")
            s.note(r, NOTE, f"[{d.get('tag', '')}] {d.get('rationale', '')}")
            r += 1
    else:
        s.rows["drv_rev_growth"] = r
        s.text(r, 2, "Revenue growth")
        for i, yy in enumerate(fcst):
            s.put(r, col[yy], ctx.drv("revenue_growth")["values"][i], "pct", "input")
        d = ctx.drv("revenue_growth")
        s.note(r, NOTE, f"[{d.get('tag', '')}] {d.get('rationale', '')}")
        r += 1
    for key, label, fmt in DRIVER_ROWS:
        s.rows[f"drv_{key}"] = r
        s.text(r, 2, label)
        for i, yy in enumerate(fcst):
            s.put(r, col[yy], ctx.drv(key)["values"][i], fmt, "input")
        d = ctx.drv(key)
        s.note(r, NOTE, f"[{d.get('tag', '')}] pillar {d.get('pillar', '') or '-'}: {d.get('rationale', '')}")
        r += 1
    r += 1

    # --- revenue build and projected cash flows
    fin = c.fin
    year_header(r, "PROJECTED CASH FLOWS (MXN millions)")
    r += 1
    if seg_keys:
        for sg in seg_keys:
            key = f"seg_{sg}"
            s.rows[key] = r
            s.text(r, 2, f"  {seg_labels.get(sg, sg)}")
            for yy in hist:
                s.put(r, col[yy], fin.get(key, {}).get(yy, 0.0), "money", "input")
            for yy in fcst:
                prev = L(col[yy] - 1)
                s.put(r, col[yy], f"={prev}{r}*(1+{L(col[yy])}{s.rows[f'drv_seg_{sg}']}+{s.a('g_shift')})", "money")
            s.note(r, NOTE, "[sourced] segment revenue, data/financials.csv (history)")
            r += 1
    s.rows["revenue"] = r
    s.text(r, 2, "Total revenue", font=F_BOLD)
    for yy in hist:
        s.put(r, col[yy], fin["revenue"][yy], "money", "input", bold=True)
    for yy in fcst:
        cl = L(col[yy])
        if seg_keys:
            first_seg, last_seg = s.rows[f"seg_{seg_keys[0]}"], s.rows[f"seg_{seg_keys[-1]}"]
            s.put(r, col[yy], f"=SUM({cl}{first_seg}:{cl}{last_seg})", "money", bold=True)
        else:
            s.put(r, col[yy], f"={L(col[yy] - 1)}{r}*(1+{cl}{s.rows['drv_rev_growth']}+{s.a('g_shift')})", "money", bold=True)
    s.note(r, NOTE, "[sourced] history data/financials.csv; forecast = sum of segments (or growth driver)")
    r += 1

    def line(key: str, label: str, hist_fn, fcst_fn, fmt="money", bold=False, note=""):
        nonlocal r
        s.rows[key] = r
        s.text(r, 2, label, font=F_BOLD if bold else F_NORM)
        for yy in hist:
            v = hist_fn(yy)
            if v is not None:
                s.put(r, col[yy], v, fmt, "input" if not (isinstance(v, str) and v.startswith("=")) else "calc", bold=bold)
        for yy in fcst:
            v = fcst_fn(yy)
            if v is not None:
                s.put(r, col[yy], v, fmt, bold=bold)
        if note:
            s.note(r, NOTE, note)
        r += 1

    R_ = lambda k, yy: yref(k, yy)  # noqa: E731
    D_ = lambda k, yy: f"{L(col[yy])}{s.rows['drv_' + k]}"  # noqa: E731
    prev_col = lambda yy: L(col[yy] - 1)  # noqa: E731

    line("rev_growth", "  Annual growth",
         lambda yy: (f"={R_('revenue', yy)}/{prev_col(yy)}{s.rows['revenue']}-1" if yy != hist[0] else None),
         lambda yy: f"={R_('revenue', yy)}/{prev_col(yy)}{s.rows['revenue']}-1", "pct")
    line("cogs", "Cost of sales",
         lambda yy: fin["cogs"][yy],
         lambda yy: f"={R_('revenue', yy)}*(1-({D_('gross_margin', yy)}+{s.a('gm_shift')}))",
         note="Forecast = revenue x (1 - (gross margin + scenario shift))")
    line("gross_profit", "Gross profit",
         lambda yy: f"={R_('revenue', yy)}-{R_('cogs', yy)}", lambda yy: f"={R_('revenue', yy)}-{R_('cogs', yy)}",
         bold=True)
    line("gm", "  Gross margin", lambda yy: f"={R_('gross_profit', yy)}/{R_('revenue', yy)}",
         lambda yy: f"={R_('gross_profit', yy)}/{R_('revenue', yy)}", "pct")
    line("opex", "Operating expenses (net, incl. PTU)", lambda yy: fin["opex"][yy],
         lambda yy: f"={R_('revenue', yy)}*{D_('opex_pct_revenue', yy)}")
    line("ebit", "EBIT", lambda yy: f"={R_('gross_profit', yy)}-{R_('opex', yy)}",
         lambda yy: f"={R_('gross_profit', yy)}-{R_('opex', yy)}", bold=True)
    line("ebit_m", "  EBIT margin", lambda yy: f"={R_('ebit', yy)}/{R_('revenue', yy)}",
         lambda yy: f"={R_('ebit', yy)}/{R_('revenue', yy)}", "pct")
    line("da", "D&A (incl. right-of-use depreciation)", lambda yy: fin["da"][yy],
         lambda yy: f"={R_('revenue', yy)}*{D_('da_pct_revenue', yy)}")
    line("ebitda", "EBITDA (post-IFRS 16) = EBIT + D&A", lambda yy: f"={R_('ebit', yy)}+{R_('da', yy)}",
         lambda yy: f"={R_('ebit', yy)}+{R_('da', yy)}", bold=True,
         note="Engine definition. Company Adjusted EBITDA adds back one-offs (bridge in data/adjustments.md).")
    line("ebitda_m", "  EBITDA margin", lambda yy: f"={R_('ebitda', yy)}/{R_('revenue', yy)}",
         lambda yy: f"={R_('ebitda', yy)}/{R_('revenue', yy)}", "pct")

    def hist_etr(yy):
        ebt = (fin["revenue"][yy] - fin["cogs"][yy] - fin["opex"][yy] - fin.get("interest_expense", {}).get(yy, 0.0)
               + fin.get("interest_income", {}).get(yy, 0.0) + fin.get("other_financial_net", {}).get(yy, 0.0))
        tax = fin.get("income_tax", {}).get(yy, 0.0)
        return tax / ebt if ebt > 0 else 0.0

    line("tax_rate_used", "  Tax rate on EBIT", lambda yy: hist_etr(yy), lambda yy: f"={D_('tax_rate', yy)}", "pct",
         note="History: effective rate = income tax / pre-tax income (display only); forecast: tax driver")
    line("tax", "Less: taxes on EBIT", lambda yy: f"=-{R_('ebit', yy)}*{R_('tax_rate_used', yy)}",
         lambda yy: f"=-{R_('ebit', yy)}*{R_('tax_rate_used', yy)}")
    line("nopat", "NOPAT (unlevered net income)", lambda yy: f"={R_('ebit', yy)}+{R_('tax', yy)}",
         lambda yy: f"={R_('ebit', yy)}+{R_('tax', yy)}", bold=True)
    line("plus_da", "Plus: D&A", lambda yy: f"={R_('da', yy)}", lambda yy: f"={R_('da', yy)}")
    line("capex", "Less: capital expenditures (PP&E + intangibles)", lambda yy: -fin["capex"][yy],
         lambda yy: f"=-{R_('revenue', yy)}*({D_('capex_pct_revenue', yy)}+{s.a('capex_shift')})")
    line("new_leases", "Less: new leases (IFRS 16; engine proxy = lease principal repaid)",
         lambda yy: -fin.get("lease_principal_paid", {}).get(yy, 0.0),
         lambda yy: f"=-{R_('revenue', yy)}*{D_('lease_principal_pct_revenue', yy)}",
         note="Engine proxy (keeps the lease liability flat). The published basis below also deducts the growth "
              "in the lease liability.")
    nwc_row = r  # change in NWC row is written after the schedule (placeholder now)
    s.rows["dnwc_cf"] = r
    s.text(r, 2, "Plus / (less): change in operating working capital")
    r += 1
    s.rows["fcff"] = r
    s.text(r, 2, "Unlevered free cash flow (FCFF)", font=F_BOLD)
    r += 1
    line("fcff_g", "  Annual growth", lambda yy: None,
         lambda yy: (f"=IFERROR({R_('fcff', yy)}/{prev_col(yy)}{s.rows['fcff']}-1,0)" if yy != fcst[0] else None), "pct")
    line("fcff_m", "  FCFF margin", lambda yy: None, lambda yy: f"={R_('fcff', yy)}/{R_('revenue', yy)}", "pct")
    r += 1

    # --- working-capital schedule
    year_header(r, "WORKING CAPITAL SCHEDULE (engine chart; non-operating balances shown separately and held flat in "
                   "the forecast, so only operating items move FCFF)")
    r += 1
    owc = ctx.extras.get("operating_wc") or {}

    def owc_v(k: str, yy: int) -> float:
        d = owc.get(k) or {}
        return float(d.get(yy, d.get(str(yy), 0.0)))

    wc_items = [
        ("wc_ar", "Trade receivables", "receivables", lambda yy: f"={R_('revenue', yy)}*{D_('dso', yy)}/365"),
        ("wc_inv", "Inventories", "inventory", lambda yy: f"={R_('cogs', yy)}*{D_('dio', yy)}/365"),
        ("wc_oca", "Other current assets (total)", "other_current_assets",
         lambda yy: f"={R_('revenue', yy)}*{D_('other_ca_pct_revenue', yy)}"),
        ("wc_ap", "Trade payables", "payables", lambda yy: f"={R_('cogs', yy)}*{D_('dpo', yy)}/365"),
        ("wc_ocl", "Other current liabilities, non-debt (total)", "other_current_liabilities",
         lambda yy: f"={R_('revenue', yy)}*{D_('other_cl_pct_revenue', yy)}"),
    ]
    for key, label, fin_key, fn in wc_items:
        line(key, label, lambda yy, fk=fin_key: fin.get(fk, {}).get(yy, 0.0), fn)
        if key == "wc_oca" and owc:
            line("wc_oca_nonop", "  of which non-operating: derivatives, collateral deposits, held for sale (flat)",
                 lambda yy: (fin.get("other_current_assets", {}).get(yy, 0.0) - owc_v("other_receivables", yy)
                             - owc_v("prepayments", yy)),
                 lambda yy: f"=${L(last_h)}{s.rows['wc_oca_nonop']}",
                 note="[calc] history = total less other receivables and prepayments (data/adjustments.md operating "
                      "working-capital table); forecast held at FY2025")
        if key == "wc_ocl" and owc:
            line("wc_ocl_nonop", "  of which non-operating: income and other taxes, accrued interest, derivatives (flat)",
                 lambda yy: (fin.get("other_current_liabilities", {}).get(yy, 0.0) - owc_v("related_party_payables", yy)
                             - owc_v("ptu_payable", yy) - owc_v("operating_accruals", yy)),
                 lambda yy: f"=${L(last_h)}{s.rows['wc_ocl_nonop']}",
                 note="[calc] history = total less related-party payables, PTU and operating accruals (note 14); "
                      "forecast held at FY2025")
    line("nwc", "Net working capital - engine basis (incl. non-operating balances)",
         lambda yy: f"={R_('wc_ar', yy)}+{R_('wc_inv', yy)}+{R_('wc_oca', yy)}-{R_('wc_ap', yy)}-{R_('wc_ocl', yy)}",
         lambda yy: f"={R_('wc_ar', yy)}+{R_('wc_inv', yy)}+{R_('wc_oca', yy)}-{R_('wc_ap', yy)}-{R_('wc_ocl', yy)}")
    if owc:
        line("nwc_op", "Operating working capital = engine basis - non-operating assets + non-operating liabilities",
             lambda yy: f"={R_('nwc', yy)}-{R_('wc_oca_nonop', yy)}+{R_('wc_ocl_nonop', yy)}",
             lambda yy: f"={R_('nwc', yy)}-{R_('wc_oca_nonop', yy)}+{R_('wc_ocl_nonop', yy)}", bold=True,
             note="Ties to data/adjustments.md: -14,860 / -15,353 / -8,599 / -6,279 / -6,232 (Dec-21 to Dec-25); "
                  "see 'Checks'")
    line("dnwc", "(Increase) / decrease in working capital (cash-flow sign; used in FCFF)",
         lambda yy: (f"=-({R_('nwc', yy)}-{prev_col(yy)}{s.rows['nwc']})" if yy != hist[0] else None),
         lambda yy: f"=-({R_('nwc', yy)}-{prev_col(yy)}{s.rows['nwc']})",
         note="Engine: change_nwc = -(dAR + dInv + dOCA - dAP - dOCL). Non-operating balances are flat in the forecast "
              "(drivers encode them as a declining share of revenue), so this equals the operating change on the base "
              "revenue path; under the g_shift scenario they move slightly with revenue.")
    if owc:
        line("dnwc_op", "  Memo: (increase) / decrease in operating working capital",
             lambda yy: (f"=-({R_('nwc_op', yy)}-{prev_col(yy)}{s.rows['nwc_op']})" if yy != hist[0] else None),
             lambda yy: f"=-({R_('nwc_op', yy)}-{prev_col(yy)}{s.rows['nwc_op']})")
    r += 1
    # fill the FCF rows that depend on the schedule
    for yy in hist[1:] + fcst:
        s.put(nwc_row, col[yy], f"={R_('dnwc', yy)}", "money")
    s.note(nwc_row, NOTE, "From the working-capital schedule below")
    for yy in hist[1:] + fcst:
        cl = L(col[yy])
        s.put(s.rows["fcff"], col[yy],
              f"={cl}{s.rows['nopat']}+{cl}{s.rows['plus_da']}+{cl}{s.rows['capex']}+{cl}{s.rows['new_leases']}"
              f"+{cl}{s.rows['dnwc_cf']}", "money", bold=True)
    s.note(s.rows["fcff"], NOTE, "Engine: FCFF = NOPAT + D&A - capex - new leases + change in NWC. History: display only.")

    # --- discounting
    year_header(r, "DISCOUNTING (to 31-Dec-2025; engine convention)")
    r += 1
    line("t", "Forecast year number", lambda yy: None, lambda yy: fcst.index(yy) + 1, "year")
    for yy in fcst:
        ws.cell(row=s.rows["t"], column=col[yy]).font = F_INPUT
    line("period", "Discount period = year number - mid-year offset", lambda yy: None,
         lambda yy: f"={R_('t', yy)}-{s.a('offset')}", "num")
    line("df", "Discount factor = 1 / (1 + WACC) ^ period", lambda yy: None,
         lambda yy: f"=1/(1+{s.a('wacc')})^{R_('period', yy)}", "factor")
    line("pv", "Present value of FCFF", lambda yy: None, lambda yy: f"={R_('fcff', yy)}*{R_('df', yy)}", bold=True)
    r += 1

    fF, fL = L(first_f), L(last_f)
    fcff_rng = f"{fF}{s.rows['fcff']}:{fL}{s.rows['fcff']}"
    fcff_N = f"{fL}{s.rows['fcff']}"
    ebitda_N = f"{fL}{s.rows['ebitda']}"
    nopat_N = f"{fL}{s.rows['nopat']}"
    N = n

    # --- terminal value and EV
    s.header(r, "PRESENT VALUE OF THE ENTERPRISE", LASTC)
    r += 1
    tv_rows = [
        ("sum_pv", f"Sum of PV of FCFF, FY{fcst[0]}E-FY{fcst[-1]}E", f"=SUM({fF}{s.rows['pv']}:{fL}{s.rows['pv']})", "money"),
        ("tv_g", "Terminal value - Gordon growth = FCFF_N x (1 + g) / (WACC - g)",
         f"={fcff_N}*(1+{s.a('g')})/({s.a('wacc')}-{s.a('g')})", "money"),
        ("pv_tv_g", "PV of terminal value - Gordon (discounted N - offset: value as of the last cash flow's date)",
         f"={{tv_g}}/(1+{s.a('wacc')})^({N}-{s.a('offset')})", "money"),
        ("ev_g", "Enterprise value - Gordon", "={sum_pv}+{pv_tv_g}", "money"),
        ("tv_x", "Terminal value - exit multiple = final-year EBITDA x multiple", f"={ebitda_N}*{s.a('exit_mult')}", "money"),
        ("pv_tv_x", "PV of terminal value - exit (a year-end value, discounted N)",
         f"={{tv_x}}/(1+{s.a('wacc')})^{N}", "money"),
        ("ev_x", "Enterprise value - exit multiple", "={sum_pv}+{pv_tv_x}", "money"),
        ("tv_share_g", "Terminal value share of EV - Gordon", "={pv_tv_g}/{ev_g}", "pct"),
        ("tv_share_x", "Terminal value share of EV - exit", "={pv_tv_x}/{ev_x}", "pct"),
    ]
    for key, label, formula, fmt in tv_rows:
        formula = formula.format(**{k: s.a(k) for k in s.names if "{" + k + "}" in formula})
        s.text(r, 2, label, font=F_BOLD if key.startswith("ev_") else F_NORM)
        s.put(r, 4, formula, fmt, key=key, bold=key.startswith("ev_"))
        r += 1
    r += 1

    # --- equity bridge
    b = core.bridge(c)
    s.header(r, f"EQUITY BRIDGE AND VALUE PER SHARE (engine basis: balances at 31-Dec-{hist[-1]})", LASTC)
    r += 1
    bridge_inputs = [
        ("b_debt_s", "Short-term debt", b["debt_short"]),
        ("b_debt_l", "Long-term debt", b["debt_long"]),
        ("b_leases", "Lease liabilities (IFRS 16)", b["lease_liabilities"]),
        ("b_cash", "Cash and equivalents", b["cash"]),
    ]
    for key, label, v in bridge_inputs:
        s.text(r, 2, label)
        s.put(r, 4, v, "money", "input", key=key)
        s.note(r, 5, f"[sourced] data/financials.csv FY{hist[-1]}")
        r += 1
    s.text(r, 2, "Net debt incl. leases")
    s.put(r, 4, f"={s.a('b_debt_s')}+{s.a('b_debt_l')}+{s.a('b_leases')}-{s.a('b_cash')}", "money", key="net_debt")
    r += 1
    s.text(r, 2, "Non-controlling interests (book)")
    s.put(r, 4, b["nci"], "money", "input", key="nci")
    s.note(r, 5, f"[sourced] data/financials.csv FY{hist[-1]}")
    r += 1
    nci_cfg = ctx.extras.get("nci_claim") or {}
    if nci_cfg:
        s.subheader(r, "NCI claim: BBU independent-operator structured entities (consolidated; their earnings sit in "
                       "EBITDA and FCFF)", LASTC)
        r += 1
        s.text(r, 2, "NCI income, annualised (1H26 x 2)")
        s.put(r, 4, float(nci_cfg.get("income", 0.0)), "money", "input", key="nci_income")
        s.note(r, 5, "[sourced] CIQ quarterly 1Q26 436 + 2Q26 471; FY2025 1,813 (EEFF25 p.13)")
        r += 1
        s.text(r, 2, "Growth of the NCI claim (MXN, US-linked)")
        s.put(r, 4, float(nci_cfg.get("growth", 0.0)), "pct2", "input", key="nci_g")
        s.note(r, 5, "[assumption] North America terminal growth in MXN; the zero-growth case is shown below")
        r += 1
        s.text(r, 2, "NCI at market value = NCI income / (cost of equity - growth)")
        s.put(r, 4, f"={s.a('nci_income')}/({s.a('ke')}-{s.a('nci_g')})", "money", key="nci_mkt")
        s.note(r, 5, "Live at the cost of equity above. " + str(nci_cfg.get("evidence", "")))
        r += 1
        s.text(r, 2, "NCI at market value, zero growth = NCI income / cost of equity (sensitivity)")
        s.put(r, 4, f"={s.a('nci_income')}/{s.a('ke')}", "money", key="nci_mkt_g0")
        r += 1
    comps = (ctx.extras.get("bridge_items") or {})
    comp_keys = []
    for ck, item in comps.items():
        s.text(r, 2, "  Debt-like component: " + str(item.get("label", ck)))
        if ck == "nci_uplift" and nci_cfg:
            s.put(r, 4, f"={s.a('nci_mkt')}-{s.a('nci')}", "money", key=f"dl_{ck}")
        else:
            s.put(r, 4, float(item.get("value", 0.0)), "money", "input", key=f"dl_{ck}")
        s.note(r, 5, f"[{item.get('tag', '')}] {item.get('source', '')}")
        comp_keys.append(f"dl_{ck}")
        r += 1
    if comp_keys:
        s.text(r, 2, "Less: debt-like items (sum of components)")
        s.put(r, 4, "=" + "+".join(s.a(k) for k in comp_keys), "money", key="debt_like")
        s.note(r, 5, ctx.src("debt_like_items", "valuation/valuation.yaml"))
        r += 1
        s.text(r, 2, "Memo: debt-like items in valuation/valuation.yaml (engine input)")
        s.put(r, 4, c.debt_like, "money", "input", key="debt_like_yaml")
        r += 1
    for key, label, v, note in [
        ("non_op", "Plus: non-operating assets", c.non_op, ctx.src("non_operating_assets", "valuation/valuation.yaml")),
        ("debt_like", "Less: debt-like items", c.debt_like, ctx.src("debt_like_items", "valuation/valuation.yaml")),
        ("shares", "Diluted shares, millions (engine: FY weighted average)", b["shares"],
         f"[sourced] data/financials.csv FY{hist[-1]}; outstanding shares at the valuation date in 'Dated_Check'"),
        ("price", "Current share price (MXN)", c.price, ctx.src("share_price", "valuation/valuation.yaml")),
    ]:
        if key == "debt_like" and comp_keys:
            continue
        s.text(r, 2, label)
        s.put(r, 4, v, "price" if key == "price" else ("shares" if key == "shares" else "money"), "input", key=key)
        s.note(r, 5, note)
        r += 1
    dps = None
    if ctx.summary:
        dps = (ctx.summary.get("valuation") or {}).get("dps_next")
    s.text(r, 2, "Next-year dividend per share (engine dps_next; see publication gate)")
    s.put(r, 4, float(dps) if dps is not None else 0.0, "price", "input", key="dps_next")
    s.note(r, 5, "[calc] engine output (first forecast year). The FY2026 dividend was already paid in 2026: "
                 "valuation/valuation-outputs.yaml carries the corrected next-12-month dividend.")
    r += 1
    s.text(r, 2, "Roll-forward factor = (1 + cost of equity) ^ years since FYE")
    s.put(r, 4, f"=(1+{s.a('ke')})^{s.a('stub')}", "num3", key="roll")
    r += 1
    s.subheader(r, "Gordon growth method", LASTC)
    r += 1
    for key, label, formula, fmt in [
        ("eq_g", "Equity value at FYE = EV + non-op - net debt - NCI - debt-like",
         f"={s.a('ev_g')}+{s.a('non_op')}-{s.a('net_debt')}-{s.a('nci')}-{s.a('debt_like')}", "money"),
        ("vps_fye_g", "Value per share at FYE", "={eq_g}/" + s.a("shares"), "price"),
        ("vps_g", "Value per share today (rolled forward)", "={vps_fye_g}*" + s.a("roll"), "price"),
        ("prem_g", "Premium / (discount) to current price - value today", "={vps_g}/" + s.a("price") + "-1", "pct"),
        ("tgt12_g", "12-month target - Gordon = value today x (1 + ke) - next dividend",
         "={vps_g}*(1+" + s.a("ke") + ")-" + s.a("dps_next"), "price"),
        ("up12_g", "Upside to 12-month target - Gordon", "={tgt12_g}/" + s.a("price") + "-1", "pct"),
    ]:
        formula = formula.format(**{k: s.a(k) for k in s.names if "{" + k + "}" in formula})
        s.text(r, 2, label, font=F_BOLD if key in ("vps_g", "tgt12_g") else F_NORM)
        s.put(r, 4, formula, fmt, key=key, bold=key in ("vps_g", "tgt12_g"))
        r += 1
    s.subheader(r, "Exit multiple method", LASTC)
    r += 1
    for key, label, formula, fmt in [
        ("eq_x", "Equity value at FYE", f"={s.a('ev_x')}+{s.a('non_op')}-{s.a('net_debt')}-{s.a('nci')}-{s.a('debt_like')}",
         "money"),
        ("vps_fye_x", "Value per share at FYE", "={eq_x}/" + s.a("shares"), "price"),
        ("vps_x", "Value per share today (rolled forward)", "={vps_fye_x}*" + s.a("roll"), "price"),
        ("prem_x", "Premium / (discount) to current price - value today", "={vps_x}/" + s.a("price") + "-1", "pct"),
    ]:
        formula = formula.format(**{k: s.a(k) for k in s.names if "{" + k + "}" in formula})
        s.text(r, 2, label, font=F_BOLD if key == "vps_x" else F_NORM)
        s.put(r, 4, formula, fmt, key=key, bold=key == "vps_x")
        r += 1
    per_share = f"/{s.a('shares')}*{s.a('roll')}"
    claims_peer_ev = "+".join(s.a(k) for k in comp_keys if k != "dl_nci_uplift") or "0"
    s.text(r, 2, "Debt-like claims outside the peer-EV definition (all components except the NCI uplift)")
    s.put(r, 4, f"={claims_peer_ev}", "money", key="claims_outside_peer")
    s.note(r, 5, "Peer EVs (CIQ basis: market cap + debt + leases + book NCI - cash) exclude pensions and other "
                 "employee claims, MEPP, provisions and derivatives. The NCI uplift stays: Bimbo's EBITDA includes the "
                 "IO earnings, while Flowers' distributors are not consolidated.")
    r += 1
    s.text(r, 2, "Value per share today - exit, engine basis on peer-EV-consistent claims (v1 reconciliation)")
    s.put(r, 4, f"={s.a('vps_x')}+{s.a('claims_outside_peer')}{per_share}", "price", key="vps_x_peer")
    r += 2

    # --- cross-checks
    s.header(r, "CROSS-CHECKS (terminal economics)", LASTC)
    r += 1
    reinv = (f"=(-{fL}{s.rows['capex']}-{fL}{s.rows['new_leases']}-{fL}{s.rows['plus_da']}-{fL}{s.rows['dnwc_cf']})"
             f"/{nopat_N}")
    for key, label, formula, fmt, note in [
        ("wacc_minus_g", "WACC - g", f"={s.a('wacc')}-{s.a('g')}", "pct2", "Engine warns below 2 points."),
        ("g_vs_gdp", "g minus long-term nominal GDP growth (must be <= 0)", f"={s.a('g')}-{s.a('lt_gdp')}", "pct2", ""),
        ("impl_mult", "Exit multiple implied by the Gordon value (TV carried to year-end / final EBITDA)",
         f"={s.a('tv_g')}*(1+{s.a('wacc')})^{s.a('offset')}/{ebitda_N}", "mult", ""),
        ("impl_g_x", "Growth implied by the exit multiple",
         f"=({s.a('tv_x')}/(1+{s.a('wacc')})^{s.a('offset')}*{s.a('wacc')}-{fcff_N})"
         f"/({s.a('tv_x')}/(1+{s.a('wacc')})^{s.a('offset')}+{fcff_N})", "pct2", ""),
        ("tv_ronic", "Return on new capital implied by the terminal value = g / (1 - FCFF_N / NOPAT_N)",
         f'=IF(AND({nopat_N}>0,{fcff_N}<{nopat_N}),{s.a("g")}/(1-{fcff_N}/{nopat_N}),"n.m.")', "pct",
         "Domain guard: undefined when NOPAT <= 0 or FCFF >= NOPAT (no net reinvestment)."),
        ("reinv", "Final-year net reinvestment rate = (capex + new leases - D&A + increase in NWC) / NOPAT", reinv, "pct",
         "Steady state: reinvestment rate x RONIC should equal g."),
        ("ronic_ss", "RONIC consistent with g and the final-year reinvestment = g / reinvestment rate",
         f'=IF({{reinv}}>0,{s.a("g")}/{{reinv}},"n.m.")', "pct", "Compare with WACC and with ROIC history."),
    ]:
        formula = formula.replace("{reinv}", s.a("reinv") if "reinv" in s.names else "0")
        s.text(r, 2, label)
        s.put(r, 4, formula, fmt, key=key)
        if note:
            s.note(r, 5, note)
        r += 1
    r += 1

    # --- published basis: leases grow with the business; exit on peer-EV claims
    year_header(r, "PUBLISHED BASIS - LEASES GROW WITH THE BUSINESS (lease liability held at the FY2025 share of revenue; "
                   "new leases = principal + growth in the liability)")
    r += 1
    line("ls_liab", "Lease liabilities (IFRS 16), closing",
         lambda yy: fin.get("lease_liabilities", {}).get(yy),
         lambda yy: f"={prev_col(yy)}{s.rows['ls_liab']}*{R_('revenue', yy)}/{prev_col(yy)}{s.rows['revenue']}",
         note="[sourced] history data/financials.csv; forecast = prior year x revenue growth (constant lease intensity)")
    line("ls_int", "  Lease liabilities / revenue",
         lambda yy: f"={R_('ls_liab', yy)}/{R_('revenue', yy)}",
         lambda yy: f"={R_('ls_liab', yy)}/{R_('revenue', yy)}", "pct2",
         note="History 6.7-9.2% (FY21-25), mean 8.0%: the forecast holds FY2025's 8.15% (test b)")
    line("ls_prin", "Lease principal repaid (engine: = new leases)",
         lambda yy: fin.get("lease_principal_paid", {}).get(yy),
         lambda yy: f"={R_('revenue', yy)}*{D_('lease_principal_pct_revenue', yy)}",
         note="[calc] history derived as payments less accrued lease interest; the filings do not split the cash "
              "interest inside lease payments (data/adjustments.md #6)")
    line("ls_growth", "Growth in the lease liability = new leases - principal", lambda yy: None,
         lambda yy: f"={R_('ls_liab', yy)}-{prev_col(yy)}{s.rows['ls_liab']}")
    line("ls_add", "New leases (right-of-use additions, organic) = principal + growth in the liability",
         lambda yy: None, lambda yy: f"={R_('ls_prin', yy)}+{R_('ls_growth', yy)}",
         note="History of new leases + modifications / revenue: 1.11% / 2.13% / 1.65% (FY23-25, 1.63% average)")
    line("fcff_pub", "FCFF - published basis = FCFF (engine basis) - growth in the lease liability", lambda yy: None,
         lambda yy: f"={R_('fcff', yy)}-{R_('ls_growth', yy)}", bold=True)
    line("pv_pub", "PV of FCFF - published basis", lambda yy: None, lambda yy: f"={R_('fcff_pub', yy)}*{R_('df', yy)}")
    lease_N = f"{fL}{s.rows['ls_liab']}"
    fcff_pub_rng = f"{fF}{s.rows['fcff_pub']}:{fL}{s.rows['fcff_pub']}"
    r += 1
    s.subheader(r, "Published valuation", LASTC)
    r += 1
    nb = "+" + s.a("non_op") + "-" + s.a("net_debt") + "-" + s.a("nci")
    pub_rows = [
        ("pub_sum_pv", "Sum of PV of FCFF - published basis", f"=SUM({fF}{s.rows['pv_pub']}:{fL}{s.rows['pv_pub']})",
         "money", ""),
        ("pub_fcff_next", "FCFF in year N+1 = FCFF_N (engine basis) x (1 + g) - g x lease liability_N",
         f"={fcff_N}*(1+{s.a('g')})-{s.a('g')}*{lease_N}", "money",
         "The lease book keeps growing at g after the forecast, so g x L_N is deducted every year."),
        ("pub_tv_g", "Terminal value - Gordon, published = FCFF_N+1 / (WACC - g)",
         f"={{pub_fcff_next}}/({s.a('wacc')}-{s.a('g')})", "money", ""),
        ("pub_pv_tv_g", "PV of terminal value - Gordon, published (discounted N - offset)",
         f"={{pub_tv_g}}/(1+{s.a('wacc')})^({N}-{s.a('offset')})", "money", ""),
        ("pub_ev_g", "Enterprise value - Gordon, published", "={pub_sum_pv}+{pub_pv_tv_g}", "money", ""),
        ("pub_ev_x", "Enterprise value - exit, published (exit TV unchanged: EBITDA is the same)",
         "={pub_sum_pv}+" + s.a("pv_tv_x"), "money", ""),
        ("pub_tv_share_g", "Terminal value share of EV - Gordon, published", "={pub_pv_tv_g}/{pub_ev_g}", "pct", ""),
        ("pub_tv_share_x", "Terminal value share of EV - exit, published", "=" + s.a("pv_tv_x") + "/{pub_ev_x}", "pct", ""),
        ("pub_eq_g", "Equity value at FYE - Gordon = EV + non-op - net debt - NCI (book) - all debt-like items",
         "={pub_ev_g}" + nb + "-" + s.a("debt_like"), "money", ""),
        ("pub_vps_g", "PUBLISHED value per share today - Gordon", "={pub_eq_g}/" + s.a("shares") + "*" + s.a("roll"),
         "price", ""),
        ("pub_eq_x", "Equity value at FYE - exit = EV + non-op - net debt - NCI (book) - NCI uplift (peer-EV claims)",
         "={pub_ev_x}" + nb + "-" + s.a("debt_like") + "+" + s.a("claims_outside_peer"), "money", ""),
        ("pub_vps_x", "PUBLISHED value per share today - exit multiple", "={pub_eq_x}/" + s.a("shares") + "*" + s.a("roll"),
         "price", ""),
        ("pub_prem_g", "Premium / (discount) to the current price - Gordon", "={pub_vps_g}/" + s.a("price") + "-1", "pct", ""),
        ("pub_prem_x", "Premium / (discount) to the current price - exit", "={pub_vps_x}/" + s.a("price") + "-1", "pct", ""),
    ]
    for key, label, formula, fmt, note in pub_rows:
        formula = formula.format(**{k: s.a(k) for k in s.names if "{" + k + "}" in formula})
        bold = key in ("pub_vps_g", "pub_vps_x")
        s.text(r, 2, label, font=F_BOLD if bold else F_NORM)
        s.put(r, 4, formula, fmt, key=key, bold=bold)
        if note:
            s.note(r, 5, note)
        r += 1
    s.text(r, 2, "Next unpaid dividend per share (12 months)")
    s.put(r, 4, 1.05, "price", "input", key="next_div")
    s.note(r, 5, "[sourced] consensus DPS FY2027E, CIQ 2026-09-20; the FY2026 dividend (MXN 1.075) was paid on 13-May-2026")
    r += 1
    for key, label, formula, fmt in [
        ("pub_tgt12_g", "12-month value - Gordon, published = value today x (1 + ke) - next unpaid dividend",
         "={pub_vps_g}*(1+" + s.a("ke") + ")-" + s.a("next_div"), "price"),
        ("pub_up12_g", "Upside to the 12-month value - Gordon", "={pub_tgt12_g}/" + s.a("price") + "-1", "pct"),
        ("pub_ronic", "Return on new capital implied by the published terminal value = g / (1 - FCFF_N+1 / NOPAT_N+1)",
         f'=IF({{pub_fcff_next}}<{nopat_N}*(1+{s.a("g")}),{s.a("g")}/(1-{{pub_fcff_next}}/({nopat_N}*(1+{s.a("g")}))),"n.m.")',
         "pct"),
        ("pub_impl_mult", "Exit multiple implied by the published Gordon value",
         f"={{pub_tv_g}}*(1+{s.a('wacc')})^{s.a('offset')}/{ebitda_N}", "mult"),
        ("pub_impl_g_x", "Growth implied by the exit multiple, published = (X x WACC - FCFF_N) / (X + FCFF_N - L_N)",
         f"=({s.a('tv_x')}/(1+{s.a('wacc')})^{s.a('offset')}*{s.a('wacc')}-{fcff_N})"
         f"/({s.a('tv_x')}/(1+{s.a('wacc')})^{s.a('offset')}+{fcff_N}-{lease_N})", "pct2"),
        ("lease_test_a", "Test (a): FY2035 growth in the lease liability / (g x lease liability FY2035)",
         f"={fL}{s.rows['ls_growth']}/({s.a('g')}*{lease_N})", "num3"),
    ]:
        formula = formula.format(**{k: s.a(k) for k in s.names if "{" + k + "}" in formula})
        s.text(r, 2, label)
        s.put(r, 4, formula, fmt, key=key)
        r += 1
    r += 1
    s.subheader(r, "Walk from the engine basis to the published values (MXN per share today)", LASTC)
    r += 1
    for key, label, formula in [
        ("walk_eng_g", "Engine basis - Gordon", "=" + s.a("vps_g")),
        ("walk_lease_g", "Lease growth (Gordon: explicit years and terminal value)", "={pub_vps_g}-" + s.a("vps_g")),
        ("walk_eng_x", "Engine basis - exit", "=" + s.a("vps_x")),
        ("walk_claims_x", "Peer-EV-consistent claims (add back the claims outside the peer EV)",
         "=" + s.a("vps_x_peer") + "-" + s.a("vps_x")),
        ("walk_lease_x", "Lease growth (exit: explicit years only)", "={pub_vps_x}-" + s.a("vps_x_peer")),
    ]:
        formula = formula.format(**{k: s.a(k) for k in s.names if "{" + k + "}" in formula})
        s.text(r, 2, label)
        s.put(r, 4, formula, "price", key=key)
        r += 1
    r += 1
    s.subheader(r, "Bridge judgements on the published basis (MXN per share today)", LASTC)
    r += 1
    for key, label, formula in [
        ("pub_g_ex_nci", "Gordon without the NCI uplift (NCI at book only)",
         "={pub_vps_g}+" + s.a("dl_nci_uplift") + per_share if "dl_nci_uplift" in s.names else "={pub_vps_g}"),
        ("pub_x_ex_nci", "Exit without the NCI uplift",
         "={pub_vps_x}+" + s.a("dl_nci_uplift") + per_share if "dl_nci_uplift" in s.names else "={pub_vps_x}"),
        ("pub_g_nci_g0", "Gordon with the NCI claim at zero growth (NCI income / ke)",
         "={pub_vps_g}+(" + s.a("nci_mkt") + "-" + s.a("nci_mkt_g0") + ")" + per_share if "nci_mkt" in s.names
         else "={pub_vps_g}"),
    ]:
        formula = formula.format(**{k: s.a(k) for k in s.names if "{" + k + "}" in formula})
        s.text(r, 2, label)
        s.put(r, 4, formula, "price", key=key)
        r += 1
    r += 1
    s.subheader(r, "Lower-return terminal cases - Gordon, published: TV = NOPAT_N+1 x (1 - g / RONIC) / (WACC - g)", LASTC)
    r += 1
    for j, hdr in enumerate(["RONIC", "TV", "Value/share", "vs published"]):
        s.text(r, 3 + j, hdr, font=F_BOLD)
    r += 1
    ronic_cases = [("RONIC = WACC (growth adds no value)", "=" + s.a("wacc")),
                   ("RONIC 12.5%", 0.125), ("RONIC 15.0%", 0.15),
                   ("RONIC implied by the published terminal value", "=" + s.a("pub_ronic"))]
    for idx, (label, rv) in enumerate(ronic_cases):
        s.text(r, 2, label)
        s.put(r, 3, rv, "pct", "input" if isinstance(rv, float) else "calc", key=f"ronic_r{idx}")
        tv_f = f"{nopat_N}*(1+{s.a('g')})*(1-{s.a('g')}/C{r})/({s.a('wacc')}-{s.a('g')})"
        s.put(r, 4, f"={tv_f}", "money")
        s.put(r, 5, f"=({s.a('pub_sum_pv')}+D{r}/(1+{s.a('wacc')})^({N}-{s.a('offset')}){nb}-{s.a('debt_like')})"
                    f"/{s.a('shares')}*{s.a('roll')}", "price", key=f"ronic_v{idx}")
        s.put(r, 6, f"=E{r}-{s.a('pub_vps_g')}", "price")
        r += 1
    s.note(r, 2, "The value-driver formula prices growth by its return: at RONIC = WACC growth adds nothing. The last row "
                 "reproduces the published value (check). Total terminal net reinvestment = 1 - FCFF_N+1 / NOPAT_N+1.")
    r += 2

    # --- sensitivity analysis
    s.header(r, "SENSITIVITY ANALYSIS - PUBLISHED BASIS (roll-forward factor and NCI claim held at the base cost of "
                "equity; the engine-basis grid is in the model workbook)", LASTC)
    r += 1
    w_steps = (-0.02, -0.01, 0.0, 0.01, 0.02)
    g_steps = (-0.01, -0.005, 0.0, 0.005, 0.01)
    m_steps = (-2.0, -1.0, 0.0, 1.0, 2.0)
    bridge_expr = f"+{s.a('non_op')}-{s.a('net_debt')}-{s.a('nci')}-{s.a('debt_like')}"
    bridge_x_expr = bridge_expr + f"+{s.a('claims_outside_peer')}"

    def grid(title: str, col_steps, col_base: str, col_fmt: str, cell_formula, cell_fmt: str):
        nonlocal r
        s.subheader(r, title, LASTC)
        r += 1
        s.text(r, 2, "WACC  \\  column variable")
        for j, st in enumerate(col_steps):
            s.put(r, 4 + j, f"={col_base}+({st})", col_fmt)
        hdr = r
        r += 1
        for i, st in enumerate(w_steps):
            s.put(r, 3, f"={s.a('wacc')}+({st})", "pct2")
            for j in range(len(col_steps)):
                wcell = f"$C{r}"
                ccell = f"{L(4 + j)}${hdr}"
                s.put(r, 4 + j, cell_formula(wcell, ccell), cell_fmt)
            r += 1
        r += 1

    npv = lambda wv: f"NPV({wv},{fcff_pub_rng})*(1+{wv})^{s.a('offset')}"  # noqa: E731
    grid("Value per share today - Gordon, published: WACC (rows) x terminal growth (columns); n.m. where WACC - g < 1%",
         g_steps, s.a("g"), "pct2",
         lambda wv, gv: (f'=IF({wv}-{gv}>=0.01,({npv(wv)}+({fcff_N}*(1+{gv})-{gv}*{lease_N})/({wv}-{gv})'
                         f'/(1+{wv})^({N}-{s.a("offset")}){bridge_expr})/{s.a("shares")}*{s.a("roll")},"n.m.")'), "price")
    grid("Value per share today - exit multiple, published (peer-EV claims): WACC (rows) x exit EV/EBITDA (columns)",
         m_steps, s.a("exit_mult"), "mult",
         lambda wv, mv: f"=({npv(wv)}+{ebitda_N}*{mv}/(1+{wv})^{N}{bridge_x_expr})/{s.a('shares')}*{s.a('roll')}",
         "price")
    grid("Enterprise value - exit multiple, published: WACC (rows) x exit EV/EBITDA (columns)",
         m_steps, s.a("exit_mult"), "mult",
         lambda wv, mv: f"={npv(wv)}+{ebitda_N}*{mv}/(1+{wv})^{N}", "money")
    grid("Implied perpetual growth of the exit value, published: WACC (rows) x exit EV/EBITDA (columns)",
         m_steps, s.a("exit_mult"), "mult",
         lambda wv, mv: (f"=({ebitda_N}*{mv}/(1+{wv})^{s.a('offset')}*{wv}-{fcff_N})"
                         f"/({ebitda_N}*{mv}/(1+{wv})^{s.a('offset')}+{fcff_N}-{lease_N})"), "pct2")
    grid("Premium / (discount) to the current price - exit multiple, published: WACC (rows) x exit EV/EBITDA (columns)",
         m_steps, s.a("exit_mult"), "mult",
         lambda wv, mv: (f"=({npv(wv)}+{ebitda_N}*{mv}/(1+{wv})^{N}{bridge_x_expr})/{s.a('shares')}*{s.a('roll')}"
                         f"/{s.a('price')}-1"), "pct")

    # --- key outputs block (top)
    ko = key_out_row
    for i, (label, ref, fmt) in enumerate([
        ("PUBLISHED value per share today - Gordon (MXN)", s.a("pub_vps_g"), "price"),
        ("PUBLISHED value per share today - exit multiple, peer-EV-consistent claims (MXN)", s.a("pub_vps_x"), "price"),
        ("PUBLISHED 12-month value - Gordon = value x (1 + ke) - next unpaid dividend (MXN)", s.a("pub_tgt12_g"), "price"),
        ("Current share price (MXN)", s.a("price"), "price"),
        ("Premium / (discount) today - Gordon / exit", None, "pct"),
        ("WACC / cost of equity", None, "pct2"),
        ("Terminal growth / exit multiple", None, "pct2"),
        ("Enterprise value - Gordon / exit, published (MXN m)", None, "money"),
        ("Terminal value share of EV - Gordon / exit, published", None, "pct"),
        ("Return on new capital implied by the terminal value - published / engine basis", None, "pct"),
        ("Gordon at RONIC = WACC / at RONIC 12.5% (lower-return cases, MXN)", None, "price"),
        ("Engine basis, reconciliation only - Gordon / exit (MXN)", None, "price"),
        ("Team target price", None, "price"),
    ]):
        rr = ko + i
        s.text(rr, 2, label, font=F_BOLD if i < 3 else F_NORM)
        if ref:
            s.put(rr, 4, f"={ref}", fmt, bold=i < 3)
            ws.cell(row=rr, column=4).fill = FILL_KEY
    pairs = {
        ko + 4: (s.a("pub_prem_g"), s.a("pub_prem_x"), "pct", "pct"),
        ko + 5: (s.a("wacc"), s.a("ke"), "pct2", "pct2"),
        ko + 6: (s.a("g"), s.a("exit_mult"), "pct2", "mult"),
        ko + 7: (s.a("pub_ev_g"), s.a("pub_ev_x"), "money", "money"),
        ko + 8: (s.a("pub_tv_share_g"), s.a("pub_tv_share_x"), "pct", "pct"),
        ko + 9: (s.a("pub_ronic"), s.a("tv_ronic"), "pct", "pct"),
        ko + 10: (s.a("ronic_v0"), s.a("ronic_v1"), "price", "price"),
        ko + 11: (s.a("vps_g"), s.a("vps_x"), "price", "price"),
    }
    for rr, (a1, a2, f1, f2) in pairs.items():
        s.put(rr, 4, f"={a1}", f1)
        s.put(rr, 5, f"={a2}", f2)
    s.text(ko + 12, 4, "null - team decision (CFA rules: the team sets the target and rating)", font=F_NOTE)
    ws.freeze_panes = "C3"

    # =================================================================== other sheets
    rev_names = _reverse_dcf(wb, ctx, s, fcff_N, N, lease_N)
    _levers(wb, ctx)
    _scenarios(wb, ctx)
    _dated_check(wb, ctx, s, col, fcst)
    _engine_recon(wb, ctx, s, col, fcst, rev_names)
    _inputs_sources(wb, ctx)
    _checks(wb, ctx, s, col, hist, fcst)
    _cover_dcf(cover, ctx)
    wb.save(out)


def _reverse_dcf(wb, ctx: Ctx, s: Sheet, fcff_N: str, N: int, lease_N: str) -> None:
    ws = wb.create_sheet("Reverse_DCF")
    t = Sheet(ws)
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 80
    ws.sheet_view.showGridLines = False
    D = "'DCF'!"
    t.text(2, 2, "Reverse Discounted Cash Flow - what the market price implies", font=F_TITLE)
    t.text(3, 2, "Layout after the S&P 'Reverse_DCF' tab. The template imputes growth with a VBA macro; here the "
                 "procedure is Excel Goal Seek on a single input cell (no macros).", font=F_SUB)
    r = 5
    t.header(r, "1. Terminal growth implied by the market price - PUBLISHED basis (headline)", 5)
    r += 1
    for key, label, formula, fmt in [
        ("pmkt_ev", "Market EV at FYE = price / roll-forward x shares + net debt + NCI + all debt-like - non-op",
         f"={D}{s.a('price')}/{D}{s.a('roll')}*{D}{s.a('shares')}+{D}{s.a('net_debt')}+{D}{s.a('nci')}"
         f"+{D}{s.a('debt_like')}-{D}{s.a('non_op')}", "money"),
        ("psum_pv", "PV of explicit FCFF, published basis", f"={D}{s.a('pub_sum_pv')}", "money"),
        ("pimpl_tv", "Terminal value implied by the market (at the Gordon discount date)",
         f"=({{pmkt_ev}}-{{psum_pv}})*(1+{D}{s.a('wacc')})^({N}-{D}{s.a('offset')})", "money"),
        ("pimpl_g", "Terminal growth implied by the market price = (TV x WACC - FCFF_N) / (TV + FCFF_N - L_N)",
         f'=IF({{pmkt_ev}}>{{psum_pv}},({{pimpl_tv}}*{D}{s.a("wacc")}-{D}{fcff_N})/({{pimpl_tv}}+{D}{fcff_N}-{D}{lease_N}),'
         f'"n.m.")', "pct2"),
        ("pbase_g", "Terminal growth used in the DCF", f"={D}{s.a('g')}", "pct2"),
    ]:
        formula = formula.format(**{k: t.a(k) for k in t.names if "{" + k + "}" in formula})
        t.text(r, 2, label)
        t.put(r, 4, formula, fmt, key=key)
        r += 1
    r += 1
    t.header(r, "1b. Terminal growth implied by the market price - engine basis (reconciliation with the engine)", 5)
    r += 1
    rows = [
        ("mkt_ev", "Market EV at FYE = price / roll-forward x shares + net debt + NCI + debt-like - non-op",
         f"={D}{s.a('price')}/{D}{s.a('roll')}*{D}{s.a('shares')}+{D}{s.a('net_debt')}+{D}{s.a('nci')}"
         f"+{D}{s.a('debt_like')}-{D}{s.a('non_op')}", "money"),
        ("sum_pv", "PV of explicit FCFF (from the DCF)", f"={D}{s.a('sum_pv')}", "money"),
        ("impl_tv", "Terminal value implied by the market (at the Gordon discount date)",
         f"=({{mkt_ev}}-{{sum_pv}})*(1+{D}{s.a('wacc')})^({N}-{D}{s.a('offset')})", "money"),
        ("impl_g", "Terminal growth implied by the market price",
         f'=IF({{mkt_ev}}>{{sum_pv}},({{impl_tv}}*{D}{s.a("wacc")}-{D}{fcff_N})/({{impl_tv}}+{D}{fcff_N}),"n.m.")', "pct2"),
        ("base_g", "Terminal growth used in the DCF", f"={D}{s.a('g')}", "pct2"),
    ]
    for key, label, formula, fmt in rows:
        formula = formula.format(**{k: t.a(k) for k in t.names if "{" + k + "}" in formula})
        t.text(r, 2, label)
        t.put(r, 4, formula, fmt, kind="link" if formula.startswith(f"={D}") and "(" not in formula[1:4] else "calc", key=key)
        r += 1
    t.note(r, 2, "Guard: defined only when the market EV exceeds the PV of the explicit cash flows.")
    r += 2
    t.header(r, "2. Revenue growth shift implied by the market price (all other drivers fixed)", 5)
    r += 1
    sol_g = sol_x = sol_ge = None
    try:
        sol_g = core.solve_growth_shift(ctx.project, ctx.case.price, "gordon", basis="published")
        sol_x = core.solve_growth_shift(ctx.project, ctx.case.price, "exit", basis="published")
        sol_ge = core.solve_growth_shift(ctx.project, ctx.case.price, "gordon")
    except Exception as exc:  # pragma: no cover - reference values only
        print(f"reverse DCF solve skipped: {exc}")
    ctx.rev_solutions = {"gordon_pub": sol_g, "exit_pub": sol_x, "gordon_engine": sol_ge}
    for label, v, fmt, note in [
        ("Procedure", None, None, "Data > What-If Analysis > Goal Seek: set 'DCF'!" + s.a("pub_vps_g").replace("$", "")
         + " (published Gordon value today) to the current price by changing 'DCF'!" + s.a("g_shift").replace("$", "")
         + ". Reset the shift to 0 afterwards."),
        ("Reference solution - Gordon, published (Python bisection, valuation/build/dcf_core.py)", sol_g, "pct2",
         "Uniform shift added to every forecast year's growth driver; n.a. if no solution in -10%..+10%."),
        ("Reference solution - exit multiple, published", sol_x, "pct2", ""),
        ("Reference solution - Gordon, engine basis", sol_ge, "pct2", "Reconciliation only."),
    ]:
        t.text(r, 2, label)
        if v is not None:
            t.put(r, 4, v, fmt, "input")
        elif fmt:
            t.text(r, 4, "n.a.")
        t.note(r, 5, note)
        r += 1
    return {k: f"'Reverse_DCF'!{v}" for k, v in t.names.items()}


def _levers(wb, ctx: Ctx) -> None:
    """Static lever grids computed with the Python replica (rebuild the workbook to refresh)."""
    ws = wb.create_sheet("Levers")
    t = Sheet(ws)
    ws.column_dimensions["B"].width = 44
    for cc in "CDEFGH":
        ws.column_dimensions[cc].width = 13
    ws.sheet_view.showGridLines = False
    t.text(2, 2, "Story levers - published value per share today (Gordon and exit), MXN", font=F_TITLE)
    t.text(3, 2, "STATIC values computed by valuation/build/dcf_core.py from the same inputs (published basis: leases "
                 "grow with revenue; exit on peer-EV claims); rebuild to refresh. Live single-scenario levers: the shift "
                 "inputs on the 'DCF' sheet.", font=F_SUB)
    r = 5
    steps = (-0.01, -0.005, 0.0, 0.005, 0.01)
    for method, key in (("Gordon", "price_gordon"), ("exit multiple", "price_exit")):
        t.header(r, f"Value per share today - {method}: capex/revenue shift (rows) x gross margin shift (columns)", 8)
        r += 1
        t.text(r, 2, "Capex shift (rows) / GM shift (columns)")
        for j, gm in enumerate(steps):
            t.put(r, 3 + j, gm, "pct2", "input")
        r += 1
        for cx in steps:
            t.put(r, 2, cx, "pct2", "input")
            for j, gm in enumerate(steps):
                v = core.value(core.build_case(ctx.project, gm_shift=gm, capex_shift=cx), basis="published")
                t.put(r, 3 + j, v[key], "price")
            r += 1
        r += 1
    t.header(r, "Terminal economics by capex shift (Gordon, published)", 8)
    r += 1
    for j, hdr in enumerate(["Capex shift", "Value / share", "RONIC implied by TV", "Terminal net reinvestment / NOPAT",
                             "FY2035 FCFF margin"]):
        t.text(r, 2 + j, hdr, font=F_BOLD)
    r += 1
    for cx in steps:
        v = core.value(core.build_case(ctx.project, capex_shift=cx), basis="published")
        f = v["lines"]
        t.put(r, 2, cx, "pct2", "input")
        t.put(r, 3, v["price_gordon"], "price")
        t.put(r, 4, v["tv_ronic"] if v["tv_ronic"] is not None else "n.m.", "pct")
        t.put(r, 5, 1 - v["fcff_next"] / v["nopat_next"], "pct")
        t.put(r, 6, f["fcff_pub"][-1] / f["revenue"][-1], "pct")
        r += 1
    r += 1
    t.header(r, "Revenue growth shift (all segments, all years) - published value per share today", 8)
    r += 1
    for j, hdr in enumerate(["Growth shift", "Gordon", "Exit"]):
        t.text(r, 2 + j, hdr, font=F_BOLD)
    r += 1
    for gs in steps:
        v = core.value(core.build_case(ctx.project, growth_shift=gs), basis="published")
        t.put(r, 2, gs, "pct2", "input")
        t.put(r, 3, v["price_gordon"], "price")
        t.put(r, 4, v["price_exit"], "price")
        r += 1


def _crp_by_basis(ctx: Ctx) -> dict[str, float]:
    """Company CRP for each weight basis from the one shared country -> region -> company build (dcf_core)."""
    return core.crp_from_countries(ctx.extras)["company"]


def _peer_results(ctx: Ctx) -> dict:
    """Verified peer results from the RECALCULATED peer workbook v3 (cached values), checked against the calculator
    spec's expected medians. A missing workbook or missing caches stop the build instead of dropping scenarios."""
    from openpyxl import load_workbook
    p = ctx.project / "valuation" / "BIMBOA_PeerAnalysis_v3.xlsx"
    if not p.is_file():
        raise SystemExit("valuation/BIMBOA_PeerAnalysis_v3.xlsx is missing: build and recalculate the peer workbook first")
    wbp = load_workbook(p, data_only=True)
    dp = wbp["Detail_PeerAnalysis"]
    out = {"core_ev_ltm": dp["B50"].value, "core_ev_ntm": dp["C50"].value, "all_ev_ltm": dp["B53"].value,
           "all_ev_ntm": dp["C53"].value}
    for k, v in out.items():
        if not isinstance(v, (int, float)):
            raise SystemExit(f"peer workbook v3 has no cached value for {k}: recalculate it in Excel first")
    spec_p = ctx.project / "valuation" / "inputs" / "peer_extras_v3.yaml"
    if not spec_p.is_file():
        raise SystemExit("valuation/inputs/peer_extras_v3.yaml (independent spec) is missing")
    exp = ((core.load_yaml(spec_p).get("expected") or {}).get("medians") or {})
    for k, v in out.items():
        e = exp.get(k)
        if e is None:
            raise SystemExit(f"the independent spec has no expected value for {k}")
        if abs(v - float(e)) > 1e-6 * max(1.0, abs(float(e))):
            raise SystemExit(f"peer median {k}: workbook {v} differs from the independent spec {e}")
    pa = wbp["PeerAnalysis"]
    for row in range(1, pa.max_row + 1):
        if str(pa.cell(row, 2).value or "").startswith("Peer median (excluding Bimbo)"):
            out["capint_peer"] = [pa.cell(row, c).value for c in (6, 7, 8)]
            out["capint_bimbo"] = [pa.cell(row + 1, c).value for c in (6, 7, 8)]
    for k in ("capint_peer", "capint_bimbo"):
        vals = out.get(k)
        if not vals or not all(isinstance(x, (int, float)) for x in vals):
            raise SystemExit(f"peer workbook v3 capital-intensity row {k} missing or not recalculated: {vals!r}")
    return out


def _rec(d: dict, key: str) -> float:
    """Value of a source record ({value, tag, source}) or a bare number; finite and present, else stop the build."""
    import math
    x = d.get(key)
    x = x.get("value") if isinstance(x, dict) else x
    if not isinstance(x, (int, float)) or not math.isfinite(float(x)):
        raise SystemExit(f"input {key!r} is missing or not a finite number: {x!r}")
    return float(x)


def _bimbo_multiple(ctx: Ctx, ke: float) -> float:
    """Bimbo's own LTM EV/EBITDA on the peer-EV claim basis at the Jun-26 bridge: (market cap + debt + leases - cash +
    NCI at market - associates) / LTM company Adj. EBITDA 60,919 (Capital IQ quarters 3Q25-2Q26). NCI at market comes
    from the shared, guarded dcf_core.nci_claim (same value as the DCF bridge)."""
    dbj = ctx.extras.get("dated_bridge") or {}
    nci_mkt = core.nci_claim(ctx.case, ke)["market"]
    claims = (_rec(dbj, "debt_jun") + _rec(dbj, "leases_jun") - _rec(dbj, "cash_jun") + nci_mkt - _rec(dbj, "non_op_jun"))
    shares = _rec(dbj, "shares_out")
    if shares <= 0:
        raise SystemExit("shares_out must be positive")
    return (ctx.case.price * shares + claims) / 60919.0


def _scenarios(wb, ctx: Ctx) -> None:
    """Scenario table on the published basis (static values from valuation/build/dcf_core.py; rebuild to refresh).
    One change at a time against the base, so each row isolates one debate for the team."""
    import dataclasses
    ws = wb.create_sheet("Scenarios")
    t = Sheet(ws)
    ws.sheet_view.showGridLines = False
    for k_, v_ in {"B": 64, "C": 11, "D": 11, "E": 11, "F": 11, "G": 11, "H": 11, "I": 11, "J": 11, "K": 70}.items():
        ws.column_dimensions[k_].width = v_
    t.text(2, 2, "Scenarios - published value per share today (MXN), one change at a time against the base", font=F_TITLE)
    t.text(3, 2, "STATIC values computed by valuation/build/dcf_core.py from the same inputs as the DCF sheet (published "
                 "basis); rebuild to refresh. The base is the assistant's proposal; every other row is an alternative the "
                 "team can adopt. Target price and rating stay with the team.", font=F_SUB)
    c0 = ctx.case
    va = ctx.extras.get("wacc_variants") or {}
    rf_rows = va.get("rf") or {}
    mbono = float((rf_rows.get("mbono10") or {}).get("value", 0.0926))
    dspread = float((rf_rows.get("def_spread") or {}).get("value", 0.0162))
    crp = _crp_by_basis(ctx)
    price = c0.price

    def run(case):
        v = core.value(case, basis="published")
        return v

    def with_wacc(**kw):
        w = dict(c0.wacc_in)
        w.update(kw)
        return dataclasses.replace(c0, wacc_in=w)

    def with_drivers(changes: dict, **kw):
        d = {k: list(v) for k, v in c0.drivers.items()}
        for k, fn in changes.items():
            d[k] = [fn(i, x) for i, x in enumerate(d[k])]
        return dataclasses.replace(c0, drivers=d, **kw)

    # North America FX: drivers.yaml builds the FY2027+ 3.0% growth in MXN as ~1.5% USD growth + ~1.5% a year of peso
    # depreciation. The PPP rows REPLACE that 1.5% with the inflation-differential depreciation (no double count),
    # recompose g with the scenario's own FY2035 North America revenue share, and move the US-linked NCI growth too.
    na_key = "north_america"
    base_dep = 0.015
    dep_target = 1.03 / 1.02 - 1           # Banxico 3% target vs Fed 2% target (same 3% as the 5.0% g cap)
    dep_survey = 1.037 / 1.02 - 1          # Banxico survey long-run inflation 3.70% (2031-34 median) vs 2%
    na_path = c0.drivers[f"seg_{na_key}_growth"]

    def ppp(dep):
        conv = lambda x: (1 + x) / (1 + base_dep) * (1 + dep) - 1  # noqa: E731
        d = {k: list(v) for k, v in c0.drivers.items()}
        d[f"seg_{na_key}_growth"] = [x if i == 0 else conv(x) for i, x in enumerate(na_path)]
        case = dataclasses.replace(c0, drivers=d)
        na = c0.h("seg_" + na_key)
        for x in d[f"seg_{na_key}_growth"]:
            na *= 1 + x
        w_na = na / core.forecast(case)["revenue"][-1]
        g_new = c0.g + w_na * (conv(na_path[-1]) - na_path[-1])
        extra = {**c0.extra, "nci_claim": {**(c0.extra.get("nci_claim") or {}),
                                           "growth": conv(float((c0.extra.get("nci_claim") or {}).get("growth", 0.03)))}}
        return dataclasses.replace(case, g=g_new, extra=extra), w_na, conv(na_path[-1])

    ppp_t, w_t, gna_t = ppp(dep_target)
    ppp_s, w_s, gna_s = ppp(dep_survey)
    peers = _peer_results(ctx)
    peer_capex = (peers.get("capint_peer") or [None, None, None])[-1]
    bimbo_mult = _bimbo_multiple(ctx, core.wacc(c0)["ke"])

    def capex_tail(level):
        return with_drivers({"capex_pct_revenue": (lambda i, x: x if i < 5 else level)})

    groups = [
        ("BASE", [("Base: route B, CRP weighted by FY2025 Adj. EBITDA (additive); proposals for drivers and terminal",
                   c0, "valuation.yaml / drivers.yaml")]),
        ("COST OF CAPITAL", [
            ("Route A local: M-Bono 10y as risk-free, no country risk premium", with_wacc(risk_free=mbono,
             country_risk_premium=0.0), "project guide's preferred local route"),
            ("Route B, CRP weighted by FY2025 sales", with_wacc(risk_free=mbono - dspread,
             country_risk_premium=crp["revenue_fy25"]), f"CRP {crp['revenue_fy25']:.2%}"),
            ("Route B, CRP weighted by FY2023-25 average Adj. EBITDA", with_wacc(risk_free=mbono - dspread,
             country_risk_premium=crp["ebitda_3y"]), f"CRP {crp['ebitda_3y']:.2%}"),
            ("Route B, CRP weighted by FY2023-25 average sales", with_wacc(risk_free=mbono - dspread,
             country_risk_premium=crp["revenue_3y"]), f"CRP {crp['revenue_3y']:.2%}"),
            ("Route B, EBITDA weights, Argentina excluded", with_wacc(risk_free=mbono - dspread,
             country_risk_premium=crp["ebitda_fy25_xar"]), f"CRP {crp['ebitda_fy25_xar']:.2%}"),
            ("Route B, CRP scaled by beta: ke = rf + beta_l x (ERP + CRP)", with_wacc(
                equity_risk_premium=c0.wacc_in["equity_risk_premium"] + c0.wacc_in["country_risk_premium"],
                country_risk_premium=0.0), "Damodaran operations-weighted ERP"),
            ("Beta: Bimbo 5-year regression vs IPC, unlevered", with_wacc(beta_unlevered=_beta_case(ctx, "unlevered")),
             "valuation/inputs/beta_regression.json"),
            ("Beta: same, Blume-adjusted", with_wacc(beta_unlevered=_beta_case(ctx, "blume_unlevered")), "0.67 x b + 0.33"),
            ("Beta: Bimbo 3-year regression, unlevered", with_wacc(beta_unlevered=_beta_case(ctx, "unlevered", "total_return_36m")),
             "2023-24 de-rating window"),
        ]),
        ("NORTH AMERICA FX ASSUMPTION (base: ~1.5% a year peso depreciation inside the 3.0% MXN growth)", [
            ("PPP at inflation targets (3% MX vs 2% US): depreciation 0.98% a year replaces 1.5%", ppp_t,
             f"NA growth FY2027+ {gna_t:.2%}; NA share FY2035 {w_t:.1%}; g {ppp_t.g:.2%}; NCI growth moved too"),
            ("PPP at Banxico survey inflation (3.70% vs 2%): depreciation 1.67% a year replaces 1.5%", ppp_s,
             f"NA growth FY2027+ {gna_s:.2%}; NA share FY2035 {w_s:.1%}; g {ppp_s.g:.2%}"),
        ]),
        ("CAPEX SENSITIVITY (FY2031-35 capex / revenue, which also sets the terminal value; base path 5.6% to 6.0%)", [
            ("Capex held at 5.5% (the FY2030 level)", capex_tail(0.055), "no further ramp"),
            ("Capex 5.0% (Bimbo consensus FY2027-28: about 5.0%)", capex_tail(0.050),
             "near-term consensus, not a perpetual reinvestment estimate"),
            ("Capex 4.5%", capex_tail(0.045), "near FY2025 D&A excl. right-of-use (4.4%)"),
        ] + ([(f"Capex at the peer consensus median FY3 ({float(peer_capex):.1%})", capex_tail(float(peer_capex)),
               "peers' FY3 consensus; RONIC shows whether it is sustainable")] if isinstance(peer_capex, (int, float)) else [])),
        ("EXIT MULTIPLE (exit method only)", [
            (f"Core bakery median EV/EBITDA LTM ({peers['core_ev_ltm']:.2f}x)",
             dataclasses.replace(c0, exit_multiple=float(peers["core_ev_ltm"])), "BIMBOA_PeerAnalysis_v3 (recalculated)"),
            (f"All eligible peers median EV/EBITDA LTM ({peers['all_ev_ltm']:.2f}x)",
             dataclasses.replace(c0, exit_multiple=float(peers["all_ev_ltm"])), "BIMBOA_PeerAnalysis_v3 (recalculated)"),
            (f"Bimbo's own LTM multiple on the same claim basis ({bimbo_mult:.2f}x)",
             dataclasses.replace(c0, exit_multiple=bimbo_mult), "Jun-26 bridge, NCI at market at the base ke"),
        ]),
    ]
    r = 5
    hdr = ["Scenario", "WACC", "ke", "g", "Gordon", "Exit", "vs price (G)", "vs price (X)", "RONIC (TV)", "What changes"]
    base_cells = {}
    for gname, rows in groups:
        t.header(r, gname, 11)
        r += 1
        for j, h in enumerate(hdr):
            t.text(r, 2 + j, h, font=F_BOLD)
        r += 1
        for label, case, note in rows:
            v = run(case)
            t.text(r, 2, label)
            t.put(r, 3, v["wacc"], "pct2")
            t.put(r, 4, v["ke"], "pct2")
            t.put(r, 5, case.g, "pct2")
            t.put(r, 6, v["price_gordon"], "price", key=None)
            t.put(r, 7, v["price_exit"], "price")
            t.put(r, 8, v["price_gordon"] / price - 1, "pct")
            t.put(r, 9, v["price_exit"] / price - 1, "pct")
            t.put(r, 10, v["tv_ronic"] if v["tv_ronic"] is not None else "n.m.", "pct")
            t.note(r, 11, note)
            if gname == "BASE":
                base_cells = {"g": f"'Scenarios'!$F${r}", "x": f"'Scenarios'!$G${r}"}
            r += 1
        r += 1
    t.note(r, 2, f"Price MXN {price:.2f} (BMV close 16-Oct-2026). RONIC = return on new capital implied by the Gordon "
                 "terminal value; 'n.m.' when terminal reinvestment is zero or negative. Capex rows are sensitivities: "
                 "lower capex with D&A held at its path raises the implied RONIC; a sustainable terminal needs capex, "
                 "non-lease depreciation, lease investment and working capital reconciled (see the lower-return table on "
                 "the DCF sheet for the value-driver cross-check).")
    t.note(r + 1, 2, "The company's impairment-test discount rates (RA25 p.265) are post-tax book rates in five "
                     "currencies; they appear only as context in the WACC workbook cross-checks, with no implied value.")
    t.note(r + 2, 2, "The Gordon terminal FCFF grows the FY2035 FCFF by g (engine convention), so the FY2035 working-"
                     "capital movement carries into the terminal year; a rebuilt steady-state terminal year is not used.")
    wb._scen_cells = base_cells


def _beta_case(ctx: Ctx, field: str, case: str = "total_return_60m") -> float:
    import json as _json
    p = ctx.project / "valuation" / "inputs" / "beta_regression.json"
    d = _json.loads(p.read_text(encoding="utf-8"))
    return float(d["cases"][case][field])


def _dated_check(wb, ctx: Ctx, s: Sheet, col: dict, fcst: list[int]) -> None:
    ws = wb.create_sheet("Dated_Check")
    t = Sheet(ws)
    ws.column_dimensions["B"].width = 72
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 90
    ws.sheet_view.showGridLines = False
    D = "'DCF'!"
    db = ctx.extras.get("dated_bridge") or {}
    t.text(2, 2, "Valuation-date cross-check (anchored at the latest balance sheet, 30-Jun-2026)", font=F_TITLE)
    t.text(3, 2, "Independent of the engine's FYE roll-forward. Still a roll-forward approximation from 30-Jun to the "
                 "valuation date, not an observed October balance sheet.", font=F_SUB)
    r = 5
    t.header(r, "1H26 actual FCFF (same definition as the forecast FCFF)", 5)
    r += 1

    def inp(key, label, default=0.0, fmt="money"):
        nonlocal r
        item = db.get(key) or {}
        v = item.get("value", default) if isinstance(item, dict) else item
        t.text(r, 2, label)
        t.put(r, 4, float(v) if v is not None else 0.0, fmt, "input", key=key)
        t.note(r, 5, (f"[{item.get('tag', 'unverified')}] {item.get('source', '')}" if isinstance(item, dict)
                      else "[unverified]"))
        r += 1

    inp("ebit_1h", "EBIT 1H26 (operating income)")
    inp("one_off_1h", "Less: one-off gain in 1H26 EBIT (La Moderna stake sale, ~50 bp of 1Q26 sales)")
    t.text(r, 2, "Tax rate (FY2026 forecast driver)")
    t.put(r, 4, f"={D}{L(col[fcst[0]])}{s.rows['drv_tax_rate']}", "pct", "link", key="tax_1h")
    r += 1
    inp("da_1h", "D&A 1H26")
    inp("capex_1h", "Capex 1H26 (positive)")
    inp("lease_1h", "Lease principal paid 1H26 (positive)")
    inp("dnwc_1h", "Change in operating NWC 1H26, cash-flow sign (Dec-25 to Jun-26 balance sheets; income taxes excluded)")
    t.text(r, 2, "FCFF 1H26 = (EBIT - one-off) x (1 - tax) + D&A - capex - lease principal + change in NWC", font=F_BOLD)
    t.put(r, 4, f"=({t.a('ebit_1h')}-{t.a('one_off_1h')})*(1-{t.a('tax_1h')})+{t.a('da_1h')}-{t.a('capex_1h')}-{t.a('lease_1h')}+{t.a('dnwc_1h')}",
          "money", key="fcff_1h", bold=True)
    r += 1
    t.text(r, 2, "FY2026E FCFF, published basis (DCF)")
    t.put(r, 4, f"={D}{L(col[fcst[0]])}{s.rows['fcff_pub']}", "money", "link", key="fcff_fy26")
    r += 1
    inp("one_off_in_fy26", "FY2026E forecast includes the 1H26 one-off gain? (1 = yes, 0 = no)", 1.0, "num")
    t.text(r, 2, "FY2026E FCFF on the same one-off basis as 1H26 = FY2026E - flag x one-off x (1 - tax)")
    t.put(r, 4, f"={t.a('fcff_fy26')}-{t.a('one_off_in_fy26')}*{t.a('one_off_1h')}*(1-{t.a('tax_1h')})", "money",
          key="fcff_fy26_same")
    r += 1
    t.text(r, 2, "2H26E FCFF = FY2026E (same basis) - 1H26 actual", font=F_BOLD)
    t.put(r, 4, f"={t.a('fcff_fy26_same')}-{t.a('fcff_1h')}", "money", key="fcff_2h", bold=True)
    t.note(r, 5, "Both halves exclude the gain: the La Moderna proceeds are already in the 30-Jun-2026 cash. The 1H26 "
                 "lease line is principal; the FY2026 growth in the lease book (published basis) falls in 2H26.")
    r += 2
    t.header(r, "Present value at 30-Jun-2026 (mid-period convention, same WACC)", 5)
    r += 1
    t.text(r, 2, "Discount time of 2H26 = 0.5 x (1 - offset)")
    t.put(r, 4, f"=0.5*(1-{D}{s.a('offset')})", "num3", key="t_2h")
    r += 1
    t.text(r, 2, "PV of 2H26E FCFF")
    t.put(r, 4, f"={t.a('fcff_2h')}/(1+{D}{s.a('wacc')})^{t.a('t_2h')}", "money", key="pv_2h")
    r += 1
    pv_terms = []
    for yy in fcst[1:]:
        k = f"pv_{yy}"
        t.text(r, 2, f"PV of FY{yy}E FCFF, published (time = {yy - fcst[0]} + 0.5 - offset)")
        t.put(r, 4, f"={D}{L(col[yy])}{s.rows['fcff_pub']}/(1+{D}{s.a('wacc')})^({yy - fcst[0]}+0.5-{D}{s.a('offset')})",
              "money", key=k)
        pv_terms.append(t.a(k))
        r += 1
    n_years = fcst[-1] - fcst[0]
    fL = L(col[fcst[-1]])
    t.text(r, 2, "PV of terminal value - Gordon, published (at the final cash flow's date)")
    t.put(r, 4, f"={D}{s.a('pub_tv_g')}/(1+{D}{s.a('wacc')})^({n_years}+0.5-{D}{s.a('offset')})", "money",
          key="pv_tvg")
    r += 1
    t.text(r, 2, "PV of terminal value - exit (year-end FY2035)")
    t.put(r, 4, f"={D}{s.a('tv_x')}/(1+{D}{s.a('wacc')})^({n_years}+0.5)", "money", key="pv_tvx")
    r += 1
    t.text(r, 2, "EV at 30-Jun-2026 - Gordon", font=F_BOLD)
    t.put(r, 4, f"={t.a('pv_2h')}+{'+'.join(pv_terms)}+{t.a('pv_tvg')}", "money", key="ev_jun_g", bold=True)
    r += 1
    t.text(r, 2, "EV at 30-Jun-2026 - exit", font=F_BOLD)
    t.put(r, 4, f"={t.a('pv_2h')}+{'+'.join(pv_terms)}+{t.a('pv_tvx')}", "money", key="ev_jun_x", bold=True)
    r += 2
    t.header(r, "Bridge at 30-Jun-2026 and roll to the valuation date", 5)
    r += 1
    inp("debt_jun", "Total debt ex-leases, 30-Jun-2026")
    inp("leases_jun", "Lease liabilities, 30-Jun-2026")
    inp("cash_jun", "Cash and equivalents, 30-Jun-2026")
    inp("nci_jun", "Non-controlling interests (book), 30-Jun-2026")
    t.text(r, 2, "NCI market-value uplift, 30-Jun-2026 = NCI at market value ('DCF', live at ke) - June book")
    if "nci_mkt" in s.names:
        t.put(r, 4, f"={D}{s.a('nci_mkt')}-{t.a('nci_jun')}", "money", key="nci_up_jun")
    else:
        t.put(r, 4, f"={D}{s.a('dl_nci_uplift')}", "money", "link", key="nci_up_jun")
    r += 1
    inp("debt_like_jun_ex_nci", "Other debt-like items, 30-Jun-2026 (same policy as the Dec-25 bridge, excl. NCI)")
    inp("tax_seasonal_jun", "Seasonal income taxes payable at 30-Jun-2026 (June less Dec-25 level; paid in 2H26)")
    inp("non_op_jun", "Non-operating assets, 30-Jun-2026 (same policy as valuation.yaml)")
    t.text(r, 2, "Equity at 30-Jun-2026 - Gordon (all claims)")
    eq_x = (f"-({t.a('debt_jun')}+{t.a('leases_jun')}-{t.a('cash_jun')})-{t.a('tax_seasonal_jun')}-{t.a('nci_jun')}"
            f"-{t.a('nci_up_jun')}+{t.a('non_op_jun')}")
    t.put(r, 4, f"={t.a('ev_jun_g')}{eq_x}-{t.a('debt_like_jun_ex_nci')}", "money", key="eq_jun_g")
    r += 1
    t.text(r, 2, "Equity at 30-Jun-2026 - exit (peer-EV-consistent claims: other debt-like items added back)")
    t.put(r, 4, f"={t.a('ev_jun_x')}{eq_x}", "money", key="eq_jun_x")
    r += 1
    inp("days_roll", "Days from 30-Jun-2026 to the valuation date", 108.0, "num")
    t.text(r, 2, "Roll factor = (1 + ke) ^ (days / 365)")
    t.put(r, 4, f"=(1+{D}{s.a('ke')})^({t.a('days_roll')}/365)", "num3", key="roll_d")
    r += 1
    inp("dist_after", "Distributions paid between 30-Jun and the valuation date (dividends + buybacks)")
    inp("shares_out", "Shares outstanding at the valuation date, millions", 0.0, "shares")
    t.text(r, 2, "Value per share at the valuation date - Gordon", font=F_BOLD)
    t.put(r, 4, f"=({t.a('eq_jun_g')}*{t.a('roll_d')}-{t.a('dist_after')})/{t.a('shares_out')}", "price",
          key="vps_d_g", bold=True)
    r += 1
    t.text(r, 2, "Value per share at the valuation date - exit", font=F_BOLD)
    t.put(r, 4, f"=({t.a('eq_jun_x')}*{t.a('roll_d')}-{t.a('dist_after')})/{t.a('shares_out')}", "price",
          key="vps_d_x", bold=True)
    r += 2
    t.header(r, "Reconciliation with the FYE roll-forward (published basis, 'DCF')", 5)
    r += 1
    for key, label, formula, fmt in [
        ("rec_fye_g", "FYE roll-forward value today - Gordon, published ('DCF')", f"={D}{s.a('pub_vps_g')}", "price"),
        ("rec_d_g", "Dated value - Gordon", f"={t.a('vps_d_g')}", "price"),
        ("rec_diff_g", "Difference - Gordon (dated - FYE roll-forward)", f"={t.a('vps_d_g')}-{D}{s.a('pub_vps_g')}", "price"),
        ("rec_pct_g", "Difference - Gordon, % of the FYE roll-forward value", f"={t.a('vps_d_g')}/{D}{s.a('pub_vps_g')}-1",
         "pct"),
        ("rec_fye_x", "FYE roll-forward value today - exit, published", f"={D}{s.a('pub_vps_x')}", "price"),
        ("rec_d_x", "Dated value - exit", f"={t.a('vps_d_x')}", "price"),
        ("rec_pct_x", "Difference - exit, % of the FYE roll-forward value", f"={t.a('vps_d_x')}/{D}{s.a('pub_vps_x')}-1",
         "pct"),
    ]:
        t.text(r, 2, label)
        t.put(r, 4, formula, fmt, key=key)
        r += 1
    wb._dated_cells = {"pct_g": f"'Dated_Check'!{t.a('rec_pct_g')}", "pct_x": f"'Dated_Check'!{t.a('rec_pct_x')}"}
    expl = (db.get("reconciliation_notes") or [])
    t.subheader(r, "What explains the gap (filled from the build; tagged)", 5)
    r += 1
    for line_ in expl or ["1H26 FCFF vs half of the FY2026 forecast; 2026 dividend and buybacks paid before 30-Jun; "
                          "FX on USD debt between Dec-25 and Jun-26; acquisitions in 1H26; roll-forward from Dec-25 vs Jun-26."]:
        t.note(r, 2, str(line_))
        r += 1


def _engine_recon(wb, ctx: Ctx, s: Sheet, col: dict, fcst: list[int], rev_names: dict) -> None:
    ws = wb.create_sheet("Engine_Recon")
    t = Sheet(ws)
    ws.column_dimensions["B"].width = 52
    for cc in "CDEF":
        ws.column_dimensions[cc].width = 16
    ws.column_dimensions["G"].width = 10
    ws.sheet_view.showGridLines = False
    t.text(2, 2, "Reconciliation with the model engine (model/model-summary.json)", font=F_TITLE)
    t.text(3, 2, "Engine values are pasted from the engine's summary (constants); workbook values are live formulas. "
                 "Tolerance: relative difference below 1e-6.", font=F_SUB)
    r = 5
    for j, hdr in enumerate(["Item", "", "Engine summary", "This workbook", "Difference", "Rel. diff.", "Status",
                             "Python replica", "Rel. diff. vs replica", "Status vs replica"]):
        cc = ws.cell(row=r, column=2 + j)
        cc.value = hdr
        cc.font = F_HDR
        cc.fill = FILL_HDR
    r += 1
    summ = ctx.summary or {}
    val = summ.get("valuation") or {}
    lines = summ.get("lines") or {}
    D = "'DCF'!"
    items: list[tuple[str, float | None, str, str]] = []
    for key, label, fmt in [("wacc", "WACC", "pct2"),
                            ("ev_gordon", "EV - Gordon", "money"), ("ev_exit", "EV - exit", "money"),
                            ("price_gordon", "Value per share today - Gordon", "price"),
                            ("price_exit", "Value per share today - exit", "price"),
                            ("target_12m_gordon", "12-month target - Gordon", "price"),
                            ("implied_exit_multiple", "Implied exit multiple", "mult"),
                            ("implied_g_exit", "Growth implied by exit", "pct2"),
                            ("implied_g", "Terminal growth implied by the price", "pct2"),
                            ("tv_share_gordon", "TV share of EV - Gordon", "pct")]:
        ref = {"wacc": s.a("wacc"), "ev_gordon": s.a("ev_g"), "ev_exit": s.a("ev_x"),
               "price_gordon": s.a("vps_g"), "price_exit": s.a("vps_x"), "target_12m_gordon": s.a("tgt12_g"),
               "implied_exit_multiple": s.a("impl_mult"), "implied_g_exit": s.a("impl_g_x"),
               "implied_g": rev_names["impl_g"], "tv_share_gordon": s.a("tv_share_g")}[key]
        items.append((label, val.get(key), ref if ref.startswith("'") else D + ref, fmt))
    fcff_engine = (lines.get("fcff") or {})
    ebitda_engine = (lines.get("ebitda") or {})
    for yy in fcst:
        items.append((f"FCFF FY{yy}E", _year_value(fcff_engine, yy), f"{D}{L(col[yy])}{s.rows['fcff']}", "money"))
    for yy in fcst:
        items.append((f"EBITDA FY{yy}E", _year_value(ebitda_engine, yy), f"{D}{L(col[yy])}{s.rows['ebitda']}", "money"))
    py = ctx.py
    py_map = {"WACC": py["wacc"], "Sum of PV of FCFF": py["sum_pv"], "EV - Gordon": py["ev_gordon"],
              "EV - exit": py["ev_exit"], "Value per share today - Gordon": py["price_gordon"],
              "Value per share today - exit": py["price_exit"],
              "12-month target - Gordon": py["price_gordon"] * (1 + py["ke"]) - float(val.get("dps_next") or 0.0),
              "Implied exit multiple": py["implied_exit_multiple"], "Growth implied by exit": py["implied_g_exit"],
              "Terminal growth implied by the price": py["implied_g"], "TV share of EV - Gordon": py["tv_share_gordon"]}
    for i, yy in enumerate(fcst):
        py_map[f"FCFF FY{yy}E"] = py["lines"]["fcff"][i]
        py_map[f"EBITDA FY{yy}E"] = py["lines"]["ebitda"][i]
    extra = [("Discount factor FY" + str(yy) + "E (replica only)", py["dfs"][i], f"{D}{L(col[yy])}{s.rows['df']}", "factor")
             for i, yy in enumerate(fcst)]
    extra += [("Sum of PV of FCFF (replica only)", py["sum_pv"], D + s.a("sum_pv"), "money"),
              ("Terminal value - Gordon (replica only)", py["tv_gordon"], D + s.a("tv_g"), "money"),
              ("PV of terminal value - Gordon (replica only)", py["pv_tv_gordon"], D + s.a("pv_tv_g"), "money"),
              ("Terminal value - exit (replica only)", py["tv_exit"], D + s.a("tv_x"), "money"),
              ("PV of terminal value - exit (replica only)", py["pv_tv_exit"], D + s.a("pv_tv_x"), "money"),
              ("Net debt incl. leases (replica only)", py["bridge"]["net_debt"], D + s.a("net_debt"), "money"),
              ("Shares (replica only)", py["bridge"]["shares"], D + s.a("shares"), "shares"),
              ("Roll-forward factor (replica only)", py["roll_factor"], D + s.a("roll"), "num3")]
    pp = ctx.py_pub
    fL = L(col[fcst[-1]])
    extra += [("PUBLISHED FCFF FY" + str(fcst[-1]) + "E (replica only)", pp["lines"]["fcff_pub"][-1],
               f"{D}{fL}{s.rows['fcff_pub']}", "money"),
              ("PUBLISHED lease liability FY" + str(fcst[-1]) + "E (replica only)", pp["lease_n"],
               f"{D}{fL}{s.rows['ls_liab']}", "money"),
              ("PUBLISHED sum of PV of FCFF (replica only)", pp["sum_pv"], D + s.a("pub_sum_pv"), "money"),
              ("PUBLISHED terminal value - Gordon (replica only)", pp["tv_gordon"], D + s.a("pub_tv_g"), "money"),
              ("PUBLISHED EV - Gordon (replica only)", pp["ev_gordon"], D + s.a("pub_ev_g"), "money"),
              ("PUBLISHED EV - exit (replica only)", pp["ev_exit"], D + s.a("pub_ev_x"), "money"),
              ("Debt-like items, NCI uplift live at ke (replica only)", pp["debt_like_gordon"], D + s.a("debt_like"),
               "money"),
              ("PUBLISHED value per share today - Gordon (replica only)", pp["price_gordon"], D + s.a("pub_vps_g"),
               "price"),
              ("PUBLISHED value per share today - exit (replica only)", pp["price_exit"], D + s.a("pub_vps_x"), "price"),
              ("PUBLISHED RONIC implied by the terminal value (replica only)", pp["tv_ronic"], D + s.a("pub_ronic"),
               "pct"),
              ("PUBLISHED growth implied by the exit multiple (replica only)", pp["implied_g_exit"],
               D + s.a("pub_impl_g_x"), "pct2"),
              ("PUBLISHED Gordon at RONIC = WACC (replica only)", core.ronic_case(ctx.case, pp, pp["wacc"]),
               D + s.a("ronic_v0"), "price"),
              ("PUBLISHED terminal growth implied by the price (replica only)", pp["implied_g"],
               rev_names["pimpl_g"], "pct2")]
    first = r
    for label, eng, ref, fmt in items:
        t.text(r, 2, label)
        t.put(r, 4, float(eng) if eng is not None else "n.a.", fmt, "input")
        t.put(r, 5, f"={ref}", fmt, "link")
        t.put(r, 6, f'=IFERROR(E{r}-D{r},"n.a.")', fmt)
        t.put(r, 7, f'=IFERROR(IF(D{r}=0,E{r}-D{r},E{r}/D{r}-1),"n.a.")', "0.00E+00")
        t.put(r, 8, f'=IFERROR(IF(OR(ABS(G{r})<0.000001,ABS(F{r})<=0.0000005),"OK","CHECK"),"CHECK")', "text")
        pv = py_map.get(label)
        t.put(r, 9, float(pv) if pv is not None else "n.a.", fmt, "input")
        t.put(r, 10, f'=IFERROR(IF(I{r}=0,E{r}-I{r},E{r}/I{r}-1),"n.a.")', "0.00E+00")
        t.put(r, 11, f'=IFERROR(IF(ABS(J{r})<0.000001,"OK","CHECK"),"CHECK")', "text")
        r += 1
    for label, pv, ref, fmt in extra:
        t.text(r, 2, label)
        t.text(r, 4, "not in summary")
        t.put(r, 5, f"={ref}", fmt, "link")
        t.text(r, 8, "replica only")
        t.put(r, 9, float(pv), fmt, "input")
        t.put(r, 10, f'=IFERROR(IF(I{r}=0,E{r}-I{r},E{r}/I{r}-1),"n.a.")', "0.00E+00")
        t.put(r, 11, f'=IFERROR(IF(ABS(J{r})<0.000001,"OK","CHECK"),"CHECK")', "text")
        r += 1
    last = r - 1
    t.text(r + 1, 2, "Rows checked against the engine summary")
    t.put(r + 1, 4, len(items), "0")
    t.text(r + 2, 2, "Rows OK vs engine (must equal the count above)")
    t.put(r + 2, 4, f'=COUNTIF(H{first}:H{first + len(items) - 1},"OK")', "0")
    t.text(r + 3, 2, "Rows checked against the Python replica")
    t.put(r + 3, 4, last - first + 1, "0")
    t.text(r + 4, 2, "Rows OK vs replica (must equal the count above)")
    t.put(r + 4, 4, f'=COUNTIF(K{first}:K{last},"OK")', "0")
    wb._recon_cells = {"eng_n": f"'Engine_Recon'!$D${r + 1}", "eng_ok": f"'Engine_Recon'!$D${r + 2}",
                       "py_n": f"'Engine_Recon'!$D${r + 3}", "py_ok": f"'Engine_Recon'!$D${r + 4}"}


def _year_value(series, year: int):
    if isinstance(series, dict):
        for k in (str(year), year, f"FY{year}"):
            if k in series:
                return series[k]
    return None


def _inputs_sources(wb, ctx: Ctx) -> None:
    ws = wb.create_sheet("Inputs_Sources")
    t = Sheet(ws)
    widths = {"B": 46, "C": 14, "D": 12, "E": 12, "F": 80, "G": 12}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    t.text(2, 2, "Inputs and sources", font=F_TITLE)
    t.text(3, 2, TAG_HELP, font=F_NOTE)
    r = 5
    for j, hdr in enumerate(["Item", "Value", "Unit", "Tag", "Source", "As of"]):
        cc = ws.cell(row=r, column=2 + j)
        cc.value = hdr
        cc.font = F_HDR
        cc.fill = FILL_HDR
    r += 1
    for item in ctx.extras.get("sources") or []:
        t.text(r, 2, str(item.get("item", "")))
        v = item.get("value")
        if isinstance(v, (int, float)):
            t.put(r, 3, v, "#,##0.0000" if abs(v) < 10 else "#,##0.00", "input")
        elif v is not None:
            t.text(r, 3, str(v))
        t.text(r, 4, str(item.get("unit", "")))
        t.text(r, 5, str(item.get("tag", "")))
        t.text(r, 6, str(item.get("source", "")), wrap=False)
        t.text(r, 7, str(item.get("as_of", "")))
        r += 1
    t.note(r + 1, 2, "Drivers: model/drivers.yaml (tag, rationale and pillar per driver). History: data/financials.csv "
                     "(document and page per figure) and data/adjustments.md.")


def _checks(wb, ctx: Ctx, s: Sheet, col: dict, hist: list[int], fcst: list[int]) -> None:
    ws = wb.create_sheet("Checks")
    t = Sheet(ws)
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 12
    t.text(2, 2, "Integrity checks", font=F_TITLE)
    r = 4
    D = "'DCF'!"
    checks = [
        ("WACC - g at least 2 points", f'=IF({D}{s.a("wacc_minus_g")}>=0.02,"OK","WARN")'),
        ("g not above long-term nominal GDP growth", f'=IF({D}{s.a("g_vs_gdp")}<=0,"OK","FAIL")'),
        ("Weights add up to 100%", f'=IF(ABS({D}{s.a("we")}+{D}{s.a("wd")}-1)<0.000001,"OK","FAIL")'),
        ("Engine reconciliation: every row OK (count of OK = rows checked)",
         f'=IF({wb._recon_cells["eng_ok"]}={wb._recon_cells["eng_n"]},"OK","CHECK")'),
        ("Python replica reconciliation: every row OK",
         f'=IF({wb._recon_cells["py_ok"]}={wb._recon_cells["py_n"]},"OK","CHECK")'),
        ("Gordon TV share of EV below 85% (published)", f'=IF({D}{s.a("pub_tv_share_g")}<0.85,"OK","WARN")'),
        ("Exit multiple implied by the published Gordon value within a factor of two of the exit multiple "
         "(0.5x to 2x)",
         f'=IF(AND({D}{s.a("pub_impl_mult")}>=0.5*{D}{s.a("exit_mult")},'
         f'{D}{s.a("pub_impl_mult")}<=2*{D}{s.a("exit_mult")}),"OK","WARN")'),
        ("Lease growth cannot add value: published Gordon <= engine-basis Gordon",
         f'=IF({D}{s.a("pub_vps_g")}<={D}{s.a("vps_g")}+0.000001,"OK","CHECK")'),
        ("Lease test (a): FY2035 growth in the lease liability within 10% of g x lease liability FY2035",
         f'=IF(ABS({D}{s.a("lease_test_a")}-1)<=0.1,"OK","WARN")'),
        ("Lower-return table reproduces the published Gordon value at the implied RONIC",
         f'=IF(ABS({D}{s.a("ronic_v3")}-{D}{s.a("pub_vps_g")})<0.000001,"OK","FAIL")'),
        ("Published RONIC above WACC (terminal growth creates value)",
         f'=IF({D}{s.a("pub_ronic")}>{D}{s.a("wacc")},"OK","WARN")'),
        ("Dated cross-check within 5% of the FYE roll-forward (Gordon and exit, published)",
         f'=IF(AND(ABS({wb._dated_cells["pct_g"]})<0.05,ABS({wb._dated_cells["pct_x"]})<0.05),"OK","WARN")'),
    ]
    sc = getattr(wb, "_scen_cells", None)
    if sc:
        checks.append(("Scenarios: base row equals the published key outputs (Gordon and exit)",
                       f'=IF(AND(ABS({sc["g"]}-{D}{s.a("pub_vps_g")})<0.000001,ABS({sc["x"]}-{D}{s.a("pub_vps_x")})<0.000001),'
                       f'"OK","FAIL")'))
    if "nwc_op" in s.rows:
        owc = (ctx.extras.get("operating_wc") or {}).get("operating_nwc") or {}
        for yy in hist:
            target = owc.get(yy, owc.get(str(yy)))
            if target is None:
                continue
            checks.append((f"Operating working capital FY{yy} ties to data/adjustments.md ({target:,.0f})",
                           f'=IF(ABS({D}{L(col[yy])}{s.rows["nwc_op"]}-({target}))<0.5,"OK","FAIL")'))
    if "debt_like_yaml" in s.names:
        checks.append(("Debt-like components (NCI uplift live at ke) add up to valuation.yaml debt_like_items",
                       f'=IF(ABS({D}{s.a("debt_like")}-{D}{s.a("debt_like_yaml")})<0.5,"OK","FAIL")'))
    if ctx.case.segments:
        seg_first = s.rows[f"seg_{ctx.case.segments[0]}"]
        seg_last = s.rows[f"seg_{ctx.case.segments[-1]}"]
        for yy in hist:
            cl = L(col[yy])
            checks.append((f"Segments add up to revenue FY{yy}",
                           f'=IF(ABS(SUM({D}{cl}{seg_first}:{cl}{seg_last})-{D}{cl}{s.rows["revenue"]})<0.5,"OK","FAIL")'))
    for label, formula in checks:
        t.text(r, 2, label)
        t.put(r, 3, formula, "text")
        r += 1
    t.note(r + 1, 2, "Formula errors are scanned after recalculation with valuation/build/scan_errors.py.")


def _cover_dcf(ws, ctx: Ctx) -> None:
    t = Sheet(ws)
    ws.column_dimensions["B"].width = 120
    ws.sheet_view.showGridLines = False
    lines = [
        ("Grupo Bimbo, S.A.B. de C.V. (BMV: BIMBO A) - DCF valuation workbook", F_TITLE),
        (f"Valuation date {ctx.valuation_date}; MXN millions; IFRS. Generated {dt.date.today().isoformat()} by "
         "valuation/build/build_spg_workbooks.py from the canonical model inputs.", F_SUB),
        ("", F_NORM),
        ("Sheets", F_BOLD),
        ("DCF - assumptions, drivers, projected free cash flow, working capital, discounting, terminal value, equity "
         "bridge, cross-checks, published basis (lease growth, peer-EV exit claims, value walk, lower-return cases) and "
         "sensitivity tables (layout after S&P 'DCF Analysis').", F_NORM),
        ("Published basis = engine mechanics plus two technical corrections: (1) FCFF also deducts the growth in the "
         "lease liability, held at the FY2025 share of revenue; (2) the exit method subtracts only the claims inside the "
         "peer EV definition (net debt, NCI at market). Engine-basis values are kept for reconciliation.", F_NORM),
        ("Reverse_DCF - terminal growth and revenue-growth shift implied by the market price.", F_NORM),
        ("Levers - static grids of the story levers (capex, gross margin, growth) on the published basis.", F_NORM),
        ("Scenarios - one change at a time against the base: cost-of-capital routes and country-risk weights, Bimbo's own "
         "regression beta, North America currency consistency (PPP), terminal capex vs the peer benchmark, exit multiples "
         "at the v3 peer medians.", F_NORM),
        ("v3 (2026-10-17): the base cost of capital is route B - MXN default-free rate (M-Bono 10y minus Mexico's default "
         "spread) plus an equity country risk premium weighted by where Bimbo operates (countries -> regions -> FY2025 "
         "Adj. EBITDA weights); full build in BIMBOA_WACC_v3.xlsx. A proposal until the team confirms it.", F_NORM),
        ("Dated_Check - independent valuation anchored at the 30-Jun-2026 balance sheet, rolled to the valuation date.",
         F_NORM),
        ("Engine_Recon - line-by-line agreement with the project's model engine (model/model-summary.json).", F_NORM),
        ("Inputs_Sources - every input with value, tag, source and date.  Checks - integrity checks.", F_NORM),
        ("", F_NORM),
        ("Conventions", F_BOLD),
        ("Cash flows discounted to 31-Dec-2025 (mid-year); Gordon TV discounted N - offset, exit TV discounted N; "
         "net debt incl. IFRS 16 leases at 31-Dec-2025; value per share rolled forward to today at the cost of equity; "
         "12-month value deducts the next unpaid dividend (the FY2026 dividend was paid on 13-May-2026).",
         F_NORM),
        ("The terminal value, WACC inputs, story drivers, target price and rating are TEAM decisions; values here are "
         "tagged proposals until the team confirms them (AGENTS.md).", F_NORM),
        ("Publication gate: report, deck and pitch read valuation/valuation-outputs.yaml, not this workbook or the "
         "engine summary directly.", F_NORM),
        ("", F_NORM),
        (TAG_HELP, F_NOTE),
    ]
    for i, (txt, font) in enumerate(lines, start=2):
        t.text(i, 2, txt, font=font)


# ----------------------------------------------------------------------------- WACC workbook

def _crp_weighting(ws, s: Sheet, va: dict, r: int) -> int:
    """Route B country risk premium: countries (Damodaran) -> regions (plant-weighted) -> company (region weights).
    The selected company CRP feeds 'crp_b'; the region CRPs are checked against the calculator spec values."""
    regions = va.get("crp_exposure") or []
    countries = va.get("crp_countries") or []
    if not regions:
        return r
    s.subheader(r, "Route B - country risk premium weighted by where Bimbo operates: countries (Damodaran; dataset in the source column) "
                   "-> regions (plant-weighted) -> company (region weights)", 15)
    r += 1
    first_c = last_c = None
    if countries:
        for j, h in enumerate(["Region", "Sub-region", "Country", "Plants", "Moody's", "Default spread", "CRP", "Include",
                               "Plants x CRP"]):
            s.text(r, 5 + j, h, font=F_BOLD)
        s.text(r, 15, "Tag / source", font=F_BOLD)
        r += 1
        first_c = r
        for c in countries:
            s.text(r, 5, str(c.get("region", "")))
            s.text(r, 6, str(c.get("subregion") or ""))
            s.text(r, 7, str(c.get("country", "")))
            s.put(r, 8, float(c.get("plants") or 0), "0", "input")
            s.text(r, 9, str(c.get("rating") or ""))
            s.put(r, 10, float(c.get("default_spread") or 0.0), "pct2", "input")
            s.put(r, 11, float(c.get("crp") or 0.0), "pct2", "input")
            s.put(r, 12, "yes" if c.get("include", True) else "no", "text", "input")
            s.put(r, 13, f'=IF(L{r}="yes",H{r}*K{r},0)', "0.0000")
            s.note(r, 15, f"plants [{c.get('plants_tag', 'sourced')}]; CRP [sourced] {c.get('source', '')}")
            r += 1
        last_c = r - 1
        r += 1
    for j, h in enumerate(["Region", "Plants included", "Region CRP", "Region CRP excl. Argentina", "W: Adj. EBITDA FY25",
                           "W: sales FY25", "W: Adj. EBITDA FY23-25", "W: sales FY23-25", "Region CRP = spec"]):
        s.text(r, 5 + j, h, font=F_BOLD)
    r += 1
    first_r = r
    core.crp_from_countries({"wacc_variants": va})          # validates labels, plants and weights before writing
    for row in regions:
        code = core.CRP_REGION_CODES.get(str(row.get("region")))
        if code is None:
            raise SystemExit(f"unknown CRP region label: {row.get('region')}")
        s.text(r, 5, str(row.get("region", "")))
        if countries:
            rng = lambda col: f"${col}${first_c}:${col}${last_c}"  # noqa: E731
            cond = f'({rng("E")}="{code}")*({rng("L")}="yes")'
            condx = cond + f'*({rng("G")}<>"Argentina")'
            s.put(r, 6, f"=SUMPRODUCT({cond}*{rng('H')})", "0")
            s.put(r, 7, f"=IFERROR(SUMPRODUCT({cond}*{rng('H')}*{rng('K')})/F{r},0)", "pct2")
            s.put(r, 8, f"=IFERROR(SUMPRODUCT({condx}*{rng('H')}*{rng('K')})/SUMPRODUCT({condx}*{rng('H')}),0)", "pct2")
            s.put(r, 13, f'=IF(ABS(G{r}-({float(row.get("crp", 0.0))!r}))<0.0000001,"OK","DIFF")', "text")
        else:
            s.put(r, 7, float(row.get("crp", 0.0)), "pct2", "input")
            s.put(r, 8, float(row.get("crp_ex_argentina", row.get("crp", 0.0))), "pct2", "input")
        s.put(r, 9, float(row.get("ebitda_weight", 0.0)), "pct", "input")
        s.put(r, 10, float(row.get("revenue_weight", 0.0)), "pct", "input")
        s.put(r, 11, float(row.get("ebitda_weight_3y", row.get("ebitda_weight", 0.0))), "pct", "input")
        s.put(r, 12, float(row.get("revenue_weight_3y", row.get("revenue_weight", 0.0))), "pct", "input")
        s.note(r, 15, str(row.get("source", "")))
        r += 1
    last_r = r - 1
    s.text(r, 5, "Company CRP = SUMPRODUCT(region weights, region CRP)", font=F_BOLD)
    for j, cw in enumerate("IJKL"):
        s.put(r, 9 + j, f"=SUMPRODUCT({cw}{first_r}:{cw}{last_r},$G${first_r}:$G${last_r})", "pct2",
              key=f"crp_tot{j + 1}", bold=True)
    r += 1
    s.text(r, 5, "Company CRP excluding Argentina (sensitivity)")
    for j, cw in enumerate("IJKL"):
        s.put(r, 9 + j, f"=SUMPRODUCT({cw}{first_r}:{cw}{last_r},$H${first_r}:$H${last_r})", "pct2",
              key=f"crp_tot{j + 1}_xar")
    r += 1
    s.text(r, 5, "Weights add up to 100%")
    for j, cw in enumerate("IJKL"):
        s.put(r, 9 + j, f'=IF(ABS(SUM({cw}{first_r}:{cw}{last_r})-1)<0.000001,"OK","FAIL")', "text")
    r += 1
    cell = ws[s.a("crp_b").replace("$", "")]
    cell.value = (f"=CHOOSE({s.a('crp_basis')},{s.a('crp_tot1')},{s.a('crp_tot2')},{s.a('crp_tot3')},"
                  f"{s.a('crp_tot4')})")
    cell.font = F_NORM
    cell.fill = PatternFill(fill_type=None)
    s.note(r, 5, "Rule: consolidated operations in; Venezuela out (not consolidated; held at fair value, RA25 p.85). "
                 "Argentina is consolidated (IAS 29), so it is in; the row above shows the value without it. Within a "
                 "region, countries are weighted by plant count (no country sales disclosure) [assumption].")
    r += 1
    s.note(r, 5, "Method (base): additive, ke = rf + beta_l x ERP + CRP, the formula of the project guide's USD route "
                 "(references/wacc-latam.md). Beta-scaled sensitivity in STEP 5.")
    r += 2
    cc = va.get("crp_crosschecks") or {}
    if cc:
        s.text(r, 5, "Cross-checks (not used in the WACC)", font=F_BOLD)
        r += 1
        for label, v, fmt in [
            ("Mexico 10-year sovereign CDS, mid (Capital IQ, 16-Oct-2026)", cc.get("cds_10y"), "pct2"),
            ("Mexico 5-year sovereign CDS, mid (Capital IQ, 16-Oct-2026)", cc.get("cds_5y"), "pct2"),
            ("Damodaran rating-based default spread, Mexico, used in route B", cc.get("default_spread"), "pct2"),
            ("Equity / bond volatility multiplier implied by Damodaran (Mexico CRP / default spread)",
             cc.get("equity_bond_multiplier"), "num"),
            ("Company impairment-test discount rates (post-tax, local currency), weighted by FY2025 Adj. EBITDA",
             cc.get("impairment_rate_ebitda_fy25"), "pct2"),
            ("Same, weighted by FY2025 sales", cc.get("impairment_rate_sales_fy25"), "pct2"),
        ]:
            if v is None:
                continue
            s.text(r, 5, label)
            s.put(r, 9, float(v), fmt, "input")
            r += 1
        imp = cc.get("impairment_rates") or {}
        if imp:
            s.text(r, 5, "Impairment-test rates 2025 (RA25 p.265): " + "; ".join(
                f"{k} {float(v):.2%}" for k, v in imp.items()), font=F_NOTE)
            r += 1
        if cc.get("notes"):
            s.note(r, 5, str(cc["notes"]))
            r += 1
        r += 1
    return r


def _rates_evidence(ws, s: Sheet, ex: dict, r: int) -> int:
    """Government curves, CDS and inflation at the pricing date (Capital IQ), as evidence for the risk-free choice."""
    rows = ex.get("rates_credit") or []
    if not rows:
        return r
    s.subheader(r, "Evidence - government curves, sovereign CDS and inflation (S&P Capital IQ, 16-Oct-2026)", 15)
    r += 1
    for j, h in enumerate(["Series", "Value", "Unit", "Date"]):
        s.text(r, 5 + j, h, font=F_BOLD)
    s.text(r, 15, "Source", font=F_BOLD)
    r += 1
    for it in rows:
        s.text(r, 5, str(it.get("series", "")))
        v = it.get("value")
        if isinstance(v, (int, float)):
            s.put(r, 6, float(v), "pct2" if str(it.get("unit", "")).startswith(("yield", "rate", "spread", "%", "inflation"))
                  else "num", "input")
        else:
            s.text(r, 6, "n/a")
        s.text(r, 7, str(it.get("unit", "")))
        s.text(r, 8, str(it.get("date", "")))
        s.note(r, 15, str(it.get("source", "")))
        if "10Y" in str(it.get("series", "")) and "M-Bono" in str(it.get("series", "")) and isinstance(v, (int, float)):
            s.put(r, 9, f'=IF(ABS(F{r}-{s.a("mbono10")})<0.000001,"= rf input","DIFF")', "text")
        r += 1
    return r + 1


def _beta_evidence(ctx: Ctx, s: Sheet, r: int) -> int:
    """Bimbo's own regression beta against the IPC (valuation/inputs/beta_regression.json)."""
    import json as _json
    p = ctx.project / "valuation" / "inputs" / "beta_regression.json"
    if not p.is_file():
        return r
    d = _json.loads(p.read_text(encoding="utf-8"))
    s.text(r, 5, "Evidence - Bimbo regression beta vs the S&P/BMV IPC (monthly; levered at its own market D/E 0.743)",
           font=F_BOLD)
    r += 1
    for j, h in enumerate(["Window", "Months", "Levered beta", "R-squared", "Std. error", "95% low", "95% high",
                           "Blume", "Unlevered", "Unlevered (Blume)"]):
        s.text(r, 5 + j, h, font=F_BOLD)
    r += 1
    for name, c in (d.get("cases") or {}).items():
        s.text(r, 5, name.replace("_", " "))
        s.put(r, 6, int(c["n"]), "0")
        for j, (k, fmt) in enumerate((("beta", "num3"), ("r2", "num3"), ("se", "num3"), ("ci95_low", "num3"),
                                      ("ci95_high", "num3"), ("blume", "num3"), ("unlevered", "num3"),
                                      ("blume_unlevered", "num3"))):
            s.put(r, 7 + j, float(c[k]), fmt, "input")
        r += 1
    s.note(r, 5, f"[calc] {d.get('source', '')}. {(d.get('selected_evidence') or {}).get('note', '')}")
    return r + 2


def _debt_evidence(ex: dict, s: Sheet, r: int) -> int:
    """Outstanding bonds and the issuer rating (Capital IQ prompt 3), as evidence for the cost of debt."""
    ds = ex.get("debt_securities") or {}
    if not ds:
        return r
    s.subheader(r, "Evidence - Bimbo bonds outstanding and issuer rating (S&P Capital IQ, 16-Oct-2026)", 15)
    r += 1
    rt = ds.get("issuer_rating") or {}
    if rt:
        s.text(r, 2, f"Issuer rating: {rt.get('agency', '')} {rt.get('rating', '')} ({rt.get('type', '')}), outlook "
                     f"{rt.get('outlook', '')}; effective {rt.get('effective', '')}, last reviewed {rt.get('last_review', '')}. "
                     f"{rt.get('source', '')}", font=F_NOTE)
        r += 1
    for j, h in enumerate(["Bond", "Currency", "Coupon", "Maturity", "Amount (m)", "Note"]):
        s.text(r, 2 + j if j == 0 else 4 + j, h, font=F_BOLD)
    r += 1
    for b in (ds.get("mxn") or []) + (ds.get("usd") or []):
        s.text(r, 2, str(b.get("bond") or b.get("ticker") or b.get("issuer") or ""))
        s.text(r, 5, str(b.get("currency", "")))
        cp = b.get("coupon")
        if isinstance(cp, (int, float)):
            s.put(r, 6, float(cp), "pct2", "input")
        else:
            s.text(r, 6, str(cp if cp is not None else "n/a"))
        s.text(r, 7, str(b.get("maturity", "")))
        am = b.get("amount_m")
        if isinstance(am, (int, float)):
            s.put(r, 8, float(am), "money", "input")
        s.note(r, 9, str(b.get("note", "")))
        r += 1
    if ds.get("source"):
        s.note(r, 2, str(ds["source"]))
        r += 1
    return r + 1


def build_wacc(ctx: Ctx, out: Path) -> None:
    ex = ctx.extras
    w = ctx.case.wacc_in
    wb = Workbook()
    ws = wb.active
    ws.title = "WACC"
    s = Sheet(ws)
    ws.sheet_view.showGridLines = False
    for k, v in {"A": 2, "B": 58, "C": 14, "D": 3, "E": 34, "F": 16, "G": 10, "H": 11, "I": 14, "J": 14, "K": 11,
                 "L": 11, "M": 11, "N": 10, "O": 60}.items():
        ws.column_dimensions[k].width = v
    s.text(2, 2, "WACC Analysis", font=F_TITLE)
    s.text(3, 2, f"{ctx.company} ({ctx.ticker}) | valuation date {ctx.valuation_date} | MXN nominal | layout after "
                 "S&P Capital IQ Pro 'WACC Analysis' (template defaults overridden: local risk-free, mature-market ERP, "
                 "statutory marginal tax, MXN cost of debt).", font=F_SUB)
    s.text(4, 2, "Blue on yellow = input. " + TAG_HELP, font=F_NOTE)
    va = ex.get("wacc_variants") or {}
    r = 6
    s.header(r, "STEP 1 - SELECTIONS", 15)
    r += 1
    s.text(r, 2, "Route: 1 = A local (M-Bono, CRP 0); 2 = B Damodaran (default-free rf + exposure-weighted CRP)")
    s.put(r, 3, int(va.get("selected_route", 1)), "0", "input", key="route")
    dv = DataValidation(type="list", formula1='"1,2"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws.cell(row=r, column=3))
    r += 1
    s.text(r, 2, "Beta: 1 = bottom-up median; 2 = bottom-up mean; 3 = industry (Damodaran); 4 = Bimbo regression (unlevered)")
    s.put(r, 3, int(va.get("selected_beta", 1)), "0", "input", key="beta_sel")
    dv2 = DataValidation(type="list", formula1='"1,2,3,4"', allow_blank=False)
    ws.add_data_validation(dv2)
    dv2.add(ws.cell(row=r, column=3))
    r += 1
    s.text(r, 2, "Country-risk weights (route B): 1 = FY2025 Adj. EBITDA; 2 = FY2025 sales; 3 = FY2023-25 average "
                 "Adj. EBITDA; 4 = FY2023-25 average sales")
    s.put(r, 3, int(va.get("crp_weight_basis", 1)), "0", "input", key="crp_basis")
    dv3 = DataValidation(type="list", formula1='"1,2,3,4"', allow_blank=False)
    ws.add_data_validation(dv3)
    dv3.add(ws.cell(row=r, column=3))
    r += 2

    # ---- risk-free and country risk
    s.header(r, "STEP 2 - RISK-FREE RATE AND COUNTRY RISK", 15)
    r += 1
    rf = va.get("rf") or {}
    for key, label, default, fmt in [
        ("mbono10", "Mexico M-Bono 10-year yield", w["risk_free"], "pct2"),
        ("def_spread", "Mexico sovereign default spread (variant B)", 0.0, "pct2"),
        ("crp_b", "Equity country risk premium, exposure-weighted (variant B; build below)", 0.0, "pct2"),
        ("size", "Size premium", 0.0, "pct2"),
    ]:
        item = rf.get(key) or {}
        v = item.get("value", default) if isinstance(item, dict) else item
        s.text(r, 2, label)
        s.put(r, 3, float(v), fmt, "input", key=key)
        s.note(r, 5, f"[{item.get('tag', 'assumption')}] {item.get('source', '')}" if isinstance(item, dict) else "")
        r += 1
    s.text(r, 2, "Default-free MXN risk-free (variant B) = M-Bono 10y - default spread")
    s.put(r, 3, f"={s.a('mbono10')}-{s.a('def_spread')}", "pct2", key="rf_b")
    r += 1
    s.text(r, 2, "Risk-free used", font=F_BOLD)
    s.put(r, 3, f"=IF({s.a('route')}=1,{s.a('mbono10')},{s.a('rf_b')})", "pct2", key="rf", bold=True)
    r += 1
    s.text(r, 2, "Country risk premium used", font=F_BOLD)
    s.put(r, 3, f"=IF({s.a('route')}=1,0,{s.a('crp_b')})", "pct2", key="crp", bold=True)
    r += 1
    erp = (va.get("erp") or {})
    s.text(r, 2, "Equity risk premium (mature market)")
    s.put(r, 3, float(erp.get("value", w["equity_risk_premium"])), "pct2", "input", key="erp")
    s.note(r, 5, f"[{erp.get('tag', 'assumption')}] {erp.get('source', '')}")
    r += 1
    s.text(r, 2, "Marginal tax rate (Mexico statutory)")
    s.put(r, 3, w["tax_rate"], "pct", "input", key="tax")
    s.note(r, 5, "[sourced] Mexico statutory corporate rate 30% (Ley del ISR); effective/cash tax is a separate forecast driver.")
    r += 2

    # ---- country risk weighting (route B): countries -> regions -> company
    r = _crp_weighting(ws, s, va, r)
    r = _rates_evidence(ws, s, ex, r)

    # ---- beta table
    s.header(r, "STEP 3 - BOTTOM-UP BETA (Hamada: beta_u = beta_l / (1 + (1 - t) x D/E); debt incl. leases, market equity)", 15)
    r += 1
    hdrs = ["Peer", "Ticker", "Marginal tax", "Levered beta", "Debt incl. leases", "Market equity", "D/E", "Unlevered beta",
            "Included", "Used"]
    for j, hdr in enumerate(hdrs):
        s.text(r, 5 + j, hdr, font=F_BOLD)
    s.text(r, 15, "Source", font=F_BOLD)
    r += 1
    first = r
    for p in ex.get("beta_peers") or []:
        s.text(r, 5, str(p.get("name", "")))
        s.text(r, 6, str(p.get("ticker", "")))
        s.put(r, 7, float(p.get("tax", 0.0)), "pct", "input")
        s.put(r, 8, float(p.get("beta_levered", 0.0)), "num3", "input")
        s.put(r, 9, float(p.get("debt_incl_leases", 0.0)), "money", "input")
        s.put(r, 10, float(p.get("market_equity", 0.0)), "money", "input")
        s.put(r, 11, f"=IFERROR(I{r}/J{r},0)", "num3")
        s.put(r, 12, f"=H{r}/(1+(1-G{r})*K{r})", "num3")
        s.put(r, 13, "yes" if p.get("include", True) else "no", "text", "input")
        s.put(r, 14, f'=IF(M{r}="yes",L{r},"excl.")', "num3")
        s.note(r, 15, f"[{p.get('tag', 'sourced')}] {p.get('source', '')}")
        r += 1
    last = r - 1
    s.text(r, 5, "Median (included)", font=F_BOLD)
    s.put(r, 14, f"=MEDIAN(N{first}:N{last})" if last >= first else 0.0, "num3", key="bu_median", bold=True)
    r += 1
    s.text(r, 5, "Mean (included)", font=F_BOLD)
    s.put(r, 14, f"=AVERAGE(N{first}:N{last})" if last >= first else 0.0, "num3", key="bu_mean", bold=True)
    r += 1
    alt = ex.get("beta_alternatives") or {}
    for key, label in [("industry", "Industry unlevered beta (Damodaran, Food Processing)"),
                       ("regression_unlevered", "Bimbo regression beta, unlevered at its own market D/E")]:
        item = alt.get(key) or {}
        s.text(r, 5, label)
        s.put(r, 14, float(item.get("value", 0.0)), "num3", "input", key=f"beta_{key}")
        s.note(r, 15, f"[{item.get('tag', 'unverified')}] {item.get('source', '')}")
        r += 1
    s.text(r, 5, "Selected unlevered beta", font=F_BOLD)
    s.put(r, 14, f"=CHOOSE({s.a('beta_sel')},{s.a('bu_median')},{s.a('bu_mean')},{s.a('beta_industry')},"
                 f"{s.a('beta_regression_unlevered')})", "num3", key="beta_u", bold=True)
    r += 2
    r = _beta_evidence(ctx, s, r)

    # ---- capital structure and cost of debt
    s.header(r, "STEP 4 - CAPITAL STRUCTURE AND COST OF DEBT (MXN)", 15)
    r += 1
    cs = ex.get("capital_structure") or {}
    for key, label, default, fmt in [
        ("de_current", "Bimbo current market D/E (debt incl. leases / market cap)", 0.0, "num3"),
        ("de_peer_median", "Peer median market D/E", 0.0, "num3"),
        ("de_mgmt", "D/E implied by management's 2.0-2.5x ex-IFRS 16 comfort range plus leases", 0.0, "num3"),
    ]:
        item = cs.get(key) or {}
        s.text(r, 2, label)
        s.put(r, 3, float(item.get("value", default)), fmt, "input", key=key)
        s.note(r, 5, f"[{item.get('tag', 'calc')}] {item.get('source', '')}")
        r += 1
    s.text(r, 2, "Target D/E used (team decision)", font=F_BOLD)
    s.put(r, 3, w["target_debt_to_equity"], "num3", "input", key="de")
    s.note(r, 5, "valuation/valuation.yaml")
    r += 1
    kd_c = ex.get("kd_candidates") or []
    for i, item in enumerate(kd_c):
        s.text(r, 2, f"Pre-tax Kd candidate: {item.get('label', '')}")
        s.put(r, 3, float(item.get("value", 0.0)), "pct2", "input", key=f"kd_c{i}")
        s.note(r, 5, f"[{item.get('tag', '')}] {item.get('source', '')}")
        r += 1
    s.text(r, 2, "Pre-tax cost of debt used (team decision)", font=F_BOLD)
    s.put(r, 3, w["pre_tax_cost_of_debt"], "pct2", "input", key="kd")
    s.note(r, 5, "valuation/valuation.yaml")
    r += 2
    r = _debt_evidence(ex, s, r)

    # ---- WACC
    s.header(r, "STEP 5 - COST OF EQUITY AND WACC", 15)
    r += 1
    rows = [
        ("beta_l", "Relevered beta = beta_u x (1 + (1 - t) x D/E)", f"={s.a('beta_u')}*(1+(1-{s.a('tax')})*{s.a('de')})", "num3"),
        ("ke", "Cost of equity = rf + beta_l x ERP + CRP + size",
         f"={s.a('rf')}+{{beta_l}}*{s.a('erp')}+{s.a('crp')}+{s.a('size')}", "pct2"),
        ("we", "E / (D + E)", f"=1/(1+{s.a('de')})", "pct"),
        ("ke_part", "Cost of equity portion", "={ke}*{we}", "pct2"),
        ("kd_after", "After-tax cost of debt", f"={s.a('kd')}*(1-{s.a('tax')})", "pct2"),
        ("wd", "D / (D + E)", f"={s.a('de')}/(1+{s.a('de')})", "pct"),
        ("kd_part", "Cost of debt portion", "={kd_after}*{wd}", "pct2"),
        ("wacc", "WACC", "={ke_part}+{kd_part}", "pct2"),
    ]
    for key, label, formula, fmt in rows:
        formula = formula.format(**{k: s.a(k) for k in s.names if "{" + k + "}" in formula})
        s.text(r, 2, label, font=F_BOLD if key == "wacc" else F_NORM)
        s.put(r, 3, formula, fmt, key=key, bold=key == "wacc")
        r += 1
    s.text(r, 2, "WACC in valuation/valuation.yaml (engine input set)")
    s.put(r, 3, core.wacc(ctx.case)["wacc"], "pct2", "input", key="wacc_yaml")
    r += 1
    s.text(r, 2, "Check: workbook WACC = valuation.yaml WACC", font=F_BOLD)
    s.put(r, 3, f'=IF(ABS({s.a("wacc")}-{s.a("wacc_yaml")})<0.000001,"OK","DIFFERENT SELECTION")', "text", key="wacc_check")
    s.note(r, 5, "Differs only if the selectors or inputs on this sheet are changed from the valuation.yaml set.")
    r += 1
    s.text(r, 2, "Sensitivity: cost of equity with the CRP scaled by beta = rf + beta_l x (ERP + CRP)")
    s.put(r, 3, f"={s.a('rf')}+{s.a('beta_l')}*({s.a('erp')}+{s.a('crp')})+{s.a('size')}", "pct2", key="ke_bs")
    r += 1
    s.text(r, 2, "Sensitivity: WACC with the beta-scaled CRP")
    s.put(r, 3, f"={s.a('we')}*{s.a('ke_bs')}+{s.a('wd')}*{s.a('kd_after')}", "pct2", key="wacc_bs")
    s.note(r, 5, "Damodaran's operations-weighted ERP approach multiplies the CRP by beta; the base keeps it additive.")
    r += 2

    # ---- sensitivity
    s.header(r, "WACC SENSITIVITY ANALYSIS", 15)
    r += 1
    s.subheader(r, "WACC: D/(D+E) (rows; beta relevered at each D/E) x pre-tax cost of debt (columns)", 15)
    r += 1
    kd_steps = (-0.01, -0.005, 0.0, 0.005, 0.01)
    wd_vals = (0.25, 0.30, 0.35, 0.40, 0.45, 0.50, 0.55)
    for j, st in enumerate(kd_steps):
        s.put(r, 3 + j, f"={s.a('kd')}+({st})", "pct2")
    hdr = r
    r += 1
    for wdv in wd_vals:
        s.put(r, 2, wdv, "pct", "input")
        for j in range(len(kd_steps)):
            kd_cell = f"{L(3 + j)}${hdr}"
            wd_cell = f"$B{r}"
            de = f"({wd_cell}/(1-{wd_cell}))"
            ke = f"({s.a('rf')}+{s.a('beta_u')}*(1+(1-{s.a('tax')})*{de})*{s.a('erp')}+{s.a('crp')}+{s.a('size')})"
            s.put(r, 3 + j, f"=(1-{wd_cell})*{ke}+{wd_cell}*{kd_cell}*(1-{s.a('tax')})", "pct2")
        r += 1
    r += 1
    s.subheader(r, "WACC: unlevered beta (rows) x equity risk premium (columns), at the target D/E", 15)
    r += 1
    erp_steps = (-0.01, -0.005, 0.0, 0.005, 0.01)
    beta_steps = (-0.2, -0.1, 0.0, 0.1, 0.2)
    for j, st in enumerate(erp_steps):
        s.put(r, 3 + j, f"={s.a('erp')}+({st})", "pct2")
    hdr = r
    r += 1
    for bst in beta_steps:
        s.put(r, 2, f"={s.a('beta_u')}+({bst})", "num3")
        for j in range(len(erp_steps)):
            ke = (f"({s.a('rf')}+$B{r}*(1+(1-{s.a('tax')})*{s.a('de')})*{L(3 + j)}${hdr}+{s.a('crp')}+{s.a('size')})")
            s.put(r, 3 + j, f"={s.a('we')}*{ke}+{s.a('wd')}*{s.a('kd_after')}", "pct2")
        r += 1
    r += 1
    s.note(r, 2, "(1) Unlevered beta = levered beta / (1 + (1 - tax) x D/E). (2) Relevered beta = unlevered beta x "
                 "(1 + (1 - tax) x D/E). Source data: valuation/inputs/wacc_inputs.yaml and the wiki (see Inputs_Sources).")
    ws.freeze_panes = "A6"
    _inputs_sources(wb, ctx)
    wb.save(out)


# ----------------------------------------------------------------------------- main

def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--project", default=".")
    ap.add_argument("--version", default="1")
    args = ap.parse_args(argv)
    project = Path(args.project).resolve()
    ctx = Ctx(project)
    ctx.version = args.version
    out_dir = project / "valuation"
    dcf_path = out_dir / f"BIMBOA_DCF_v{args.version}.xlsx"
    wacc_path = out_dir / f"BIMBOA_WACC_v{args.version}.xlsx"
    for p in (dcf_path, wacc_path):
        if p.name.startswith("SPG_"):
            raise SystemExit("refusing to write over an S&P template")
    build_dcf(ctx, dcf_path)
    build_wacc(ctx, wacc_path)
    py = ctx.py
    print(f"wrote {dcf_path.name} and {wacc_path.name}")
    print(f"python replica: WACC {py['wacc']:.6%}  ke {py['ke']:.6%}  EV Gordon {py['ev_gordon']:,.1f}  "
          f"EV exit {py['ev_exit']:,.1f}  price Gordon {py['price_gordon']:.4f}  price exit {py['price_exit']:.4f}")
    pp = ctx.py_pub
    print(f"published:      EV Gordon {pp['ev_gordon']:,.1f}  EV exit {pp['ev_exit']:,.1f}  price Gordon "
          f"{pp['price_gordon']:.6f}  price exit {pp['price_exit']:.6f}  RONIC {pp['tv_ronic']:.4%}")
    print(f"reverse DCF solutions: {getattr(ctx, 'rev_solutions', {})}")
    if ctx.summary and ctx.summary.get("valuation"):
        v = ctx.summary["valuation"]
        print(f"engine:        WACC {v.get('wacc')}  EV Gordon {v.get('ev_gordon')}  price Gordon {v.get('price_gordon')}  "
              f"price exit {v.get('price_exit')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
