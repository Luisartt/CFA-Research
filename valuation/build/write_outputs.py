"""Publication gate v3 (valuation/valuation-outputs.yaml) and verification record (valuation/verification-v3.json).

Run after the v3 workbooks (EN and ES) are built, recalculated in Excel and scanned:
  python -I valuation/build/write_outputs.py --project .
It reads the workbooks' cached values, the engine summary and the Python replica (dcf_core.py). It refuses to write
when a formula has no cached result or shows an error, when a status cell is not OK, when the replica and the
workbook disagree, or when a Spanish copy's numbers differ from the English file.
"""
from __future__ import annotations

import argparse
import dataclasses
import hashlib
import json
import math
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import dcf_core as core  # noqa: E402

ERRORS = ("#REF!", "#N/A", "#DIV/0!", "#VALUE!", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
# Informational flags that document a known, explained difference (not a failed test). Anything else not OK stops.
EXPLAINED_FLAGS = {
    ("peers", "Checks!F9"): "PepsiCo EV vs the CIQ table: our claim basis adds operating leases 3,846 that CIQ excludes "
                           "(explained in Checks!G9; same as v2)",
}
OK_WORDS, BAD_WORDS = {"OK", "ok"}, {"FAIL", "CHECK", "FLAG", "WARN"}
FILES = {"dcf": "valuation/BIMBOA_DCF_v3.xlsx", "wacc": "valuation/BIMBOA_WACC_v3.xlsx",
         "peers": "valuation/BIMBOA_PeerAnalysis_v3.xlsx", "dcf_es": "valuation/BIMBOA_DCF_v3_ES.xlsx",
         "wacc_es": "valuation/BIMBOA_WACC_v3_ES.xlsx", "peers_es": "valuation/BIMBOA_PeerAnalysis_v3_ES.xlsx"}


def scan(path: Path) -> dict:
    """Formula count, missing caches, error values and status cells (OK vs FAIL/CHECK/FLAG/WARN) of one workbook."""
    wf, wv = load_workbook(path, data_only=False), load_workbook(path, data_only=True)
    out = {"formulas": 0, "missing_cache": [], "errors": [], "status_ok": 0, "status_not_ok": []}
    for ws in wf.worksheets:
        vs = wv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if not (isinstance(c.value, str) and c.value.startswith("=")):
                    continue
                out["formulas"] += 1
                v = vs[c.coordinate].value
                ref = f"{ws.title}!{c.coordinate}"
                if v is None:
                    out["missing_cache"].append(ref)
                elif isinstance(v, str) and v.strip() in ERRORS:
                    out["errors"].append(f"{ref} {v}")
                elif isinstance(v, str) and v.strip() in OK_WORDS:
                    out["status_ok"] += 1
                elif isinstance(v, str) and v.strip() in BAD_WORDS:
                    out["status_not_ok"].append(f"{ref} {v}")
    return out


def numbers(path: Path) -> dict:
    wv = load_workbook(path, data_only=True)
    return {(ws.title, c.coordinate): c.value for ws in wv.worksheets for row in ws.iter_rows() for c in row
            if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)}


def find_row(ws, label: str, col: int = 2) -> int:
    hits = [r for r in range(1, ws.max_row + 1) if str(ws.cell(r, col).value or "").strip() == label]
    if len(hits) != 1:
        raise SystemExit(f"{ws.title}: label {label!r} found {len(hits)} times")
    return hits[0]


def first_number(ws, r: int, start: int = 3):
    for c in range(start, ws.max_column + 1):
        v = ws.cell(r, c).value
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return float(v)
    raise SystemExit(f"{ws.title} row {r}: no number")


def r6(x):
    return None if x is None else round(float(x), 6)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--project", default=".")
    a = ap.parse_args(argv)
    P = Path(a.project).resolve()
    probs: list[str] = []

    # ---- 1. every workbook: hashes, caches, errors, status cells; ES numbers equal EN
    ver_files = {}
    for k, rel in FILES.items():
        p = P / rel
        if not p.is_file():
            raise SystemExit(f"missing {rel}")
        s = scan(p)
        fam = k.replace("_es", "")
        explained = [x for x in s["status_not_ok"] if (fam, x.split(" ")[0]) in EXPLAINED_FLAGS]
        s["status_not_ok"] = [x for x in s["status_not_ok"] if x not in explained]
        ver_files[k] = {"path": rel, "sha256": hashlib.sha256(p.read_bytes()).hexdigest(), "formulas": s["formulas"],
                        "missing_cache": len(s["missing_cache"]), "errors": len(s["errors"]),
                        "status_ok": s["status_ok"], "status_not_ok": s["status_not_ok"],
                        "explained_flags": {x: EXPLAINED_FLAGS[(fam, x.split(" ")[0])] for x in explained}}
        for what in ("missing_cache", "errors", "status_not_ok"):
            if s[what]:
                probs.append(f"{rel}: {what} {s[what][:5]}")
    for en, es in (("dcf", "dcf_es"), ("wacc", "wacc_es"), ("peers", "peers_es")):
        ne, ns = numbers(P / FILES[en]), numbers(P / FILES[es])
        diff = [k for k in set(ne) | set(ns) if not (k in ne and k in ns and
                                                    math.isclose(ne[k], ns[k], rel_tol=1e-12, abs_tol=1e-12))]
        ver_files[es]["numbers_equal_english"] = not diff
        ver_files[es]["numeric_cells"] = len(ns)
        if diff:
            probs.append(f"{FILES[es]}: {len(diff)} numeric cells differ from the English file, e.g. {sorted(diff)[:5]}")

    # ---- 2. DCF workbook: checks sheet, engine and replica reconciliation, dated check, scenarios
    wd = load_workbook(P / FILES["dcf"], data_only=True)
    ck = wd["Checks"]
    checks = [(str(ck.cell(r, 2).value), ck.cell(r, 3).value) for r in range(4, ck.max_row + 1)
              if ck.cell(r, 3).value in OK_WORDS | BAD_WORDS]
    er = wd["Engine_Recon"]
    rec = {k: first_number(er, find_row(er, lab), 4) for k, lab in (
        ("eng_n", "Rows checked against the engine summary"), ("eng_ok", "Rows OK vs engine (must equal the count above)"),
        ("py_n", "Rows checked against the Python replica"), ("py_ok", "Rows OK vs replica (must equal the count above)"))}
    if rec["eng_ok"] != rec["eng_n"] or rec["py_ok"] != rec["py_n"]:
        probs.append(f"Engine_Recon not fully OK: {rec}")
    wb_pub = {k: float(er.cell(find_row(er, lab), 5).value) for k, lab in (
        ("gordon", "PUBLISHED value per share today - Gordon (replica only)"),
        ("exit", "PUBLISHED value per share today - exit (replica only)"))}
    dc = wd["Dated_Check"]
    dated = {"gordon": first_number(dc, find_row(dc, "Value per share at the valuation date - Gordon")),
             "exit": first_number(dc, find_row(dc, "Value per share at the valuation date - exit"))}
    sc = wd["Scenarios"]
    scen = []
    for r in range(1, sc.max_row + 1):
        lab, g_, x_ = sc.cell(r, 2).value, sc.cell(r, 6).value, sc.cell(r, 7).value
        if isinstance(lab, str) and isinstance(g_, (int, float)) and isinstance(x_, (int, float)):
            scen.append({"scenario": lab, "wacc": r6(sc.cell(r, 3).value), "ke": r6(sc.cell(r, 4).value),
                         "g": r6(sc.cell(r, 5).value), "gordon": r6(g_), "exit": r6(x_),
                         "ronic_tv": r6(sc.cell(r, 10).value) if isinstance(sc.cell(r, 10).value, (int, float)) else None,
                         "note": str(sc.cell(r, 11).value or "")})

    # ---- 3. replica (published basis) and the cross-check with the workbook
    c = core.build_case(P)
    p0, e0 = core.value(c, basis="published"), core.value(c)
    for k, key in (("gordon", "price_gordon"), ("exit", "price_exit")):
        if abs(wb_pub[k] / p0[key] - 1) > 1e-6:
            probs.append(f"workbook published {k} {wb_pub[k]} != replica {p0[key]}")
    va = (c.extra.get("wacc_variants") or {}).get("rf") or {}
    mbono, dspread = float(va["mbono10"]["value"]), float(va["def_spread"]["value"])
    if abs(c.wacc_in["risk_free"] - (mbono - dspread)) > 1e-12:
        probs.append("valuation.yaml risk_free is not M-Bono - default spread (route B)")
    crp = core.crp_from_countries(c.extra)["company"]
    if abs(c.wacc_in["country_risk_premium"] - crp["ebitda_fy25"]) > 1e-12:
        probs.append("valuation.yaml CRP is not the FY2025 Adj. EBITDA-weighted company CRP")
    price, ke, wacc_ = c.price, p0["ke"], p0["wacc"]
    ps = p0["roll_factor"] / p0["bridge"]["shares"]
    n = p0["nci"]
    dps_next = 1.05

    def pub(case):
        return core.value(case, basis="published")

    def with_wacc(**kw):
        w = dict(c.wacc_in)
        w.update(kw)
        return dataclasses.replace(c, wacc_in=w)

    v2 = pub(with_wacc(risk_free=mbono, country_risk_premium=0.0))
    rf_only = pub(with_wacc(country_risk_premium=0.0))
    for k in ("price_gordon", "price_exit"):
        if abs(v2[k] - {"price_gordon": 35.717554, "price_exit": 46.604264}[k]) > 5e-7:
            probs.append(f"route A no longer reproduces the v2 published {k}: {v2[k]}")
    rev_g = core.solve_growth_shift(P, price, "gordon", basis="published")
    rev_x = core.solve_growth_shift(P, price, "exit", basis="published")
    lev = {name: pub(core.build_case(P, **kw))["price_gordon"] for name, kw in (
        ("capex_minus_50bp", {"capex_shift": -0.005}), ("capex_plus_50bp", {"capex_shift": 0.005}),
        ("gross_margin_plus_50bp", {"gm_shift": 0.005}), ("gross_margin_minus_50bp", {"gm_shift": -0.005}),
        ("growth_plus_50bp", {"growth_shift": 0.005}), ("growth_minus_50bp", {"growth_shift": -0.005}))}
    pe = yaml.safe_load((P / "valuation/inputs/peer_extras_v3.yaml").read_text(encoding="utf-8"))
    meds = pe["expected"]["medians"]
    wp = load_workbook(P / FILES["peers"], data_only=True)
    dp = wp["Detail_PeerAnalysis"]
    for key, cell in (("core_ev_ltm", "B50"), ("core_ev_ntm", "C50"), ("all_ev_ltm", "B53"), ("all_ev_ntm", "C53"),
                      ("all_pe_ltm", "E53"), ("all_pe_ntm", "F53")):
        if abs(float(dp[cell].value) - float(meds[key])) > 1e-6 * max(1.0, abs(float(meds[key]))):
            probs.append(f"peer median {key}: workbook {dp[cell].value} vs spec {meds[key]}")
    pck = wp["Checks"]
    sec, p_ok, p_n = False, 0, 0
    for r in range(1, pck.max_row + 1):
        lab = str(pck.cell(r, 1).value or "")
        if lab.startswith("10. v3"):
            sec = True
            continue
        if sec and isinstance(pck.cell(r, 3).value, (int, float)):
            p_n += 1
            p_ok += pck.cell(r, 4).value == "ok"
    if p_n == 0 or p_ok != p_n:
        probs.append(f"peer Checks section 10: {p_ok}/{p_n} ok")
    dbj = c.extra["dated_bridge"]
    rec_ = lambda k: float(dbj[k]["value"] if isinstance(dbj[k], dict) else dbj[k])  # noqa: E731
    claims_jun = rec_("debt_jun") + rec_("leases_jun") - rec_("cash_jun") + core.nci_claim(c, ke)["market"] - rec_("non_op_jun")
    shares_jun = rec_("shares_out")
    comps_v = lambda m: (m * 60919.0 - claims_jun) / shares_jun  # noqa: E731
    bimbo_mult = (price * shares_jun + claims_jun) / 60919.0
    summ = core.load_summary(P / "model" / "model-summary.json") or {}
    ev_ = summ.get("valuation") or {}
    engine_model = f"model/{Path(str(summ.get('model_file') or '')).name}"
    if not (P / engine_model).is_file():
        probs.append(f"engine model file {engine_model} not found")
    if any(x.get("status") != "OK" for x in (summ.get("checks") or [])) or summ.get("warnings"):
        probs.append("engine summary has checks that are not OK or warnings")
    if abs(float(ev_.get("wacc", 0)) - round(wacc_, 6)) > 1e-6:
        probs.append(f"engine summary WACC {ev_.get('wacc')} is stale vs {wacc_:.6f}: rebuild the engine model")
    if probs:
        raise SystemExit("NOT written:\n  " + "\n  ".join(probs))

    old = yaml.safe_load((P / "valuation/valuation-outputs.yaml").read_text(encoding="utf-8"))
    tv_share_x = p0["pv_tv_exit"] / p0["ev_exit"]
    gate = {
        "as_of": "2026-10-17", "version": 3, "valuation_date": "2026-10-16", "currency": "MXN",
        "share_price": old["share_price"],
        "inputs_status": "assistant proposals - not confirmed by the team",
        "sources": {"dcf": FILES["dcf"], "wacc": FILES["wacc"], "peers": FILES["peers"],
                    "engine": f"{engine_model}; model/model-summary.json",
                    "verification": "valuation/verification-v3.json"},
        "published": {
            "value_today_gordon": {"value": r6(p0["price_gordon"]), "status": "usable",
                                   "label": "value per share today, Gordon (published basis, route B proposal)"},
            "value_today_exit": {"value": r6(p0["price_exit"]), "status": "usable",
                                 "label": "value per share today, exit 6.0x, peer-EV-consistent claims"},
            "value_12m_gordon": {"value": r6(p0["price_gordon"] * (1 + ke) - dps_next), "status": "usable",
                                 "label": f"12-month value = {p0['price_gordon']:.4f} x (1 + ke {ke:.4%}) - next unpaid "
                                          f"dividend {dps_next}"},
            "next_unpaid_dividend": old["published"]["next_unpaid_dividend"],
            "premium_today_gordon": r6(p0["price_gordon"] / price - 1),
            "premium_today_exit": r6(p0["price_exit"] / price - 1),
            "upside_12m_gordon": r6((p0["price_gordon"] * (1 + ke) - dps_next) / price - 1),
            "wacc_route": {"value": "B", "status": "proposal",
                           "note": "rf = M-Bono 10y 9.26% - Mexico default spread 1.75% (Damodaran July 2026, Baa3); "
                                   "ke = rf + beta_l x ERP + CRP (CRP weighted by FY2025 Adj. EBITDA over Bimbo's countries); "
                                   "route A (M-Bono, CRP 0) kept as the alternative"},
            "risk_free": r6(c.wacc_in["risk_free"]), "country_risk_premium": r6(c.wacc_in["country_risk_premium"]),
            "wacc": round(wacc_, 10), "cost_of_equity": round(ke, 10), "terminal_growth": c.g,
            "exit_multiple": c.exit_multiple, "ev_gordon": round(p0["ev_gordon"], 2), "ev_exit": round(p0["ev_exit"], 2),
            "tv_share_gordon": r6(p0["tv_share_gordon"]), "tv_share_exit": r6(tv_share_x),
            "tv_ronic": {"value": r6(p0["tv_ronic"]),
                         "note": f"return on new capital implied by the published terminal value; WACC {wacc_:.2%}"},
            "terminal_net_reinvestment_rate": r6(c.g / p0["tv_ronic"]),
            "implied_exit_multiple_from_gordon": r6(p0["implied_exit_multiple"]),
            "implied_g_exit": r6(p0["implied_g_exit"]),
            "market_implied_terminal_growth": {"value": r6(p0["implied_g"]),
                                               "note": "above the 5.0% long-term nominal GDP cap [assumption]"},
            "reverse_dcf_growth_shift": {"gordon": r6(rev_g), "exit": r6(rev_x),
                                         "note": "uniform revenue-growth shift that equates value to price"},
            "lower_return_cases_gordon": {"ronic_equals_wacc": r6(core.ronic_case(c, p0, wacc_)),
                                          "ronic_12_5pct": r6(core.ronic_case(c, p0, 0.125)),
                                          "ronic_15pct": r6(core.ronic_case(c, p0, 0.15))},
            "bridge_judgements": {
                "without_nci_uplift_gordon": r6(p0["price_gordon"] + n["uplift"] * ps),
                "without_nci_uplift_exit": r6(p0["price_exit"] + n["uplift"] * ps),
                "nci_zero_growth_gordon": r6(p0["price_gordon"] + (n["uplift"] - n["uplift_g0"]) * ps),
                "status": "provisional",
                "note": f"NCI uplift {n['uplift']:,.2f} = NCI income 1,814 / (ke - 3%) - book 743, live at ke "
                        f"(Ps {n['uplift'] * ps:.2f} per share)"},
            "dated_cross_check": {"gordon": round(dated["gordon"], 4), "exit": round(dated["exit"], 4),
                                  "vs_fye_gordon": round(dated["gordon"] / p0["price_gordon"] - 1, 4),
                                  "vs_fye_exit": round(dated["exit"] / p0["price_exit"] - 1, 4), "status": "usable",
                                  "note": "30-Jun-2026 anchor (Dated_Check sheet); open: ebit_1h basis (~+Ps 0.33 each)"},
            "sensitivity": {"status": "usable",
                            "note": "published-basis grids in BIMBOA_DCF_v3.xlsx (NCI claim and roll factor at the base ke)"},
            "levers_gordon": {k: r6(v) for k, v in lev.items()},
            "comps_indicative": {
                f"core_bakery_{meds['core_ev_ltm']:.2f}x".replace(".", "_"): round(comps_v(float(meds["core_ev_ltm"])), 2),
                "exit_6_0x": round(comps_v(6.0), 2),
                f"all_eligible_{meds['all_ev_ltm']:.2f}x".replace(".", "_"): round(comps_v(float(meds["all_ev_ltm"])), 2),
                "bimbo_ltm_multiple_peer_claims": round(bimbo_mult, 3), "status": "provisional",
                "note": "LTM EBITDA 60,919; Jun-26 bridge; net debt incl. leases and NCI at market (at the base ke); "
                        "medians from BIMBOA_PeerAnalysis_v3.xlsx"},
            "scenarios": {"status": "usable", "note": "one change at a time against the base (DCF v3 'Scenarios' sheet)",
                          "rows": scen},
        },
        "engine_fields": {
            "wacc": {"engine": ev_.get("wacc"), "status": "usable"},
            "terminal_growth": {"engine": ev_.get("terminal_growth"), "status": "usable"},
            "exit_multiple": {"engine": ev_.get("exit_multiple"), "status": "usable"},
            "de_market": {"engine": ev_.get("de_market"), "status": "usable",
                          "note": "engine uses FY2025 debt and weighted shares; 0.743 at Jun-26 debt and outstanding shares"},
            "ev_gordon": {"engine": ev_.get("ev_gordon"), "status": "reconciliation", "use": "published.ev_gordon"},
            "ev_exit": {"engine": ev_.get("ev_exit"), "status": "reconciliation", "use": "published.ev_exit"},
            "price_gordon": {"engine": ev_.get("price_gordon"), "status": "superseded", "use": r6(p0["price_gordon"]),
                             "reason": "engine proxy (new leases = principal) keeps the lease liability flat while revenue grows"},
            "price_exit": {"engine": ev_.get("price_exit"), "status": "superseded", "use": r6(p0["price_exit"]),
                           "reason": "peer EVs exclude pensions, provisions, MEPP and derivatives; lease growth in the explicit years"},
            "dps_next": {"engine": ev_.get("dps_next"), "status": "superseded", "use": dps_next,
                         "reason": "FY2026 dividend already paid on 13-May-2026; next unpaid = consensus DPS FY2027E"},
            "target_12m_gordon": {"engine": ev_.get("target_12m_gordon"), "status": "superseded",
                                  "use": r6(p0["price_gordon"] * (1 + ke) - dps_next)},
            "tv_ronic": {"engine": ev_.get("tv_ronic"), "status": "superseded", "use": r6(p0["tv_ronic"])},
            "implied_g": {"engine": ev_.get("implied_g"), "status": "superseded", "use": r6(p0["implied_g"])},
            "price_comps_ev_ebitda": {"engine": ev_.get("price_comps_ev_ebitda"), "status": "superseded",
                                      "use": "published.comps_indicative",
                                      "reason": "engine applies FY2026 Bimbo EBITDA to NTM peer multiples of the P/E-eligible set"},
            "price_comps_pe": {"engine": ev_.get("price_comps_pe"), "status": "superseded",
                               "use": "see valuation/BIMBOA_PeerAnalysis_v3.xlsx"},
            "football": {"status": "provisional", "note": "52-week row usable; DCF and comps rows superseded by the published values"},
        },
        "v1_to_v2_walk": old["v1_to_v2_walk"],
        "v2_to_v3_walk": {
            "v2_published": [r6(v2["price_gordon"]), r6(v2["price_exit"])],
            "rf_default_spread_removed": [r6(rf_only["price_gordon"] - v2["price_gordon"]),
                                          r6(rf_only["price_exit"] - v2["price_exit"])],
            "crp_operations_weighted": [r6(p0["price_gordon"] - rf_only["price_gordon"]),
                                        r6(p0["price_exit"] - rf_only["price_exit"])],
            "route_b_cost_of_capital": [r6(p0["price_gordon"] - v2["price_gordon"]), r6(p0["price_exit"] - v2["price_exit"])],
            "v3_published": [r6(p0["price_gordon"]), r6(p0["price_exit"])],
        },
        "target_price": None,
        "rating": None,
    }
    header = (
        "# valuation/valuation-outputs.yaml - PUBLICATION GATE: the downstream source for the report, deck and pitch skills.\n"
        "# v3 (2026-10-17): route B cost of capital (proposal) with Damodaran's July 2026 country data; peer workbook v3\n"
        "#   (Capital IQ prompts 4-7, USD/MXN 18.479 at the pricing date, Gruma in the NTM and P/E medians).\n"
        "# PUBLISHED basis = valuation/BIMBOA_DCF_v3.xlsx key outputs: engine mechanics plus two technical corrections\n"
        "#   (1) FCFF also deducts the growth in the lease liability, held at the FY2025 share of revenue (8.15%);\n"
        "#   (2) the exit method subtracts only the claims inside the peer-EV definition (net debt incl. leases, NCI at market).\n"
        "# ENGINE basis = model engine summary: kept untouched for reconciliation only. Written by valuation/build/write_outputs.py\n"
        "#   after the Excel recalculation; verification record (hashes, checks) in valuation/verification-v3.json.\n"
        "# All values rest on assistant PROPOSALS (WACC route, terminal, forecast, bridge) until the team confirms them.\n"
        "# target_price and rating are the team's decisions (null). Values per share in MXN; money in MXN millions.\n")
    (P / "valuation/valuation-outputs.yaml").write_text(
        header + yaml.safe_dump(gate, sort_keys=False, allow_unicode=True, width=140), encoding="utf-8")
    dcf_s, wacc_s = ver_files["dcf"], ver_files["wacc"]
    verification = {
        "as_of": "2026-10-17", "files": ver_files,
        "formula_cells_dcf_wacc": dcf_s["formulas"] + wacc_s["formulas"],
        "formula_errors": sum(v["errors"] + v["missing_cache"] for v in ver_files.values()),
        "dcf": {"checks_ok": sum(1 for _, v in checks if v in OK_WORDS), "checks_total": len(checks),
                "engine_ok": int(rec["eng_ok"]), "engine_total": int(rec["eng_n"]),
                "replica_ok": int(rec["py_ok"]), "replica_total": int(rec["py_n"])},
        "peers": {"spec_ok": p_ok, "spec_total": p_n},
        "published": {"gordon": p0["price_gordon"], "exit": p0["price_exit"], "wacc": wacc_, "ke": ke},
        "engine_model": engine_model,
        "engine_checks": {"ok": sum(1 for x in (summ.get("checks") or []) if x.get("status") == "OK"),
                          "total": len(summ.get("checks") or [])},
    }
    (P / "valuation/verification-v3.json").write_text(json.dumps(verification, indent=1, ensure_ascii=False),
                                                      encoding="utf-8")
    print("wrote valuation/valuation-outputs.yaml (v3) and valuation/verification-v3.json")
    print(f"  published Gordon {p0['price_gordon']:.6f}  exit {p0['price_exit']:.6f}  WACC {wacc_:.6%}  ke {ke:.6%}")
    print(f"  DCF checks {verification['dcf']['checks_ok']}/{verification['dcf']['checks_total']}  engine "
          f"{verification['dcf']['engine_ok']}/{verification['dcf']['engine_total']}  replica "
          f"{verification['dcf']['replica_ok']}/{verification['dcf']['replica_total']}  peers {p_ok}/{p_n}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
