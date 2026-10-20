"""Convert a project markdown note (headings, paragraphs, lists, tables, quotes) to a Word file.

Supports the subset used in valuation/*.md: #/##/### headings, paragraphs, '-' bullets (nested by indent),
'1.' numbered items (numbers written literally, so lists never run on), indented continuation lines,
pipe tables (also indented under a list item), '>' quote blocks, and inline **bold** and `code`.

Usage: python -I valuation/build/md_to_docx.py <input.md> <output.docx> [--footer "text"]
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x59, 0x59, 0x59)
CODE = RGBColor(0x7F, 0x3F, 0x00)
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
NUMBERED = re.compile(r"^(\d+)\.\s+(.*)$")


def add_runs(par, text: str, *, size: float | None = None, color: RGBColor | None = None, bold: bool = False,
             italic: bool = False) -> None:
    text = text.replace("\\*", "*")
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            _run(par, text[pos:m.start()], size, color, bold, italic)
        tok = m.group(0)
        if tok.startswith("**"):
            add_runs(par, tok[2:-2], size=size, color=color, bold=True, italic=italic)
        else:
            r = _run(par, tok[1:-1], size, color if color is not None else CODE, bold, italic)
            r.font.name = "Consolas"
            r.font.size = Pt((size or 10) - 1)
        pos = m.end()
    if pos < len(text):
        _run(par, text[pos:], size, color, bold, italic)


def _run(par, text, size, color, bold, italic):
    r = par.add_run(text)
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    r.bold = bold or None
    r.italic = italic or None
    return r


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _row_flag(row, tag: str) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement(f"w:{tag}")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _plain(s: str) -> str:
    return s.replace("**", "").replace("`", "")


def add_table(doc, rows: list[list[str]], usable_cm: float = 17.59) -> None:
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    lens = [max(len(_plain(r[j])) if j < len(r) else 0 for r in rows) for j in range(ncol)]
    weights = [min(max(n, 10), 48) for n in lens]
    widths = [Cm(usable_cm * w / sum(weights)) for w in weights]
    keep_whole = len(rows) <= 16  # small tables stay on one page
    for i, row in enumerate(rows):
        _row_flag(t.rows[i], "cantSplit")
        if i == 0:
            _row_flag(t.rows[i], "tblHeader")  # repeats on the next page if the table breaks
        for j in range(ncol):
            cell = t.cell(i, j)
            cell.width = widths[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.keep_with_next = (i == 0) or (keep_whole and i < len(rows) - 1)
            txt = row[j] if j < len(row) else ""
            if i == 0:
                add_runs(par, txt, size=8.5, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
                shade(cell, "1F3864")
            else:
                add_runs(par, txt, size=8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def split_row(line: str) -> list[str]:
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def is_sep(line: str) -> bool:
    return bool(re.match(r"^\s*\|?\s*:?-{3,}", line))


def add_footer(doc, text: str) -> None:
    sec = doc.sections[0]
    par = sec.footer.paragraphs[0]
    add_runs(par, text + "  |  ", size=8, color=GREY)
    run = par.add_run()
    for kind, val in (("begin", None), ("instr", "PAGE"), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = val
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)
    run.font.size = Pt(8)


def convert(src: Path, out: Path, footer: str | None) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)  # Letter
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.0))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    for name, size in (("Title", 17), ("Heading 1", 13), ("Heading 2", 11.5)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = NAVY
        st.font.bold = True
    if footer:
        add_footer(doc, footer)

    lines = src.read_text(encoding="utf-8").splitlines()
    i, para_buf = 0, []

    def flush():
        if para_buf:
            p = doc.add_paragraph()
            add_runs(p, " ".join(s.strip() for s in para_buf))
            para_buf.clear()

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.lstrip()
        indent = len(line) - len(stripped)
        if not stripped:
            flush()
            i += 1
            continue
        if stripped.startswith("|"):
            flush()
            if doc.paragraphs:  # keep the lead-in line with its table
                doc.paragraphs[-1].paragraph_format.keep_with_next = True
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not is_sep(lines[i]):
                    rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        if indent == 0 and line.startswith("#"):
            flush()
            level = len(line) - len(line.lstrip("#"))
            text = line[level:].strip()
            h = doc.add_heading(level=0 if level == 1 else min(level - 1, 3))
            add_runs(h, text)
            i += 1
            continue
        if stripped.startswith(">"):
            flush()
            body = stripped[1:].strip()
            if body.startswith("- "):
                p = doc.add_paragraph(style="List Bullet 2")
                add_runs(p, body[2:], color=GREY, size=9.5)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                add_runs(p, body, color=GREY, size=9.5)
            i += 1
            continue
        if stripped.startswith("- "):
            flush()
            style = "List Bullet" if indent < 2 else ("List Bullet 2" if indent < 5 else "List Bullet 3")
            p = doc.add_paragraph(style=style)
            add_runs(p, stripped[2:])
            i += 1
            continue
        m = NUMBERED.match(stripped)
        if m:
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.65 + 0.6 * (indent // 3))
            p.paragraph_format.first_line_indent = Cm(-0.65)
            add_runs(p, f"{m.group(1)}.\t{m.group(2)}")
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.65 + 0.6 * (indent // 3)))
            i += 1
            continue
        if indent >= 3:  # continuation paragraph under a list item
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.65)
            add_runs(p, stripped)
            i += 1
            continue
        para_buf.append(stripped)
        i += 1
    flush()
    doc.save(out)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("out")
    ap.add_argument("--footer", default=None)
    a = ap.parse_args(argv)
    convert(Path(a.src), Path(a.out), a.footer)
    print("wrote", a.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
