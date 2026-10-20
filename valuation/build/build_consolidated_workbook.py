"""Build valuation/BIMBOA_DCF_consolidated_v1.xlsx from valuation/inputs/consolidated_results.json.

Live formulas: Inputs (weights, facts, judgments, blended drivers), WACC, DCF (FCFF, exit, Gordon,
12-month view) and Sensitivity. Values written from the Python replica: Compare, Walks, Risk_ESG,
Football and the replica checks (valuation/build/consolidate_models.py). Run that script first,
then this one, then valuation/build/recalc_excel.ps1 and scan_errors.py.

Usage (from the project folder): python -I valuation/build/build_consolidated_workbook.py --project .
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "valuation/BIMBOA_DCF_consolidated_v1.xlsx"
SRC = "valuation/inputs/consolidated_results.json"

NAVY = "1F3864"
F_TITLE = Font(name="Calibri", size=14, bold=True, color=NAVY)
F_SUB = Font(name="Calibri", size=10, italic=True, color="595959")
F_HDR = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_BOLD = Font(name="Calibri", size=10, bold=True)
F_NORM = Font(name="Calibri", size=10)
F_INPUT = Font(name="Calibri", size=10, color="0000FF")
F_LINK = Font(name="Calibri", size=10, color="00803C")
F_NOTE = Font(name="Calibri", size=9, italic=True, color="595959")
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_SUB = PatternFill("solid", fgColor="D9E1F2")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_KEY = PatternFill("solid", fgColor="E2EFDA")
TOP = Border(top=Side(style="thin", color="000000"))
WRAP = Alignment(wrap_text=True, vertical="top")
FMT = {"m": '#,##0.0;(#,##0.0);"-"', "m0": '#,##0;(#,##0);"-"', "p": '0.00%;(0.00%);"-"', "p1": '0.0%;(0.0%);"-"',
       "x": '0.00"x"', "ps": '0.00;(0.00);"-"', "n4": '0.0000', "n": '0.00', "d": '0', "bp": '+0.00;-0.00;"-"'}
YEARS = ["FY2026E", "FY2027E", "FY2028E", "FY2029E", "FY2030E"]
YC = ["B", "C", "D", "E", "F"]  # year columns on the Inputs driver block
DC = ["C", "D", "E", "F", "G"]  # year columns on the DCF sheet


def put(ws, addr, value, fmt=None, font=None, fill=None, wrap=False, bold=False):
    c = ws[addr]
    c.value = value
    if fmt:
        c.number_format = FMT.get(fmt, fmt)
    if font is not None:
        c.font = font
    elif isinstance(value, str) and value.startswith("="):
        c.font = F_NORM
    elif isinstance(value, (int, float)):
        c.font = F_INPUT
    else:
        c.font = F_BOLD if bold else F_NORM
    if fill is not None:
        c.fill = fill
    if wrap:
        c.alignment = WRAP
    return c


def header(ws, row, labels, start=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row, start + i, lab)
        c.font, c.fill = F_HDR, FILL_HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")


def section(ws, row, text, ncol=7):
    for i in range(1, ncol + 1):
        ws.cell(row, i).fill = FILL_SUB
    c = ws.cell(row, 1, text)
    c.font = F_BOLD


def title(ws, text, sub):
    ws["A1"].value, ws["A1"].font = text, F_TITLE
    ws["A2"].value, ws["A2"].font = sub, F_SUB


def widths(ws, ws_widths):
    for col, w in ws_widths.items():
        ws.column_dimensions[col].width = w


# ---------------------------------------------------------------- Inputs


def build_inputs(wb, R):
    ws = wb.create_sheet("Inputs")
    A, D, F = R["inputs"]["alejandro"], R["inputs"]["diego"], R["inputs"]["facts"]
    P, V3 = R["inputs"]["peers"], R["inputs"]["v3"]
    title(ws, "Consolidated DCF inputs - Grupo Bimbo (BIMBO A)",
          "MXN millions unless stated. Valuation date 16-Oct-2026; balance sheet and cash-flow anchor 30-Jun-2026. "
          "Blue = typed value; black = formula; yellow = decision cell the team can change.")
    put(ws, "A3", "Rule: facts take one sourced value; mechanics take one convention (the technically correct one); "
                  "only judgments are averaged, with weight w on the v3 set. Alejandro's model carries the v3 inputs, so "
                  "it is not a third independent set.", font=F_NOTE)
    header(ws, 5, ["Input", "v3 / Alejandro", "Diego", "Consolidated (used)", "Tag", "Bucket", "Rule and source"])
    r = 6
    put(ws, f"A{r}", "Weight on the v3 input set (w)")
    put(ws, f"D{r}", 0.5, "n4", fill=FILL_IN)
    put(ws, f"E{r}", "[assumption]"), put(ws, f"F{r}", "Method")
    put(ws, f"G{r}", "0.5 = two independent input sets (recommended). 0.6667 = one third per model, which gives v3 two "
                     "thirds because Alejandro's model carries the v3 inputs.", wrap=True)
    put(ws, "A7", "Diego exit-multiple basis (clean / raw)")
    put(ws, "D7", "clean", fill=FILL_IN, font=F_INPUT)
    put(ws, "E7", "[assumption]"), put(ws, "F7", "Method")
    put(ws, "G7", "clean = Diego's own rule (median of his four core peers) on the lease-consistent peer table (v3); "
                  "raw = his 6.4x, built on US GAAP multiples with operating-lease cost inside EBITDA and Gruma's CIQ EBITDA.",
        wrap=True)
    put(ws, "A8", "NCI treatment (market / book)")
    put(ws, "D8", "market", fill=FILL_IN, font=F_INPUT)
    put(ws, "E8", "[assumption]"), put(ws, "F8", "Mechanic (provisional)")
    put(ws, "G8", "market = NCI income / (ke - NCI growth): BBU independent operators withdraw more than their book "
                  "equity every year (valuation.md section 5). book = 721.7 at 30-Jun-2026. Open team decision.", wrap=True)
    dv1 = DataValidation(type="list", formula1='"clean,raw"', allow_blank=False)
    dv2 = DataValidation(type="list", formula1='"market,book"', allow_blank=False)
    ws.add_data_validation(dv1), ws.add_data_validation(dv2)
    dv1.add("D7"), dv2.add("D8")

    section(ws, 10, "A. Facts (one sourced value; not averaged)")
    facts = [
        (11, "Share price, 16-Oct-2026 (MXN)", 58.60, 58.60, F["price"], "ps", "[sourced]", "Fact",
         "BMV official close (PPP) 16-Oct-2026"),
        (12, "Shares outstanding (millions)", F["shares_m"], D["shares_m"], F["shares_m"], "n4", "[sourced]", "Fact",
         "BMV notice 29-Jul-2026 (4,294.9m, after the April-2026 cancellation). Diego: Capital IQ 30-Jun-2026"),
        (13, "Financial debt, 30-Jun-2026", F["fin_debt"], None, F["fin_debt"], "m", "[sourced]", "Fact",
         "2Q26 release p.6; ties to Capital IQ debt ex leases and derivatives (wiki ciq-debt-stack-and-maturity-ladder)"),
        (14, "Lease liabilities (IFRS 16), 30-Jun-2026", F["leases_bs"], None, F["leases_bs"], "m", "[sourced]", "Fact",
         "Capital IQ balance sheet 2026 FQ2 (7,113 + 26,248)"),
        (15, "Debt incl. leases", F["total_debt"], D["debt_ciq"], "=D13+D14", "m", "[calc]", "Fact",
         "[calc] reconciliation, for Diego to confirm: CIQ total debt is MXN 189,679 = 153,660 + 33,361 + 2,658 "
         "derivative adjustments (wiki ciq-debt-stack-and-maturity-ladder). CIQ's peso and dollar figures imply 17.47 "
         "MXN/US$; Diego reconverted US$10,855m at 17.96 (+5,239), the rate that wiki bimbo-wacc-market-inputs-capiq-"
         "2026-10 says CIQ used. Evidence for 17.47: US$948m cash x 17.47 = 16,565 vs company 16,563"),
        (16, "Cash and equivalents, 30-Jun-2026", F["cash"], D["cash_ciq"], F["cash"], "m", "[sourced]", "Fact",
         "Company 16,563 (net debt 137,097, 2Q26 release p.6). Diego's 17,023 = US$948m at 17.96 (+460)"),
        (17, "NCI annual income", F["nci_income"], None, F["nci_income"], "m", "[sourced]", "Fact",
         "1H26 x 2 (Alejandro Datos_Bimbo D18; valuation.md section 5)"),
        (18, "NCI income growth (for the market value)", 0.03, None, 0.03, "p", "[assumption]", "Judgment (v3)",
         "Tied to North America (about 3% in pesos); valuation/inputs/spg_inputs.yaml nci_claim"),
        (19, "NCI at book", F["nci_book"], D["nci_book"], F["nci_book"], "m", "[sourced]", "Fact",
         "30-Jun-2026 (Diego: Dec-2025 book 743)"),
        (20, "Investments in associates", F["assoc"], 0, F["assoc"], "m", "[sourced]", "Mechanic",
         "Non-operating asset; its results sit below EBIT, so the DCF does not capture it. Diego omits it"),
        (21, "Seasonal income tax payable (1H26, paid in 2H26)", F["tax_seasonal"], 0, F["tax_seasonal"], "m", "[sourced]",
         "Mechanic", "The 2H26 flow is counted in full, so the tax accrued in 1H26 that inflates June cash is deducted"),
        (22, "Other debt-like claims (Gordon only)", F["debt_like"], 0, F["debt_like"], "m", "[sourced]", "Mechanic",
         "Pensions, provisions incl. uncertain tax positions, derivatives (Alejandro Datos_Bimbo D21). Peer EVs exclude "
         "them, so the exit method does not deduct them (v3 rule)"),
        (23, "FY2025 revenue", F["rev0"], F["rev0"], F["rev0"], "m", "[sourced]", "Fact", "EEFF25; data/financials.csv"),
        (24, "Lease principal paid / revenue", 0.0139, D["lease_principal"], 0.0139, "p", "[calc]", "Mechanic",
         "FY2025 5,923 / 426,952 (model/drivers.yaml rounds to 1.39%)"),
        (25, "Lease liability / revenue (held)", R["inputs"]["lease_liability_ratio"], None,
         R["inputs"]["lease_liability_ratio"], "p", "[calc]", "Mechanic",
         "FY2025 34,790 / 426,952. New leases = principal + growth in the liability (v3 published basis). Diego's "
         "principal-only basis keeps the liability flat while revenue grows, so it understates reinvestment"),
        (26, "Stub days (30-Jun to 31-Dec-2026)", A["stub_days"], D["stub_days"], A["stub_days"], "d", "[calc]", "Mechanic",
         "Cash flows and the balance sheet on the same date (30-Jun). Diego discounts from 16-Oct with a 76-day stub "
         "but keeps the 30-Jun balance sheet, which drops about 3.5 months of cash flow"),
        (27, "Days from 30-Jun to 16-Oct-2026 (roll-forward at ke)", 108, 0, 108, "d", "[calc]", "Mechanic",
         "No dividend was paid between the two dates (the 2026 dividend was paid 13-May-2026)"),
        (28, "Next unpaid dividend per share", R["dps_next"], None, R["dps_next"], "ps", "[sourced]", "Fact",
         "Consensus DPS FY2027E (CIQ 20-Sep-2026); valuation-outputs.yaml"),
    ]
    for row, lab, v3v, dv, used, fmt, tag, bucket, note in facts:
        put(ws, f"A{row}", lab)
        if v3v is not None:
            put(ws, f"B{row}", v3v, fmt)
        if dv is not None:
            put(ws, f"C{row}", dv, fmt)
        put(ws, f"D{row}", used, fmt)
        put(ws, f"E{row}", tag), put(ws, f"F{row}", bucket), put(ws, f"G{row}", note, wrap=True)

    section(ws, 30, "B. Cost of capital (same sources in all three models; only the beta and the debt figure differ)")
    coc = [
        (31, "Risk-free: M-Bono 10y 9.26% - Mexico default spread 1.75%", A["rf"], D["rf"], A["rf"], "p", "[sourced]/[calc]",
         "Fact", "Capital IQ 16-Oct-2026; Damodaran July 2026 (Baa3)"),
        (32, "Equity risk premium (mature market)", A["rm"] - A["rf"], D["rm"] - D["rf"], A["rm"] - A["rf"], "p", "[sourced]",
         "Fact", "Damodaran implied ERP"),
        (33, "Country risk premium, weighted by FY2025 Adj. EBITDA", A["crp"], D["crp"], A["crp"], "p", "[calc]", "Fact",
         "Same method and weights in all three models (NA 28.8%, Mexico 55.4%, EAA 9.8%, LatAm 6.0%); added once to ke"),
        (34, "Pre-tax cost of debt: M-Bono + 45 bp (BIMBO 26)", A["kd"], D["kd"], A["kd"], "p", "[sourced]/[calc]", "Fact",
         "Issue spread of the Feb-2026 bond"),
        (35, "Marginal tax rate (tax shield, relevering)", A["tax_marg"], D["tax_marg"], A["tax_marg"], "p", "[sourced]",
         "Fact", "Mexican statutory rate"),
        (36, "D/E, market value, debt incl. leases", A["de"], D["de"], "=D15/(D11*D12)", "n4", "[calc]", "Fact",
         "Both chose the current market structure; Diego's differs only through his debt figure (row 15)"),
        (37, "Unlevered beta", A["beta_u"], D["beta_u"], "=$D$6*B37+(1-$D$6)*C37", "n4", "[sourced]", "Judgment",
         "Damodaran Food Processing (Global): v3 0.53 (Jan-2026 file); Diego 0.56 (cash-corrected)"),
    ]
    for row, lab, v3v, dv, used, fmt, tag, bucket, note in coc:
        put(ws, f"A{row}", lab), put(ws, f"B{row}", v3v, fmt), put(ws, f"C{row}", dv, fmt), put(ws, f"D{row}", used, fmt)
        put(ws, f"E{row}", tag), put(ws, f"F{row}", bucket), put(ws, f"G{row}", note, wrap=True)

    section(ws, 39, "C. Terminal value and tax on operating profit (judgments, averaged)")
    put(ws, "A40", "Exit EV/EBITDA on FY2030 EBITDA")
    put(ws, "B40", V3["exit_mult"], "x"), put(ws, "C40", D["exit_mult"], "x")
    put(ws, "D40", '=$D$6*B40+(1-$D$6)*IF($D$7="raw",C40,C41)', "x")
    put(ws, "E40", "[assumption]"), put(ws, "F40", "Judgment")
    put(ws, "G40", "v3: judgment above the core bakery median (5.53x) for the Mexico mix. Diego: median of his core four "
                   "(Flowers, Gruma, Yamazaki, Aryzta) = 6.37x on his table", wrap=True)
    put(ws, "A41", "Diego's rule on the lease-consistent peer table (row 52)")
    put(ws, "C41", "=B52", "x"), put(ws, "E41", "[calc]"), put(ws, "F41", "Data basis")
    put(ws, "G41", "Same four peers, EV/EBITDA LTM on one claim basis (valuation/inputs/peers.yaml)", wrap=True)
    put(ws, "A42", "Terminal growth g (nominal MXN)")
    put(ws, "B42", A["g"], "p"), put(ws, "C42", D["g"], "p"), put(ws, "D42", "=$D$6*B42+(1-$D$6)*C42", "p")
    put(ws, "E42", "[assumption]"), put(ws, "F42", "Judgment")
    put(ws, "G42", "v3 4.25% (regional mix, below 5% nominal GDP). Diego 4.0% (2% real + ~3% inflation, USD-linked mix)",
        wrap=True)
    put(ws, "A43", "Tax on EBIT (cash flows)")
    put(ws, "B43", A["tax_ebit"][0], "p"), put(ws, "C43", D["tax_ebit"][0], "p"), put(ws, "D43", "=$D$6*B43+(1-$D$6)*C43", "p")
    put(ws, "E43", "[assumption]"), put(ws, "F43", "Judgment (flagged)")
    put(ws, "G43", "v3 34% (effective rate 33-36% ex-2022: non-deductibles, foreign rates). Diego 30% (statutory). The "
                   "effective rate is the better estimate of cash taxes; the average is shown because the user asked "
                   "for averages", wrap=True)
    put(ws, "A44", "Return on new capital after 2030 (Gordon)")
    put(ws, "B44", V3["ronic"], "p"), put(ws, "D44", "=B44", "p", fill=FILL_IN)
    put(ws, "E44", "[assumption]"), put(ws, "F44", "Judgment (v3 only)")
    put(ws, "G44", "RONIC implied by the published v3 terminal value (valuation-outputs.yaml tv_ronic). Diego has no "
                   "Gordon or reinvestment judgment, so nothing to average", wrap=True)

    section(ws, 46, "D. Peer EV/EBITDA LTM used for Diego's rule")
    header(ws, 47, ["Peer", "Lease-consistent (v3 table)", "Diego's table", "Gap", "", "", "Why they differ"])
    notes = {"FLO": "US GAAP: v3 adds operating-lease cost to EBITDA because CIQ debt includes leases; Diego does not",
             "GRUMA": "Company EBITDA 1,108.9 (v3) vs Capital IQ 995.0 (Diego flagged it as open)",
             "YMZ": "Same figure; J-GAAP lease basis not normalisable (excluded from the v3 medians)",
             "ARYN": "Same figure; NCI not verified (excluded from the v3 medians)"}
    names = {"FLO": "Flowers Foods", "GRUMA": "Gruma", "YMZ": "Yamazaki Baking", "ARYN": "Aryzta"}
    for i, k in enumerate(["FLO", "GRUMA", "YMZ", "ARYN"]):
        row = 48 + i
        put(ws, f"A{row}", names[k]), put(ws, f"B{row}", P["clean_ltm"][k], "x"), put(ws, f"C{row}", P["diego_raw"][k], "x")
        put(ws, f"D{row}", f"=C{row}-B{row}", "x"), put(ws, f"G{row}", notes[k], wrap=True)
    put(ws, "A52", "Median of the four", bold=True)
    put(ws, "B52", "=MEDIAN(B48:B51)", "x"), put(ws, "C52", "=MEDIAN(C48:C51)", "x"), put(ws, "D52", "=C52-B52", "x")
    put(ws, "G52", "B52 feeds row 41 (clean). C52 reproduces Diego's 6.37x", wrap=True)

    section(ws, 54, "E. Forecast drivers (judgments, averaged year by year)")
    header(ws, 55, ["Driver"] + YEARS + ["Tag and source"])
    blocks = [
        ("Revenue growth", "growth", "v3: sum of regional builds (model/drivers.yaml). Diego: Capital IQ consensus "
                                    "FY26-28 (426,275 / 445,288 / 468,431), fade 4.7% / 4.3%"),
        ("EBITDA margin (IFRS 16)", "ebitda_m", "v3: gross margin minus opex (pillars 1-2), to 15.3%. Diego: consensus "
                                               "EBITDA / revenue, held at 14.73% after 2028"),
        ("D&A / revenue", "da_pct", "v3: 5.85-5.9%. Diego: template mechanics, D&A rises to 100% of capex + leases by "
                                    "2030 (no net investment while revenue grows). Mechanic, so the v3 path is used"),
        ("Capex excluding leases / revenue", "capex_pct", "v3: 4.6% (2026 guide) to 5.5%. Diego: consensus 5.16% / "
                                                          "5.02% / 4.90%, then 4.85% / 4.80%"),
        ("Change in working capital / revenue (+ releases cash)", "wc_pct", "Both hold FY2025 days; schedules differ"),
    ]
    row = 56
    rows_used = {}
    for lab, key, note in blocks:
        put(ws, f"A{row}", f"{lab} - v3 / Alejandro")
        put(ws, f"A{row + 1}", f"{lab} - Diego")
        put(ws, f"A{row + 2}", f"{lab} - consolidated", bold=True)
        mechanic = key == "da_pct"
        for j, col in enumerate(YC):
            put(ws, f"{col}{row}", A[key][j], "p")
            put(ws, f"{col}{row + 1}", D[key][j], "p")
            f = f"={col}{row}" if mechanic else f"=$D$6*{col}{row}+(1-$D$6)*{col}{row + 1}"
            put(ws, f"{col}{row + 2}", f, "p", fill=FILL_KEY)
        put(ws, f"G{row}", "[sourced]/[assumption] " + note, wrap=True)
        put(ws, f"G{row + 2}", "[calc] mechanic: v3 path for every w" if mechanic else "[calc] w x v3 + (1 - w) x Diego")
        rows_used[key] = row + 2
        row += 3
    put(ws, f"A{row}", "Reference, not used: new leases / revenue - v3 / Alejandro")
    put(ws, f"A{row + 1}", "Reference, not used: new leases / revenue - Diego (principal only)")
    for j, col in enumerate(YC):
        put(ws, f"{col}{row}", A["lease_pct"][j], "p"), put(ws, f"{col}{row + 1}", D["lease_pct"][j], "p")
    put(ws, f"G{row}", "The consolidated DCF recomputes new leases from rows 24-25 on its own revenue path (DCF row 16)",
        wrap=True)
    section(ws, 74, "F. FY2025 actuals (base year: linked to the IS, BS and CF sheets; shown in DCF column B for reference, not part of the valuation)")
    header(ws, 75, ["Item", "", "", "FY2025A", "Tag", "Bucket", "Source"])
    actuals = [
        (76, "FY2024 revenue (base for FY2025 growth)", f"=IS!E{IS_ROWS['rev']}", "[sourced]", "IS sheet (IA25-EEFF p.13)"),
        (77, "FY2025 Adjusted EBITDA", f"=IS!F{IS_ROWS['adj']}", "[sourced]",
         "IS sheet (IA25-EEFF p.85-86 segment note: operating income + D&A + impairment + non-cash items)"),
        (78, "FY2025 D&A", f"=IS!F{IS_ROWS['da']}", "[sourced]", "IS sheet (IA25-EEFF p.85-86, continuing operations)"),
        (79, "FY2025 capex (PP&E acquired)", f"=CF!F{CF_ROWS['capex_ppe']}", "[sourced]",
         "CF sheet (IA25-EEFF p.85-86); capex on the cash-flow statement is 22,841 (CF sheet row 6)"),
        (80, "FY2025 lease principal paid", f"=CF!F{CF_ROWS['lease']}", "[calc]", "CF sheet; numerator of row 24"),
        (81, "FY2025 change in working capital (+ releases cash)", f"=CF!F{CF_ROWS['dnwc']}", "[calc]",
         "CF sheet: decrease in operating working capital on the balance sheet (same definition as the forecast driver), not a "
         "reported cash-flow line"),
    ]
    for row, lab, val, tag, note in actuals:
        put(ws, f"A{row}", lab), put(ws, f"D{row}", val, "m", font=F_LINK), put(ws, f"E{row}", tag), put(ws, f"F{row}", "Link")
        put(ws, f"G{row}", note, wrap=True)
    widths(ws, {"A": 52, "B": 15, "C": 15, "D": 17, "E": 15, "F": 15, "G": 90})
    ws.freeze_panes = "B6"
    return rows_used


# ---------------------------------------------------------------- WACC


def build_wacc(wb, R):
    ws = wb.create_sheet("WACC")
    title(ws, "Consolidated cost of capital", "ke = rf + beta_levered x ERP + CRP; beta_levered = beta_unlevered x (1 + (1 - t) x D/E)")
    header(ws, 3, ["Item", "Value", "Formula or source"])
    rows = [
        (4, "Unlevered beta", "=Inputs!D37", "n4", "Inputs row 37 (w-weighted)"),
        (5, "D/E (market, debt incl. leases)", "=Inputs!D36", "n4", "Inputs row 36"),
        (6, "Marginal tax rate", "=Inputs!D35", "p", "Inputs row 35"),
        (7, "Levered beta", "=B4*(1+(1-B6)*B5)", "n4", "Hamada"),
        (8, "Risk-free", "=Inputs!D31", "p", ""),
        (9, "Equity risk premium", "=Inputs!D32", "p", ""),
        (10, "Country risk premium", "=Inputs!D33", "p", "Added once, not multiplied by beta"),
        (11, "Cost of equity (ke)", "=B8+B7*B9+B10", "p", ""),
        (12, "Pre-tax cost of debt", "=Inputs!D34", "p", ""),
        (13, "After-tax cost of debt", "=B12*(1-B6)", "p", ""),
        (14, "Weight of debt D/(D+E)", "=B5/(1+B5)", "p", ""),
        (15, "Weight of equity", "=1-B14", "p", ""),
        (16, "WACC", "=B15*B11+B14*B13", "p", ""),
    ]
    for row, lab, f, fmt, note in rows:
        put(ws, f"A{row}", lab, bold=row in (11, 16)), put(ws, f"B{row}", f, fmt), put(ws, f"C{row}", note)
    ws["B16"].fill = FILL_KEY
    ws["B11"].fill = FILL_KEY
    section(ws, 18, "How the three WACCs compare", 3)
    A, D = R["inputs"]["alejandro"], R["inputs"]["diego"]
    comp = [
        (19, "v3 published (valuation-outputs.yaml)", R["inputs"]["v3"]["wacc"], "Same inputs as Alejandro"),
        (20, "Alejandro (WACC!C87 of his workbook)", A["wacc"], "beta 0.53, D/E 0.743"),
        (21, "Diego (WACC!C87 of his workbook)", D["wacc"], "beta 0.56, D/E 0.774 (CIQ debt at the LTM FX rate)"),
        (22, "Simple average of the two input sets", f"=(B20+B21)/2", ""),
        (23, "Simple average of the three models", f"=(B19+B20+B21)/3", ""),
        (24, "Consolidated (this sheet)", "=B16", "beta averaged, D/E reconciled to the company debt"),
    ]
    for row, lab, v, note in comp:
        put(ws, f"A{row}", lab), put(ws, f"B{row}", v, "p"), put(ws, f"C{row}", note)
    put(ws, "A26", "All three use the same risk-free, ERP, EBITDA-weighted country risk premium and cost of debt, so the "
                   "WACCs differ by less than 10 bp: there is little to average.", font=F_NOTE)
    widths(ws, {"A": 46, "B": 14, "C": 70})


# ---------------------------------------------------------------- DCF


def build_dcf(wb, R, drv_rows):
    ws = wb.create_sheet("DCF")
    title(ws, "Consolidated DCF - FCFF, exit multiple and Gordon",
          "S&P template mechanics: FY2026 stub from 30-Jun (184 days), mid-period discounting, terminal value at "
          "end-2030; value rolled to 16-Oct-2026 at ke. MXN millions.")
    header(ws, 4, ["MXN m", "FY2025A", *YEARS])
    labels = {5: "Revenue", 6: "  growth", 7: "  EBITDA margin", 8: "EBITDA (IFRS 16)", 9: "  D&A / revenue",
              10: "D&A", 11: "EBIT", 12: "Tax on EBIT", 13: "NOPAT", 14: "  capex ex leases / revenue",
              15: "Capex ex leases", 16: "New leases (principal + growth in the liability)", 17: "  change in WC / revenue",
              18: "Change in working capital (+ releases cash)", 19: "Free cash flow to the firm (FCFF)",
              20: "Share of the year counted", 21: "Discount exponent (years from 30-Jun-2026)", 22: "Discount factor",
              23: "PV of FCFF"}
    for row, lab in labels.items():
        put(ws, f"A{row}", lab, bold=row in (5, 8, 13, 19, 23))
    put(ws, "B5", "=Inputs!D23", "m")
    for j, col in enumerate(DC):
        prev = "B" if j == 0 else DC[j - 1]
        yc = YC[j]
        put(ws, f"{col}5", f"={prev}5*(1+{col}6)", "m")
        put(ws, f"{col}6", f"=Inputs!{yc}{drv_rows['growth']}", "p")
        put(ws, f"{col}7", f"=Inputs!{yc}{drv_rows['ebitda_m']}", "p")
        put(ws, f"{col}8", f"={col}5*{col}7", "m")
        put(ws, f"{col}9", f"=Inputs!{yc}{drv_rows['da_pct']}", "p")
        put(ws, f"{col}10", f"={col}5*{col}9", "m")
        put(ws, f"{col}11", f"={col}8-{col}10", "m")
        put(ws, f"{col}12", f"={col}11*Inputs!$D$43", "m")
        put(ws, f"{col}13", f"={col}11-{col}12", "m")
        put(ws, f"{col}14", f"=Inputs!{yc}{drv_rows['capex_pct']}", "p")
        put(ws, f"{col}15", f"={col}5*{col}14", "m")
        put(ws, f"{col}16", f"=Inputs!$D$24*{col}5+Inputs!$D$25*({col}5-{prev}5)", "m")
        put(ws, f"{col}17", f"=Inputs!{yc}{drv_rows['wc_pct']}", "p")
        put(ws, f"{col}18", f"={col}5*{col}17", "m")
        put(ws, f"{col}19", f"={col}13+{col}10-{col}15-{col}16+{col}18", "m")
        put(ws, f"{col}20", "=Inputs!$D$26/365" if j == 0 else 1, "n4", font=F_NORM)
        if j == 0:
            put(ws, f"{col}21", "=Inputs!$D$26/365/2", "n4")
        elif j == 1:
            put(ws, f"{col}21", "=Inputs!$D$26/365+0.5", "n4")
        else:
            put(ws, f"{col}21", f"={DC[j - 1]}21+1", "n4")
        put(ws, f"{col}22", f"=1/(1+WACC!$B$16)^{col}21", "n4")
        put(ws, f"{col}23", f"={col}19*{col}20*{col}22", "m")
    # FY2025A base year: reported lines from Inputs rows 76-81; tax on EBIT, NOPAT and FCFF are pro forma at the DCF's tax rate
    put(ws, "B6", "=B5/Inputs!D76-1", "p")
    put(ws, "B8", "=Inputs!D77", "m"), put(ws, "B7", "=B8/B5", "p")
    put(ws, "B10", "=Inputs!D78", "m"), put(ws, "B9", "=B10/B5", "p")
    put(ws, "B11", "=B8-B10", "m")
    put(ws, "B12", "=B11*Inputs!$D$43", "m", font=F_NOTE), put(ws, "B13", "=B11-B12", "m", font=F_NOTE)
    put(ws, "B15", "=Inputs!D79", "m"), put(ws, "B14", "=B15/B5", "p")
    put(ws, "B16", "=Inputs!D80", "m")
    put(ws, "B18", "=Inputs!D81", "m"), put(ws, "B17", "=B18/B5", "p")
    put(ws, "B19", "=B13+B10-B15-B16+B18", "m", font=F_NOTE)
    put(ws, "A3", "FY2025A = base year: reported revenue, Adjusted EBITDA, D&A, capex, lease principal paid (new leases) and "
                  "working capital (Inputs rows 76-81). Tax on EBIT, NOPAT and FCFF in grey are pro forma at the DCF's tax rate, "
                  "for comparison only. The valuation counts cash flows from 30-Jun-2026.", font=F_NOTE)
    for row in (8, 13, 19, 23):
        for col in ["A", "B", *DC]:
            ws[f"{col}{row}"].border = TOP

    section(ws, 25, "Valuation at 30-Jun-2026, rolled to 16-Oct-2026", 7)
    v = [
        (26, "PV of FCFF, 2H26-FY2030", "=SUM(C23:G23)", "m"),
        (27, "Years from 30-Jun-2026 to end-2030 (terminal value)", "=Inputs!D26/365+4", "n4"),
        (29, "Exit EV/EBITDA", "=Inputs!D40", "x"),
        (30, "Terminal value (FY2030 EBITDA x multiple)", "=G8*B29", "m"),
        (31, "PV of terminal value", "=B30/(1+WACC!B16)^B27", "m"),
        (32, "Enterprise value (exit)", "=B26+B31", "m"),
        (33, "Terminal value / EV", "=B31/B32", "p1"),
        (34, "Perpetual FCFF growth implied by the multiple (template formula)", "=(B30*WACC!B16-G19)/(B30+G19)", "p"),
        (36, "(-) Debt incl. leases", "=-Inputs!D15", "m"),
        (37, "(+) Cash", "=Inputs!D16", "m"),
        (38, "(-) Non-controlling interests", '=-IF(Inputs!D8="book",Inputs!D19,Inputs!D17/(WACC!B11-Inputs!D18))', "m"),
        (39, "(+) Associates - seasonal tax payable", "=Inputs!D20-Inputs!D21", "m"),
        (40, "Equity value, exit, 30-Jun-2026", "=B32+SUM(B36:B39)", "m"),
        (41, "Per share, 30-Jun-2026", "=B40/Inputs!D12", "ps"),
        (42, "Roll-forward factor to 16-Oct-2026 = (1 + ke)^(days / 365)", "=(1+WACC!B11)^(Inputs!D27/365)", "n4"),
        (43, "VALUE PER SHARE TODAY - EXIT MULTIPLE", "=B41*B42", "ps"),
        (44, "Premium / (discount) to the MXN 58.60 price", "=B43/Inputs!D11-1", "p1"),
        (46, "Terminal growth g", "=Inputs!D42", "p"),
        (47, "Return on new capital after 2030 (RONIC)", "=Inputs!D44", "p"),
        (48, "NOPAT FY2031 = NOPAT FY2030 x (1 + g)", "=G13*(1+B46)", "m"),
        (49, "Reinvestment rate = g / RONIC", "=B46/B47", "p1"),
        (50, "FCFF FY2031 = NOPAT x (1 - reinvestment rate)", "=B48*(1-B49)", "m"),
        (51, "Terminal value = FCFF FY2031 / (WACC - g)", "=B50/(WACC!B16-B46)", "m"),
        (52, "PV of terminal value", "=B51/(1+WACC!B16)^B27", "m"),
        (53, "Enterprise value (Gordon)", "=B26+B52", "m"),
        (54, "(-) Other debt-like claims (pensions, provisions, derivatives)", "=-Inputs!D22", "m"),
        (55, "Equity value, Gordon, 30-Jun-2026", "=B53+SUM(B36:B39)+B54", "m"),
        (56, "VALUE PER SHARE TODAY - GORDON (RONIC-consistent)", "=B55/Inputs!D12*B42", "ps"),
        (57, "Terminal value / FY2030 EBITDA", "=B51/G8", "x"),
        (59, "FCFF FY2031 = FCFF FY2030 x (1 + g)", "=G19*(1+B46)", "m"),
        (60, "Value per share today - Gordon, template style", "=(B26+B59/(WACC!B16-B46)/(1+WACC!B16)^B27+SUM(B36:B39)+B54)/Inputs!D12*B42", "ps"),
        (61, "RONIC that this template-style Gordon implies", "=B46/((G15+G16-G10-G18)/G13)", "p1"),
        (62, "Value per share today - Gordon with RONIC = WACC (growth creates no value)",
         "=(B26+B48*(1-B46/WACC!B16)/(WACC!B16-B46)/(1+WACC!B16)^B27+SUM(B36:B39)+B54)/Inputs!D12*B42", "ps"),
    ]
    for row, lab, f, fmt in v:
        put(ws, f"A{row}", lab, bold=row in (32, 43, 53, 56)), put(ws, f"B{row}", f, fmt)
    for row in (43, 56):
        ws[f"B{row}"].fill = FILL_KEY
    put(ws, "A28", "Exit multiple method", bold=True)
    put(ws, "A35", "Equity bridge at 30-Jun-2026 (each claim once)", bold=True)
    put(ws, "A45", "Gordon method, RONIC-consistent: terminal reinvestment = g / RONIC", bold=True)
    put(ws, "A58", "Gordon, template style (as in Alejandro's sheet): shown for reference; its implied RONIC is too high",
        bold=True)
    section(ws, 64, "12-month view (recommendation-guide convention; the team sets target and rating)", 7)
    header(ws, 65, ["Method", "Value today", "12-month value = value x (1 + ke) - next dividend",
                    "Expected return = 12m value / price - 1 + dividend / price", "", "", ""])
    for i, (lab, ref) in enumerate([("Exit multiple", "B43"), ("Gordon, RONIC-consistent", "B56"),
                                     ("Gordon, template style", "B60"), ("Gordon, RONIC = WACC", "B62")]):
        row = 66 + i
        put(ws, f"A{row}", lab), put(ws, f"B{row}", f"={ref}", "ps")
        put(ws, f"C{row}", f"=B{row}*(1+WACC!$B$11)-Inputs!$D$28", "ps")
        put(ws, f"D{row}", f"=C{row}/Inputs!$D$11-1+Inputs!$D$28/Inputs!$D$11", "p1")
    put(ws, "A71", "Convention in references/recommendation-guide.md: BUY at 15% or more, SELL at -10% or less, HOLD in "
                   "between. HOLD needs a value today of at least price x 0.9 / (1 + ke) = D72.", font=F_NOTE)
    put(ws, "A72", "Value today needed for the HOLD floor (-10%)"), put(ws, "D72", "=Inputs!D11*0.9/(1+WACC!B11)", "ps")
    put(ws, "A73", "Value today needed for the BUY floor (+15%)"), put(ws, "D73", "=Inputs!D11*1.15/(1+WACC!B11)", "ps")
    widths(ws, {"A": 62, "B": 14, "C": 14, "D": 16, "E": 14, "F": 14, "G": 14})
    ws.freeze_panes = "B5"


# ---------------------------------------------------------------- Sensitivity


def build_sens(wb, R):
    ws = wb.create_sheet("Sensitivity")
    title(ws, "Sensitivity of the consolidated value per share today (MXN)",
          "Live formulas. NCI claim and the roll-forward factor held at the base ke, as in the S&P template grids.")
    put(ws, "A4", "Exit method: WACC (rows) x exit EV/EBITDA (columns)", bold=True)
    mults = [5.0, 5.5, "=Inputs!D40", 6.5, 7.0]
    put(ws, "A5", "WACC \\ multiple", bold=True)
    for j, m in enumerate(mults):
        put(ws, f"{get_column_letter(2 + j)}5", m, "x", fill=FILL_SUB)
    for i, dw in enumerate([-0.01, -0.005, 0.0, 0.005, 0.01]):
        row = 6 + i
        put(ws, f"A{row}", f"=WACC!$B$16{dw:+.3f}" if dw else "=WACC!$B$16", "p", fill=FILL_SUB)
        for j in range(5):
            col = get_column_letter(2 + j)
            f = (f"=(SUMPRODUCT(DCF!$C$19:$G$19*DCF!$C$20:$G$20/(1+$A{row})^DCF!$C$21:$G$21)"
                 f"+DCF!$G$8*{col}$5/(1+$A{row})^DCF!$B$27+SUM(DCF!$B$36:$B$39))/Inputs!$D$12*DCF!$B$42")
            put(ws, f"{col}{row}", f, "ps")
    ws["D8"].fill = FILL_KEY
    put(ws, "A13", "Gordon (RONIC-consistent): WACC (rows) x terminal growth g (columns)", bold=True)
    put(ws, "A14", "WACC \\ g", bold=True)
    for j, dg in enumerate([-0.01, -0.005, 0.0, 0.005, 0.01]):
        put(ws, f"{get_column_letter(2 + j)}14", f"=Inputs!$D$42{dg:+.3f}" if dg else "=Inputs!$D$42", "p", fill=FILL_SUB)
    for i, dw in enumerate([-0.01, -0.005, 0.0, 0.005, 0.01]):
        row = 15 + i
        put(ws, f"A{row}", f"=WACC!$B$16{dw:+.3f}" if dw else "=WACC!$B$16", "p", fill=FILL_SUB)
        for j in range(5):
            col = get_column_letter(2 + j)
            f = (f"=(SUMPRODUCT(DCF!$C$19:$G$19*DCF!$C$20:$G$20/(1+$A{row})^DCF!$C$21:$G$21)"
                 f"+DCF!$G$13*(1+{col}$14)*(1-{col}$14/Inputs!$D$44)/($A{row}-{col}$14)/(1+$A{row})^DCF!$B$27"
                 f"+SUM(DCF!$B$36:$B$39)+DCF!$B$54)/Inputs!$D$12*DCF!$B$42")
            put(ws, f"{col}{row}", f, "ps")
    ws["D17"].fill = FILL_KEY
    put(ws, "A21", "Centre cells (green) equal DCF!B43 and DCF!B56 (see Checks).", font=F_NOTE)
    widths(ws, {"A": 22, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12})


# ---------------------------------------------------------------- Compare


def build_compare(wb, R):
    ws = wb.create_sheet("Compare")
    A, D, V3 = R["inputs"]["alejandro"], R["inputs"]["diego"], R["inputs"]["v3"]
    C = R["consolidated"]
    run, st = C["run"], C["state"]
    title(ws, "The three models side by side", "Values read from each workbook (cached results); consolidated = this "
                                                 "workbook at the base settings. Per-share values are for 16-Oct-2026.")
    header(ws, 4, ["Item", "v3 (valuation/)", "Alejandro", "Diego", "Consolidated", "Bucket", "What differs and the rule used"])
    rev_cagr = lambda rev, rev0: (rev[-1] / rev0) ** (1 / 5) - 1  # noqa: E731
    rows = [
        ("Provenance", None),
        ("Workbook", ["BIMBOA_DCF_v3.xlsx", "Bimbo_DCF_Alejandro.xlsm", "SPG_DCF_independent_Diego.xlsm",
                      "BIMBOA_DCF_consolidated_v1.xlsx", "", "valuation/ and valuation/team-models/"]),
        ("Forecast source", ["v3 drivers (regional build)", "v3 drivers, copied (controls: FCFF = v3)",
                             "Capital IQ consensus FY26-28 + fade", "Average of the two sets", "Judgment",
                             "Alejandro is not an independent set: his FCFF FY2026-30 equal v3 to the peso"]),
        ("Horizon", ["10 years (FY2026-35)", "5 years (2H26-FY2030)", "5 years (76-day stub-FY2030)", "5 years (2H26-FY2030)",
                     "Mechanic", "Template horizon for both teammates; v3 kept as a reference row"]),
        ("Valuation date mechanics", ["Fiscal-year-end roll (0.79 y)", "30-Jun anchor, rolled at ke",
                                      "16-Oct, 30-Jun balance sheet", "30-Jun anchor, rolled at ke", "Mechanic",
                                      "Diego's mix of dates drops about 3.5 months of cash flow (Walks: -0.75)"]),
        ("Cost of capital", None),
        ("Risk-free", [0.075101, A["rf"], D["rf"], A["rf"], "Fact", "Same in all three"], "p"),
        ("Country risk premium (EBITDA-weighted)", [0.020496, A["crp"], D["crp"], A["crp"], "Fact", "Same method and weights"], "p"),
        ("Unlevered beta", [0.53, A["beta_u"], D["beta_u"], st["beta_u"], "Judgment", "Averaged"], "n4"),
        ("D/E", [0.7431, A["de"], D["de"], st["de"], "Fact", "Company debt; Diego's CIQ debt figure is reconverted at the wrong FX rate"], "n4"),
        ("Levered beta", [0.806, A["beta_l"], D["beta_l"], run["beta_l"], "", ""], "n4"),
        ("Cost of equity", [V3["ke"], A["ke"], D["ke"], run["ke"], "", ""], "p"),
        ("WACC", [V3["wacc"], A["wacc"], D["wacc"], run["wacc"], "", "Less than 10 bp apart"], "p"),
        ("Forecast", None),
        ("Revenue FY2030 (MXN m)", [None, A["rev"][-1], D["rev"][-1], run["rev"][-1], "Judgment", ""], "m"),
        ("Revenue CAGR FY2025-30", [None, rev_cagr(A["rev"], R["inputs"]["facts"]["rev0"]),
                                    rev_cagr(D["rev"], R["inputs"]["facts"]["rev0"]),
                                    rev_cagr(run["rev"], R["inputs"]["facts"]["rev0"]), "Judgment", ""], "p"),
        ("EBITDA margin FY2030", [None, A["ebitda_m"][-1], D["ebitda_m"][-1], st["ebitda_m"][-1], "Judgment",
                                  "The largest judgment gap (Walks: -3.3 per share)"], "p"),
        ("Capex + new leases / revenue FY2030", [None, A["capex_lease_pct"][-1], D["capex_lease_pct"][-1],
                                                 (run["capex"][-1] + run["leases"][-1]) / run["rev"][-1], "Judgment + mechanic",
                                                 "Capex averaged; leases recomputed with the v3 mechanics"], "p"),
        ("Tax on EBIT", [0.34, A["tax_ebit"][0], D["tax_ebit"][0], st["tax_ebit"], "Judgment (flagged)", "Averaged"], "p"),
        ("FCFF FY2026-30, sum (MXN m)", [None, sum(A["fcff"]), sum(D["fcff"]), sum(run["fcff"]), "", ""], "m"),
        ("Terminal value", None),
        ("Exit EV/EBITDA", [V3["exit_mult"], A["exit_mult"], D["exit_mult"], st["exit_mult"], "Judgment",
                            "Diego's rule on clean data gives 5.83x; average with v3 6.0x"], "x"),
        ("Perpetual growth implied by the multiple", [0.050788, A["implied_g"], D["implied_g"], run["implied_g_exit"], "", ""], "p"),
        ("Terminal growth g (Gordon)", [V3["g"], A["g"], D["g"], st["g"], "Judgment", "Diego: his reverse-DCF g"], "p"),
        ("RONIC behind the Gordon value", [V3["ronic"], A["ronic_gordon"], None, st["ronic"], "Judgment",
                                           "Alejandro's template-style Gordon implies 20%"], "p1"),
        ("Equity bridge (MXN m)", None),
        ("Debt incl. leases", [None, -A["bridge_debt"], D["debt_ciq"], st["debt"], "Fact", "See Inputs row 15"], "m"),
        ("Cash", [None, A["bridge_cash"], D["cash_ciq"], st["cash"], "Fact", ""], "m"),
        ("Non-controlling interests", [None, A["nci_market"], D["nci_book"], run["nci"], "Mechanic (provisional)",
                                       "Market value in v3 and Alejandro; book in Diego"], "m"),
        ("Associates - seasonal tax", [None, A["assoc_net"], 0, st["assoc_net"], "Mechanic", ""], "m"),
        ("Other debt-like claims (Gordon only)", [None, A["other_claims"], 0, st["other_claims"], "Mechanic", ""], "m"),
        ("Shares (millions)", [4294.9, A["shares"], D["shares_m"], st["shares"], "Fact", ""], "n"),
        ("Results (MXN per share, 16-Oct-2026)", None),
        ("Exit multiple", [V3["exit"], R["results"]["alejandro"]["ps_exit"], R["results"]["diego"]["ps_exit"],
                           run["ps_exit"], "", ""], "ps"),
        ("Gordon (as each model computes it)", [V3["gordon"], R["results"]["alejandro"]["ps_gordon"], None,
                                                run["ps_gordon_ronic"], "", "Consolidated: RONIC-consistent"], "ps"),
        ("Exit vs price", [V3["exit"] / 58.6 - 1, R["results"]["alejandro"]["ps_exit"] / 58.6 - 1,
                           R["results"]["diego"]["ps_exit"] / 58.6 - 1, run["ps_exit"] / 58.6 - 1, "", ""], "p1"),
    ]
    r = 5
    for item in rows:
        lab, vals = item[0], item[1]
        fmt = item[2] if len(item) > 2 else None
        if vals is None:
            section(ws, r, lab, 7)
            r += 1
            continue
        put(ws, f"A{r}", lab)
        for j, v in enumerate(vals):
            col = "BCDEFG"[j]
            if v is None:
                put(ws, f"{col}{r}", "n/a", font=F_NOTE)
            elif isinstance(v, str):
                put(ws, f"{col}{r}", v, wrap=True)
            else:
                put(ws, f"{col}{r}", v, fmt)
        r += 1
    widths(ws, {"A": 40, "B": 20, "C": 20, "D": 22, "E": 20, "F": 18, "G": 70})
    ws.freeze_panes = "B5"


# ---------------------------------------------------------------- Walks


def build_walks(wb, R):
    ws = wb.create_sheet("Walks")
    title(ws, "Why the values differ", "Values from the Python replica (valuation/build/consolidate_models.py), which "
                                       "reproduces both workbooks to the centavo (Checks). One change at a time, in order.")
    put(ws, "A4", "Walk 1 - Alejandro to Diego, exit method (MXN per share, 16-Oct-2026)", bold=True)
    header(ws, 5, ["Step", "Value", "Change", "Type"])
    kinds = ["Start", "Mechanic", "Fact", "Fact (FX reconversion, to confirm)", "Fact", "Mechanic", "Mechanic (provisional)",
             "Judgment (beta) + fact (debt)",
             "Judgment", "Judgment", "Mechanic (template D&A)", "Judgment", "Mechanic", "Judgment", "Judgment (flagged)",
             "Judgment + data basis"]
    r = 6
    for i, row in enumerate(R["walk"]):
        put(ws, f"A{r}", row["step"], wrap=True), put(ws, f"B{r}", row["value"], "ps")
        if row["delta"] is not None:
            put(ws, f"C{r}", row["delta"], "bp")
        put(ws, f"D{r}", kinds[i] if i < len(kinds) else "")
        r += 1
    put(ws, f"A{r}", "Diego's workbook (SPG_DCF_independent_Diego.xlsm, DCF!I155)", bold=True)
    put(ws, f"B{r}", R["results"]["diego"]["ps_exit"], "ps")
    put(ws, f"C{r}", f"=B{r}-B{r - 1}", "bp")
    put(ws, f"D{r}", "gap must be 0.00")
    r += 2
    put(ws, f"A{r}", "Walk 2 - v3 set to Diego's set on the same facts and mechanics (where the analysts really differ)", bold=True)
    r += 1
    header(ws, r, ["Step", "Exit", "Change", "Gordon (RONIC)", "Change"])
    r += 1
    for row in R["judgment_walk"]:
        put(ws, f"A{r}", row["step"], wrap=True), put(ws, f"B{r}", row["exit"], "ps")
        put(ws, f"D{r}", row["gordon_ronic"], "ps")
        if row["d_exit"] is not None:
            put(ws, f"C{r}", row["d_exit"], "bp"), put(ws, f"E{r}", row["d_gordon"], "bp")
        r += 1
    put(ws, f"A{r + 1}", "Reading: once facts and mechanics are reconciled the two sets are only about Ps 1.3 apart on the "
                         "exit method. Margin is the main disagreement (v3 higher); revenue, capex and the tax rate offset "
                         "it (Diego higher).", font=F_NOTE)
    widths(ws, {"A": 95, "B": 12, "C": 12, "D": 16, "E": 12})


# ---------------------------------------------------------------- Risk and ESG


def build_risk(wb, R):
    ws = wb.create_sheet("Risk_ESG")
    title(ws, "Risks and ESG translated into value (consolidated base)",
          "Each row changes one input against the base. Values from the Python replica. ESG evidence from Diego's General "
          "ESG Report v2 (team-models/BIMBOA_ESG_Report_EN_Diego.pdf); figures he tagged [unverified] are not used.")
    header(ws, 4, ["Lever", "Theme and evidence", "Exit (MXN/share)", "Gordon RONIC (MXN/share)"])
    evidence = {
        "EBITDA margin -50 bp every year": "Wheat is the main input: raw materials ~64% of cost of sales [sourced/calc]; "
                                           "the CFO tied 60 bp of 2022 EBITDA margin to the wheat cycle [guidance]; the "
                                           "2026 wheat move (+17%) is [unverified]",
        "Capex +50 bp of revenue every year": "101 work centres in high or extremely high water stress [sourced]; "
                                              "withdrawal +15.1% since FY2019 [calc]; a separate sustainability capex "
                                              "envelope since 2025, amount not disclosed [sourced]",
        "Revenue growth -50 bp every year": "Front-of-pack labelling phase 3 from 2028-01-01 and school-food rules from "
                                            "2025-03-29 [sourced]; consumer shifts on sugar and refined flour",
        "Exit multiple -0.5x": "Controlled company: 72.03% control group and 38.9% independent directors [sourced]; "
                               "minority protection 1 of 10 on his compiled rubric [assumption]; governance is the weak "
                               "pillar in every rating source",
        "One-off claim of MXN 1bn": "Canada Bread price-fixing class actions open with no provision; the CAD 50m fine "
                                    "was paid in 2023 [sourced]; uncertain tax positions 8,537 [sourced] already "
                                    "deducted in Gordon",
        "Recurring cost of MXN 1bn a year from FY2027 (pre-tax)": "Mexican carbon market, pilot 2020-25, rules "
                                                                  "unpublished [sourced]; extended producer "
                                                                  "responsibility under the law of 2026-01-19 [sourced]",
        "Bimbo 23L coupon step-up (Scope 3 target missed in 2030)": "+25 bp on MXN 12,000m from 21-Nov-2031 to "
                                                                    "20-May-2033 if the -28% Scope 3 target is not "
                                                                    "verified [sourced]; 2025 Scope 3 +46.7% vs the bond "
                                                                    "base [calc]",
    }
    r = 5
    for row in R["consolidated"]["levers"]:
        put(ws, f"A{r}", row["lever"], wrap=True)
        put(ws, f"B{r}", evidence.get(row["lever"], row["theme"]), wrap=True)
        put(ws, f"C{r}", row["d_exit"], "bp"), put(ws, f"D{r}", row["d_gordon_ronic"], "bp")
        r += 1
    r += 1
    put(ws, f"A{r}", "Rules of thumb (live)", bold=True)
    put(ws, f"A{r + 1}", "Each MXN 1bn of claim at 30-Jun-2026, per share today")
    put(ws, f"C{r + 1}", "=-1000/Inputs!D12*DCF!B42", "bp")
    put(ws, f"A{r + 2}", "No ESG adjustment to the discount rate: Diego's report finds no source that links Bimbo's ESG "
                         "scores to its cost of debt or equity, and the 23L step-up is about 0.05% of FY2025 EBITDA. "
                         "Adding a premium would double-count risks already in the cash flows.", font=F_NOTE, wrap=True)
    widths(ws, {"A": 46, "B": 80, "C": 16, "D": 22})


# ---------------------------------------------------------------- Football field


def build_football(wb, R):
    ws = wb.create_sheet("Football")
    C, V3 = R["consolidated"], R["inputs"]["v3"]
    g = C["grids"]
    title(ws, "Football field (MXN per share, 16-Oct-2026)", "Price MXN 58.60. Ranges: consolidated grids at +/-0.5 pp of "
                                                             "WACC and +/-0.5x or +/-0.5 pp of g. The team sets the target.")
    header(ws, 4, ["Method", "Low", "Point", "High", "Note"])
    ex, go = g["exit"], g["gordon_ronic"]
    sub = lambda grid: [grid[i][j] for i in (1, 2, 3) for j in (1, 2, 3)]  # noqa: E731
    rows = [
        ("52-week range", 53.03, None, 68.67, "Capital IQ, 06-Oct-2026"),
        ("Consolidated - exit multiple", min(sub(ex)), C["run"]["ps_exit"], max(sub(ex)),
         "5.92x; range: WACC +/-0.5 pp and 5.5x-6.5x"),
        ("Consolidated - Gordon, RONIC 14.6%", min(sub(go)), C["run"]["ps_gordon_ronic"], max(sub(go)),
         "g 4.125%; range: WACC and g +/-0.5 pp"),
        ("Consolidated - Gordon, template style", None, C["run"]["ps_gordon_naive"], None,
         "Implies RONIC of about 33%: not credible"),
        ("v3 (10 years) - exit 6.0x / Gordon", V3["gordon"], None, V3["exit"], "Published v3"),
        ("Alejandro - exit 6.0x / Gordon", R["results"]["alejandro"]["ps_gordon"], None, R["results"]["alejandro"]["ps_exit"],
         "v3 inputs, 5 years"),
        ("Diego - exit 6.4x", None, R["results"]["diego"]["ps_exit"], None, "Before fixing facts and mechanics"),
        ("Comps (v3 table) - core 5.53x / all eligible 8.27x", V3["comps_core"], None, V3["comps_all"],
         "LTM EBITDA 60,919; peer-EV claims"),
    ]
    r = 5
    for lab, lo, pt, hi, note in rows:
        put(ws, f"A{r}", lab)
        for col, v in zip("BCD", (lo, pt, hi)):
            if v is not None:
                put(ws, f"{col}{r}", v, "ps")
        put(ws, f"E{r}", note)
        r += 1
    put(ws, f"A{r}", "Current price"), put(ws, f"C{r}", 58.60, "ps")
    widths(ws, {"A": 52, "B": 11, "C": 11, "D": 11, "E": 48})


# ---------------------------------------------------------------- Checks and README


def build_checks(wb, R):
    ws = wb.create_sheet("Checks")
    title(ws, "Checks", "Replica = valuation/build/consolidate_models.py. Live checks hold at the base settings "
                        "(w = 0.5, clean, market).")
    header(ws, 4, ["Check", "Replica", "Reference", "Difference", "Result"])
    r = 5
    for k, c in R["checks"].items():
        put(ws, f"A{r}", k), put(ws, f"B{r}", c["replica"], "n4"), put(ws, f"C{r}", c["workbook"], "n4")
        put(ws, f"D{r}", f"=B{r}-C{r}", "0.00E+00")
        tol = 0.005 if k.endswith(("exit", "gordon")) else (1 if ("ev" in k or "fcff" in k) else 1e-6)
        put(ws, f"E{r}", f'=IF(ABS(D{r})<={tol},"OK","CHECK")')
        r += 1
    put(ws, f"A{r}", "walk Alejandro -> Diego closes"), put(ws, f"B{r}", R["walk"][-1]["value"], "n4")
    put(ws, f"C{r}", R["results"]["diego"]["ps_exit"], "n4"), put(ws, f"D{r}", f"=B{r}-C{r}", "0.00E+00")
    put(ws, f"E{r}", f'=IF(ABS(D{r})<=0.000001,"OK","CHECK")')
    r += 2
    run = R["consolidated"]["run"]
    live = [
        ("WACC (live) vs replica", "=WACC!B16", run["wacc"], 1e-9),
        ("Exit value per share (live) vs replica", "=DCF!B43", run["ps_exit"], 0.005),
        ("Gordon RONIC (live) vs replica", "=DCF!B56", run["ps_gordon_ronic"], 0.005),
        ("Gordon template style (live) vs replica", "=DCF!B60", run["ps_gordon_naive"], 0.005),
        ("Gordon RONIC = WACC (live) vs replica", "=DCF!B62", run["ps_gordon_ronic_wacc"], 0.005),
        ("Exit grid centre = DCF!B43", "=Sensitivity!D8", "=DCF!B43", 0.005),
        ("Gordon grid centre = DCF!B56", "=Sensitivity!D17", "=DCF!B56", 0.005),
        ("FY2030 revenue (live) vs replica", "=DCF!G5", run["rev"][-1], 0.01),
        ("FCFF FY2026-30 sum (live) vs replica", "=SUM(DCF!C19:G19)", sum(run["fcff"]), 0.01),
        ("Income statement: net income, computed vs reported (sum of |diff|)", f"=SUMPRODUCT(ABS(IS!B{IS_ROWS['ni_chk']}:F{IS_ROWS['ni_chk']}))", 0, 0.5),
        ("Balance sheet: total assets, lines vs reported (sum of |diff|)", f"=SUMPRODUCT(ABS(BS!B{BS_ROWS['ta_chk']}:F{BS_ROWS['ta_chk']}))", 0, 0.5),
        ("Balance sheet balances: assets - liabilities - equity (sum of |diff|)", f"=SUMPRODUCT(ABS(BS!B{BS_ROWS['bal_chk']}:F{BS_ROWS['bal_chk']}))", 0, 0.5),
        ("History: regional net sales add up to the income statement (sum of |diff|)", f"=SUMPRODUCT(ABS(History!B{REFS['seg_sales_chk']}:F{REFS['seg_sales_chk']}))", 0, 0.5),
        ("History: regional Adjusted EBITDA adds up to the income statement (sum of |diff|)", f"=SUMPRODUCT(ABS(History!B{REFS['seg_ebitda_chk']}:F{REFS['seg_ebitda_chk']}))", 0, 0.5),
        ("DCF FY2025A net sales = income statement", "=DCF!B5", f"=IS!F{IS_ROWS['rev']}", 0.5),
        ("DCF FY2025A Adjusted EBITDA = income statement", "=DCF!B8", f"=IS!F{IS_ROWS['adj']}", 0.5),
    ]
    for lab, a, b, tol in live:
        put(ws, f"A{r}", lab), put(ws, f"B{r}", a, "n4"), put(ws, f"C{r}", b, "n4")
        put(ws, f"D{r}", f"=B{r}-C{r}", "0.00E+00")
        put(ws, f"E{r}", f'=IF(ABS(D{r})<={tol},"OK","CHECK")')
        r += 1
    put(ws, f"A{r + 1}", "All checks", bold=True)
    put(ws, f"E{r + 1}", f'=IF(COUNTIF(E5:E{r - 1},"CHECK")=0,"OK","CHECK")')
    widths(ws, {"A": 44, "B": 16, "C": 16, "D": 14, "E": 10})
    return f"E{r + 1}"


def build_readme(wb, R):
    ws = wb.active
    ws.title = "README"
    title(ws, "Grupo Bimbo (BIMBO A) - consolidated DCF v1 (assistant proposal, 2026-10-19)",
          "Built from the team's three models: v3 (valuation/), Alejandro and Diego (valuation/team-models/). "
          "Target price and rating stay with the team.")
    lines = [
        ("What this is", "One DCF on averaged judgments. Facts (debt, cash, shares) take one sourced value and mechanics "
                         "(valuation date, lease reinvestment, bridge) take one convention, because averaging a data "
                         "error with a correct figure only halves the error."),
        ("Why 50/50", "Alejandro's model carries the v3 inputs (his FCFF and WACC equal v3), so there are two input sets, "
                      "not three. Set Inputs!D6 to 0.6667 to give each model one third."),
        ("Sheets", "Inputs (decision cells in yellow) -> IS, BS, CF and History (audited FY2021-FY2025, from data/financials.csv) -> "
                   "WACC -> DCF -> Sensitivity. Compare, Walks, Risk_ESG and Football explain the numbers; Checks ties everything to the "
                   "Python replica and the statements to their reported totals."),
        ("Colour code", "Blue = typed value; black = formula; yellow = decision cell; green = key result."),
        ("Rebuild", "python -I valuation/build/consolidate_models.py --project . ; python -I "
                    "valuation/build/build_consolidated_workbook.py --project . ; recalc_excel.ps1 ; scan_errors.py"),
    ]
    r = 4
    for lab, txt in lines:
        put(ws, f"A{r}", lab, bold=True), put(ws, f"B{r}", txt, wrap=True)
        r += 1
    r += 1
    section(ws, r, "Key results (live)", 3)
    r += 1
    key = [
        ("WACC", "=WACC!B16", "p"), ("Cost of equity", "=WACC!B11", "p"), ("Exit EV/EBITDA", "=Inputs!D40", "x"),
        ("Terminal growth", "=Inputs!D42", "p"),
        ("Value per share today - exit multiple", "=DCF!B43", "ps"),
        ("Value per share today - Gordon, RONIC-consistent", "=DCF!B56", "ps"),
        ("Value per share today - Gordon, template style (reference)", "=DCF!B60", "ps"),
        ("12-month value, exit (guide convention)", "=DCF!C66", "ps"),
        ("Expected return, exit (guide convention)", "=DCF!D66", "p1"),
        ("Value today needed for the HOLD floor", "=DCF!D72", "ps"),
        ("All checks", None, None),
    ]
    for lab, f, fmt in key:
        put(ws, f"A{r}", lab)
        if f:
            put(ws, f"B{r}", f, fmt)
        r += 1
    widths(ws, {"A": 56, "B": 110})
    return f"B{r - 1}"


# ---------------------------------------------------------------- Financial statements and history (audited FY2021-FY2025)

FIN_CSV = "data/financials.csv"
HC = ["B", "C", "D", "E", "F"]  # FY2021A ... FY2025A on the statement sheets
HYEARS = ["FY2021A", "FY2022A", "FY2023A", "FY2024A", "FY2025A"]
IS_ROWS = {"rev": 5, "cogs": 6, "gp": 7, "opex": 8, "ebit": 9, "intexp": 10, "intinc": 11, "ofr": 12, "pbt": 13, "tax": 14, "ni": 15,
           "ni_rep": 16, "ni_chk": 17, "nci": 18, "ni_sh": 19, "shares": 20, "eps": 21, "da": 22, "ebitda": 23, "adj": 24}
BS_ROWS = {"cash": 5, "ar": 6, "inv": 7, "oca": 8, "tca": 9, "ppe": 10, "intang": 11, "onca": 12, "ta": 13, "ta_rep": 14, "ta_chk": 15,
           "ap": 16, "ocl": 17, "std": 18, "tcl": 19, "ltd": 20, "lease": 21, "oncl": 22, "tl": 23, "eq_p": 24, "nci": 25, "te": 26,
           "tle": 27, "bal_chk": 28, "debt": 29, "nd": 30, "nd_ebitda": 31, "nwc": 32}
CF_ROWS = {"cfo": 5, "capex_cash": 6, "lease": 7, "fcf": 8, "div": 9, "fcf_m": 10, "capex_ppe": 11, "capex_ppe_pct": 12, "dnwc": 13}
REFS: dict = {}  # row numbers of the History checks, read by build_checks

ADJ_EBITDA = [47372, 53446, 54942, 55473, 59456]
SEG_EBITDA = {"North America": [21452, 22611, 20217, 15812, 17117], "Mexico": [20771, 23320, 27484, 30762, 31630],
              "Europe, Asia and Africa": [2704, 2626, 2932, 4212, 5830], "Latin America": [1926, 3433, 3531, 3603, 3552],
              "Eliminations and other": [519, 1456, 778, 1084, 1327]}
CAPEX_PPE = [20671, 28669, 34754, 29402, 22530]
EXFX_GROUP = [0.093, 0.206, 0.107, 0.005, 0.016]
EXFX_REGION = {"North America": [None, None, 0.067, -0.046, -0.040], "Mexico": [None, None, 0.115, 0.040, 0.024],
               "Europe, Asia and Africa": [None, None, 0.235, 0.079, 0.117], "Latin America": [None, None, 0.169, 0.100, 0.101]}
SRC_SEG = "RA22 p.313 (FY21 restated); RA23 p.358 (FY22); IA25-EEFF p.85-86 (FY23-25)"


def _short_doc(path: str) -> str:
    p = path.replace("\\", "/").split("/")[-1]
    if "Reporte Anual 2023" in p:
        return "RA23 XBRL"
    if "IA25_GB_EEFF" in p:
        return "IA25-EEFF"
    return p[:24]


def load_financials(project):
    import csv

    data, tag, src = {}, {}, {}
    with open(project / FIN_CSV, encoding="utf-8") as f:
        for r in csv.DictReader(f):
            k, y = r["line_item"], int(r["year"])
            data.setdefault(k, {})[y] = float(r["value"])
            tag[k] = r["tag"]
            src.setdefault(k, {}).setdefault((_short_doc(r["source_doc"]), r["page"]), []).append(y)
    text = {}
    for k, v in src.items():
        parts = []
        for (d, p), ys in sorted(v.items(), key=lambda kv: kv[1][0]):
            span = f"FY{ys[0] % 100}" + (f"-{ys[-1] % 100}" if len(ys) > 1 else "")
            parts.append(f"{d} p.{p} ({span})")
        text[k] = "; ".join(parts)
    return data, tag, text


def _typed(ws, r, label, vals, fmt, tag, source, bold=False):
    put(ws, f"A{r}", label, bold=bold)
    for j, v in enumerate(vals):
        if v is not None:
            put(ws, f"{HC[j]}{r}", v, fmt)
    put(ws, f"G{r}", tag), put(ws, f"H{r}", source)


def _calc(ws, r, label, tmpl, fmt, start=0, bold=False, note=""):
    put(ws, f"A{r}", label, bold=bold)
    for j in range(start, 5):
        put(ws, f"{HC[j]}{r}", tmpl.format(c=HC[j], p=HC[j - 1] if j else ""), fmt)
    put(ws, f"G{r}", "[calc]")
    if note:
        put(ws, f"H{r}", note)


def _stmt_sheet(wb, name, heading, sub):
    ws = wb.create_sheet(name)
    title(ws, heading, sub)
    header(ws, 4, ["MXN m", *HYEARS, "Tag", "Source / note"])
    widths(ws, {"A": 56, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12, "G": 11, "H": 95})
    ws.freeze_panes = "B5"
    return ws


def build_financials(wb, project):
    data, tag, text = load_financials(project)
    vals = lambda k: [data[k][y] for y in range(2021, 2026)]  # noqa: E731
    tg = lambda k: "[" + tag[k] + "]"  # noqa: E731
    I, B, C = IS_ROWS, BS_ROWS, CF_ROWS

    # ------------------------------------------------------------ IS
    ws = _stmt_sheet(wb, "IS", "Income statement, FY2021A-FY2025A (audited, MXN millions)",
                     "Continuing operations. Blue = typed from the filings (data/financials.csv); black = formula. FY2021 is restated "
                     "for Ricolino (sold 1-Nov-2022), which sits in 'other financial result' as discontinued operations after tax.")
    _typed(ws, I["rev"], "Net sales", vals("revenue"), "m0", tg("revenue"), text["revenue"], bold=True)
    _typed(ws, I["cogs"], "Cost of sales", vals("cogs"), "m0", tg("cogs"), text["cogs"])
    _calc(ws, I["gp"], "Gross profit", "={c}%d-{c}%d" % (I["rev"], I["cogs"]), "m0", bold=True)
    _typed(ws, I["opex"], "Operating expenses, net", vals("opex"), "m0", tg("opex"),
           text["opex"] + ". Distribution and selling + administration + integration + other (income) expenses, net")
    _calc(ws, I["ebit"], "EBIT (operating income)", "={c}%d-{c}%d" % (I["gp"], I["opex"]), "m0", bold=True,
          note="FY2022 includes a non-cash gain on multi-employer pension plans of about 19,010 (RA23 EEFF notes 21-22)")
    _typed(ws, I["intexp"], "Interest expense", vals("interest_expense"), "m0", tg("interest_expense"), text["interest_expense"])
    _typed(ws, I["intinc"], "Interest income", vals("interest_income"), "m0", tg("interest_income"), text["interest_income"])
    _typed(ws, I["ofr"], "Other financial result, net", vals("other_financial_net"), "m0", tg("other_financial_net"),
           text["other_financial_net"] + ". Associates + FX + monetary position + discontinued operations after tax "
           "(Ricolino: 1,254 in FY2021; 16,988 in FY2022)")
    _calc(ws, I["pbt"], "Pre-tax income (incl. discontinued operations after tax)",
          "={c}%d-{c}%d+{c}%d+{c}%d" % (I["ebit"], I["intexp"], I["intinc"], I["ofr"]), "m0")
    _typed(ws, I["tax"], "Income tax", vals("income_tax"), "m0", tg("income_tax"), text["income_tax"] + ". Continuing operations")
    _calc(ws, I["ni"], "Net income", "={c}%d-{c}%d" % (I["pbt"], I["tax"]), "m0", bold=True)
    _typed(ws, I["ni_rep"], "Net income as reported", vals("net_income_reported"), "m0", tg("net_income_reported"),
           text["net_income_reported"] + ". Consolidated, before the NCI split")
    _calc(ws, I["ni_chk"], "Check: computed - reported", "={c}%d-{c}%d" % (I["ni"], I["ni_rep"]), "m0")
    _typed(ws, I["nci"], "Attributable to non-controlling interests", vals("nci_income"), "m0", tg("nci_income"),
           text["nci_income"] + ". FY2021-22 include the NCI share of discontinued operations")
    _calc(ws, I["ni_sh"], "Net income to shareholders", "={c}%d-{c}%d" % (I["ni"], I["nci"]), "m0", bold=True)
    _typed(ws, I["shares"], "Diluted shares (millions)", vals("shares_diluted"), "m", tg("shares_diluted"), text["shares_diluted"])
    _calc(ws, I["eps"], "Diluted EPS (MXN)", "={c}%d/{c}%d" % (I["ni_sh"], I["shares"]), "ps")
    _typed(ws, I["da"], "D&A (continuing)", vals("da"), "m0", tg("da"), text["da"] + ". Segment note, consistent with EBIT")
    _calc(ws, I["ebitda"], "EBITDA (EBIT + D&A, IFRS 16)", "={c}%d+{c}%d" % (I["ebit"], I["da"]), "m0", bold=True)
    _typed(ws, I["adj"], "Adjusted EBITDA (company definition)", ADJ_EBITDA, "m0", "[sourced]",
           SRC_SEG + ". Operating income + D&A + impairment + non-cash items (mainly multi-employer pension plans)", bold=True)

    # ------------------------------------------------------------ BS
    ws = _stmt_sheet(wb, "BS", "Balance sheet, FY2021A-FY2025A (audited, MXN millions, year-end)",
                     "Blue = typed from the filings (data/financials.csv); black = formula. Lines are grouped as in the model; "
                     "net debt includes lease liabilities.")
    _typed(ws, B["cash"], "Cash and equivalents", vals("cash"), "m0", tg("cash"), text["cash"])
    _typed(ws, B["ar"], "Accounts receivable", vals("receivables"), "m0", tg("receivables"), text["receivables"])
    _typed(ws, B["inv"], "Inventories", vals("inventory"), "m0", tg("inventory"), text["inventory"])
    _typed(ws, B["oca"], "Other current assets", vals("other_current_assets"), "m0", tg("other_current_assets"), text["other_current_assets"])
    _calc(ws, B["tca"], "Total current assets", "=SUM({c}%d:{c}%d)" % (B["cash"], B["oca"]), "m0", bold=True)
    _typed(ws, B["ppe"], "PP&E and right-of-use assets, net", vals("ppe_net"), "m0", tg("ppe_net"), text["ppe_net"])
    _typed(ws, B["intang"], "Intangibles and goodwill", vals("intangibles_goodwill"), "m0", tg("intangibles_goodwill"),
           text["intangibles_goodwill"])
    _typed(ws, B["onca"], "Other non-current assets", vals("other_noncurrent_assets"), "m0", tg("other_noncurrent_assets"),
           text["other_noncurrent_assets"])
    _calc(ws, B["ta"], "Total assets", "={c}%d+{c}%d+{c}%d+{c}%d" % (B["tca"], B["ppe"], B["intang"], B["onca"]), "m0", bold=True)
    _typed(ws, B["ta_rep"], "Total assets as reported", vals("total_assets_reported"), "m0", tg("total_assets_reported"),
           text["total_assets_reported"])
    _calc(ws, B["ta_chk"], "Check: lines - reported", "={c}%d-{c}%d" % (B["ta"], B["ta_rep"]), "m0")
    _typed(ws, B["ap"], "Accounts payable", vals("payables"), "m0", tg("payables"), text["payables"])
    _typed(ws, B["ocl"], "Other current liabilities", vals("other_current_liabilities"), "m0", tg("other_current_liabilities"),
           text["other_current_liabilities"])
    _typed(ws, B["std"], "Short-term debt (current portion)", vals("debt_short"), "m0", tg("debt_short"), text["debt_short"])
    _calc(ws, B["tcl"], "Total current liabilities", "=SUM({c}%d:{c}%d)" % (B["ap"], B["std"]), "m0", bold=True)
    _typed(ws, B["ltd"], "Long-term debt", vals("debt_long"), "m0", tg("debt_long"), text["debt_long"])
    _typed(ws, B["lease"], "Lease liabilities (IFRS 16, current + non-current)", vals("lease_liabilities"), "m0", tg("lease_liabilities"),
           text["lease_liabilities"])
    _typed(ws, B["oncl"], "Other non-current liabilities", vals("other_noncurrent_liabilities"), "m0", tg("other_noncurrent_liabilities"),
           text["other_noncurrent_liabilities"])
    _calc(ws, B["tl"], "Total liabilities", "={c}%d+{c}%d+{c}%d+{c}%d" % (B["tcl"], B["ltd"], B["lease"], B["oncl"]), "m0", bold=True)
    _typed(ws, B["eq_p"], "Shareholders' equity (parent)", vals("equity_parent"), "m0", tg("equity_parent"),
           text["equity_parent"] + ". FY2021-22 include perpetual notes classed as equity (8,867; 8,098)")
    _typed(ws, B["nci"], "Non-controlling interests", vals("nci_equity"), "m0", tg("nci_equity"), text["nci_equity"])
    _calc(ws, B["te"], "Total equity", "={c}%d+{c}%d" % (B["eq_p"], B["nci"]), "m0", bold=True)
    _calc(ws, B["tle"], "Total liabilities and equity", "={c}%d+{c}%d" % (B["tl"], B["te"]), "m0", bold=True)
    _calc(ws, B["bal_chk"], "Check: assets - (liabilities + equity)", "={c}%d-{c}%d" % (B["ta"], B["tle"]), "m0")
    _calc(ws, B["debt"], "Financial debt incl. leases", "={c}%d+{c}%d+{c}%d" % (B["std"], B["ltd"], B["lease"]), "m0")
    _calc(ws, B["nd"], "Net debt incl. leases", "={c}%d-{c}%d" % (B["debt"], B["cash"]), "m0", bold=True)
    _calc(ws, B["nd_ebitda"], "Net debt / Adjusted EBITDA (x)", "={c}%d/IS!{c}%d" % (B["nd"], I["adj"]), "x")
    _calc(ws, B["nwc"], "Operating working capital (receivables + inventories + other current assets - payables - other current liabilities)",
          "={c}%d+{c}%d+{c}%d-{c}%d-{c}%d" % (B["ar"], B["inv"], B["oca"], B["ap"], B["ocl"]), "m0",
          note="Balance-sheet basis, same definition as the forecast driver; data/adjustments.md uses a narrower one (excludes financing and tax items)")

    # ------------------------------------------------------------ CF
    ws = _stmt_sheet(wb, "CF", "Cash flow, FY2021A-FY2025A (audited, MXN millions; abridged)",
                     "Only the lines the valuation uses. Blue = typed from the filings; black = formula. FY2021-22 CFO includes Ricolino "
                     "until its sale.")
    _typed(ws, C["cfo"], "Cash flow from operations", vals("cfo"), "m0", tg("cfo"), text["cfo"], bold=True)
    _typed(ws, C["capex_cash"], "Capex, cash basis (PP&E + intangibles)", vals("capex"), "m0", tg("capex"), text["capex"])
    _typed(ws, C["lease"], "Lease principal paid (IFRS 16)", vals("lease_principal_paid"), "m0", tg("lease_principal_paid"),
           text["lease_principal_paid"] + ". Financing line 'lease payments' less lease interest accrued")
    _calc(ws, C["fcf"], "Free cash flow (CFO - capex - lease principal)", "={c}%d-{c}%d-{c}%d" % (C["cfo"], C["capex_cash"], C["lease"]),
          "m0", bold=True)
    _typed(ws, C["div"], "Dividends paid (to the parent)", vals("dividends_paid"), "m0", tg("dividends_paid"), text["dividends_paid"])
    _calc(ws, C["fcf_m"], "FCF margin (% of net sales)", "={c}%d/IS!{c}%d" % (C["fcf"], I["rev"]), "p1")
    _typed(ws, C["capex_ppe"], "Capex, PP&E acquired (segment note; the basis of the capex guidance)", CAPEX_PPE, "m0", "[sourced]", SRC_SEG)
    _calc(ws, C["capex_ppe_pct"], "Capex, PP&E acquired / net sales", "={c}%d/IS!{c}%d" % (C["capex_ppe"], I["rev"]), "p1")
    _calc(ws, C["dnwc"], "Change in operating working capital (+ releases cash)", "=BS!{p}%d-BS!{c}%d" % (B["nwc"], B["nwc"]), "m0",
          start=1, note="Decrease in operating working capital on the balance sheet (BS row %d)" % B["nwc"])

    # ------------------------------------------------------------ History
    ws = _stmt_sheet(wb, "History", "Historical growth, margins, returns and segments, FY2021A-FY2025A",
                     "Formulas on the IS, BS and CF sheets; blue = typed from the filings or releases. FY2021 is restated for Ricolino, so "
                     "FY2021 growth rates are not available on the same basis.")
    r = 5
    section(ws, r, "A. Growth, % y/y", 8)
    r += 1
    H = {}
    for key, lab, row_ref in (("g_rev", "Net sales (MXN)", I["rev"]), ("g_adj", "Adjusted EBITDA (MXN)", I["adj"]),
                              ("g_ebit", "EBIT", I["ebit"]), ("g_ni", "Net income to shareholders", I["ni_sh"]),
                              ("g_eps", "Diluted EPS", I["eps"])):
        _calc(ws, r, lab, "=IS!{c}%d/IS!{p}%d-1" % (row_ref, row_ref), "p1", start=1,
              note="FY2022 EBIT and net income include the pension and Ricolino one-offs" if key == "g_ebit" else "")
        H[key] = r
        r += 1
    _typed(ws, r, "Net sales ex-FX (company)", EXFX_GROUP, "p1", "[sourced]",
           "4T21 p.2 (FY21, as originally reported, incl. Ricolino); 4T22 p.4; 4T23 p.1; 4T24 p.2; 4T25 p.3")
    H["g_exfx"] = r
    r += 2
    section(ws, r, "B. Compound annual growth (CAGR) and simple average", 8)
    r += 1
    for lab, rowkey, n, first in (("Net sales CAGR FY2021-FY2025 (4 years)", "rev", 4, "B"), ("Net sales CAGR FY2022-FY2025 (3 years)", "rev", 3, "C"),
                                  ("Adjusted EBITDA CAGR FY2021-FY2025 (4 years)", "adj", 4, "B"),
                                  ("Adjusted EBITDA CAGR FY2022-FY2025 (3 years)", "adj", 3, "C")):
        put(ws, f"A{r}", lab)
        put(ws, f"B{r}", f"=(IS!F{I[rowkey]}/IS!{first}{I[rowkey]})^(1/{n})-1", "p1")
        put(ws, f"G{r}", "[calc]")
        r += 1
    put(ws, f"A{r}", "Net sales: simple average of the annual rates FY2022-FY2025"), put(ws, f"B{r}", f"=AVERAGE(C{H['g_rev']}:F{H['g_rev']})", "p1")
    put(ws, f"G{r}", "[calc]"), put(ws, f"H{r}", "FY2022 (+17.7%) lifts the average; see the CAGR from FY2022")
    r += 2
    section(ws, r, "C. Margins, % of net sales", 8)
    r += 1
    for lab, tmpl, note in (("Gross margin", "=IS!{c}%d/IS!{c}%d" % (I["gp"], I["rev"]), ""),
                            ("EBIT margin", "=IS!{c}%d/IS!{c}%d" % (I["ebit"], I["rev"]), ""),
                            ("EBITDA margin (EBIT + D&A)", "=IS!{c}%d/IS!{c}%d" % (I["ebitda"], I["rev"]), ""),
                            ("Adjusted EBITDA margin", "=IS!{c}%d/IS!{c}%d" % (I["adj"], I["rev"]), ""),
                            ("Net margin (to shareholders)", "=IS!{c}%d/IS!{c}%d" % (I["ni_sh"], I["rev"]), "FY2022 includes the Ricolino gain"),
                            ("FCF margin", "=CF!{c}%d" % C["fcf_m"], ""),
                            ("D&A / net sales", "=IS!{c}%d/IS!{c}%d" % (I["da"], I["rev"]), ""),
                            ("Capex (PP&E acquired) / net sales", "=CF!{c}%d" % C["capex_ppe_pct"], "")):
        _calc(ws, r, lab, tmpl, "p1", note=note)
        r += 1
    r += 1
    section(ws, r, "D. Returns, leverage and liquidity", 8)
    r += 1
    _calc(ws, r, "Return on average equity (parent)", "=IS!{c}%d/AVERAGE(BS!{p}%d,BS!{c}%d)" % (I["ni_sh"], B["eq_p"], B["eq_p"]), "p1", start=1,
          note="FY2022 includes the one-offs")
    r += 1
    _calc(ws, r, "ROIC = EBIT x (1 - effective tax rate) / opening (equity + net debt)",
          "=IS!{c}%d*(1-IS!{c}%d/IS!{c}%d)/(BS!{p}%d+BS!{p}%d)" % (I["ebit"], I["tax"], I["pbt"], B["te"], B["nd"]), "p1", start=1)
    r += 1
    _calc(ws, r, "Effective tax rate", "=IS!{c}%d/IS!{c}%d" % (I["tax"], I["pbt"]), "p1")
    r += 1
    _calc(ws, r, "Net debt incl. leases / Adjusted EBITDA (x)", "=BS!{c}%d" % B["nd_ebitda"], "x")
    r += 1
    _calc(ws, r, "Interest cover: EBIT / interest expense (x)", "=IS!{c}%d/IS!{c}%d" % (I["ebit"], I["intexp"]), "x")
    r += 1
    _calc(ws, r, "Current ratio (x)", "=BS!{c}%d/BS!{c}%d" % (B["tca"], B["tcl"]), "x")
    r += 2

    section(ws, r, "E. Net sales by region before eliminations, MXN m", 8)
    r += 1
    seg_keys = [("North America", "seg_north_america"), ("Mexico", "seg_mexico"), ("Europe, Asia and Africa", "seg_europe_asia_africa"),
                ("Latin America", "seg_latin_america"), ("Eliminations", "seg_eliminations")]
    seg_rows = {}
    for lab, k in seg_keys:
        _typed(ws, r, lab, vals(k), "m0", tg(k), text[k])
        seg_rows[lab] = r
        r += 1
    first_seg, last_seg = seg_rows["North America"], seg_rows["Eliminations"]
    _calc(ws, r, "Consolidated net sales", "=SUM({c}%d:{c}%d)" % (first_seg, last_seg), "m0", bold=True)
    tot_sales = r
    r += 1
    _calc(ws, r, "Check: segments - income statement", "={c}%d-IS!{c}%d" % (tot_sales, I["rev"]), "m0")
    REFS["seg_sales_chk"] = r
    r += 1
    section(ws, r, "Growth y/y by region, MXN", 8)
    r += 1
    for lab, _ in seg_keys[:4]:
        _calc(ws, r, lab, "={c}%d/{p}%d-1" % (seg_rows[lab], seg_rows[lab]), "p1", start=1)
        r += 1
    section(ws, r, "Ex-FX growth by region (company; not available for FY2021-FY2022 by region)", 8)
    r += 1
    for lab in ("North America", "Mexico", "Europe, Asia and Africa", "Latin America"):
        _typed(ws, r, lab, EXFX_REGION[lab], "p1", "[sourced]", "4T23 p.2; 4T24 p.2; RA25 p.132")
        r += 1
    r += 1
    section(ws, r, "F. Adjusted EBITDA by region, MXN m", 8)
    r += 1
    eb_rows = {}
    for lab, v in SEG_EBITDA.items():
        _typed(ws, r, lab, v, "m0", "[sourced]", SRC_SEG + ". Regional results exclude intercompany royalties")
        eb_rows[lab] = r
        r += 1
    _calc(ws, r, "Consolidated Adjusted EBITDA", "=SUM({c}%d:{c}%d)" % (eb_rows["North America"], eb_rows["Eliminations and other"]), "m0",
          bold=True)
    tot_eb = r
    r += 1
    _calc(ws, r, "Check: segments - income statement", "={c}%d-IS!{c}%d" % (tot_eb, I["adj"]), "m0")
    REFS["seg_ebitda_chk"] = r
    r += 1
    section(ws, r, "Adjusted EBITDA margin by region (EBITDA / region net sales)", 8)
    r += 1
    for lab in ("North America", "Mexico", "Europe, Asia and Africa", "Latin America"):
        _calc(ws, r, lab, "={c}%d/{c}%d" % (eb_rows[lab], seg_rows[lab]), "p1")
        r += 1
    section(ws, r, "Share of Adjusted EBITDA (Mexico includes eliminations and other, as in the company's presentation)", 8)
    r += 1
    _calc(ws, r, "North America", "={c}%d/{c}%d" % (eb_rows["North America"], tot_eb), "p1")
    r += 1
    _calc(ws, r, "Mexico (incl. eliminations and other)", "=({c}%d+{c}%d)/{c}%d" % (eb_rows["Mexico"], eb_rows["Eliminations and other"], tot_eb), "p1")
    r += 1
    _calc(ws, r, "Europe, Asia and Africa", "={c}%d/{c}%d" % (eb_rows["Europe, Asia and Africa"], tot_eb), "p1")
    r += 1
    _calc(ws, r, "Latin America", "={c}%d/{c}%d" % (eb_rows["Latin America"], tot_eb), "p1")
    return H


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--project", default=".")
    args = ap.parse_args(argv)
    project = Path(args.project).resolve()
    R = json.loads((project / SRC).read_text(encoding="utf-8"))
    wb = Workbook()
    all_checks_cell = build_readme(wb, R)
    drv = build_inputs(wb, R)
    build_financials(wb, project)
    build_wacc(wb, R)
    build_dcf(wb, R, drv)
    build_sens(wb, R)
    build_compare(wb, R)
    build_walks(wb, R)
    build_risk(wb, R)
    build_football(wb, R)
    checks_cell = build_checks(wb, R)
    put(wb["README"], all_checks_cell, f"=Checks!{checks_cell}")
    out = project / OUT
    wb.save(out)
    print("wrote", out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
