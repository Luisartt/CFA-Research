"""Scan recalculated workbooks for Excel error values in formula cells.

Usage: python -I scan_errors.py <file.xlsx> [<file.xlsx> ...]
Exit code 1 if any formula cell shows an error value or has no cached value.
"""

from __future__ import annotations

import sys

import openpyxl

ERRORS = ("#REF!", "#N/A", "#DIV/0!", "#VALUE!", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")


def scan(path: str) -> int:
    formulas = openpyxl.load_workbook(path, data_only=False)
    values = openpyxl.load_workbook(path, data_only=True)
    problems = 0
    n_formulas = 0
    for ws in formulas.worksheets:
        wv = values[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    n_formulas += 1
                    cached = wv[cell.coordinate].value
                    if cached is None:
                        problems += 1
                        print(f"{path} | {ws.title}!{cell.coordinate}: no cached value | {cell.value[:90]}")
                    elif isinstance(cached, str) and cached.strip() in ERRORS:
                        problems += 1
                        print(f"{path} | {ws.title}!{cell.coordinate}: {cached} | {cell.value[:90]}")
    print(f"{path}: {n_formulas} formula cells, {problems} problems")
    return problems


def main(argv: list[str]) -> int:
    total = sum(scan(p) for p in argv)
    return 1 if total else 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
