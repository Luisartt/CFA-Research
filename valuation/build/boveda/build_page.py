"""Assemble the Boveda Vision page: template + model.js + data.json (+ walk and audit from the gate file).

python -I build_page.py --project <project root>
Writes boveda-vision.html (publish) and preview.html (local look) next to this script.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

import yaml

HERE = Path(__file__).resolve().parent


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--project", default=".")
    ap.add_argument("--js-test", default="ver prueba JS", help="result of node test_model.js, e.g. 39 / 39")
    a = ap.parse_args()
    project = Path(a.project).resolve()
    data = json.loads((HERE / "data.json").read_text(encoding="utf-8"))
    data.pop("reference", None)
    gate = yaml.safe_load((project / "valuation" / "valuation-outputs.yaml").read_text(encoding="utf-8"))
    ver = json.loads((project / "valuation" / "verification-v3.json").read_text(encoding="utf-8"))
    w1, w2 = gate["v1_to_v2_walk"], gate["v2_to_v3_walk"]
    data["walk"] = {"steps": [
        {"label": "v1 auditada (motor)", "d": w1["v1_engine"], "total": True},
        {"label": "Capital de trabajo sin partidas financieras", "d": w1["working_capital_fix"]},
        {"label": "Provisiones fiscales y beneficios por país", "d": w1["bridge_benefits_and_provisions"]},
        {"label": "v2 motor", "d": w1["v2_engine"], "total": True},
        {"label": "Múltiplo: solo reclamos que restan los pares", "d": w1["peer_ev_claims_exit"]},
        {"label": "Arrendamientos crecen con las ventas", "d": w1["lease_growth"]},
        {"label": "v2 publicada (ruta A)", "d": w1["v2_published"], "total": True},
        {"label": "Costo de capital: ruta B (rf sin incumplimiento + prima país)", "d": w2["route_b_cost_of_capital"]},
        {"label": "v3 publicada (ruta B, propuesta)", "d": w2["v3_published"], "total": True},
    ]}
    data["refs"] = {
        "consensus": {"low": 51.0, "high": 87.0, "mean": 70.17, "median": 69.5, "n": 12, "date": "20-sep-2026",
                      "source": "[sourced] CIQ Estimate Highlights, data/capitaliq/consensus.csv"},
        "thesis": {"low": 59.9, "high": 75.2, "rec": "COMPRA", "date": "2026-10-08",
                   "source": "[calc] output/bimbo-tesis-3-pilares.md, Pilar 3: P/E 18.84x y EV/EBITDA 7.68x promedio a 5 años"},
    }
    dated = gate["published"]["dated_cross_check"]
    data["audit"] = {
        "rounds": [
            {"title": "Ronda 1 · auditoría en frío de la v1",
             "text": "Recalculó todo en un Python independiente y reprodujo el valor a precisión completa. Encontró 14 problemas "
                     "(3 altos, 9 medios, 2 bajos): capital de trabajo con partidas financieras, reclamos del múltiplo, NCI, "
                     "beneficios a empleados, arrendamientos, etiquetas y convenciones.", "pass": False},
            {"title": "Ronda 2 · revisión de las correcciones",
             "text": "Confirmó las correcciones y dejó el NCI como juicio provisional. Encontró un problema nuevo: impuestos "
                     "estacionales dentro del capital de trabajo del 1S26.", "pass": False},
            {"title": "Ronda 3 · AUDIT RESULT: PASS",
             "text": "Codex verificó la verificación fechada al 30-jun-2026 (Gordon 35.94 y múltiplo 46.47 en ese momento) y "
                     "confirmó que no hay doble conteo de impuestos.", "pass": True},
            {"title": "Ronda 4 · plan de integración v3: APPROVED",
             "text": "Codex revisó en tres vueltas el plan para integrar los datos de Capital IQ y la ruta B. Encontró el doble "
                     "conteo de la depreciación del peso en el escenario de paridad, libros sin resultados en caché después "
                     "de recalcular y controles débiles de la prima país; todo quedó corregido antes de aprobar.", "pass": True},
        ],
        "update": (f"v3 del 17-oct: tasa libre de riesgo en pesos sin riesgo de incumplimiento (M-Bono 9.26% menos el "
                   f"diferencial de México 1.75%, Damodaran jul-2026) más una prima país de "
                   f"{gate['published']['country_risk_premium']:.2%} ponderada por el EBITDA ajustado FY2025 de cada país. "
                   f"Pares v3 al tipo de cambio del 16-oct (18.479) y Gruma dentro de las medianas NTM y P/U. La verificación "
                   f"fechada al 30-jun da Gordon {dated['gordon']:.2f} ({dated['vs_fye_gordon'] * 100:+.1f}%) y múltiplo "
                   f"{dated['exit']:.2f} ({dated['vs_fye_exit'] * 100:+.1f}%) contra el traslado desde diciembre."),
        "checks": [[f"Controles del Excel DCF v3", f"{ver['dcf']['checks_ok']} / {ver['dcf']['checks_total']} OK"],
                   ["Conciliación con el motor del modelo", f"{ver['dcf']['engine_ok']} / {ver['dcf']['engine_total']}"],
                   ["Conciliación con la réplica en Python", f"{ver['dcf']['replica_ok']} / {ver['dcf']['replica_total']}"],
                   [f"Errores de fórmula ({ver['formula_cells_dcf_wacc']:,} celdas DCF y WACC)", str(ver["formula_errors"])],
                   ["Pares v3 contra el cálculo independiente", f"{ver['peers']['spec_ok']} / {ver['peers']['spec_total']}"],
                   ["Motor de esta página contra Python", a.js_test]],
        "pending": ["Ruta del costo de capital: la B es la propuesta base y la A queda como alternativa; decide el equipo",
                    "Pares: cierre oficial de BMV para Gruma; NCI en libros de Aryzta y Arca; arrendamientos operativos de "
                    "Yamazaki (J-GAAP); consenso de Gruma en su moneda de origen",
                    "Verificación fechada: el EBIT del 1S26 mezcla bases (CIQ antes de reestructura contra 16,738.3 de BMV); "
                    "corregirlo sube ~0.33 por acción la verificación, no el valor publicado",
                    "Acciones en circulación: BMV muestra 4,284.6 millones contra 4,294.9 millones al 29-jul (recompras después de julio)",
                    "Resueltos el 17-oct: M-Bono 9.26%, CDS y tasas de EUA con extracto de Capital IQ (prompt 2); beta por "
                    "regresión de Bimbo contra el IPC; datos de mercado, utilidad bruta y capex de consenso de los pares"],
        "files": [ver["files"][k]["path"] for k in ("dcf", "wacc", "peers", "dcf_es", "wacc_es", "peers_es")] +
                 ["valuation/valuation.md", "valuation/valuation-outputs.yaml", "valuation/verification-v3.json",
                  ver["engine_model"], "model/review.md"],
    }
    # Financials and peers from the recalculated peer workbook v3 (template front sheets)
    from openpyxl import load_workbook
    pw = load_workbook(project / "valuation" / "BIMBOA_PeerAnalysis_v3.xlsx", data_only=True)
    ts, pa = pw["TearSheet"], pw["PeerAnalysis"]

    def num(v):
        return v if isinstance(v, (int, float)) else None

    def ts_row(r):
        return [num(ts.cell(r, c).value) for c in range(5, 12)]

    data["fin"] = {
        "years": [ts.cell(27, c).value for c in range(5, 12)],
        "flags": [ts.cell(26, c).value for c in range(5, 12)],
        "rows": [{"label": "Ventas", "v": ts_row(28), "fmt": "m"}, {"label": "Crecimiento anual", "v": ts_row(29), "fmt": "p"},
                 {"label": "Margen bruto", "v": ts_row(57), "fmt": "p"}, {"label": "EBITDA (ajustado)", "v": ts_row(34), "fmt": "m"},
                 {"label": "Margen EBITDA", "v": ts_row(60), "fmt": "p"}, {"label": "Utilidad neta", "v": ts_row(36), "fmt": "m"},
                 {"label": "UPA (MXN)", "v": ts_row(37), "fmt": "d"}, {"label": "Dividendo por acción (MXN)", "v": ts_row(39), "fmt": "d"},
                 {"label": "Capex", "v": ts_row(50), "fmt": "m"},
                 {"label": "Deuda neta (A: con arrendamientos; E: consenso)", "v": ts_row(47), "fmt": "m"}],
        "source": "valuation/BIMBOA_PeerAnalysis_v3.xlsx, hoja TearSheet: real = CIQ y reportes; E = promedio de consenso CIQ (20-sep-2026)",
    }
    peers = []
    for i in range(10):                       # claim-basis block: 9 peers then Bimbo (rows 64-73)
        r, rm, rp = 64 + i, 13 + i, 50 + i
        peers.append({"name": pa.cell(r, 2).value, "group": pa.cell(r, 4).value, "ticker": pa.cell(r, 5).value,
                      "ev_ltm": num(pa.cell(r, 6).value), "ev_ntm": num(pa.cell(r, 7).value), "ev_sales": num(pa.cell(r, 8).value),
                      "pe_ltm": num(pa.cell(r, 9).value), "pe_ntm": num(pa.cell(r, 10).value), "nd_ebitda": num(pa.cell(r, 11).value),
                      "margin": num(pa.cell(r, 12).value), "elig": pa.cell(r, 13).value, "mcap": num(pa.cell(rm, 9).value),
                      "fcf_yield": num(pa.cell(rp, 9).value), "ps_ltm": num(pa.cell(rp, 15).value)})
    stats = []
    for r, label in ((74, "Mediana del núcleo (elegibles)"), (75, "Mediana de todos los pares (elegibles)"), (76, "Promedio de todos los pares")):
        stats.append({"label": label, "ev_ltm": num(pa.cell(r, 6).value), "ev_ntm": num(pa.cell(r, 7).value),
                      "ev_sales": num(pa.cell(r, 8).value), "pe_ltm": num(pa.cell(r, 9).value), "pe_ntm": num(pa.cell(r, 10).value),
                      "nd_ebitda": num(pa.cell(r, 11).value), "margin": num(pa.cell(r, 12).value)})
    data["peer_table"] = {"rows": peers, "stats": stats,
                          "source": "valuation/BIMBOA_PeerAnalysis_v3.xlsx, hoja PeerAnalysis (base de reclamos, normalizada por arrendamientos)"}
    blob = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    tpl = (HERE / "page.html").read_text(encoding="utf-8")
    model = (HERE / "model.js").read_text(encoding="utf-8")
    assert "__DATA_JSON__" in tpl and "/*__MODEL_JS__*/" in tpl
    page = tpl.replace("__DATA_JSON__", blob).replace("/*__MODEL_JS__*/", model)
    (HERE / "boveda-vision.html").write_text(page, encoding="utf-8")
    preview = ("<!doctype html><html lang=\"es\"><head><meta charset=\"utf-8\"><meta name=\"viewport\" "
               "content=\"width=device-width, initial-scale=1, viewport-fit=cover\"><style>body{margin:0}"
               "[hidden]{display:none!important}</style></head><body>" + page + "</body></html>")
    (HERE / "preview.html").write_text(preview, encoding="utf-8")
    print("wrote boveda-vision.html", len(page.encode("utf-8")), "bytes")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
