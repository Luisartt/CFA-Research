"""Build analisis-dcf-2026-10-26/evidencia/BIMBOA_FX_premium_evidence_v1.xlsx: is the 0.63 pp gap a premium or expected peso depreciation?

Audit finding V-1 (2026-10-26). The MXN default-free rate of the WACC (M-Bono 10y less Mexico's default spread) is
2.30% a year above the US Treasury 10y. Expected depreciation by relative PPP at the Banxico survey's long-run
inflation (3.70%) vs 2% is 1.67%. This workbook tests, with public data and live formulas, how much of the 2.30% is
expected depreciation (belongs in the cash flows) and how much is premium (belongs only in the discount rate).

Every downloaded value is typed as stored in its source file [sourced]; every result is a live Excel formula [calc].
Sources (all public, downloaded 2026-10-26 into analisis-dcf-2026-10-26/datos/fx_premium/):
  FRED: DEXMXUS (daily and monthly average), DGS10, DFII10, T10YIE, EXPINF10YR (Cleveland Fed),
        IRLTLT01MXM156N and IRLTLT01USM156N (OECD 10y government yields), CPIAUCSL, MEXCPIALLMINMEI (OECD)
  Banxico SIE cuadro CF107: SF44071 Bono tasa fija 10 anos, SF43924 Udibono 10 anos (real), auction yields
  Banxico, Encuesta sobre las Expectativas de los Especialistas, septiembre 2026 (published 1-Oct-2026)

Run from the repo root:  python analisis-dcf-2026-10-26/build/build_fx_premium_evidence.py
then recalculate:        powershell -File valuation/build/recalc_excel.ps1 <Windows path of the workbook>
"""
import csv
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "analisis-dcf-2026-10-26" / "datos" / "fx_premium"
OUT = ROOT / "analisis-dcf-2026-10-26" / "evidencia" / "BIMBOA_FX_premium_evidence_v1.xlsx"
VAL_DATE = dt.date(2026, 10, 16)  # valuation date of the published DCF (valuation/valuation-outputs.yaml)

TITLE = Font(name="Arial", size=12, bold=True)
BOLD = Font(name="Arial", size=9, bold=True)
BLACK = Font(name="Arial", size=9)
BLUE = Font(name="Arial", size=9, color="0000FF")
NOTE = Font(name="Arial", size=8, italic=True, color="595959")
HDR = Font(name="Arial", size=9, bold=True, color="FFFFFF")
FILL_HDR = PatternFill("solid", fgColor="1F4E78")
FILL_KEY = PatternFill("solid", fgColor="FFF2CC")
PCT = '0.00%;(0.00%)'


def fred(name):
    """FRED csv -> {date: float}; blank or '.' = missing."""
    out = {}
    with open(SRC / name, encoding="utf-8") as f:
        r = csv.reader(f)
        next(r)
        for d, v in r:
            if v not in ("", "."):
                out[dt.date.fromisoformat(d)] = float(v)
    return out


def banxico_auctions():
    rows = []
    with open(SRC / "banxico_CF107_bono10_udibono10.csv", encoding="latin-1") as f:
        lines = list(csv.reader(f))
    title = next(l for l in lines if l and l[0].startswith("T"))
    assert "Bono tasa fija 10" in title[1] and "Udibonos a 10" in title[2], title
    for l in lines:
        if len(l) == 3 and l[0][:2].isdigit() and "/" in l[0]:
            d = dt.datetime.strptime(l[0], "%d/%m/%Y").date()
            b = float(l[1]) if l[1] != "N/E" else None
            u = float(l[2]) if l[2] != "N/E" else None
            if b is not None or u is not None:
                rows.append((d, b, u))
    return rows


def hdr(ws, row, labels, widths=None):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h)
        c.font, c.fill = HDR, FILL_HDR
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate(widths or [], 1):
        ws.column_dimensions[chr(64 + j)].width = w


def main():
    wb = Workbook()
    rd = wb.active
    rd.title = "README"

    # ---------------- Data_daily (from 2000) ----------------
    daily = {k: fred(f"fred_{k}.csv") for k in ("DEXMXUS", "DGS10", "DFII10", "T10YIE")}
    dates = sorted(d for d in set().union(*[set(v) for v in daily.values()]) if d >= dt.date(2000, 1, 1))
    dd = wb.create_sheet("Data_daily")
    dd["A1"] = "FRED daily series as stored [sourced]; percent yields as stored (5.09 = 5.09%). Files: analisis-dcf-2026-10-26/datos/fx_premium/fred_*.csv"
    dd["A1"].font = NOTE
    hdr(dd, 2, ["date", "DEXMXUS (MXN per USD)", "DGS10 UST 10y %", "DFII10 TIPS 10y real %", "T10YIE US 10y breakeven %"],
        [11, 14, 12, 14, 16])
    for i, d in enumerate(dates, 3):
        dd.cell(i, 1, d).number_format = "yyyy-mm-dd"
        for j, k in enumerate(("DEXMXUS", "DGS10", "DFII10", "T10YIE"), 2):
            if d in daily[k]:
                dd.cell(i, j, daily[k][d]).font = BLUE
    dlast = 2 + len(dates)

    # ---------------- Data_monthly ----------------
    m = {"FX": fred("fred_DEXMXUS_monthly_avg.csv"), "MX10": fred("fred_IRLTLT01MXM156N.csv"),
         "US10": fred("fred_IRLTLT01USM156N.csv"), "USCPI": fred("fred_CPIAUCSL.csv"),
         "MXCPI": fred("fred_MEXCPIALLMINMEI.csv"), "EXPINF": fred("fred_EXPINF10YR.csv")}
    months = []
    y, mo = 2000, 1
    while (y, mo) <= (2026, 10):
        months.append(dt.date(y, mo, 1))
        y, mo = (y + 1, 1) if mo == 12 else (y, mo + 1)
    dm = wb.create_sheet("Data_monthly")
    dm["A1"] = ("FRED monthly series as stored [sourced]: DEXMXUS monthly average (FRED aggregation), OECD 10y government yields "
                "(IRLTLT01MXM156N, IRLTLT01USM156N), CPIAUCSL, MEXCPIALLMINMEI (ends Jul-2024), EXPINF10YR (Cleveland Fed 10y expected inflation)")
    dm["A1"].font = NOTE
    keys = ("FX", "MX10", "US10", "USCPI", "MXCPI", "EXPINF")
    hdr(dm, 2, ["month", "USD/MXN avg", "Mexico 10y %", "US 10y %", "US CPI index", "Mexico CPI index", "US 10y expected inflation % (Cleveland Fed)"],
        [11, 11, 11, 11, 12, 13, 18])
    mrow = {}
    for i, d in enumerate(months, 3):
        mrow[d] = i
        dm.cell(i, 1, d).number_format = "yyyy-mm"
        for j, k in enumerate(keys, 2):
            if d in m[k]:
                dm.cell(i, j, m[k][d]).font = BLUE

    # ---------------- Data_auctions ----------------
    au = banxico_auctions()
    da = wb.create_sheet("Data_auctions")
    da["A1"] = ("Banxico SIE cuadro CF107, weekly auction results [sourced]: SF44071 Bono tasa fija 10 anos yield %, SF43924 Udibono "
                "10 anos real yield %. File analisis-dcf-2026-10-26/datos/fx_premium/banxico_CF107_bono10_udibono10.csv (consulted 26-Oct-2026)")
    da["A1"].font = NOTE
    hdr(da, 2, ["auction date", "Bono M 10y %", "Udibono 10y real %"], [12, 12, 14])
    for i, (d, b, u) in enumerate(au, 3):
        da.cell(i, 1, d).number_format = "yyyy-mm-dd"
        if b is not None:
            da.cell(i, 2, b).font = BLUE
        if u is not None:
            da.cell(i, 3, u).font = BLUE
    alast = 2 + len(au)

    # ---------------- Inputs ----------------
    ip = wb.create_sheet("Inputs")
    ip["A1"] = "Inputs typed from documents (not downloadable as series) and the WACC inputs being tested"
    ip["A1"].font = TITLE
    hdr(ip, 3, ["Item", "Value", "Tag", "Source"], [62, 11, 11, 100])
    inputs = [
        ("vd", "Valuation date", VAL_DATE, "[sourced]", "valuation/valuation-outputs.yaml valuation_date (published DCF v3)"),
        ("mb", "M-Bono 10y used in the WACC (%)", 9.26, "[sourced]", "S&P Capital IQ 16-Oct-2026; valuation/valuation.yaml wacc.risk_free comment"),
        ("ds", "Mexico default spread removed in route B (%)", 1.75, "[sourced]", "Damodaran July 2026, Baa3; valuation/valuation.yaml"),
        ("pi_mx", "Mexico expected inflation, 5-8 years (2031-2034), median of the survey (%)", 3.70, "[sourced]",
         "Banxico, Encuesta sobre las Expectativas de los Especialistas, septiembre 2026, Cuadro 5 p.6 (media 3.69)"),
        ("pi_mx14", "Mexico expected inflation, 1-4 years (2027-2030), median of the survey (%)", 3.80, "[sourced]",
         "Same survey, Cuadro 5 p.6 (media 3.77)"),
        ("pi_fed", "US inflation target (%)", 2.00, "[sourced]", "Federal Reserve longer-run goal (PCE)"),
        ("sv26", "Survey: USD/MXN at end-2026, median", 17.50, "[sourced]", "Same survey, Cuadro 11 p.13 (media 17.57)"),
        ("sv27", "Survey: USD/MXN at end-2027, median", 18.04, "[sourced]", "Same survey, Cuadro 11 p.13 (media 18.11)"),
        ("sv28", "Survey: USD/MXN at end-2028, median", 18.64, "[sourced]", "Same survey, Cuadro 11 p.13 (media 18.71)"),
        ("cq27", "Capital IQ: USD/MXN annual average 2027", 18.21, "[sourced]",
         "BIMBOA/research/Informacion_BIMBO_macro.txt L1128 (macro workbook sheet FX); same path the statements use"),
        ("cq30", "Capital IQ: USD/MXN annual average 2030", 19.24, "[sourced]", "Same source"),
    ]
    rn = {}
    for i, (k, lab, v, tag, src) in enumerate(inputs, 4):
        rn[k] = i
        ip.cell(i, 1, lab).font = BLACK
        c = ip.cell(i, 2, v)
        c.font = BLUE
        c.number_format = "yyyy-mm-dd" if isinstance(v, dt.date) else "0.00"
        ip.cell(i, 3, tag).font = NOTE
        ip.cell(i, 4, src).font = NOTE

    def I(k):
        return "Inputs!$B$%d" % rn[k]

    # ---------------- Decomposition at the valuation date ----------------
    de = wb.create_sheet("Decomposition")
    de["A1"] = "Test 1-3: what the market prices at the valuation date (all rates 10 years, annual)"
    de["A1"].font = TITLE
    de["A2"] = ("Identity: (1 + MXN nominal) / (1 + USD nominal) = [(1 + MXN real) / (1 + USD real)] x [(1 + MXN breakeven) / (1 + USD breakeven)]. "
                "Under relative PPP, expected depreciation = expected inflation differential; everything else in the rate gap is premium.")
    de["A2"].font = NOTE
    hdr(de, 4, ["Line", "Value", "Tag", "Formula / note"], [70, 11, 9, 100])
    D = "Data_daily!$A$3:$A$%d" % dlast
    col = lambda c: "Data_daily!$%s$3:$%s$%d" % (c, c, dlast)
    A_D, A_B, A_U = "Data_auctions!$A$3:$A$%d" % alast, "Data_auctions!$B$3:$B$%d" % alast, "Data_auctions!$C$3:$C$%d" % alast
    M_A, M_G = "Data_monthly!$A$3:$A$%d" % (2 + len(months)), "Data_monthly!$G$3:$G$%d" % (2 + len(months))
    rows = [
        ("h1", "Market inputs", None, "", ""),
        ("ust", "UST 10y nominal", "=INDEX(%s,MATCH(%s,%s,0))/100" % (col("C"), I("vd"), D), "[sourced]", "FRED DGS10 on the valuation date"),
        ("tips", "TIPS 10y real", "=INDEX(%s,MATCH(%s,%s,0))/100" % (col("D"), I("vd"), D), "[sourced]", "FRED DFII10 on the valuation date"),
        ("bei_fred", "US 10y breakeven inflation as published (DGS10 - DFII10, additive)", "=INDEX(%s,MATCH(%s,%s,0))/100" % (col("E"), I("vd"), D),
         "[sourced]", "FRED T10YIE on the valuation date; cross-check only"),
        ("fx0", "USD/MXN spot", "=INDEX(%s,MATCH(%s,%s,0))" % (col("B"), I("vd"), D), "[sourced]", "FRED DEXMXUS on the valuation date"),
        ("mb", "M-Bono 10y nominal (WACC input)", "=%s/100" % I("mb"), "[sourced]", "Capital IQ, as in the WACC"),
        ("mb_ax", "  cross-check: Bono M 10y, last Banxico auction on or before the date",
         "=LOOKUP(2,1/((%s<=%s)*(%s<>\"\")),%s)/100" % (A_D, I("vd"), A_B, A_B), "[sourced]", "Banxico SIE SF44071"),
        ("udi", "Udibono 10y real, last Banxico auction on or before the date",
         "=LOOKUP(2,1/((%s<=%s)*(%s<>\"\")),%s)/100" % (A_D, I("vd"), A_U, A_U), "[sourced]", "Banxico SIE SF43924"),
        ("udi_d", "  date of that Udibono auction", "=LOOKUP(2,1/((%s<=%s)*(%s<>\"\")),%s)" % (A_D, I("vd"), A_U, A_D),
         "[sourced]", "Auctions are every few weeks; the secondary-market yield moves daily"),
        ("ds", "Mexico default spread (route B)", "=%s/100" % I("ds"), "[sourced]", "Damodaran"),
        ("pi_mx", "Mexico expected inflation 5-8 years (survey median)", "=%s/100" % I("pi_mx"), "[sourced]", "Banxico survey Sep-2026"),
        ("pi_us", "US expected inflation 10y, latest month on or before the date (Cleveland Fed)",
         "=LOOKUP(2,1/((%s<=%s)*(%s<>\"\")),%s)/100" % (M_A, I("vd"), M_G, M_G), "[sourced]", "FRED EXPINF10YR (model of yields, inflation, swaps and surveys)"),
        ("pi_fed", "US inflation target", "=%s/100" % I("pi_fed"), "[sourced]", "Fed"),
        ("b1", "", None, "", ""),
        ("h2", "The gap being tested", None, "", ""),
        ("gap_nom", "Nominal gap M-Bono vs UST", "=(1+B{mb})/(1+B{ust})-1", "[calc]", "Includes Mexico's default risk"),
        ("gap", "Default-free gap used by the WACC (route B): the 2.30%", "=(1+B{mb}-B{ds})/(1+B{ust})-1", "[calc]",
         "(1 + M-Bono - default spread) / (1 + UST) - 1"),
        ("b2", "", None, "", ""),
        ("h3", "Test 1-2: split of the nominal gap into real and inflation parts", None, "", ""),
        ("bei_mx", "Mexico 10y breakeven inflation", "=(1+B{mb})/(1+B{udi})-1", "[calc]", "(1 + M-Bono) / (1 + Udibono) - 1 = expected inflation + inflation risk premium"),
        ("bei_us", "US 10y breakeven inflation (multiplicative, same convention as Mexico)", "=(1+B{ust})/(1+B{tips})-1", "[calc]",
         "(1 + UST) / (1 + TIPS) - 1; T10YIE above uses the additive convention"),
        ("bei_gap", "Breakeven gap Mexico vs US", "=(1+B{bei_mx})/(1+B{bei_us})-1", "[calc]", "Upper bound of the expected inflation differential (still includes inflation risk premia)"),
        ("real_gap", "Real gap Udibono vs TIPS", "=(1+B{udi})/(1+B{tips})-1", "[calc]", "Credit risk + real-rate premium; not depreciation under relative PPP"),
        ("chk", "Check: real gap x breakeven gap = nominal gap (must be 0)", "=(1+B{real_gap})*(1+B{bei_gap})-(1+B{gap_nom})", "[calc]", "Identity"),
        ("b3", "", None, "", ""),
        ("h4", "Test 3: expected inflation (surveys and models) -> expected depreciation and premium", None, "", ""),
        ("irp_mx", "Mexico inflation risk premium = breakeven vs survey", "=(1+B{bei_mx})/(1+B{pi_mx})-1", "[calc]", "Premium, not expected inflation"),
        ("irp_us", "US inflation risk premium = breakeven vs Cleveland Fed", "=(1+B{bei_us})/(1+B{pi_us})-1", "[calc]", "Negative = TIPS liquidity premium dominates"),
        ("dep_cl", "Expected depreciation, relative PPP (survey 3.70% vs Cleveland Fed US)", "=(1+B{pi_mx})/(1+B{pi_us})-1", "[calc]", "Market-based US expectation"),
        ("dep_fed", "Expected depreciation, relative PPP (survey 3.70% vs Fed target 2%)", "=(1+B{pi_mx})/(1+B{pi_fed})-1", "[calc]", "The 1.67% of the debate"),
        ("prem_cl", "Premium inside the 2.30% (with Cleveland Fed US)", "=(1+B{gap})/(1+B{dep_cl})-1", "[calc]", "Gap not explained by expected inflation"),
        ("prem_fed", "Premium inside the 2.30% (with Fed target): the 0.63 of the debate", "=(1+B{gap})/(1+B{dep_fed})-1", "[calc]", ""),
    ]
    drn = {k: 5 + i for i, (k, *_) in enumerate(rows)}
    for k, lab, f, tag, note in rows:
        r = drn[k]
        c = de.cell(r, 1, lab)
        c.font = BOLD if k.startswith("h") else BLACK
        if f is not None:
            v = de.cell(r, 2, f.format(**{kk: vv for kk, vv in drn.items()}) if "{" in f else f)
            v.number_format = "yyyy-mm-dd" if k == "udi_d" else ("0.0000" if k == "fx0" else ("0.000000" if k == "chk" else PCT))
            if k in ("gap", "dep_fed", "dep_cl", "prem_fed", "prem_cl"):
                v.fill = FILL_KEY
        de.cell(r, 3, tag).font = NOTE
        de.cell(r, 4, note).font = NOTE

    # ---------------- Forecasts ----------------
    fc = wb.create_sheet("Forecasts")
    fc["A1"] = "Test 4: what forecasters expect for USD/MXN (no premium inside a forecast)"
    fc["A1"].font = TITLE
    hdr(fc, 3, ["Line", "Value", "Tag", "Formula / note"], [70, 11, 9, 90])
    frows = [
        ("sv_a", "Survey: depreciation a year, end-2026 to end-2028 (median)", "=(%s/%s)^(1/2)-1" % (I("sv28"), I("sv26")), "Short horizon (2 years)"),
        ("cq_a", "Capital IQ: depreciation a year, 2027 to 2030 average", "=(%s/%s)^(1/3)-1" % (I("cq30"), I("cq27")), "Path used by the statements"),
        ("sv_p", "Premium inside the 2.30% if the survey path is the expectation", "=(1+Decomposition!B%d)/(1+B4)-1" % drn["gap"], "Negative = forecasters expect more depreciation than the gap"),
        ("cq_p", "Premium inside the 2.30% if the Capital IQ path is the expectation", "=(1+Decomposition!B%d)/(1+B5)-1" % drn["gap"], ""),
        ("sv_s", "Survey: depreciation a year, spot on the valuation date to end-2028 (median)",
         "=(%s/Decomposition!B%d)^(365/(DATE(2028,12,31)-%s))-1" % (I("sv28"), drn["fx0"], I("vd")),
         "Same survey, measured from today's spot: the end-2026 median (17.50) is below spot, so the path from spot is flatter"),
        ("sv_sp", "Premium inside the 2.30% if the spot-to-2028 survey path is the expectation", "=(1+Decomposition!B%d)/(1+B8)-1" % drn["gap"], ""),
    ]
    for i, (k, lab, f, note) in enumerate(frows, 4):
        fc.cell(i, 1, lab).font = BLACK
        c = fc.cell(i, 2, f)
        c.number_format = PCT
        c.fill = FILL_KEY
        fc.cell(i, 3, "[calc]").font = NOTE
        fc.cell(i, 4, note).font = NOTE

    # ---------------- History ----------------
    hi = wb.create_sheet("History")
    hi["A1"] = "Test 5: did the peso depreciate by the rate gap, or by the inflation gap? Rolling 10-year windows, monthly starts"
    hi["A1"].font = TITLE
    hi["A2"] = ("Mexico yield = last Banxico Bono M 10y auction on or before the start month's end (the OECD monthly series has empty months); "
                "US yield = OECD IRLTLT01USM156N. The yield gap uses full government yields, so it still contains Mexico's default premium; no default-spread history is "
                "available, so this test bounds the premium from above. Inflation test ends with windows starting Jul-2014 (Mexico CPI on FRED ends Jul-2024).")
    hi["A2"].font = NOTE
    hdr(hi, 4, ["start month", "USD/MXN start", "USD/MXN +10y", "realized depreciation a year", "yield gap at start (Bono M 10y last auction vs US 10y)",
                "realized inflation gap a year", "yield gap - depreciation", "depreciation - inflation gap"],
        [11, 11, 11, 13, 14, 13, 13, 13])
    starts = [d for d in months if dt.date(2001, 7, 1) <= d <= dt.date(2016, 9, 1)]
    for i, d in enumerate(starts, 5):
        r0 = mrow[d]
        r1 = r0 + 120
        hi.cell(i, 1, "=Data_monthly!A%d" % r0).number_format = "yyyy-mm"
        hi.cell(i, 2, "=Data_monthly!B%d" % r0).number_format = "0.0000"
        hi.cell(i, 3, "=Data_monthly!B%d" % r1).number_format = "0.0000"
        hi.cell(i, 4, "=(C%d/B%d)^(1/10)-1" % (i, i)).number_format = PCT
        hi.cell(i, 5, "=(1+LOOKUP(2,1/((%s<=EOMONTH(A%d,0))*(%s<>\"\")),%s)/100)/(1+Data_monthly!D%d/100)-1"
                % (A_D, i, A_B, A_B, r0)).number_format = PCT
        hi.cell(i, 6, '=IF(AND(ISNUMBER(Data_monthly!F%d),ISNUMBER(Data_monthly!F%d)),((Data_monthly!F%d/Data_monthly!F%d)/(Data_monthly!E%d/Data_monthly!E%d))^(1/10)-1,"")'
                % (r1, r0, r1, r0, r1, r0)).number_format = PCT
        hi.cell(i, 7, "=E%d-D%d" % (i, i)).number_format = PCT
        hi.cell(i, 8, '=IF(F%d="","",D%d-F%d)' % (i, i, i)).number_format = PCT
    hl = 4 + len(starts)
    s0 = hl + 2
    summ = [
        ("Windows", "=COUNT(D5:D%d)" % hl, "0"),
        ("Average realized depreciation a year", "=AVERAGE(D5:D%d)" % hl, PCT),
        ("Average yield gap at start", "=AVERAGE(E5:E%d)" % hl, PCT),
        ("Average yield gap - depreciation (premium incl. default risk)", "=AVERAGE(G5:G%d)" % hl, PCT),
        ("Share of windows where the peso depreciated less than the yield gap", "=COUNTIF(G5:G%d,\">0\")/COUNT(G5:G%d)" % (hl, hl), "0%"),
        ("Windows with inflation data", "=COUNT(F5:F%d)" % hl, "0"),
        ("Average realized inflation gap a year", "=AVERAGE(F5:F%d)" % hl, PCT),
        ("Average depreciation - inflation gap (relative PPP error)", "=AVERAGE(H5:H%d)" % hl, PCT),
    ]
    hi.cell(s0, 1, "Summary").font = BOLD
    for j, (lab, f, fmt) in enumerate(summ, s0 + 1):
        hi.cell(j, 1, lab).font = BLACK
        c = hi.cell(j, 4, f)
        c.number_format = fmt
        c.fill = FILL_KEY
    hs = {lab: s0 + 1 + j for j, (lab, *_) in enumerate(summ)}

    # ---------------- Summary ----------------
    su = wb.create_sheet("Summary", 1)
    su["A1"] = "Is the 0.63 pp a premium? Evidence at a glance (live links)"
    su["A1"].font = TITLE
    su["A2"] = ("Gap tested = MXN default-free 10y vs UST 10y (route B of the WACC). Expected depreciation belongs in the cash flows and g; "
                "premium belongs only in the discount rate.")
    su["A2"].font = NOTE
    hdr(su, 4, ["Test", "Expected depreciation a year", "Premium inside the gap", "What it measures / limit"], [58, 16, 16, 90])
    srows = [
        ("Gap in the WACC (route B)", "=Decomposition!B%d" % drn["gap"], None, "The number to split"),
        ("1-3a. Relative PPP: survey 3.70% vs Fed target 2%", "=Decomposition!B%d" % drn["dep_fed"], "=Decomposition!B%d" % drn["prem_fed"],
         "Long-run survey; US side is a target, not an expectation"),
        ("1-3b. Relative PPP: survey 3.70% vs Cleveland Fed expected inflation", "=Decomposition!B%d" % drn["dep_cl"], "=Decomposition!B%d" % drn["prem_cl"],
         "Both sides are expectations"),
        ("1-2. Breakeven gap (upper bound of expected inflation gap)", "=Decomposition!B%d" % drn["bei_gap"], "=(1+Summary!B5)/(1+B8)-1",
         "Still contains inflation risk premia; Udibono from the last auction"),
        ("4a. Banxico survey FX path 2026-2028", "=Forecasts!B4", "=Forecasts!B6", "2-year horizon only"),
        ("4a'. Banxico survey FX path, spot to end-2028", "=Forecasts!B8", "=Forecasts!B9", "Same survey measured from spot"),
        ("4b. Capital IQ FX path 2027-2030", "=Forecasts!B5", "=Forecasts!B7", "Forecast already used by the statements"),
        ("5a. History: realized depreciation vs yield gap (2001-2016 starts)", "=History!D%d" % hs["Average realized depreciation a year"],
         "=History!D%d" % hs["Average yield gap - depreciation (premium incl. default risk)"], "Gap includes default premium (upper bound)"),
        ("5b. History: realized depreciation vs inflation gap (PPP check)", "=History!D%d" % hs["Average realized inflation gap a year"],
         "=History!D%d" % hs["Average depreciation - inflation gap (relative PPP error)"],
         "Column B = realized inflation gap; column C = realized depreciation in excess of it (real depreciation), not a premium"),
    ]
    for i, (lab, a, b, note) in enumerate(srows, 5):
        su.cell(i, 1, lab).font = BOLD if i == 5 else BLACK
        c = su.cell(i, 2, a)
        c.number_format = PCT
        if b:
            c2 = su.cell(i, 3, b)
            c2.number_format = PCT
        su.cell(i, 4, note).font = NOTE

    # ---------------- README ----------------
    lines = [
        ("Grupo Bimbo - Is the 0.63 pp gap a premium or expected peso depreciation? (audit finding V-1, 2026-10-26)", TITLE),
        ("Built by analisis-dcf-2026-10-26/build/build_fx_premium_evidence.py from public data in analisis-dcf-2026-10-26/datos/fx_premium/ (downloaded 2026-10-26 with the team's approval).", NOTE),
        ("Blue = typed from a source file [sourced]. Black = live formula [calc]. Yellow = key results.", NOTE),
        ("", None),
        ("Sheets: Summary (results), Decomposition (tests 1-3 at the valuation date), Forecasts (test 4), History (test 5), Inputs (typed from documents), "
         "Data_daily, Data_monthly, Data_auctions (downloaded series).", None),
        ("Rule: expected depreciation goes in the cash flows and in g; premium stays only in the discount rate. The team decides the base.", None),
    ]
    for i, (txt, f) in enumerate(lines, 1):
        c = rd.cell(i, 1, txt)
        if f:
            c.font = f
    rd.column_dimensions["A"].width = 140
    wb.save(OUT)
    print("wrote", OUT, "| daily rows", len(dates), "| auctions", len(au), "| windows", len(starts))


if __name__ == "__main__":
    main()
