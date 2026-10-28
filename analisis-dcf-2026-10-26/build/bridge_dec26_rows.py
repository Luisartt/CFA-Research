"""Equity bridge at 31-Dec-2026E for the two V1V3 workbooks (team decisions 2026-10-27, stub option B).

Net debt and leases link live to the projected FY2026E balance sheet. NCI at market, debt-like items, associates and
shares are held at their 30-Sep-2026 values (team 2026-10-27). The debt-like total is computed with live formulas in
analisis-dcf-2026-10-26/puente/BIMBOA_bridge_3Q26_v1.xlsx and read from that workbook at build time, so a rebuild picks up
any change there. Used only by build_statements_workbook_V1V3.py and build_statements_workbook_auditfix_V1V3.py.
"""
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
BRIDGE = ROOT / "analisis-dcf-2026-10-26" / "puente" / "BIMBOA_bridge_3Q26_v1.xlsx"
BRIDGE_REL = "analisis-dcf-2026-10-26/puente/BIMBOA_bridge_3Q26_v1.xlsx"
XBRL = "BMV 3Q26 XBRL, analisis-dcf-2026-10-26/datos/bmv_3t26/Reporte_Definitivo_BMV_XBRL_Espanol_Sep_26.pdf"
DL_LABEL = "Debt-like items excluding NCI"
PCT = {"bke", "bgn"}
RATIO = {"bdisc"}


def bridge_debt_like():
    """Debt-like items excluding NCI at 30-Sep-2026 and its cell, from the recalculated bridge workbook."""
    ws = load_workbook(BRIDGE, data_only=True)["Bridge"]
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == DL_LABEL:
            v = ws.cell(r, 2).value
            assert isinstance(v, (int, float)), "recalculate %s first" % BRIDGE_REL
            return v, "Bridge!B%d" % r
    raise KeyError(DL_LABEL)


def rows(bsn, bs_rownum, col):
    """Assumptions rows (key, label, kind, payload, tag, note); col is the FY2026E balance sheet column."""
    bs = lambda k: "%s!%s%d" % (bsn, col, bs_rownum[k])
    dl, cell = bridge_debt_like()
    return [
        ("bl_b", "", "b", None, "", ""),
        ("h_b", "Equity bridge at 31-Dec-2026E (stub option B, team 2026-10-27); DCF not run", "h", None, "", ""),
        ("bdebt", "Gross debt at 31-Dec-2026E", "vf", "=B{debt26}", "[calc]", "Balance sheet FY2026E"),
        ("blease", "Lease liabilities at 31-Dec-2026E", "vf", "=%s+%s" % (bs("leasec"), bs("leasenc")), "[calc]",
         "Balance sheet FY2026E; in net debt because the FCFF deducts new leases"),
        ("bcash", "Cash at 31-Dec-2026E", "vf", "=" + bs("cash"), "[calc]", "Balance sheet FY2026E"),
        ("bnd", "Net debt including leases at 31-Dec-2026E", "vf", "=B{bdebt}+B{blease}-B{bcash}", "[calc]",
         "169,383 reported at 30-Sep-2026 (" + BRIDGE_REL + ")"),
        ("bn9m", "NCI net income 9M26", "v", 1365.514, "[sourced]", XBRL + " [310000] p.13"),
        ("bke", "Cost of equity (v3 route B, reference)", "v", 0.1294365113, "[calc]",
         "valuation/valuation-outputs.yaml cost_of_equity; team confirms the WACC"),
        ("bgn", "NCI growth in perpetuity", "v", 0.0, "[assumption]", "Team 2026-10-26: 0%, the NCI book base shrinks"),
        ("bncim", "NCI at market value = 9M26 x 4 / 3 / (ke - g)", "vf", "=B{bn9m}*4/3/(B{bke}-B{bgn})", "[calc]",
         "Team 2026-10-26/27: 9M26 run-rate, same value at 31-Dec-2026"),
        ("bdl", "Debt-like items excluding NCI, held at 30-Sep-2026", "v", round(dl, 2), "[calc]",
         "Team 2026-10-27: held at the 30-Sep-2026 value; live formula in %s %s, read at build time" % (BRIDGE_REL, cell)),
        ("bassoc", "Investments in associates, held at 30-Sep-2026 (added to EV)", "v", 6532.927, "[sourced]",
         XBRL + " [210000] p.11; team 2026-10-27: held, not the projected balance"),
        ("bclaims", "Claims at 31-Dec-2026E = net debt incl. leases + NCI at market + debt-like - associates", "vf",
         "=B{bnd}+B{bncim}+B{bdl}-B{bassoc}", "[calc]", "Equity value at 31-Dec-2026 = DCF enterprise value - this line"),
        ("bsh_is", "Issued shares (millions), 30-Sep-2026", "v", 4304.744119, "[sourced]", XBRL + " [700000] p.24"),
        ("bsh_tr", "Repurchased shares held (millions), 30-Sep-2026", "v", 20.184108, "[sourced]", XBRL + " [700000] p.24"),
        ("bsh", "Shares outstanding at 31-Dec-2026E (millions)", "vf", "=B{bsh_is}-B{bsh_tr}", "[calc]",
         "Team 2026-10-23: no buybacks in 4Q26, so the 30-Sep-2026 count holds"),
        ("bdays", "Days from the valuation date (16-Oct-2026) to 31-Dec-2026", "vf", "=DATE(2026,12,31)-DATE(2026,10,16)", "[calc]",
         "Valuation date: 16-Oct-2026 BMV close (valuation/valuation.yaml)"),
        ("bdisc", "Discount factor to 16-Oct-2026 = 1 / (1 + ke) ^ (days / 365)", "vf", "=1/(1+B{bke})^(B{bdays}/365)", "[calc]",
         "Equity value at 16-Oct-2026 = (EV at 31-Dec-2026 - claims) x this factor; no dividends fall in the 76 days"),
    ]


ES = {
    "Equity bridge at 31-Dec-2026E (stub option B, team 2026-10-27); DCF not run":
        "Puente al valor del capital al 31-dic-2026E (stub opción B, equipo 2026-10-27); DCF no corrido",
    "Gross debt at 31-Dec-2026E": "Deuda bruta al 31-dic-2026E",
    "Balance sheet FY2026E": "Balance FY2026E",
    "Lease liabilities at 31-Dec-2026E": "Pasivos por arrendamiento al 31-dic-2026E",
    "Balance sheet FY2026E; in net debt because the FCFF deducts new leases":
        "Balance FY2026E; en la deuda neta porque el FCFF resta los arrendamientos nuevos",
    "Cash at 31-Dec-2026E": "Caja al 31-dic-2026E",
    "Net debt including leases at 31-Dec-2026E": "Deuda neta con arrendamientos al 31-dic-2026E",
    "169,383 reported at 30-Sep-2026 (" + BRIDGE_REL + ")": "169,383 reportada al 30-sep-2026 (" + BRIDGE_REL + ")",
    "NCI net income 9M26": "Utilidad de los minoritarios 9M26",
    XBRL + " [310000] p.13": XBRL + " [310000] p.13",
    "Cost of equity (v3 route B, reference)": "Costo del capital (v3 ruta B, referencia)",
    "valuation/valuation-outputs.yaml cost_of_equity; team confirms the WACC":
        "valuation/valuation-outputs.yaml cost_of_equity; el equipo confirma el WACC",
    "NCI growth in perpetuity": "Crecimiento perpetuo de los minoritarios",
    "Team 2026-10-26: 0%, the NCI book base shrinks": "Equipo 2026-10-26: 0%, la base en libros de los minoritarios se encoge",
    "NCI at market value = 9M26 x 4 / 3 / (ke - g)": "Minoritarios a valor de mercado = 9M26 x 4 / 3 / (ke - g)",
    "Team 2026-10-26/27: 9M26 run-rate, same value at 31-Dec-2026":
        "Equipo 2026-10-26/27: utilidad 9M26 anualizada, mismo valor al 31-dic-2026",
    "Debt-like items excluding NCI, held at 30-Sep-2026": "Partidas tipo deuda sin minoritarios, a su valor del 30-sep-2026",
    "Investments in associates, held at 30-Sep-2026 (added to EV)": "Inversión en asociadas, a su valor del 30-sep-2026 (se suma)",
    XBRL + " [210000] p.11; team 2026-10-27: held, not the projected balance":
        XBRL + " [210000] p.11; equipo 2026-10-27: se mantiene, no el saldo proyectado",
    "Claims at 31-Dec-2026E = net debt incl. leases + NCI at market + debt-like - associates":
        "Total a restar al 31-dic-2026E = deuda neta con arrendamientos + minoritarios a mercado + partidas tipo deuda - asociadas",
    "Equity value at 31-Dec-2026 = DCF enterprise value - this line":
        "Valor del capital al 31-dic-2026 = valor de la empresa del DCF - este renglón",
    "Issued shares (millions), 30-Sep-2026": "Acciones emitidas (millones), 30-sep-2026",
    XBRL + " [700000] p.24": XBRL + " [700000] p.24",
    "Repurchased shares held (millions), 30-Sep-2026": "Acciones recompradas en tesorería (millones), 30-sep-2026",
    "Shares outstanding at 31-Dec-2026E (millions)": "Acciones en circulación al 31-dic-2026E (millones)",
    "Team 2026-10-23: no buybacks in 4Q26, so the 30-Sep-2026 count holds":
        "Equipo 2026-10-23: sin recompras en 4T26, así que se mantiene el número del 30-sep-2026",
    "Days from the valuation date (16-Oct-2026) to 31-Dec-2026": "Días de la fecha de valuación (16-oct-2026) al 31-dic-2026",
    "Valuation date: 16-Oct-2026 BMV close (valuation/valuation.yaml)":
        "Fecha de valuación: cierre BMV del 16-oct-2026 (valuation/valuation.yaml)",
    "Discount factor to 16-Oct-2026 = 1 / (1 + ke) ^ (days / 365)": "Factor de descuento al 16-oct-2026 = 1 / (1 + ke) ^ (días / 365)",
    "Equity value at 16-Oct-2026 = (EV at 31-Dec-2026 - claims) x this factor; no dividends fall in the 76 days":
        "Valor del capital al 16-oct-2026 = (valor de la empresa al 31-dic-2026 - total a restar) x este factor; no hay dividendos en los 76 días",
}


def es_strings():
    """Spanish strings, including the build-time note of the debt-like row."""
    _, cell = bridge_debt_like()
    en = "Team 2026-10-27: held at the 30-Sep-2026 value; live formula in %s %s, read at build time" % (BRIDGE_REL, cell)
    es = "Equipo 2026-10-27: a su valor del 30-sep-2026; fórmula viva en %s %s, leída al construir" % (BRIDGE_REL, cell)
    return dict(ES, **{en: es})
