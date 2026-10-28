"""Audit-fix copy of the statements workbook (2026-10-26): Spanish strings and the Audit_fixes trace sheet.

Used only by build_statements_workbook_auditfix.py. Adds to statements_es at import; the original builder never imports it.
"""
from openpyxl.styles import Alignment

import statements_es

statements_es.SHEETS["Audit_fixes"] = "Ajustes_auditoría"

_H05 = "(leases grow with sales, valuation basis v2 2026-10-17)"
_H05_ES = "(los arrendamientos crecen con las ventas, base de valuación v2 2026-10-17)"
statements_es.ES.update({
    # H-01
    "Audit fix H-01: integration by nature FY2026E (1H26 x 2), so the IS and the expenses by nature agree":
        "Ajuste de auditoría H-01: integración por naturaleza FY2026E (1S26 x 2), para que el ER y los gastos por naturaleza coincidan",
    "Audit fix H-01: FY2026E general + other expenses (option A) less D&A; non-recurring lines from 1H26 x 2, the rest by the FY2025 mix of the recurring lines":
        "Ajuste de auditoría H-01: gastos generales + otros gastos FY2026E (opción A) menos D&A; los renglones no recurrentes salen del 1S26 x 2 "
        "y el resto se reparte con la mezcla FY2025 de los renglones recurrentes",
    "Multi-employer pension plan revaluation 1H26": "Revaluación de planes de pensiones multipatronales 1S26",
    "Integration expenses 1H26": "Gastos de integración 1S26",
    "Restructuring expenses 1H26": "Gastos de reestructura 1S26",
    "Audit fix H-01: BMV Jun-26 p.81, expenses by nature (MXN thousands -14,032)":
        "Ajuste de auditoría H-01: BMV jun-26 p.81, gastos por naturaleza (miles de pesos -14,032)",
    "Audit fix H-01: BMV Jun-26 p.81, expenses by nature (MXN thousands 366,681)":
        "Ajuste de auditoría H-01: BMV jun-26 p.81, gastos por naturaleza (miles de pesos 366,681)",
    "Audit fix H-01: BMV Jun-26 p.82, expenses by nature (MXN thousands 1,708,999; 1H25 1,763,831)":
        "Ajuste de auditoría H-01: BMV jun-26 p.82, gastos por naturaleza (miles de pesos 1,708,999; 1S25 1,763,831)",
    "Audit fix H-01: 1H26 actual x 2 (second half = first half); replaces the FY2025 mix":
        "Ajuste de auditoría H-01: real 1S26 x 2 (segundo semestre = primero); sustituye la mezcla FY2025",
    "Audit fix H-01: non-D&A lines less the non-recurring lines, by the FY2025 mix of the recurring lines":
        "Ajuste de auditoría H-01: renglones sin D&A menos los no recurrentes, con la mezcla FY2025 de los renglones recurrentes",
    # H-03
    "Audit fix H-03: FY2026E share of the PP&E part x PP&E depreciation (rate x opening PP&E); the 9M26 non-cash items are not projected":
        "Ajuste de auditoría H-03: participación FY2026E en la parte de PP&E x depreciación de PP&E (tasa x PP&E inicial); "
        "las partidas no monetarias del 9M26 no se proyectan",
    "Audit fixes H-03, H-05: right-of-use depreciation grows with net sales; amortization held at FY2025; PP&E part = its FY2026E share x PP&E depreciation":
        "Ajustes de auditoría H-03 y H-05: la depreciación de derechos de uso crece con las ventas; amortización en el nivel FY2025; "
        "parte de PP&E = su participación FY2026E x depreciación de PP&E",
    # H-04
    "Average cost of debt reported by Bimbo (incl. derivatives), reference":
        "Costo promedio de la deuda reportado por Bimbo (incl. derivados), referencia",
    "Gross debt at 31-Dec-2025": "Deuda bruta al 31-dic-2025",
    "Balance sheet FY2025 (EEFF25)": "Balance general FY2025 (EEFF25)",
    "FX factor on interest FY2026E": "Factor cambiario de intereses FY2026E",
    "Same formula as FY2027E, with the FY2026E USD/MXN average": "Misma fórmula que FY2027E, con el promedio USD/MXN FY2026E",
    "Cost of debt applied FY2027E-FY2030E (FY2026E run-rate, all-in net)":
        "Costo de la deuda aplicado FY2027E-FY2030E (ritmo FY2026E, neto total)",
    "Audit fix H-04: the rate that, in the FY2027E formula, reproduces FY2026E net interest (9M26 + 4Q26 = 3Q26) less lease interest; "
    "net of interest income, which the IS sets to 0 from FY2026E":
        "Ajuste de auditoría H-04: la tasa que, en la fórmula FY2027E, reproduce el interés neto FY2026E (9M26 + 4T26 = 3T26) menos el interés "
        "de arrendamientos; neta de ingresos por intereses, que el ER deja en 0 desde FY2026E",
    "Audit fixes H-04, H-05: average gross debt x cost of debt (FY2026E run-rate) x FX factor + lease interest growing with net sales":
        "Ajustes de auditoría H-04 y H-05: deuda bruta promedio x costo de la deuda (ritmo FY2026E) x factor cambiario + interés de "
        "arrendamientos que crece con las ventas",
    # H-05
    "FY2025 (EEFF25 p.41); held in FY2026E; audit fix H-05: from FY2027E right-of-use assets, lease liabilities, their depreciation and interest grow with net sales":
        "FY2025 (EEFF25 p.41); se mantiene en FY2026E; ajuste de auditoría H-05: desde FY2027E los derechos de uso, el pasivo por "
        "arrendamiento, su depreciación y su interés crecen con las ventas",
    "Right-of-use depreciation FY2027E": "Depreciación de derechos de uso FY2027E",
    "Lease interest FY2027E": "Interés de arrendamientos FY2027E",
    "Audit fix H-05: FY2025 level x net sales FY2027E / FY2026E " + _H05:
        "Ajuste de auditoría H-05: nivel FY2025 x ventas FY2027E / FY2026E " + _H05_ES,
    "Audit fix H-05: FY2025 level x net sales of the year / FY2026E":
        "Ajuste de auditoría H-05: nivel FY2025 x ventas del año / FY2026E",
    # README
    "MXN millions. AUDIT-FIX COPY built by analisis-dcf-2026-10-26/build/build_statements_workbook_auditfix.py; the macro sheets are copied from "
    "analisis-dcf-2026-10-26/datos/nominal_gdp_growth_2016-2030_auditfix.xlsx (built by analisis-dcf-2026-10-26/build/build_macro_gdp_auditfix.py).":
        "Millones de pesos. COPIA CON AJUSTES DE AUDITORÍA construida con analisis-dcf-2026-10-26/build/build_statements_workbook_auditfix.py; las hojas "
        "macro se copian de analisis-dcf-2026-10-26/datos/nominal_gdp_growth_2016-2030_auditfix.xlsx (construido con analisis-dcf-2026-10-26/build/build_macro_gdp_auditfix.py).",
    "Audit fixes H-01 to H-05 applied (sheet Audit_fixes). The team's original workbook is BIMBOA_statements_2021-2030_v1.xlsx; "
    "the Methodology sheet still describes the original methods where Audit_fixes says otherwise.":
        "Ajustes de auditoría H-01 a H-05 aplicados (hoja Ajustes_auditoría). El libro original del equipo es "
        "BIMBOA_estados_2021-2030_v1_ES.xlsx; la hoja Metodología sigue describiendo los métodos originales donde Ajustes_auditoría dice otra cosa.",
})
for _y in (2028, 2029, 2030):
    statements_es.ES["Right-of-use depreciation FY%dE" % _y] = "Depreciación de derechos de uso FY%dE" % _y
    statements_es.ES["Lease interest FY%dE" % _y] = "Interés de arrendamientos FY%dE" % _y


def remap_macro(old, new):
    """Spanish strings that cite the macro file: add the same entry with the audit-fix file name."""
    for k in [k for k in statements_es.ES if old in k]:
        statements_es.ES[k.replace(old, new)] = statements_es.ES[k].replace(old, new)


# (finding, original method, this copy, Assumptions keys to show live, source) - each as (en, es)
FIXES = [
    (("H-01", "H-01"),
     ("FY2026E expenses by nature split with the FY2025 mix: restructuring 3,846, integration 532, MEPP 27 (4,465 non-recurring with the "
      "Other share), all removed from FY2027E; integration on the IS 539.8 (a second figure).",
      "Gastos FY2026E por naturaleza repartidos con la mezcla FY2025: reestructura 3,846, integración 532, MEPP 27 (4,465 no recurrentes con la "
      "parte de Otros), todos eliminados desde FY2027E; integración en el ER 539.8 (una segunda cifra)."),
     ("Restructuring, integration and MEPP FY2026E = 1H26 actual x 2; the rest of the non-D&A lines is split by the FY2025 mix of the "
      "recurring lines; the IS integration line uses the same figure.",
      "Reestructura, integración y MEPP FY2026E = real 1S26 x 2; el resto de los renglones sin D&A se reparte con la mezcla FY2025 de los "
      "renglones recurrentes; la línea de integración del ER usa la misma cifra."),
     ("p_rs", "p_ig", "p_mp", "a_int", "em27"),
     ("BMV Jun-26 p.81-82 [sourced]; second half = first half [assumption]",
      "BMV jun-26 p.81-82 [sourced]; segundo semestre = primero [assumption]")),
    (("H-02", "H-02"),
     ("KOR, KAZ, MAR and UKR 2026 nominal GDP growth from the stored Capital IQ levels (22.4% / 28.1% / 10.5% / 22.3%, likely series "
      "breaks): EAA GDP 2026 7.86%, EAA organic ratio 0.561.",
      "Crecimiento del PIB nominal 2026 de KOR, KAZ, MAR y UKR con los niveles guardados de Capital IQ (22.4% / 28.1% / 10.5% / 22.3%, "
      "probables saltos de serie): PIB EAA 2026 7.86%, ratio orgánico EAA 0.561."),
     ("Their 2026 growth = (1 + real GDP growth) x (1 + CPI) - 1, both Capital IQ forecasts (macro sheet Data, columns L-M). 2027-2030 "
      "GDP is unchanged; only the EAA ratio moves.",
      "Su crecimiento 2026 = (1 + crecimiento real) x (1 + inflación) - 1, ambos pronósticos de Capital IQ (hoja macro Datos, columnas L-M). "
      "El PIB 2027-2030 no cambia; solo se mueve el ratio de EAA."),
     ("rt_eaa", "go_eaa", "s27"),
     ("Capital IQ realGDPGrowth and consumerPriceInflation [sourced]; CPI as proxy for the GDP deflator [assumption]",
      "Capital IQ realGDPGrowth y consumerPriceInflation [sourced]; inflación como aproximación del deflactor del PIB [assumption]")),
    (("H-03", "H-03"),
     ("PP&E part of D&A moved with PP&E depreciation from the FY2026E level, which includes the 9M26 impairment and other non-cash items "
      "(EBITDA - EBIT): PP&E depreciation applied FY2027E 17,399 vs the decided rate x opening PP&E 16,731.",
      "La parte de PP&E de la D&A se movía con la depreciación de PP&E desde el nivel FY2026E, que incluye deterioros y otras partidas no "
      "monetarias del 9M26 (UAFIDA - UAIIR): depreciación de PP&E aplicada FY2027E 17,399 contra la decidida, tasa x PP&E inicial, 16,731."),
     ("From FY2027E the PP&E part of D&A = rate x opening PP&E (10.54%), split between cost of sales and expenses by its FY2026E share; "
      "no non-cash items projected. FY2026E unchanged.",
      "Desde FY2027E la parte de PP&E de la D&A = tasa x PP&E inicial (10.54%), repartida entre costo de ventas y gastos con su participación "
      "FY2026E; sin partidas no monetarias proyectadas. FY2026E sin cambio."),
     ("ppd27", "dp27", "q27_da", "da27"),
     ("PP&E roll-forward EEFF25 p.38 [sourced]; team decision on the rate (memory.md 2026-10-23)",
      "Movimiento de PP&E EEFF25 p.38 [sourced]; decisión del equipo sobre la tasa (memory.md 2026-10-23)")),
    (("H-04", "H-04"),
     ("Interest FY2027E-FY2030E = average debt x 6.5% (Bimbo's reported cost incl. derivatives) x FX factor + lease interest.",
      "Intereses FY2027E-FY2030E = deuda promedio x 6.5% (costo reportado por Bimbo, incl. derivados) x factor cambiario + interés de arrendamientos."),
     ("Rate = the FY2026E run-rate that reproduces FY2026E net interest in the same formula (all-in, net of interest income). The auditor's "
      "7.47% is gross; the 955 of other financial costs FY2025 were offset by 897 of interest income, which the model sets to 0.",
      "Tasa = el ritmo FY2026E que reproduce el interés neto FY2026E con la misma fórmula (total, neto de ingresos por intereses). El 7.47% "
      "del auditor es bruto; los 955 de otros costos financieros FY2025 los compensaban 897 de ingresos por intereses, que el modelo deja en 0."),
     ("kd_rep", "kd", "int27"),
     ("Releases 1T26-3T26 net interest [sourced]; 4Q26 = 3Q26 [assumption]; no FCFF effect",
      "Intereses netos de los reportes 1T26-3T26 [sourced]; 4T26 = 3T26 [assumption]; sin efecto en el FCFF")),
    (("H-05", "H-05"),
     ("Right-of-use assets, lease liabilities, right-of-use depreciation (6,142) and lease interest (2,137) held at FY2025 in every year "
      "(decision 2026-10-23), against the valuation basis v2 (2026-10-17): leases grow with revenue.",
      "Derechos de uso, pasivo por arrendamiento, depreciación de derechos de uso (6,142) e interés de arrendamientos (2,137) fijos en el "
      "nivel FY2025 todos los años (decisión 2026-10-23), contra la base de valuación v2 (2026-10-17): crecen con los ingresos."),
     ("From FY2027E all four grow with net sales (year / FY2026E). New leases are a non-cash addition; lease payments = depreciation + "
      "interest less the asset growth not matched by liability growth, so Checks close. FY2026E unchanged.",
      "Desde FY2027E los cuatro crecen con las ventas (año / FY2026E). Los nuevos arrendamientos entran como adición no monetaria; pagos de "
      "arrendamiento = depreciación + interés menos el crecimiento del activo que no acompaña al pasivo, para que Cuadres cierre. FY2026E sin cambio."),
     ("rou27", "li27", "rou30", "li30"),
     ("EEFF25 p.41 and p.43 [sourced]; growth with sales [assumption], valuation basis v2",
      "EEFF25 p.41 y p.43 [sourced]; crecimiento con ventas [assumption], base de valuación v2")),
]
HEAD = [("Finding", "Hallazgo"), ("Original workbook", "Libro original"), ("This copy", "Esta copia"),
        ("Rows in Assumptions (live value)", "Renglones en Supuestos (valor vivo)"), ("Source and tag", "Fuente y etiqueta"),
        ("Value now", "Valor actual")]


def build_fix_sheet(wb, rn, lang, sn, t, TITLE, NOTE, HDR, FILL_HDR, BOLD, BLACK):
    """One row per audit fix, with live links to the Assumptions rows it changes."""
    i = 0 if lang == "en" else 1
    ws = wb.create_sheet(sn("Audit_fixes"), 1)
    ws["A1"] = ("Grupo Bimbo - Audit fixes H-01 to H-05 applied in this copy (AI proposal, 2026-10-26)",
                "Grupo Bimbo - Ajustes de auditoría H-01 a H-05 aplicados en esta copia (propuesta de la IA, 2026-10-26)")[i]
    ws["A1"].font = TITLE
    ws["A2"] = ("Audit: model/audit_statements_v1_2026-10-25_ES.md. Change log with before/after figures: analisis-dcf-2026-10-26/notas/audit_fixes_v1_2026-10-26_ES.md. "
                "The team decides which fixes to keep.",
                "Auditoría: model/audit_statements_v1_2026-10-25_ES.md. Bitácora con cifras antes/después: analisis-dcf-2026-10-26/notas/audit_fixes_v1_2026-10-26_ES.md. "
                "El equipo decide qué ajustes conserva.")[i]
    ws["A2"].font = NOTE
    for j, h in enumerate(HEAD, 1):
        c = ws.cell(4, j, h[i])
        c.font = HDR
        c.fill = FILL_HDR
    an = sn("Assumptions")
    wrap = Alignment(wrap_text=True, vertical="top")
    r = 5
    for fid, orig, new, keys, src in FIXES:
        for j, v in enumerate((fid[i], orig[i], new[i], None, src[i]), 1):
            c = ws.cell(r, j, v)
            c.font = BOLD if j == 1 else BLACK
            c.alignment = wrap
        for n_, k in enumerate(keys):
            rr = r + n_
            ws.cell(rr, 4, "=%s!A%d" % (an, rn[k])).alignment = wrap
            c = ws.cell(rr, 6, "=%s!B%d" % (an, rn[k]))
            c.number_format = '#,##0.0000' if k in ("rt_eaa", "kd", "kd_rep", "go_eaa", "em27") else '#,##0'
        r += max(len(keys), 1) + 1
    for col, w in zip("ABCDEF", (9, 60, 60, 46, 40, 14)):
        ws.column_dimensions[col].width = w
